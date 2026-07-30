"""Build the no-write Selection1 historical claim package and review workbook."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.historical_selection1_claim_backfill import (  # noqa: E402
    build_selection1_canonical_package,
    claim_fields_from_snapshot,
)


def _write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _claim_columns(snapshot: dict[str, Any] | None, period: str | None) -> dict[str, Any]:
    fields = claim_fields_from_snapshot(snapshot, period)
    return {
        "salesperson": fields.get("salesperson"),
        "claim_flag": fields.get("claim_flag"),
        "daily_sales": fields.get("daily_sales"),
        "reject_reason": fields.get("reject_reason"),
        "feedback_summary": fields.get("feedback_summary"),
        "note": fields.get("note"),
    }


def _append_header(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _fit_columns(sheet) -> None:
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
        sheet.column_dimensions[letter].width = max(width, 12)


def _write_review_workbook(path: Path, package: dict[str, Any], source_name: str) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    _append_header(summary, ["项目", "数量", "说明"])
    counts = package["counts"]
    summary.append(["源数据行", counts["raw_rows"], "开发0414期至开发0721期 + 财根拆分期；旧开发0727期-财根本次未纳入"])
    summary.append(["可入库档案行", counts["archive_rows"], "同源完全重复或空值互补已合并"])
    summary.append(["可入库认领/不认领事实", counts["claim_facts"], "有销售员即保留一条认领或不认领关系"])
    summary.append(["认领", counts["claims"], "正数认领单销"])
    summary.append(["不认领", counts["rejections"], "否、拒绝理由、单销栏文字、空值或0"])
    summary.append(["待业务核对组", counts["unresolved_groups"], "同源同身份存在非空业务字段冲突，未纳入本次入库"])
    summary.append(["待业务核对原始行", counts["unresolved_rows"], "请按来源行核对后反馈"])

    facts = workbook.create_sheet("可入库认领事实")
    fact_headers = [
        "业务期", "国家", "站点", "主SKU", "子SKU", "销售员", "认领结果", "认领单销",
        "不认领理由", "销售反馈", "来源文件", "来源Sheet", "来源行", "认领列",
    ]
    _append_header(facts, fact_headers)
    for item in package["claim_facts"]:
        facts.append([
            item.get("business_period"), item.get("country"), item.get("site"), item.get("main_sku"), item.get("sub_sku"),
            item.get("salesperson_name"), "已认领" if item.get("outcome") == "claim" else "不认领",
            item.get("claim_daily_sales"), item.get("reject_reason"), item.get("feedback_summary"),
            item.get("source_file"), item.get("source_sheet"), item.get("source_row"), item.get("claim_column"),
        ])

    review = workbook.create_sheet("待业务核对")
    review_headers = ["业务期", "国家", "站点", "主SKU", "子SKU", "来源Sheet", "来源行", "冲突字段", "处理意见"]
    _append_header(review, review_headers)
    for item in package["unresolved"]:
        review.append([
            item.get("business_period"), item.get("country"), item.get("site"), item.get("main_sku"), item.get("sub_sku"),
            item.get("source_sheet"), ", ".join(str(row) for row in item.get("source_rows") or []),
            json.dumps(item.get("conflicting_fields") or {}, ensure_ascii=False),
            "核对后保留正确值；未确认前本组不导入",
        ])

    raw = workbook.create_sheet("核对原始行")
    raw_headers = [
        "业务期", "国家", "站点", "主SKU", "子SKU", "主SKU名称", "子SKU名称", "来源Sheet", "来源行",
        "主销售员", "是否认领", "认领单销", "不认领理由", "销售反馈", "备注",
    ]
    _append_header(raw, raw_headers)
    for group in package["unresolved"]:
        period = group.get("business_period")
        for row in group.get("source_records") or []:
            claim = _claim_columns(row.get("snapshot"), period)
            raw.append([
                period, group.get("country"), group.get("site"), group.get("main_sku"), group.get("sub_sku"),
                row.get("main_sku_name"), row.get("sub_sku_name"), group.get("source_sheet"), row.get("source_row"),
                claim["salesperson"], claim["claim_flag"], claim["daily_sales"], claim["reject_reason"],
                claim["feedback_summary"], claim["note"],
            ])

    excluded = workbook.create_sheet("本次不纳入")
    _append_header(excluded, ["来源文件", "范围", "原因"])
    excluded.append([source_name, "旧开发0727期-财根", "已由财根拆分期替代，本次不导入、不恢复认领"])

    for sheet in workbook.worksheets:
        _fit_columns(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a no-write canonical Selection1 historical claim package.")
    parser.add_argument("--input", type=Path, required=True, help="selection1 parsed JSON output")
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    source = json.loads(args.input.read_text(encoding="utf-8"))
    package = build_selection1_canonical_package(source.get("rows") or [])
    args.output_dir.mkdir(parents=True, exist_ok=True)
    _write_json(args.output_dir / "archive_rows.json", package["archive_rows"])
    _write_json(args.output_dir / "claim_facts.json", package["claim_facts"])
    _write_json(args.output_dir / "unresolved.json", package["unresolved"])
    _write_json(args.output_dir / "unresolved_claim_facts.json", package["unresolved_claim_facts"])
    _write_json(
        args.output_dir / "import_rows.json",
        {
            "source_file": source.get("source_file"),
            "source_sha256": source.get("source_sha256"),
            "rows": package["archive_rows"],
        },
    )
    manifest = {
        "source": {key: source.get(key) for key in ("source_file", "source_sha256")},
        "scope": "开发0414期至开发0721期 + 开发0624/0701/0708/0715期-财根",
        "excluded": ["旧开发0727期-财根"],
        "counts": package["counts"],
        "files": {
            "archive_rows": "archive_rows.json",
            "import_rows": "import_rows.json",
            "claim_facts": "claim_facts.json",
            "unresolved": "unresolved.json",
            "unresolved_claim_facts": "unresolved_claim_facts.json",
        },
    }
    _write_json(args.output_dir / "manifest.json", manifest)
    _write_review_workbook(args.output_dir / "选品1历史认领-待业务核对.xlsx", package, str(source.get("source_file") or ""))
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
