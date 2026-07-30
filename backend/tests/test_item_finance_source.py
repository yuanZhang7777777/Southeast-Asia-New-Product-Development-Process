from __future__ import annotations

import json
from hashlib import sha256
from pathlib import Path
from re import sub
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.item_finance_source import country_from_shop, item_id_text, read_item_finance_period, reconcile_item_candidates, validate_period_inputs


def test_resets_malformed_sheet_dimension_before_reading_rows(tmp_path: Path) -> None:
    source = tmp_path / "period.xlsx"
    build_workbook(source, [(123, "MAIN-A", "Shopee-1TH", "2026-04-16")])
    rewrite_dimension(source, "A1:A2")

    records = read_item_finance_period("2026-04-16", source)

    assert [(record["item_id"], record["country"]) for record in records] == [("123", "TH")]


def test_reads_display_values_from_merged_cells(tmp_path: Path) -> None:
    source = tmp_path / "merged.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "finance"
    worksheet.append(["ITEMID", "主SKU", "店铺", "审核时间"])
    worksheet.append(["I-1", "MAIN-A", "Shopee-1TH", "2026-04-16"])
    worksheet.append([None, "MAIN-B", "Shopee-2TH", "2026-04-17"])
    worksheet.append([None, "MAIN-C", "Shopee-3TH", "2026-04-18"])
    worksheet.merge_cells("A2:A4")
    workbook.save(source)

    records = read_item_finance_period("2026-04-16", source)

    assert [(record["source_reference"]["source_row"], record["item_id"], record["main_sku"]) for record in records] == [
        (2, "I-1", "MAIN-A"),
        (3, "I-1", "MAIN-B"),
        (4, "I-1", "MAIN-C"),
    ]


def test_rejects_source_without_all_required_columns(tmp_path: Path) -> None:
    source = tmp_path / "missing.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["ITEMID", "主SKU", "店铺"])
    workbook.save(source)

    with pytest.raises(ValueError, match="审核时间"):
        read_item_finance_period("2026-04-16", source)


def test_extracts_country_from_real_shop_suffixes() -> None:
    assert country_from_shop("Shopee-569TH") == "TH"
    assert country_from_shop("Shopee-11VN") == "VN"
    assert country_from_shop("Shopee-101PH") == "PH"
    assert country_from_shop("Shopee-569TH extra") is None

def test_preserves_item_id_as_text_without_float_suffix_or_scientific_notation() -> None:
    assert item_id_text("00123") == "00123"
    assert item_id_text(123) == "123"
    assert item_id_text(123.0) == "123"
    assert item_id_text(1.23e10) == "12300000000"


def test_excludes_blank_summary_rows_and_invalid_country_suffixes(tmp_path: Path) -> None:
    source = tmp_path / "period.xlsx"
    build_workbook(
        source,
        [
            ("I-1", "MAIN", "Shopee-1TH", "2026-04-16"),
            ("I-2", "", "Shopee-1TH", "2026-04-16"),
            ("I-3", "MAIN", "", "2026-04-16"),
            ("I-4", "MAIN", "Shopee-1TH extra", "2026-04-16"),
            ("I-5", "MAIN", "Alpha TX", "2026-04-16"),
        ],
    )

    records = read_item_finance_period("2026-04-16", source)

    assert [(record["item_id"], record["country"]) for record in records] == [("I-1", "TH")]


def test_deduplicates_across_periods_and_keeps_all_source_traceability(tmp_path: Path) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    build_workbook(first, [("I-1", "MAIN", "Shopee-1TH", "2026-04-16")])
    build_workbook(second, [("I-1", "MAIN", "Shopee-1TH", "2026-04-23")])

    result = reconcile_item_candidates({"records": []}, [("2026-04-16", first), ("2026-04-23", second)])

    candidate = result["candidates"][0]
    assert candidate["first_period"] == "2026-04-16"
    assert candidate["last_period"] == "2026-04-23"
    assert candidate["row_count"] == 2
    assert candidate["audit_times"] == ["2026-04-16", "2026-04-23"]
    assert candidate["source_references"] == [
        {
            "period": "2026-04-16",
            "source_file": first.name,
            "source_sha256": sha256(first.read_bytes()).hexdigest(),
            "source_sheet": "finance",
            "source_row": 2,
        },
        {
            "period": "2026-04-23",
            "source_file": second.name,
            "source_sha256": sha256(second.read_bytes()).hexdigest(),
            "source_sheet": "finance",
            "source_row": 2,
        },
    ]


def test_classifies_watchlist_records_and_unions_ambiguous_main_sku_candidates(tmp_path: Path) -> None:
    source = tmp_path / "period.xlsx"
    build_workbook(
        source,
        [
            ("I-1", "MAIN-A", "Shopee-1TH", "2026-04-16"),
            ("I-2", "MAIN-B", "Shopee-2TH", "2026-04-16"),
            ("I-3", "MAIN-C", "Shopee-3TH", "2026-04-16"),
        ],
    )
    watchlist = {
        "records": [
            {"country": "TH", "child_sku": "SUB-1", "main_skus": ["MAIN-A"]},
            {"country": "TH", "child_sku": "SUB-2", "main_skus": ["MAIN-B", "MAIN-C"]},
            {"country": "TH", "child_sku": "SUB-3", "main_skus": ["MISSING"]},
        ]
    }

    result = reconcile_item_candidates(watchlist, [("2026-04-16", source)])

    assert [(record["child_sku"], record["classification"]) for record in result["matches"]] == [
        ("SUB-1", "unique"),
        ("SUB-2", "multiple"),
        ("SUB-3", "unmatched"),
    ]
    assert [candidate["item_id"] for candidate in result["matches"][1]["candidates"]] == ["I-2", "I-3"]


def test_reports_item_owner_and_main_sku_shop_conflicts(tmp_path: Path) -> None:
    source = tmp_path / "period.xlsx"
    build_workbook(
        source,
        [
            ("I-1", "MAIN-A", "Shopee-1TH", "2026-04-16"),
            ("I-1", "MAIN-B", "Shopee-2TH", "2026-04-16"),
            ("I-2", "MAIN-A", "Shopee-1TH", "2026-04-16"),
        ],
    )

    result = reconcile_item_candidates({"records": []}, [("2026-04-16", source)])

    assert result["anomalies"]["item_multiple_owners"] == [
        {
            "item_id": "I-1",
            "owners": [
                {"country": "TH", "main_sku": "MAIN-A", "shop": "Shopee-1TH"},
                {"country": "TH", "main_sku": "MAIN-B", "shop": "Shopee-2TH"},
            ],
        }
    ]
    assert result["anomalies"]["main_sku_shop_multiple_items"] == [
        {"country": "TH", "main_sku": "MAIN-A", "shop": "Shopee-1TH", "item_ids": ["I-1", "I-2"]}
    ]


def build_workbook(path: Path, rows: list[tuple[object, object, object, object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "finance"
    worksheet.append(["ITEMID", "主SKU", "店铺", "审核时间"])
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def rewrite_dimension(path: Path, dimension: str) -> None:
    with ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    files["xl/worksheets/sheet1.xml"] = sub(
        rb'<dimension ref="[^"]+"/>', f'<dimension ref="{dimension}"/>'.encode(), files["xl/worksheets/sheet1.xml"]
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)


def test_accepts_matching_manifest_and_filter_period(tmp_path: Path) -> None:
    source, manifest = build_manifested_workbook(tmp_path, "0416-0422")

    validate_period_inputs([("0416-0422", source)], {"0416-0422": manifest})


def test_rejects_missing_manifest_by_default(tmp_path: Path) -> None:
    source, _ = build_manifested_workbook(tmp_path, "0416-0422")

    with pytest.raises(ValueError, match="missing manifest"):
        validate_period_inputs([("0416-0422", source)], {})


def test_rejects_manifest_period_and_hash_mismatches(tmp_path: Path) -> None:
    source, manifest = build_manifested_workbook(tmp_path, "0416-0422")
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    payload["period"] = "0423-0429"
    manifest.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(ValueError, match="period"):
        validate_period_inputs([("0416-0422", source)], {"0416-0422": manifest})

    payload["period"] = "0416-0422"
    payload["sha256"] = "0" * 64
    manifest.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError, match="sha256"):
        validate_period_inputs([("0416-0422", source)], {"0416-0422": manifest})


def test_rejects_wrong_filter_sheet_period(tmp_path: Path) -> None:
    source, manifest = build_manifested_workbook(tmp_path, "0416-0422", filter_period="0423-0429")

    with pytest.raises(ValueError, match="filter period"):
        validate_period_inputs([("0416-0422", source)], {"0416-0422": manifest})


def build_manifested_workbook(tmp_path: Path, period: str, filter_period: str | None = None) -> tuple[Path, Path]:
    source = tmp_path / "item_finance.xlsx"
    workbook = Workbook()
    data = workbook.active
    data.title = "ItemID财务数据八部"
    data.append(["ITEMID", "主SKU", "店铺", "审核时间"])
    data.append(["I-1", "MAIN", "Shopee-1TH", "2026-04-16"])
    filters = workbook.create_sheet("ItemID财务数据八部_过滤条件")
    filters.cell(row=1, column=2).value = f'【时间周期】属于"{filter_period or period}"'
    workbook.save(source)
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "status": "ok",
                "period": period,
                "payload_widget_name": "ItemID财务数据八部",
                "period_replaced_count": 1,
                "file": source.name,
                "size": source.stat().st_size,
                "sha256": sha256(source.read_bytes()).hexdigest(),
                "sheets": [
                    {"name": "ItemID财务数据八部", "row_count": 1},
                    {"name": "ItemID财务数据八部_过滤条件", "row_count": 0},
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return source, manifest
