import base64
import json
import os
import sys
import uuid
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_workflow.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest  # noqa: E402
from cryptography.hazmat.primitives import serialization  # noqa: E402
from cryptography.hazmat.primitives.asymmetric import padding, rsa  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from app import finebi_auto_pull as fap  # noqa: E402
from app import historical_finebi_import as hfi  # noqa: E402
from app import models  # noqa: E402
from app.config import Settings, get_settings  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402

client = TestClient(app)

TEST_PRIVATE_KEY = rsa.generate_private_key(public_exponent=65537, key_size=2048)
TEST_PUBLIC_KEY_B64 = base64.b64encode(
    TEST_PRIVATE_KEY.public_key().public_bytes(
        serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo
    )
).decode("ascii")
# 生产登录页格式：Dec.system = JSON.parse('{\"encryptionType\":0,...}')
DEC_SYSTEM_JSON = json.dumps(
    {"encryptionType": 0, "encryptionKey": TEST_PUBLIC_KEY_B64}, separators=(",", ":")
)
DEC_SYSTEM_ESCAPED = DEC_SYSTEM_JSON.replace("\\", "\\\\").replace('"', '\\"')
LOGIN_PAGE = (
    "<html><head><title>数据决策系统</title></head><body>"
    f"<script>Dec.system = JSON.parse('{DEC_SYSTEM_ESCAPED}');</script>"
    "<form>login</form></body></html>"
)


def decrypt_password(value: str) -> str:
    return "".join(
        TEST_PRIVATE_KEY.decrypt(base64.b64decode(part), padding.PKCS1v15()).decode("utf-8")
        for part in value.split("---")
    )


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def teardown_function() -> None:
    for name in ("AUTH_REQUIRED", "AUTH_SECRET_KEY", "FINEBI_AUTO_PULL_ENABLED"):
        os.environ.pop(name, None)
    os.environ["AUTH_REQUIRED"] = "false"
    get_settings.cache_clear()


def finebi_workbook_bytes(rows: list[list], headers: list[str] | None = None) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ItemID财务数据八部"
    worksheet.append(
        headers
        or ["ITEMID", "主SKU", "店铺", "审核时间", "订单特性", "总收入", "商品数量", "订单量", "一次毛利",
            "SKU一次毛利率", "SKU合单率", "商品成本", "预估重量"]
    )
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def finebi_row(item: str, main_sku: str, shop: str, revenue: float, orders: int, gross: float) -> list:
    return [item, main_sku, shop, "2026-07-23", "普通", revenue, 1, orders, gross, 0.2, 0.5, 10, 0.3]


class FakeFineBI:
    def __init__(
        self,
        *,
        login_ok: bool = True,
        export_status: int = 200,
        download_status: int = 200,
        download_body: bytes = b"",
        download_content_type: str = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ) -> None:
        self.login_ok = login_ok
        self.export_status = export_status
        self.download_status = download_status
        self.download_body = download_body
        self.download_content_type = download_content_type
        self.calls: list[tuple[str, str]] = []
        self.login_payload: dict | None = None
        self.login_headers: dict | None = None
        self.export_url: str | None = None
        self.export_body: bytes | None = None
        self.download_url: str | None = None
        self._cookies: set[str] = set()

    def request(self, method: str, url: str, *, data=None, headers=None, timeout=120) -> fap.HttpResponse:
        self.calls.append((method, url))
        if url.endswith(fap.LOGIN_PATH) and method == "GET":
            return fap.HttpResponse(200, {"content-type": "text/html"}, LOGIN_PAGE.encode("utf-8"))
        if url.endswith(fap.LOGIN_PATH) and method == "POST":
            self.login_payload = json.loads(data.decode("utf-8"))
            self.login_headers = dict(headers or {})
            if self.login_ok:
                self._cookies.add("fine_auth_token")
                return fap.HttpResponse(200, {"content-type": "application/json"}, b'{"errorCode":""}')
            return fap.HttpResponse(200, {"content-type": "application/json"}, b'{"errorCode":"INVALID_USER","errorMsg":"denied"}')
        if f"{fap.DOWNLOAD_PATH}/" in url and method == "GET":
            self.download_url = url
            return fap.HttpResponse(
                self.download_status, {"content-type": self.download_content_type}, self.download_body
            )
        if fap.EXPORT_PATH in url and method == "POST":
            self.export_url = url
            self.export_body = data
            # 生产实测：创建导出可能返回 Content-Length: 0 的空 body
            return fap.HttpResponse(self.export_status, {"content-length": "0"}, b"")
        raise AssertionError(f"unexpected request {method} {url}")

    def cookie_names(self) -> set[str]:
        return set(self._cookies)


def make_settings(tmp_path: Path, *, enabled: bool = True, payload: str = '{"demo": true}') -> Settings:
    payload_file = tmp_path / "itemid_finance.json"
    payload_file.write_text(payload, encoding="utf-8")
    return Settings(
        finebi_base_url="https://finebi.internal.test:8443",
        finebi_username="ops-user",
        finebi_password="ops-pass",
        finebi_report_id="RPT-1",
        finebi_payload_file=str(payload_file),
        finebi_auto_pull_enabled=enabled,
    )


def test_pull_week_success_writes_file_and_backs_up_existing(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    fake = FakeFineBI(download_body=finebi_workbook_bytes([finebi_row("10000000001", "MAINA", "Shopee-101PH", 100.0, 3, 20.0)]))
    target_dir = tmp_path / "finebi_live"
    target_dir.mkdir()
    (target_dir / "0723-0729.xlsx").write_bytes(b"old-content")

    path = fap.pull_week("0723-0729", settings=settings, target_dir=target_dir, session=fake)

    assert path == target_dir / "0723-0729.xlsx"
    assert (target_dir / "0723-0729.xlsx.bak").read_bytes() == b"old-content"
    workbook = load_workbook(path, read_only=True)
    assert workbook.active.cell(row=2, column=1).value == "10000000001"
    workbook.close()
    # 登录 payload：密码按 Dec.system 公钥 RSA 加密（可用私钥解回），固定字段与生产脚本一致
    assert fake.login_payload["username"] == "ops-user"
    assert fake.login_payload["password"] != "ops-pass"
    assert decrypt_password(fake.login_payload["password"]) == "ops-pass"
    assert fake.login_payload["validity"] == -2
    assert fake.login_payload["sliderToken"] == ""
    assert fake.login_payload["origin"] == ""
    assert fake.login_payload["encrypted"] is True
    # 登录 headers 与生产脚本一致
    assert fake.login_headers["Content-Type"] == "application/json"
    assert fake.login_headers["X-Requested-With"] == "XMLHttpRequest"
    assert fake.login_headers["Origin"] == "https://finebi.internal.test:8443"
    assert fake.login_headers["Referer"] == "https://finebi.internal.test:8443/webroot/decision/login"
    assert fake.login_headers["transEncryptLevel"] == "1"
    assert "Chrome" in fake.login_headers["User-Agent"]
    # 导出 URL 携带 reportId/entryType/operationId，body 是配置文件原文
    query = parse_qs(urlsplit(fake.export_url).query)
    assert query["reportId"] == ["RPT-1"]
    assert query["entryType"] == ["6"]
    assert query["operationId"][0]
    assert fake.export_body == b'{"demo": true}'
    # 下载 URL 用 operationId 路径 + 运行时 sessionID
    assert f"{fap.DOWNLOAD_PATH}/{query['operationId'][0]}?" in fake.download_url


def test_operation_and_session_ids_regenerated_each_run(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    body = finebi_workbook_bytes([finebi_row("10000000001", "MAINA", "Shopee-101PH", 100.0, 3, 20.0)])
    seen: list[tuple[str, str]] = []
    for _ in range(2):
        fake = FakeFineBI(download_body=body)
        fap.pull_week("0716-0722", settings=settings, target_dir=tmp_path / "live", session=fake)
        operation_id = parse_qs(urlsplit(fake.export_url).query)["operationId"][0]
        session_id = parse_qs(urlsplit(fake.download_url).query)["sessionID"][0]
        uuid.UUID(operation_id)
        uuid.UUID(session_id)
        seen.append((operation_id, session_id))
    assert seen[0][0] != seen[1][0]
    assert seen[0][1] != seen[1][1]


def test_login_failure_raises_with_error_json_and_page_snippet(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    fake = FakeFineBI(login_ok=False)
    with pytest.raises(fap.FineBIPullError) as excinfo:
        fap.pull_week("0723-0729", settings=settings, target_dir=tmp_path / "live", session=fake)
    message = str(excinfo.value)
    assert "登录失败" in message
    assert "fine_auth_token" in message
    assert "denied" in message
    assert "INVALID_USER" in message
    assert "数据决策系统" in message
    assert not (tmp_path / "live").exists()


def test_parse_dec_system_unescapes_json_parse_literal() -> None:
    assert fap._parse_dec_system(LOGIN_PAGE) == {
        "encryptionType": 0,
        "encryptionKey": TEST_PUBLIC_KEY_B64,
    }


def test_missing_dec_system_raises_with_page_snippet() -> None:
    with pytest.raises(fap.FineBIPullError) as excinfo:
        fap._parse_dec_system("<html><title>数据决策系统</title><body>no config</body></html>")
    assert "Dec.system" in str(excinfo.value)
    assert "数据决策系统" in str(excinfo.value)


def test_unsupported_encryption_type_rejected() -> None:
    with pytest.raises(fap.FineBIPullError, match="Unsupported encryption type"):
        fap._encrypt_password("pass", {"encryptionType": 1, "encryptionKey": TEST_PUBLIC_KEY_B64})
    with pytest.raises(fap.FineBIPullError, match="Unsupported encryption type"):
        fap._encrypt_password("pass", {"encryptionType": 0})


def test_password_chunked_every_50_chars_joined_with_triple_dash(monkeypatch) -> None:
    monkeypatch.setattr(fap, "_load_public_key", lambda key: "PUB")
    monkeypatch.setattr(fap, "_rsa_encrypt", lambda key, chunk: b"<" + chunk + b">")
    password = "".join(chr(ord("a") + index % 26) for index in range(120))
    result = fap._encrypt_password(password, {"encryptionType": 0, "encryptionKey": "KEY"})
    parts = [base64.b64decode(part).decode("utf-8") for part in result.split("---")]
    assert parts == [f"<{password[0:50]}>", f"<{password[50:100]}>", f"<{password[100:120]}>"]
    short = fap._encrypt_password("abc", {"encryptionType": 0, "encryptionKey": "KEY"})
    assert "---" not in short
    assert base64.b64decode(short) == b"<abc>"


def test_public_key_pem_wraps_base64_at_64_chars() -> None:
    pem = fap._public_key_pem(TEST_PUBLIC_KEY_B64)
    assert pem.startswith("-----BEGIN PUBLIC KEY-----\n")
    assert pem.endswith("\n-----END PUBLIC KEY-----\n")
    body_lines = pem.strip().splitlines()[1:-1]
    assert all(len(line) <= 64 for line in body_lines)
    assert "".join(body_lines) == TEST_PUBLIC_KEY_B64
    # 转出的 PEM 能被 cryptography 加载
    assert fap._load_public_key(TEST_PUBLIC_KEY_B64).key_size == 2048


def test_export_created_with_empty_body_is_not_failure(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    fake = FakeFineBI(download_body=finebi_workbook_bytes([finebi_row("10000000001", "MAINA", "Shopee-101PH", 1.0, 1, 1.0)]))
    path = fap.pull_week("0723-0729", settings=settings, target_dir=tmp_path / "live", session=fake)
    assert path.is_file()


def test_download_returning_html_fails_and_writes_nothing(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    fake = FakeFineBI(
        download_body="<html><body>会话已过期，请重新登录</body></html>".encode("utf-8"),
        download_content_type="text/html; charset=utf-8",
    )
    with pytest.raises(fap.FineBIPullError) as excinfo:
        fap.pull_week("0723-0729", settings=settings, target_dir=tmp_path / "live", session=fake)
    assert "下载失败" in str(excinfo.value)
    assert "会话已过期" in str(excinfo.value)
    assert not (tmp_path / "live" / "0723-0729.xlsx").exists()


def test_header_mismatch_rejected_before_saving(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    fake = FakeFineBI(
        download_body=finebi_workbook_bytes([], headers=["ITEMID", "店铺", "总收入", "订单量", "一次毛利"])
    )
    with pytest.raises(fap.FineBIPullError) as excinfo:
        fap.pull_week("0723-0729", settings=settings, target_dir=tmp_path / "live", session=fake)
    assert "表头" in str(excinfo.value)
    assert "主SKU" in str(excinfo.value)
    assert not (tmp_path / "live" / "0723-0729.xlsx").exists()


def test_invalid_week_label_rejected(tmp_path: Path) -> None:
    settings = make_settings(tmp_path)
    with pytest.raises(ValueError):
        fap.pull_week("2026-07", settings=settings, target_dir=tmp_path / "live", session=FakeFineBI())
    with pytest.raises(ValueError):
        fap.pull_week("1345-1352", settings=settings, target_dir=tmp_path / "live", session=FakeFineBI())


def test_pull_and_import_ingests_week_even_when_in_excluded_periods(tmp_path: Path) -> None:
    assert "0723-0729" in hfi.EXCLUDED_PERIODS
    settings = make_settings(tmp_path)
    target_dir = tmp_path / "finebi_live"
    body = finebi_workbook_bytes([finebi_row("10000000001", "MAINX", "Shopee-101PH", 100.0, 3, 20.0)])
    with SessionLocal() as db:
        db.add(
            models.NewProductOpportunity(
                source_type="selection1_developer_claim_feedback",
                main_sku="MAINX",
                sub_sku="MAINX-1",
                country="PH",
            )
        )
        db.commit()

    with SessionLocal() as db:
        report = fap.pull_and_import(
            db, "0723-0729", imported_by="tester", settings=settings, target_dir=target_dir,
            session=FakeFineBI(download_body=body), year=2026,
        )
        db.commit()

    assert report["periods"] == ["0723-0729"]
    assert report["excluded_periods"] == []
    assert report["apply_counts"]["listings_created"] == 1
    assert report["apply_counts"]["weeks_created"] == 1
    with SessionLocal() as db:
        week = db.query(models.ItemObservationPeriod).one()
        batch = db.query(models.ImportBatch).one()
        assert week.source_snapshot["period"] == "0723-0729"
        assert week.metrics_origin == "finebi_live"
        assert batch.imported_by == "tester"

    # 不传 include_period 时，默认行为保持不变：0723-0729 仍被排除
    _, periods, skipped = hfi.read_finebi_periods(target_dir, 2026)
    assert periods == []
    assert skipped == ["0723-0729"]

    # 幂等：重复拉取不产生新纪录
    with SessionLocal() as db:
        second = fap.pull_and_import(
            db, "0723-0729", imported_by="tester", settings=settings, target_dir=target_dir,
            session=FakeFineBI(download_body=body), year=2026,
        )
        db.commit()
    assert second["apply_counts"]["listings_created"] == 0
    assert second["apply_counts"]["listings_reused"] == 1
    assert second["apply_counts"]["weeks_created"] == 0
    assert (target_dir / "0723-0729.xlsx.bak").is_file()


def admin_headers() -> dict[str, str]:
    with SessionLocal() as db:
        db.add(models.RoleMapping(name="Admin", role="super_admin", dingtalk_user_id="dt-admin", enabled=True))
        db.commit()
    login = client.post("/auth/dingtalk/login", json={"dingtalk_user_id": "dt-admin"})
    assert login.status_code == 200
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def enable_auth(finebi_enabled: bool) -> None:
    os.environ["AUTH_REQUIRED"] = "true"
    os.environ["AUTH_SECRET_KEY"] = "test-auth-secret"
    os.environ["FINEBI_AUTO_PULL_ENABLED"] = "true" if finebi_enabled else "false"
    get_settings.cache_clear()


def test_endpoint_returns_403_with_switch_name_when_disabled() -> None:
    enable_auth(finebi_enabled=False)
    headers = admin_headers()
    response = client.post("/admin/finebi/pull", headers=headers, json={"week_label": "0723-0729"})
    assert response.status_code == 403
    assert "FINEBI_AUTO_PULL_ENABLED" in response.json()["detail"]


def test_endpoint_requires_super_admin() -> None:
    enable_auth(finebi_enabled=True)
    with SessionLocal() as db:
        db.add(models.RoleMapping(name="Manager A", role="manager", dingtalk_user_id="dt-manager", enabled=True))
        db.commit()
    login = client.post("/auth/dingtalk/login", json={"dingtalk_user_id": "dt-manager"})
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    response = client.post("/admin/finebi/pull", headers=headers, json={"week_label": "0723-0729"})
    assert response.status_code == 403


def test_endpoint_rejects_bad_week_label_without_touching_network() -> None:
    enable_auth(finebi_enabled=True)
    headers = admin_headers()
    response = client.post("/admin/finebi/pull", headers=headers, json={"week_label": "bad-label"})
    assert response.status_code == 400
    assert "MMDD-MMDD" in response.json()["detail"]


def test_endpoint_returns_pull_and_import_report(monkeypatch) -> None:
    enable_auth(finebi_enabled=True)
    headers = admin_headers()
    captured: dict[str, object] = {}

    def fake_pull_and_import(db, week_label, imported_by=None, *, settings=None, **kwargs):
        captured["week_label"] = week_label
        captured["imported_by"] = imported_by
        assert settings.finebi_auto_pull_enabled is True
        return {"week_label": week_label, "apply_counts": {"listings_created": 1}}

    monkeypatch.setattr(fap, "pull_and_import", fake_pull_and_import)
    response = client.post("/admin/finebi/pull", headers=headers, json={"week_label": "0723-0729"})
    assert response.status_code == 200
    assert response.json() == {"week_label": "0723-0729", "apply_counts": {"listings_created": 1}}
    assert captured == {"week_label": "0723-0729", "imported_by": "Admin"}


def test_endpoint_maps_protocol_error_to_502(monkeypatch) -> None:
    enable_auth(finebi_enabled=True)
    headers = admin_headers()

    def failing_pull(*args, **kwargs):
        raise fap.FineBIPullError("FineBI 下载失败：返回的不是 Excel/ZIP 文件")

    monkeypatch.setattr(fap, "pull_and_import", failing_pull)
    response = client.post("/admin/finebi/pull", headers=headers, json={"week_label": "0723-0729"})
    assert response.status_code == 502
    assert "下载失败" in response.json()["detail"]
