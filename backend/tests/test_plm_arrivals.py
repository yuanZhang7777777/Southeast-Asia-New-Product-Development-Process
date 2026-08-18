from pathlib import Path

from openpyxl import Workbook

from app.config import Settings, get_settings
from app.plm_arrivals import parse_plm_arrival_preview

GROUP_EIGHT = "\u96c6\u56e2\u516b\u90e8"
GROUP_ONE = "\u96c6\u56e2\u4e00\u90e8"
SALES_A = "\u9500\u552eA"
SALES_B = "\u9500\u552eB"
SALES_C = "\u9500\u552eC"
PH = "\u83f2\u5f8b\u5bbe"
TH = "\u6cf0\u56fd"
PH_WAREHOUSE = "\u83f2\u5f8b\u5bbe\u4ed3"
TH_WAREHOUSE = "\u6cf0\u56fd\u4ed3"


def test_parse_plm_arrival_preview_uses_first_listing_window_without_group_filter(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)

    preview = parse_plm_arrival_preview(workbook_path, "2026-07-12", bloc_name=GROUP_EIGHT)

    assert preview["row_count"] == 2
    assert preview["new_arrival_count"] == 2
    assert preview["restock_count"] == 0
    assert preview["unknown_count"] == 0
    assert preview["by_salesperson"] == [
        {"salesperson_name": SALES_A, "new_arrival_count": 1, "restock_count": 0, "unknown_count": 0, "total_count": 1},
        {"salesperson_name": SALES_C, "new_arrival_count": 1, "restock_count": 0, "unknown_count": 0, "total_count": 1},
    ]
    assert [item["arrival_type"] for item in preview["items"]] == ["new_arrival", "new_arrival"]
    assert [item["source_row"] for item in preview["items"]] == [2, 5]
    assert preview["items"][0]["product_name"] == "new"
    assert preview["items"][1]["product_name"] == "other group"
    assert preview["items"][1]["bloc_name"] == GROUP_ONE


def test_parse_plm_arrival_preview_includes_two_days_around_target_day(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm-first-listing.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品名称", "子SKU", "主SKU", "销售员", "集团", "海外仓", "国家", "最后一次入库时间", "首次上架时间"])
    sheet.append(["target minus two", "OLD-1", "M1", SALES_A, GROUP_EIGHT, PH_WAREHOUSE, PH, "2026-07-12 10:00:00", "2026-07-10 00:00:00"])
    sheet.append(["first listing target", "NEW-1", "M2", SALES_A, GROUP_EIGHT, PH_WAREHOUSE, PH, "2026-07-11 10:00:00", "2026-07-12 00:00:00"])
    sheet.append(["target plus two", "NEW-2", "M3", SALES_A, GROUP_EIGHT, PH_WAREHOUSE, PH, "2026-07-12 10:00:00", "2026-07-14 00:00:00"])
    sheet.append(["target minus three", "OLD-2", "M4", SALES_A, GROUP_EIGHT, PH_WAREHOUSE, PH, "2026-07-12 10:00:00", "2026-07-09 00:00:00"])
    workbook.save(workbook_path)

    preview = parse_plm_arrival_preview(workbook_path, "2026-07-12", bloc_name=GROUP_EIGHT)

    assert preview["row_count"] == 3
    assert preview["new_arrival_count"] == 3
    assert [item["product_name"] for item in preview["items"]] == [
        "target minus two",
        "first listing target",
        "target plus two",
    ]


def test_plm_arrival_preview_endpoint_reads_cached_workbook(tmp_path: Path) -> None:
    from fastapi.testclient import TestClient

    from app.main import app

    workbook_path = tmp_path / "plm-2026-07-12.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["\u5546\u54c1\u540d\u79f0", "\u5b50SKU", "\u4e3bSKU", "\u9500\u552e\u5458", "\u96c6\u56e2", "\u6d77\u5916\u4ed3", "\u56fd\u5bb6", "\u6700\u540e\u4e00\u6b21\u5165\u5e93\u65f6\u95f4", "\u9996\u6b21\u4e0a\u67b6\u65f6\u95f4"])
    sheet.append(["new", "NEW-1", "M1", SALES_A, GROUP_EIGHT, PH_WAREHOUSE, PH, "2026-07-12 10:00:00", "2026-07-12 00:00:00"])
    workbook.save(workbook_path)

    app.dependency_overrides[get_settings] = lambda: Settings(plm_cache_dir=str(tmp_path), plm_bloc_name=GROUP_EIGHT)
    try:
        response = TestClient(app).get("/arrival/plm-preview?date=2026-07-12")
    finally:
        app.dependency_overrides.pop(get_settings, None)

    assert response.status_code == 200
    assert response.json()["new_arrival_count"] == 1
    assert response.json()["items"][0]["product_name"] == "new"


def build_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["\u5546\u54c1\u540d\u79f0", "\u5b50SKU", "\u4e3bSKU", "\u9500\u552e\u5458", "\u96c6\u56e2", "\u6d77\u5916\u4ed3", "\u56fd\u5bb6", "\u6700\u540e\u4e00\u6b21\u5165\u5e93\u65f6\u95f4", "\u9996\u6b21\u4e0a\u67b6\u65f6\u95f4"])
    sheet.append(["new", "NEW-1", "M1", SALES_A, GROUP_EIGHT, PH_WAREHOUSE, PH, "2026-07-12 10:00:00", "2026-07-12 00:00:00"])
    sheet.append(["restock", "OLD-1", "M2", SALES_A, GROUP_EIGHT, PH_WAREHOUSE, PH, "2026-07-12 11:00:00", "2026-07-01 00:00:00"])
    sheet.append(["unknown", "UNK-1", "M3", SALES_B, GROUP_EIGHT, TH_WAREHOUSE, TH, "2026-07-12 12:00:00", None])
    sheet.append(["other group", "OTHER-1", "M4", SALES_C, GROUP_ONE, TH_WAREHOUSE, TH, "2026-07-12 12:00:00", "2026-07-12 00:00:00"])
    sheet.append(["old day", "OLD-DAY", "M5", SALES_A, GROUP_EIGHT, PH_WAREHOUSE, PH, "2026-07-11 12:00:00", "2026-07-01 00:00:00"])
    workbook.save(path)
