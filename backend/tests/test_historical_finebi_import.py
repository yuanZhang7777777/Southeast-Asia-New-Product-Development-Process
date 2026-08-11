import os
import sys
from datetime import date
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_workflow.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402

from app import historical_finebi_import as hfi  # noqa: E402
from app import models, services  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def build_workbench(path: Path, rows: list[list]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "精品流程"
    for _ in range(3):
        worksheet.append([None] * 38)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def workbench_row(
    main_sku: str,
    shop: str,
    item: str,
    *,
    salesperson: str = "运营A",
    country: str = "PH",
    week1_metrics: list | None = None,
    week1_review: list | None = None,
    summary: str | None = None,
) -> list:
    row = [None] * 38
    row[0] = "2026-06-01"
    row[1] = main_sku
    row[2] = country
    row[3] = salesperson
    row[4] = shop
    row[5] = item
    row[6] = "标题优化"
    if week1_metrics:
        row[12:16] = week1_metrics
    if week1_review:
        row[16:19] = week1_review
    if summary:
        row[37] = summary
    return row


def build_finebi(path: Path, rows: list[list]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ItemID财务数据八部"
    worksheet.append(
        ["ITEMID", "主SKU", "店铺", "审核时间", "订单特性", "总收入", "商品数量", "订单量", "一次毛利",
         "SKU一次毛利率", "SKU合单率", "商品成本", "预估重量"]
    )
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def finebi_row(item: str, main_sku: str, shop: str, revenue: float, orders: int, gross: float) -> list:
    return [item, main_sku, shop, "2026-04-03", "普通", revenue, 1, orders, gross, 0.2, 0.5, 10, 0.3]


def run_import(tmp_path: Path, workbench_rows: list[list], finebi_files: dict[str, list[list]], apply: bool = True):
    workbench_path = tmp_path / "刊登监控.xlsx"
    build_workbench(workbench_path, workbench_rows)
    finebi_dir = tmp_path / "finebi_live"
    finebi_dir.mkdir(exist_ok=True)
    for period, rows in finebi_files.items():
        build_finebi(finebi_dir / f"05_ItemID财务数据八部_{period}.xlsx", rows)
    workbench, warnings = hfi.read_workbench(workbench_path)
    finebi, periods, skipped = hfi.read_finebi_periods(finebi_dir, 2026)
    with SessionLocal() as db:
        plan = hfi.build_plan(db, workbench, finebi, 2026)
        counts = None
        if apply:
            counts = hfi.apply_plan(db, plan, source_label="test")
            db.commit()
    return plan, counts, {"warnings": warnings, "periods": periods, "skipped": skipped}


def test_shared_item_aggregation_review_and_idempotency(tmp_path: Path) -> None:
    workbench_rows = [
        workbench_row("MAINA", "Shopee-101PH", "10000000001",
                      week1_review=["2026-06-08", "利润款", "调价"], summary="总体表现稳定"),
        workbench_row("MAINB", "Shopee-101PH", "10000000001"),
    ]
    finebi_files = {
        "0402-0408": [
            finebi_row("10000000001", "MAINA", "Shopee-101PH", 100.0, 3, 20.0),
            finebi_row("10000000001", "MAINB", "Shopee-101PH", 50.0, 2, 10.0),
        ],
        "0409-0415": [finebi_row("10000000001", "MAINA", "Shopee-101PH", 80.0, 4, 16.0)],
        "0723-0729": [finebi_row("10000000001", "MAINA", "Shopee-101PH", 999.0, 99, 999.0)],
    }
    plan, counts, meta = run_import(tmp_path, workbench_rows, finebi_files)

    assert meta["skipped"] == []
    assert plan["summary"]["scoped_items"] == 1
    assert plan["summary"]["shared_items"] == 1
    assert counts["listings_created"] == 1
    assert counts["bindings_created"] == 2

    with SessionLocal() as db:
        listing = db.query(models.ListingRecord).one()
        bindings = db.query(models.ListingSkuBinding).order_by(models.ListingSkuBinding.main_sku).all()
        weeks = (
            db.query(models.ItemObservationPeriod)
            .order_by(models.ItemObservationPeriod.week_number)
            .all()
        )
        assert listing.source_type == "history_finebi"
        assert listing.is_shared_item is True
        assert listing.representative_rule == "lexical_first"
        assert listing.main_sku == "MAINA"
        assert listing.first_period_start == date(2026, 4, 2)
        assert [binding.main_sku for binding in bindings] == ["MAINA", "MAINB"]
        assert {binding.binding_source for binding in bindings} == {"history_finebi"}
        assert [week.week_number for week in weeks] == [1, 2, 3]
        week1, week2, week3 = weeks
        assert (week1.order_count, week1.total_revenue, week1.gross_profit_amount) == (5, 150.0, 30.0)
        assert week1.period_start == date(2026, 4, 2)
        assert week1.record_source == "history_finebi"
        assert week1.metrics_origin == "finebi_live"
        assert week1.product_positioning == "利润款"
        assert week2.order_count == 4
        assert week3.period_start == date(2026, 7, 23)
        assert week3.four_week_summary == "总体表现稳定"
        assert db.query(models.NotificationLog).count() == 0
        assert db.query(models.SalesClaimForecast).count() == 0

    plan2, counts2, _ = run_import(tmp_path, workbench_rows, finebi_files)
    assert counts2["listings_created"] == 0
    assert counts2["listings_reused"] == 1
    assert counts2["bindings_created"] == 0
    assert counts2["weeks_created"] == 0
    assert counts2["weeks_updated"] == 0
    with SessionLocal() as db:
        assert db.query(models.ListingRecord).count() == 1
        assert db.query(models.ListingSkuBinding).count() == 2
        assert db.query(models.ItemObservationPeriod).count() == 3

    with SessionLocal() as db:
        workbench_view = services.list_listing_workbench(db)
        summary_view = services.listing_summary(db, "MAINA")
    assert workbench_view["listing_records"] == []
    assert workbench_view["period_rows"] == []
    assert len(summary_view["listing_records"]) == 1
    assert summary_view["listing_records"][0]["is_shared_item"] is True
    assert {row["record_source"] for row in summary_view["period_rows"]} == {"history_finebi"}
    assert len(summary_view["period_rows"]) == 3


def test_finebi_metrics_use_latest_four_periods_with_data(tmp_path: Path) -> None:
    workbench_rows = [workbench_row("MAINF", "Shopee-105PH", "10000000006")]
    finebi_files = {
        "0402-0408": [finebi_row("10000000006", "MAINF", "Shopee-105PH", 0.0, 0, 0.0)],
        "0409-0415": [finebi_row("10000000006", "MAINF", "Shopee-105PH", 10.0, 1, 2.0)],
        "0416-0422": [finebi_row("10000000006", "MAINF", "Shopee-105PH", 20.0, 2, 4.0)],
        "0423-0429": [finebi_row("10000000006", "MAINF", "Shopee-105PH", 30.0, 3, 6.0)],
        "0430-0506": [finebi_row("10000000006", "MAINF", "Shopee-105PH", 40.0, 4, 8.0)],
        "0723-0729": [finebi_row("10000000006", "MAINF", "Shopee-105PH", 999.0, 99, 999.0)],
    }
    plan, counts, meta = run_import(tmp_path, workbench_rows, finebi_files)

    assert meta["skipped"] == []
    assert len(plan["items"][0]["finebi_weeks"]) == 4
    assert [week["period"] for week in plan["items"][0]["finebi_weeks"]] == [
        "0416-0422",
        "0423-0429",
        "0430-0506",
        "0723-0729",
    ]
    assert counts["weeks_created"] == 4
    with SessionLocal() as db:
        weeks = db.query(models.ItemObservationPeriod).order_by(models.ItemObservationPeriod.week_number).all()
        assert [week.period_start for week in weeks] == [
            date(2026, 4, 16),
            date(2026, 4, 23),
            date(2026, 4, 30),
            date(2026, 7, 23),
        ]
        assert weeks[-1].order_count == 99


def test_sparse_finebi_periods_keep_real_zero_week_without_placeholders(tmp_path: Path) -> None:
    workbench_rows = [workbench_row("MAINZ", "Shopee-108PH", "10000000009")]
    finebi_files = {
        "0716-0722": [finebi_row("10000000009", "MAINZ", "Shopee-108PH", 0.0, 0, 0.0)],
        "0723-0729": [finebi_row("10000000009", "MAINZ", "Shopee-108PH", 20.0, 2, 4.0)],
    }
    plan, counts, _ = run_import(tmp_path, workbench_rows, finebi_files)

    assert [week["period"] for week in plan["items"][0]["finebi_weeks"]] == ["0716-0722", "0723-0729"]
    assert counts["weeks_created"] == 2
    with SessionLocal() as db:
        weeks = db.query(models.ItemObservationPeriod).order_by(models.ItemObservationPeriod.week_number).all()
        assert [week.week_number for week in weeks] == [1, 2]
        assert weeks[0].period_start == date(2026, 7, 16)
        assert weeks[0].order_count == 0
        assert weeks[1].period_start == date(2026, 7, 23)
def test_workbench_only_item_creates_binding_without_observation_period(tmp_path: Path) -> None:
    workbench_rows = [
        workbench_row("MAINC", "Shopee-102TH", "10000000002", country="TH",
                      week1_metrics=[6, 300.0, 60.0, 0.2], week1_review=["2026-06-08", "引流款", "加广告"]),
    ]
    plan, counts, _ = run_import(tmp_path, workbench_rows, {})

    assert plan["summary"]["workbench_fallback_week_rows"] == 0
    assert counts["weeks_created"] == 0
    with SessionLocal() as db:
        listing = db.query(models.ListingRecord).one()
        assert db.query(models.ItemObservationPeriod).count() == 0
        assert listing.country == "TH"
        assert listing.first_period_start is None


def test_finebi_only_item_scoped_by_known_main_sku(tmp_path: Path) -> None:
    with SessionLocal() as db:
        db.add(
            models.NewProductOpportunity(
                source_type="selection1_developer_claim_feedback",
                main_sku="MAIND",
                sub_sku="MAIND-1",
                country="VN",
            )
        )
        db.commit()
        opportunity_id = db.query(models.NewProductOpportunity).one().id

    finebi_files = {
        "0402-0408": [
            finebi_row("10000000003", "MAIND", "Shopee-103VN", 10.0, 1, 2.0),
            finebi_row("10000000004", "UNKNOWN-SKU", "Shopee-103VN", 10.0, 1, 2.0),
        ],
    }
    plan, counts, _ = run_import(tmp_path, [], finebi_files)

    assert plan["summary"]["scoped_items"] == 1
    assert plan["summary"]["finebi_items_out_of_scope"] == 1
    with SessionLocal() as db:
        listing = db.query(models.ListingRecord).one()
        binding = db.query(models.ListingSkuBinding).one()
        assert listing.item == "10000000003"
        assert listing.salesperson_name == ""
        assert binding.opportunity_id == opportunity_id


def test_history_import_attaches_to_existing_platform_listing(tmp_path: Path) -> None:
    with SessionLocal() as db:
        listing = models.ListingRecord(
            id=models.new_id(),
            source_group_key="seed-task",
            source_claim_ids=[],
            source_type="selection1",
            main_sku="MAINE",
            salesperson_name="运营B",
            shop="Shopee-104PH",
            item="10000000005",
            listing_strategy="平台策略",
            first_period_start=date(2026, 7, 16),
            first_period_end=date(2026, 7, 22),
            representative_rule="single_binding",
        )
        db.add(listing)
        db.add(
            models.ItemObservationPeriod(
                id=models.new_id(),
                listing_record_id=listing.id,
                week_number=1,
                period_start=date(2026, 7, 16),
                period_end=date(2026, 7, 22),
            )
        )
        db.commit()
        platform_listing_id = listing.id

    finebi_files = {"0402-0408": [finebi_row("10000000005", "MAINE", "Shopee-104PH", 10.0, 1, 2.0)]}
    plan, counts, _ = run_import(tmp_path, [], finebi_files)

    assert plan["summary"]["scoped_items"] == 1
    assert counts["listings_created"] == 0
    assert counts["listings_reused"] == 1
    with SessionLocal() as db:
        assert db.query(models.ListingRecord).count() == 1
        history_weeks = (
            db.query(models.ItemObservationPeriod)
            .filter_by(listing_record_id=platform_listing_id, record_source="history_finebi")
            .all()
        )
        platform_weeks = (
            db.query(models.ItemObservationPeriod)
            .filter_by(listing_record_id=platform_listing_id, record_source="platform")
            .all()
        )
        assert len(history_weeks) == 1
        assert len(platform_weeks) == 1
        assert platform_weeks[0].period_start == date(2026, 7, 16)


def test_candidate_json_import_deduplicates_shared_item_metrics() -> None:
    raw_plan = {
        "push_listing_candidates": [
            {
                "can_enter_observation": True,
                "binding_key": "Shopee-150PH|25239908809|MAINA",
                "listing_key": "Shopee-150PH|25239908809",
                "country_zh": "菲律宾",
                "push_salesperson": "运营A",
                "main_sku": "MAINA",
                "shop": "Shopee-150PH",
                "item": "25239908809",
                "strategy_raw": {"ads": "是"},
                "selected_platform_period": "开发0526期",
                "push_source": {"row": 10},
            },
            {
                "can_enter_observation": True,
                "binding_key": "Shopee-150PH|25239908809|MAINB",
                "listing_key": "Shopee-150PH|25239908809",
                "country_zh": "菲律宾",
                "push_salesperson": "运营A",
                "main_sku": "MAINB",
                "shop": "Shopee-150PH",
                "item": "25239908809",
                "strategy_raw": {"ads": "是"},
                "selected_platform_period": "开发0526期",
                "push_source": {"row": 11},
            },
        ],
        "candidate_observation_periods": [
            {
                "shop": "Shopee-150PH",
                "item": "25239908809",
                "main_sku": "MAINA",
                "main_skus": ["MAINA", "MAINB"],
                "binding_key": "Shopee-150PH|25239908809|MAINA",
                "binding_keys": ["Shopee-150PH|25239908809|MAINA", "Shopee-150PH|25239908809|MAINB"],
                "week_number": 1,
                "period_code": "0716-0722",
                "period_start": "2026-07-16",
                "period_end": "2026-07-22",
                "orders": 3,
                "revenue": 100.0,
                "gross_profit": 20.0,
                "finebi_source": {"file_name": "05_ItemID财务数据八部_0716-0722.xlsx", "rows": [8], "main_skus": ["MAINA", "MAINB"]},
            }
        ],
    }

    with SessionLocal() as db:
        plan = hfi.build_plan_from_candidate_json(db, raw_plan)
        counts = hfi.apply_plan(db, plan, source_label="candidate-json")
        db.commit()

    assert counts["listings_created"] == 1
    assert counts["bindings_created"] == 2
    assert counts["weeks_created"] == 1
    with SessionLocal() as db:
        listing = db.query(models.ListingRecord).one()
        bindings = db.query(models.ListingSkuBinding).order_by(models.ListingSkuBinding.main_sku).all()
        weeks = db.query(models.ItemObservationPeriod).all()
        assert listing.business_period == "开发0526期"
        assert listing.is_shared_item is True
        assert [binding.main_sku for binding in bindings] == ["MAINA", "MAINB"]
        assert len(weeks) == 1
        assert weeks[0].source_snapshot["binding_keys"] == [
            "Shopee-150PH|25239908809|MAINA",
            "Shopee-150PH|25239908809|MAINB",
        ]


def test_candidate_json_zero_data_binding_creates_listing_without_period() -> None:
    raw_plan = {
        "push_listing_candidates": [
            {
                "can_enter_observation": True,
                "binding_key": "Shopee-311PH|54013530289|YZH007",
                "listing_key": "Shopee-311PH|54013530289",
                "country_zh": "菲律宾",
                "push_salesperson": "陈丽妹",
                "main_sku": "YZH007",
                "shop": "Shopee-311PH",
                "item": "54013530289",
                "strategy_raw": {},
                "selected_platform_period": "开发0526期",
                "rejected_platform_periods": ["选品2-财根0526期"],
                "push_source": {"row": 442},
            }
        ],
        "candidate_observation_periods": [],
    }

    with SessionLocal() as db:
        plan = hfi.build_plan_from_candidate_json(db, raw_plan)
        counts = hfi.apply_plan(db, plan, source_label="candidate-json")
        db.commit()

    assert counts["listings_created"] == 1
    assert counts["bindings_created"] == 1
    assert counts["weeks_created"] == 0
    with SessionLocal() as db:
        listing = db.query(models.ListingRecord).one()
        assert listing.business_period == "开发0526期"
        assert listing.first_period_start is None
        assert db.query(models.ItemObservationPeriod).count() == 0


def test_history_period_with_blank_push_positioning_does_not_default_to_secondary(tmp_path: Path) -> None:
    with SessionLocal() as db:
        opportunity = models.NewProductOpportunity(
            source_type="selection1_developer_claim_feedback",
            batch="开发0703期",
            country="PH",
            main_sku="MAING",
            sub_sku="MAING-1",
        )
        db.add(opportunity)
        db.flush()
        claim = models.SalesClaimForecast(
            opportunity_id=opportunity.id,
            salesperson_name="运营C",
            claim_result="claim",
            product_positioning="稳定款",
        )
        db.add(claim)
        db.flush()
        listing = models.ListingRecord(
            id=models.new_id(),
            source_group_key="seed-task",
            source_claim_ids=[claim.id],
            source_type="selection1",
            business_period="开发0703期",
            country="PH",
            site="PH",
            main_sku="MAING",
            salesperson_name="运营C",
            shop="Shopee-106PH",
            item="10000000007",
            listing_strategy="平台策略",
            first_period_start=date(2026, 7, 16),
            first_period_end=date(2026, 7, 22),
            representative_rule="single_binding",
        )
        db.add(listing)
        db.add(
            models.ItemObservationPeriod(
                id=models.new_id(),
                listing_record_id=listing.id,
                week_number=1,
                period_start=date(2026, 7, 16),
                period_end=date(2026, 7, 22),
            )
        )
        db.commit()

    finebi_files = {"0402-0408": [finebi_row("10000000007", "MAING", "Shopee-106PH", 10.0, 1, 2.0)]}
    run_import(tmp_path, [], finebi_files)

    with SessionLocal() as db:
        summary = services.listing_summary(db, "MAING")
    rows_by_source = {row["record_source"]: row for row in summary["period_rows"]}
    assert rows_by_source["platform"]["default_product_positioning"] == "稳定款"
    assert rows_by_source["history_finebi"]["product_positioning"] is None
    assert rows_by_source["history_finebi"]["default_product_positioning"] is None


def test_history_push_listing_period_does_not_default_to_secondary_positioning() -> None:
    with SessionLocal() as db:
        opportunity = models.NewProductOpportunity(
            source_type="selection1_developer_claim_feedback",
            batch="开发0703期",
            country="PH",
            main_sku="MAINH",
            sub_sku="MAINH-1",
        )
        db.add(opportunity)
        db.flush()
        claim = models.SalesClaimForecast(
            opportunity_id=opportunity.id,
            salesperson_name="运营D",
            claim_result="claim",
            product_positioning="稳定款",
        )
        db.add(claim)
        db.flush()
        listing = models.ListingRecord(
            id=models.new_id(),
            source_group_key="history-push",
            source_claim_ids=[claim.id],
            source_type="history_push_20260727",
            business_period="开发0703期",
            country="PH",
            site="PH",
            main_sku="MAINH",
            salesperson_name="运营D",
            shop="Shopee-107PH",
            item="10000000008",
            listing_strategy="历史策略",
            first_period_start=date(2026, 7, 16),
            first_period_end=date(2026, 7, 22),
            representative_rule="history_main_sku_group",
        )
        db.add(listing)
        db.add(
            models.ItemObservationPeriod(
                id=models.new_id(),
                listing_record_id=listing.id,
                week_number=1,
                period_start=date(2026, 7, 16),
                period_end=date(2026, 7, 22),
                record_source="platform",
            )
        )
        db.commit()

    with SessionLocal() as db:
        summary = services.listing_summary(db, "MAINH")

    row = summary["period_rows"][0]
    assert row["product_positioning"] is None
    assert row["default_product_positioning"] is None
