import os
import sys
from base64 import b64decode
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_workflow.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image  # noqa: E402

from app import models  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.routers import opportunities as opportunities_router  # noqa: E402


client = TestClient(app)
UPLOAD_ROOT = Path(__file__).resolve().parents[1] / ".uploaded_sources"
PNG_BYTES = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def test_selection1_import_creates_batch_and_links_snapshots(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1.xlsx"
    build_selection1_fixture(workbook_path)

    response = client.post("/opportunities/import/selection1", json={"source_file": str(workbook_path), "source_sheet": "W27"})

    assert response.status_code == 200
    body = response.json()
    assert body["imported_count"] == 1
    assert body["import_batch_id"]

    with SessionLocal() as db:
        batch = db.get(models.ImportBatch, body["import_batch_id"])
        opportunity = db.query(models.NewProductOpportunity).one()
        snapshot = db.query(models.SourceRecordSnapshot).one()

    assert batch is not None
    assert batch.created_count == 1
    assert opportunity.import_batch_id == batch.id
    assert opportunity.current_status == "pending_assignment"
    assert snapshot.import_batch_id == batch.id
    assert snapshot.column_range == "A:BX,CC:CH"
    assert snapshot.payload["cells"]["M"] == "箱规 62*42*25cm"
    assert snapshot.payload["cells"]["AW"] == 107.98
    assert snapshot.payload["cells"]["BX"] == 20
    assert db.query(models.FlowTask).count() == 0


def test_selection1_business_period_is_separate_from_source_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_business_period.xlsx"
    build_selection1_fixture(workbook_path, sheet_name="选品1原始数据")

    response = client.post(
        "/opportunities/import/selection1",
        json={
            "source_file": str(workbook_path),
            "source_sheet": "选品1原始数据",
            "business_period": "2026年第29期",
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["source_sheet"] == "选品1原始数据"
    assert body["business_period"] == "2026年第29期"

    with SessionLocal() as db:
        batch = db.get(models.ImportBatch, body["import_batch_id"])
        opportunity = db.query(models.NewProductOpportunity).one()
        snapshot = db.query(models.SourceRecordSnapshot).one()

    assert batch is not None
    assert batch.source_sheet == "选品1原始数据"
    assert batch.business_period == "2026年第29期"
    assert opportunity.source_sheet == "选品1原始数据"
    assert opportunity.batch == "2026年第29期"
    assert snapshot.source_file == workbook_path.name
    assert snapshot.source_sheet == "选品1原始数据"
    assert snapshot.source_row == 3


def test_selection1_whitespace_business_period_defaults_to_source_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_whitespace_period.xlsx"
    build_selection1_fixture(workbook_path)

    response = client.post(
        "/opportunities/import/selection1",
        json={"source_file": str(workbook_path), "source_sheet": "W27", "business_period": "   "},
    )

    assert response.status_code == 200
    assert response.json()["business_period"] == "W27"
    with SessionLocal() as db:
        batch = db.get(models.ImportBatch, response.json()["import_batch_id"])
        opportunity = db.query(models.NewProductOpportunity).one()

    assert batch is not None
    assert batch.business_period == "W27"
    assert opportunity.batch == "W27"


def test_selection1_reimport_matches_by_business_identity_not_source_location(tmp_path: Path) -> None:
    first_path = tmp_path / "selection1_first.xlsx"
    second_path = tmp_path / "selection1_second.xlsx"
    build_selection1_fixture(first_path, sheet_name="选品1原始数据", sub_sku="SUB-A", sub_sku_name="Old name")
    build_selection1_fixture(
        second_path,
        sheet_name="选品1原始数据",
        row_index=8,
        sub_sku="SUB-A",
        sub_sku_name="Updated name",
        extra_rows=[{"sub_sku": "SUB-B", "sub_sku_name": "New name", "row_index": 9}],
    )
    payload = {"source_sheet": "选品1原始数据", "business_period": "2026年第29期"}

    first = client.post("/opportunities/import/selection1", json={**payload, "source_file": str(first_path)})
    assert first.status_code == 200
    with SessionLocal() as db:
        opportunity = db.query(models.NewProductOpportunity).filter_by(sub_sku="SUB-A").one()
        original_id = opportunity.id
        services_claim = models.SalesClaimForecast(
            opportunity_id=original_id,
            salesperson_name="Operator A",
            claim_result="claim",
            claim_daily_sales=3,
            source_column="platform",
            downstream_status="stocking_draft",
        )
        review = models.ReviewRecord(opportunity_id=original_id, review_status="approved", reviewer_name="Manager A")
        db.add_all([services_claim, review])
        db.commit()
        claim_id = services_claim.id
        review_id = review.id

    second = client.post("/opportunities/import/selection1", json={**payload, "source_file": str(second_path)})

    assert second.status_code == 200
    assert second.json()["created_count"] == 1
    assert second.json()["updated_count"] == 1
    with SessionLocal() as db:
        opportunities = db.query(models.NewProductOpportunity).order_by(models.NewProductOpportunity.sub_sku).all()
        sub_a = db.query(models.NewProductOpportunity).filter_by(sub_sku="SUB-A").one()
        claim = db.get(models.SalesClaimForecast, claim_id)
        review = db.get(models.ReviewRecord, review_id)
        snapshots = db.query(models.SourceRecordSnapshot).filter_by(opportunity_id=original_id).order_by(models.SourceRecordSnapshot.source_row).all()

    assert [item.sub_sku for item in opportunities] == ["SUB-A", "SUB-B"]
    assert sub_a.id == original_id
    assert sub_a.source_file == second_path.name
    assert sub_a.source_sheet == "选品1原始数据"
    assert sub_a.source_row == 8
    assert sub_a.sub_sku_name == "Updated name"
    assert sub_a.current_status == "pending_assignment"
    assert claim is not None
    assert claim.opportunity_id == original_id
    assert claim.downstream_status == "stocking_draft"
    assert review is not None
    assert review.opportunity_id == original_id
    assert [snapshot.source_file for snapshot in snapshots] == [first_path.name, second_path.name]
    assert [snapshot.source_row for snapshot in snapshots] == [3, 8]


def test_selection1_upload_import_uses_browser_file(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_upload.xlsx"
    build_selection1_fixture(workbook_path)

    with workbook_path.open("rb") as handle:
        response = client.post(
            "/opportunities/import/selection1/upload",
            data={"source_sheet": "W27"},
            files={"file": ("selection1.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert response.status_code == 200
    body = response.json()
    assert body["imported_count"] == 1
    assert "selection1" in body["source_file"]


def test_selection1_uploads_are_saved_in_backend_persistent_volume() -> None:
    backend_root = Path(__file__).resolve().parents[1]

    assert opportunities_router.UPLOAD_ROOT == backend_root / ".private_uploads" / "source-workbooks"


def test_selection1_upload_accepts_oss_signed_filename(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_upload.xlsx"
    build_selection1_fixture(workbook_path)

    with workbook_path.open("rb") as handle:
        response = client.post(
            "/opportunities/import/selection1/upload",
            data={"source_sheet": "W27"},
            files={
                "file": (
                    "selection1.xlsx%3FExpires=1783214557&OSSAccessKeyId=masked&Signature=masked",
                    handle,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["imported_count"] == 1
    assert body["source_file"].endswith("selection1.xlsx")


def test_excel_sheet_upload_returns_first_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "multi_sheet.xlsx"
    workbook = Workbook()
    workbook.active.title = "latest"
    workbook.create_sheet("history")
    workbook.save(workbook_path)

    with workbook_path.open("rb") as handle:
        response = client.post(
            "/opportunities/excel-sheets/upload",
            files={"file": ("multi_sheet.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert response.status_code == 200
    assert response.json() == {"sheets": ["latest", "history"], "default_sheet": "latest"}


def test_selection1_import_extracts_embedded_product_image(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_image.xlsx"
    image_path = tmp_path / "product.png"
    image_path.write_bytes(PNG_BYTES)
    build_selection1_fixture(workbook_path, image_path=image_path)

    response = client.post("/opportunities/import/selection1", json={"source_file": str(workbook_path), "source_sheet": "W27"})

    assert response.status_code == 200
    with SessionLocal() as db:
        opportunity = db.query(models.NewProductOpportunity).one()

    assert opportunity.image_url
    assert opportunity.image_url.startswith("/uploaded-sources/product-images/")
    assert (UPLOAD_ROOT / opportunity.image_url.removeprefix("/uploaded-sources/")).exists()
    assert client.get(opportunity.image_url).status_code == 200


def test_selection1_import_skips_repeated_header_rows(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_headers.xlsx"
    build_selection1_fixture(workbook_path, include_repeated_header=True)

    response = client.post("/opportunities/import/selection1", json={"source_file": str(workbook_path), "source_sheet": "W27"})

    assert response.status_code == 200
    assert response.json()["imported_count"] == 1
    with SessionLocal() as db:
        assert db.query(models.NewProductOpportunity).filter_by(main_sku="MAINSKU").count() == 0


def test_selection1_import_keeps_product_when_reason_contains_sales_total(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_reason_contains_total.xlsx"
    build_selection1_fixture(workbook_path, reason="多家竞对店铺月销合计3000以上")

    response = client.post("/opportunities/import/selection1", json={"source_file": str(workbook_path), "source_sheet": "W27"})

    assert response.status_code == 200
    assert response.json()["imported_count"] == 1
    with SessionLocal() as db:
        assert db.query(models.NewProductOpportunity).filter_by(sub_sku="SUB-001").count() == 1


def test_selection1_reimport_keeps_previous_source_snapshots(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_reimport.xlsx"
    build_selection1_fixture(workbook_path)
    payload = {"source_file": str(workbook_path), "source_sheet": "W27"}

    first = client.post("/opportunities/import/selection1", json=payload)
    second = client.post("/opportunities/import/selection1", json=payload)

    assert first.status_code == 200
    assert second.status_code == 200
    with SessionLocal() as db:
        snapshots = db.query(models.SourceRecordSnapshot).order_by(models.SourceRecordSnapshot.created_at).all()

    assert len(snapshots) == 2
    assert snapshots[0].import_batch_id == first.json()["import_batch_id"]
    assert snapshots[1].import_batch_id == second.json()["import_batch_id"]


def test_selection1_import_matches_market_and_pricing_by_headers_when_columns_shift(tmp_path: Path) -> None:
    workbook_path = tmp_path / "selection1_shifted.xlsx"
    build_selection1_shifted_fixture(workbook_path)

    response = client.post("/opportunities/import/selection1", json={"source_file": str(workbook_path), "source_sheet": "W27"})

    assert response.status_code == 200
    with SessionLocal() as db:
        opportunity = db.query(models.NewProductOpportunity).one()
        market_items = db.query(models.MarketResearchItem).order_by(models.MarketResearchItem.research_type).all()

    assert response.json()["market_research_count"] == 2
    assert {item.research_type for item in market_items} == {"新晋", "最低价"}
    assert {item.competitor_url for item in market_items} == {"https://lowest.example/item", "https://new.example/item"}
    assert {item.reference_daily_sales for item in market_items} == {18}
    assert {item.reference_price for item in market_items} == {129}
    assert opportunity.snapshot["headers_by_column"]["BA"] == ["最低价链接"]
    assert opportunity.snapshot["pricing_snapshot"]["一次毛利额RMB"] == 12.5
    assert opportunity.snapshot["pricing_snapshot"]["推广期利润率"] == "8.5%"


def build_selection1_fixture(
    path: Path,
    image_path: Path | None = None,
    include_repeated_header: bool = False,
    sheet_name: str = "W27",
    row_index: int = 3,
    sub_sku: str = "SUB-001",
    sub_sku_name: str | None = None,
    reason: str | None = None,
    extra_rows: list[dict[str, object]] | None = None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append([None] * 86)
    headers = [None] * 86
    for column_index, title in {
        1: "site",
        8: "main_sku",
        10: "sub_sku",
        82: "salesperson",
        83: "claim_result",
        84: "claim_daily_sales",
    }.items():
        headers[column_index - 1] = title
    worksheet.append(headers)
    row = [None] * 86
    for column_index, value in {
        1: "TH",
        8: "MAIN-001",
        10: "SUB-001",
        13: "箱规 62*42*25cm",
        49: 107.98,
        76: 20,
        82: "Operator A",
        83: "yes",
        84: 1,
    }.items():
        row[column_index - 1] = value
    row[9] = sub_sku
    row[8] = sub_sku_name
    row[11] = reason
    while worksheet.max_row < row_index - 1:
        worksheet.append([None] * 86)
    worksheet.append(row)
    for extra in extra_rows or []:
        extra_row = list(row)
        extra_row[9] = extra["sub_sku"]
        extra_row[8] = extra.get("sub_sku_name")
        target_row = int(extra.get("row_index", worksheet.max_row + 1))
        while worksheet.max_row < target_row - 1:
            worksheet.append([None] * 86)
        worksheet.append(extra_row)
    if include_repeated_header:
        repeated = [None] * 86
        repeated[7] = "MainSKU"
        repeated[9] = "SubSKU"
        worksheet.append(repeated)
    if image_path:
        worksheet.add_image(Image(str(image_path)), "F3")
    workbook.save(path)


def build_selection1_shifted_fixture(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "W27"
    worksheet.append([None] * 100)
    headers = [None] * 100
    for column_index, title in {
        1: "站点",
        8: "主SKU",
        10: "子SKU",
        53: "最低价链接",
        54: "售价1(PHP）",
        55: "月销1",
        56: "新晋链接",
        57: "售价3(PHP）",
        58: "月销3",
        59: "参考单销",
        60: "稳定期定价 / （PHP）",
        61: "一次毛利额 / （人民币）",
        62: "推广期定价",
        63: "推广期利润率",
    }.items():
        headers[column_index - 1] = title
    worksheet.append(headers)
    row = [None] * 100
    for column_index, value in {
        1: "PH",
        8: "MAIN-SHIFT",
        10: "SUB-SHIFT",
        53: "https://lowest.example/item",
        54: 99,
        55: 300,
        56: "https://new.example/item",
        57: 109,
        58: 90,
        59: 18,
        60: 129,
        61: 12.5,
        62: 119,
        63: "8.5%",
    }.items():
        row[column_index - 1] = value
    worksheet.append(row)
    workbook.save(path)
