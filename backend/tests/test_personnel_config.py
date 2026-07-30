import os
import sys
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_workflow.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402

from app import models  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def test_personnel_config_imports_assignment_profiles(tmp_path: Path) -> None:
    workbook_path = tmp_path / "集团八部销售员信息表.xlsx"
    build_personnel_fixture(workbook_path)

    from app.personnel_importer import import_personnel_config

    with SessionLocal() as db:
        result = import_personnel_config(db, workbook_path)
        db.commit()
        profiles = db.query(models.OperatorAssignmentProfile).order_by(models.OperatorAssignmentProfile.operator_name).all()
        role_mappings = db.query(models.RoleMapping).order_by(models.RoleMapping.name).all()

    assert result == {"created_count": 2, "updated_count": 0, "skipped_count": 0}
    assert [
        (
            profile.operator_name,
            profile.role,
            profile.operator_level,
            profile.business_type,
            profile.key_site,
            profile.key_category1,
            profile.key_category2,
            profile.enabled,
        )
        for profile in profiles
    ] == [
        ("冯卓宏", "组员", "P1", "精品（海外仓）", "PH", "汽摩配", "家居厨卫", True),
        ("陈丽妹", "组员", "P1", "精品（海外仓）", "PH", "家居厨卫", "商办工业", True),
    ]
    assert [(item.name, item.role, item.enabled) for item in role_mappings] == [
        ("冯卓宏", "operator", True),
        ("陈丽妹", "operator", True),
    ]


def test_operator_profile_admin_crud_exposes_only_assignment_fields() -> None:
    from fastapi.testclient import TestClient

    from app.main import app

    with SessionLocal() as db:
        user = models.User(name="陈丽妹", enabled=True)
        db.add(user)
        db.flush()
        db.add(models.RoleMapping(user_id=user.id, name="陈丽妹", role="operator", enabled=True))
        db.commit()

    client = TestClient(app)
    create_response = client.post(
        "/admin/operator-profiles",
        json={
            "operator_name": "陈丽妹",
            "key_site": "PH",
            "key_category1": "家居厨卫",
            "key_category2": "商办工业",
            "enabled": True,
            "role": "组员",
            "operator_level": "P1",
            "business_type": "精品（海外仓）",
        },
    )

    assert create_response.status_code == 200
    body = create_response.json()
    assert set(body) == {
        "id",
        "operator_name",
        "key_site",
        "key_category1",
        "key_category2",
        "key_categories",
        "assignment_priority",
        "display_order",
        "enabled",
    }
    assert body["operator_name"] == "陈丽妹"

    update_response = client.patch(
        f"/admin/operator-profiles/{body['id']}",
        json={"key_site": "TH", "key_category1": "汽摩配", "enabled": False},
    )
    assert update_response.status_code == 200
    assert update_response.json()["key_site"] == "TH"
    assert update_response.json()["enabled"] is False

    list_response = client.get("/admin/operator-profiles")
    assert list_response.status_code == 200
    assert list_response.json()[0]["key_category1"] == "汽摩配"

    delete_response = client.delete(f"/admin/operator-profiles/{body['id']}")
    assert delete_response.status_code == 204
    assert client.get("/admin/operator-profiles").json() == []


def test_operator_profile_rejects_more_than_six_level1_category_groups() -> None:
    from fastapi.testclient import TestClient

    from app.main import app

    with SessionLocal() as db:
        user = models.User(name="陈丽妹", enabled=True)
        db.add(user)
        db.flush()
        db.add(models.RoleMapping(user_id=user.id, name="陈丽妹", role="operator", enabled=True))
        db.commit()

    response = TestClient(app).post(
        "/admin/operator-profiles",
        json={
            "operator_name": "陈丽妹",
            "enabled": True,
            "key_categories": [{"level1": f"类目{i}", "level2": ""} for i in range(7)],
        },
    )

    assert response.status_code == 422
    assert "最多配置 6 个一级类目组" in response.text


def test_company_categories_import_from_company_category_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "海外仓新品主攻类目.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    company_sheet = workbook.create_sheet("公司类目")
    company_sheet.append(["一级类目", "二级类目"])
    company_sheet.append(["家居厨卫", "收纳整理"])
    company_sheet.append(["家居厨卫", "厨房工具"])
    company_sheet.append(["汽配与摩配", "摩托车配件"])
    workbook.save(workbook_path)

    from app.company_category_importer import import_company_categories

    with SessionLocal() as db:
        result = import_company_categories(db, workbook_path)
        db.commit()
        categories = db.query(models.CompanyCategory).order_by(
            models.CompanyCategory.level1,
            models.CompanyCategory.level2,
        ).all()

    assert result == {"created_count": 3, "updated_count": 0, "skipped_count": 0}
    assert [(item.level1, item.level2, item.enabled) for item in categories] == [
        ("家居厨卫", "厨房工具", True),
        ("家居厨卫", "收纳整理", True),
        ("汽配与摩配", "摩托车配件", True),
    ]


def test_operator_profile_admin_crud_saves_multi_category_selection() -> None:
    from fastapi.testclient import TestClient

    from app.main import app

    with SessionLocal() as db:
        user = models.User(name="陈丽妹", enabled=True)
        db.add(user)
        db.flush()
        db.add(models.RoleMapping(user_id=user.id, name="陈丽妹", role="operator", enabled=True))
        db.commit()

    client = TestClient(app)
    create_response = client.post(
        "/admin/operator-profiles",
        json={
            "operator_name": "陈丽妹",
            "key_site": "VN",
            "key_categories": [
                {"level1": "家居厨卫"},
                {"level1": "汽配与摩配", "level2": "摩托车配件"},
            ],
            "enabled": True,
        },
    )

    assert create_response.status_code == 200
    body = create_response.json()
    assert body["key_categories"] == [
        {"level1": "家居厨卫", "level2": None},
        {"level1": "汽配与摩配", "level2": "摩托车配件"},
    ]

    update_response = client.patch(
        f"/admin/operator-profiles/{body['id']}",
        json={"key_categories": [{"level1": "办公文教用品"}]},
    )
    assert update_response.status_code == 200
    assert update_response.json()["key_categories"] == [{"level1": "办公文教用品", "level2": None}]


def test_assignment_matches_any_configured_category_at_same_weight() -> None:
    from types import SimpleNamespace

    from app.assignment_rules import preview_main_sku_assignment_groups

    opportunities = [
        SimpleNamespace(
            id="opp-1",
            source_type="selection1",
            batch="B1",
            main_sku="MAIN-1",
            site="PH",
            country="PH",
            category_level1="家居厨卫",
            category_level2=None,
        )
    ]
    profiles = [
        SimpleNamespace(
            operator_name="未命中",
            key_site="PH",
            key_categories=[{"level1": "办公文教用品"}],
            assignment_priority=0,
            display_order=1,
            enabled=True,
        ),
        SimpleNamespace(
            operator_name="命中",
            key_site="PH",
            key_categories=[{"level1": "汽配与摩配"}, {"level1": "家居厨卫", "level2": "收纳整理"}],
            assignment_priority=0,
            display_order=2,
            enabled=True,
        ),
    ]

    result = preview_main_sku_assignment_groups(opportunities, profiles, initial_loads={"未命中": 0, "命中": 0})

    assert result[0].suggested_assignee == "命中"
    assert result[0].match_reason == "一级类目命中"


def test_assignment_exact_level2_match_beats_lower_load_level1_match() -> None:
    from types import SimpleNamespace

    from app.assignment_rules import preview_main_sku_assignment_groups

    opportunities = [
        SimpleNamespace(
            id="opp-1",
            source_type="selection1",
            batch="开发0728期",
            main_sku="MAIN-1",
            site="PH",
            country="PH",
            category_level1="家居厨卫",
            category_level2="收纳整理",
        )
    ]
    profiles = [
        SimpleNamespace(operator_name="一级低负载", key_site="PH", key_categories=[{"level1": "家居厨卫"}], assignment_priority=0, display_order=1, enabled=True),
        SimpleNamespace(operator_name="二级高负载", key_site="PH", key_categories=[{"level1": "家居厨卫", "level2": "收纳整理"}], assignment_priority=0, display_order=2, enabled=True),
    ]

    result = preview_main_sku_assignment_groups(opportunities, profiles, initial_loads={"一级低负载": 0, "二级高负载": 9})

    assert result[0].suggested_assignee == "二级高负载"
    assert result[0].match_reason == "二级类目命中"


def test_assignment_without_category_falls_back_to_lowest_load_enabled_operator() -> None:
    from types import SimpleNamespace

    from app.assignment_rules import preview_main_sku_assignment_groups

    opportunities = [
        SimpleNamespace(
            id="opp-1",
            source_type="selection2",
            batch="选品2-财根0711期",
            main_sku="CG-MAIN",
            site="PH",
            country="PH",
            category_level1=None,
            category_level2=None,
        )
    ]
    profiles = [
        SimpleNamespace(operator_name="PH高负载", key_site="PH", key_categories=[{"level1": "家居厨卫"}], assignment_priority=0, display_order=1, enabled=True),
        SimpleNamespace(operator_name="TH低负载", key_site="TH", key_categories=[{"level1": "汽摩配"}], assignment_priority=0, display_order=2, enabled=True),
    ]

    result = preview_main_sku_assignment_groups(opportunities, profiles, initial_loads={"PH高负载": 5, "TH低负载": 1})

    assert result[0].suggested_assignee == "TH低负载"
    assert result[0].match_reason == "无类目-均衡分配"


def test_assignment_honors_sixth_category_and_balances_matching_operators() -> None:
    from types import SimpleNamespace

    from app.assignment_rules import preview_main_sku_assignment_groups

    sixth_category = [
        {"level1": "类目1"},
        {"level1": "类目2"},
        {"level1": "类目3"},
        {"level1": "类目4"},
        {"level1": "类目5"},
        {"level1": "目标类目"},
    ]
    profiles = [
        SimpleNamespace(operator_name="运营甲", key_site="PH", key_categories=sixth_category, assignment_priority=0, display_order=1, enabled=True),
        SimpleNamespace(operator_name="运营乙", key_site="PH", key_categories=sixth_category, assignment_priority=0, display_order=2, enabled=True),
    ]
    opportunities = [
        SimpleNamespace(id=f"opp-{index}", source_type="selection1", batch="开发0728期", main_sku=f"MAIN-{index}", site="PH", country="PH", category_level1="目标类目", category_level2=None)
        for index in range(6)
    ]

    suggestions = preview_main_sku_assignment_groups(opportunities, profiles, initial_loads={"运营甲": 0, "运营乙": 0})

    assert [item.suggested_assignee for item in suggestions].count("运营甲") == 3
    assert [item.suggested_assignee for item in suggestions].count("运营乙") == 3
def test_operator_profiles_keep_append_order_and_priority() -> None:
    from fastapi.testclient import TestClient

    from app.main import app

    with SessionLocal() as db:
        users = [models.User(name="Beta", enabled=True), models.User(name="Alpha", enabled=True)]
        db.add_all(users)
        db.flush()
        db.add_all(
            [
                models.RoleMapping(user_id=users[0].id, name="Beta", role="operator", enabled=True),
                models.RoleMapping(user_id=users[1].id, name="Alpha", role="operator", enabled=True),
            ]
        )
        db.commit()

    client = TestClient(app)
    first = client.post(
        "/admin/operator-profiles",
        json={"operator_name": "Beta", "key_site": "PH", "assignment_priority": 1, "enabled": True},
    )
    second = client.post(
        "/admin/operator-profiles",
        json={"operator_name": "Alpha", "key_site": "PH", "assignment_priority": 9, "enabled": True},
    )

    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["display_order"] < second.json()["display_order"]
    assert second.json()["assignment_priority"] == 9

    listed = client.get("/admin/operator-profiles").json()
    assert [item["operator_name"] for item in listed] == ["Beta", "Alpha"]

    with SessionLocal() as db:
        role_names = [
            item.name
            for item in db.query(models.RoleMapping)
            .filter(models.RoleMapping.role == "operator")
            .order_by(models.RoleMapping.name)
            .all()
        ]
    assert role_names == ["Alpha", "Beta"]


def build_personnel_fixture(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(["销售员", "入职日期", "岗位", "运营分层", "业务类型", "重点站点", "重点品类1", "重点品类2"])
    worksheet.append(["陈丽妹", "2024-03-06", "组员", "P1", "精品（海外仓）", "PH", "家居厨卫", "商办工业"])
    worksheet.append(["冯卓宏", "2025-03-10", "组员", "P1", "精品（海外仓）", "PH", "汽摩配", "家居厨卫"])
    workbook.save(path)
