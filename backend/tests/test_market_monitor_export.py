import os
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_workflow.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app import models  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.field_mapping import normalize_header  # noqa: E402
from app.main import app  # noqa: E402


client = TestClient(app)


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def test_market_monitor_export_maps_selection1_source_and_secondary_research() -> None:
    claim_time = datetime(2026, 7, 13, 10, 30, tzinfo=timezone.utc)
    make_market_monitor_claim(
        salesperson_name="Alice",
        daily_sales=4.5,
        feedback_summary="first research conclusion",
        secondary_research_at=claim_time,
    )

    response = client.get("/market-monitor/export?business_period=2026-W29")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook["PH精品"]
    assert sheet["J2"].value == 12.5
    assert sheet["P2"].value == "desk organizer"
    assert sheet["AD2"].value == "利润款"
    assert sheet["AE2"].value == 123
    assert sheet["AF2"].value == 8.8
    assert sheet["AG2"].value == 0.25
    assert sheet["AH2"].value == 99
    assert sheet["AI2"].value == 0.3
    assert sheet["AJ2"].value == 4.5
    assert sheet["AK2"].value == "first research conclusion"
    assert sheet["AL2"].value == datetime(2026, 7, 13, 18, 30)
    assert sheet["AM2"].value == "https://shopee.ph/secondary"
    assert sheet["AN2"].value == "secondary conclusion"
    assert sheet["AO2"].value == "引流款"

    traceability = workbook["source_traceability"]
    trace_headers = [cell.value for cell in traceability[1]]
    trace_row = dict(zip(trace_headers, [cell.value for cell in traceability[2]]))
    assert trace_row["source_file"] == "selection1.xlsx"
    assert trace_row["source_sheet"] == "开发0623期"
    assert trace_row["source_row"] == 12
    assert trace_row["business_period"] == "2026-W29"
    assert trace_row["source_snapshot_id"]


def test_market_monitor_export_outputs_one_row_per_responsibility() -> None:
    opportunity_id = make_market_monitor_claim(
        salesperson_name="Alice",
        daily_sales=4.5,
        feedback_summary="alice first research",
        secondary_conclusion="alice secondary",
        positioning="引流款",
    )
    make_claim_row(
        opportunity_id,
        salesperson_name="Bob",
        daily_sales=2,
        feedback_summary="bob first research",
        secondary_conclusion="bob secondary",
        positioning="利润款",
    )
    make_claim_row(
        opportunity_id,
        salesperson_name="Ignored Source",
        daily_sales=8,
        source_column="CC:CH",
    )
    make_claim_row(
        opportunity_id,
        salesperson_name="Not Submitted",
        daily_sales=9,
        secondary_research_submitted_at=None,
    )

    response = client.get("/market-monitor/export?business_period=2026-W29")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook["PH精品"]
    rows = sorted(sheet.iter_rows(min_row=2, values_only=True), key=lambda row: row[3])

    assert [(row[3], row[5], row[35], row[36], row[39], row[40]) for row in rows] == [
        ("Alice", "SUB-MM", 4.5, "alice first research", "alice secondary", "引流款"),
        ("Bob", "SUB-MM", 2, "bob first research", "bob secondary", "利润款"),
    ]


def make_market_monitor_claim(
    salesperson_name: str,
    daily_sales: float,
    feedback_summary: str,
    secondary_research_at: datetime | None = None,
    secondary_conclusion: str = "secondary conclusion",
    positioning: str = "引流款",
) -> str:
    with SessionLocal() as db:
        opportunity = models.NewProductOpportunity(
            source_type="selection1_developer_claim_feedback",
            source_file="selection1.xlsx",
            source_sheet="开发0623期",
            source_row=12,
            batch="2026-W29",
            country="PH",
            site="PH",
            main_sku="MAIN-MM",
            sub_sku="SUB-MM",
            main_sku_name="Main product",
            sub_sku_name="Sub product",
            category_level1="Home",
            keyword="fallback keyword",
            snapshot={},
        )
        db.add(opportunity)
        db.flush()
        db.add(
            models.SourceRecordSnapshot(
                opportunity_id=opportunity.id,
                source_file=opportunity.source_file,
                source_sheet=opportunity.source_sheet,
                source_row=opportunity.source_row,
                column_range="A:BX,CC:CH",
                payload={
                    "fields_by_column": {
                        "K": "利润款",
                        "AJ": 123,
                        "AL": 8.8,
                        "AM": 0.25,
                        "AO": 99,
                        "AP": 0.3,
                    },
                    "fields_by_header": {
                        normalize_header("关键词"): "desk organizer",
                        normalize_header("商品成本-含税（元）"): 12.5,
                    },
                },
            )
        )
        make_claim_row(
            opportunity.id,
            salesperson_name=salesperson_name,
            daily_sales=daily_sales,
            feedback_summary=feedback_summary,
            secondary_research_at=secondary_research_at,
            secondary_conclusion=secondary_conclusion,
            positioning=positioning,
            db=db,
        )
        db.commit()
        return opportunity.id


def make_claim_row(
    opportunity_id: str,
    salesperson_name: str,
    daily_sales: float,
    feedback_summary: str = "first research",
    secondary_research_at: datetime | None = None,
    secondary_conclusion: str = "secondary conclusion",
    positioning: str = "引流款",
    source_column: str = "platform",
    secondary_research_submitted_at: datetime | None = datetime(2026, 7, 13, 11, 0, tzinfo=timezone.utc),
    db=None,
) -> None:
    close = False
    if db is None:
        db = SessionLocal()
        close = True
    db.add(
        models.SalesClaimForecast(
            opportunity_id=opportunity_id,
            salesperson_name=salesperson_name,
            claim_result="claim",
            claim_daily_sales=daily_sales,
            feedback_summary=feedback_summary,
            source_column=source_column,
            secondary_research_at=secondary_research_at or datetime(2026, 7, 13, 10, 0, tzinfo=timezone.utc),
            secondary_competitor_url="https://shopee.ph/secondary",
            secondary_conclusion=secondary_conclusion,
            product_positioning=positioning,
            secondary_research_submitted_at=secondary_research_submitted_at,
        )
    )
    if close:
        db.commit()
        db.close()
