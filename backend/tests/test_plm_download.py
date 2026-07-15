import json
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import plm_download


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["子SKU", "海外仓可发"])
    worksheet.append(["SKU-001", 8])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_beijing_window_uses_the_full_requested_day() -> None:
    start, end = plm_download.beijing_window_ms("2026-07-08")

    assert start == 1783440000000
    assert end == 1783526399000


def test_select_new_export_ignores_old_unready_and_stale_rows() -> None:
    ready = plm_download.select_new_export(
        [
            {"id": "old", "status": "2", "downloadUrl": "old.xlsx", "fileName": "库存-1783590768018.xlsx"},
            {"id": "new-unready", "status": "1", "downloadUrl": "wait.xlsx", "fileName": "库存-1783590768019.xlsx"},
            {"id": "new-stale", "status": "2", "downloadUrl": "stale.xlsx", "fileName": "库存-1783590000000.xlsx"},
            {"id": "new-ready", "status": "2", "downloadUrl": "ready.xlsx", "fileName": "库存-1783590768020.xlsx"},
        ],
        existing_ids={"old"},
        created_after_ms=1783590700000,
    )

    assert ready and ready["id"] == "new-ready"


def test_download_percent_encodes_chinese_file_name(monkeypatch) -> None:
    opened: list[str] = []

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *_):
            return None

        def read(self) -> bytes:
            return b"xlsx"

    def fake_urlopen(url: str, timeout: int):
        opened.append(url)
        return Response()

    monkeypatch.setattr(plm_download, "urlopen", fake_urlopen)

    assert plm_download._download_bytes("http://minio.example/真仓库存明细数据.xlsx") == b"xlsx"
    assert opened == [
        "http://minio.example/%E7%9C%9F%E4%BB%93%E5%BA%93%E5%AD%98%E6%98%8E%E7%BB%86%E6%95%B0%E6%8D%AE.xlsx"
    ]


def test_first_token_prefers_access_token_over_short_token() -> None:
    token = plm_download._first_token({"data": {"token": "short-token", "accessToken": "long-access-token"}})

    assert token == "long-access-token"


def test_download_plm_export_uses_group_eight_and_reuses_valid_cache(monkeypatch, tmp_path) -> None:
    calls: list[tuple[str, dict, dict]] = []

    def fake_post_json(url: str, body: dict, headers: dict | None = None) -> dict:
        calls.append((url, body, headers or {}))
        if url.endswith("/api/system/login"):
            return {"success": True, "data": "test-token"}
        if url.endswith("/api/hz-inventory/overseas/details/exportSummaryExcel"):
            return {"success": True, "now": 1783590768000}
        if len([call for call in calls if call[0].endswith("/selectPage")]) == 1:
            return {"success": True, "data": {"list": [{"id": "old"}]}}
        return {
            "success": True,
            "data": {
                "list": [
                    {"id": "old", "status": "2", "downloadUrl": "old.xlsx"},
                    {
                        "id": "new",
                        "status": "2",
                        "downloadUrl": "http://minio.example/plm.xlsx",
                        "fileName": "真仓库存明细数据-1783590768018.xlsx",
                    },
                ]
            },
        }

    monkeypatch.setattr(plm_download, "_post_json", fake_post_json)
    monkeypatch.setattr(plm_download, "_download_bytes", lambda *_: _workbook_bytes())

    downloaded = plm_download.download_plm_export(
        "2026-07-08",
        base_url="http://plm.example",
        username="test-user",
        password="test-password",
        bloc_name="集团八部",
        cache_dir=tmp_path,
    )

    export_call = next(call for call in calls if call[0].endswith("/exportSummaryExcel"))
    assert export_call[1]["blocNameList"] == ["集团八部"]
    assert export_call[1]["latestStorageTimeStart"] == 1783440000000
    assert export_call[1]["latestStorageTimeEnd"] == 1783526399000
    assert export_call[2] == {"authorization": "test-token"}
    assert downloaded.name == "plm-2026-07-08.xlsx"
    assert load_workbook(downloaded, read_only=True).active["A2"].value == "SKU-001"

    call_count = len(calls)
    assert plm_download.download_plm_export(
        "2026-07-08",
        base_url="http://plm.example",
        username="test-user",
        password="test-password",
        bloc_name="集团八部",
        cache_dir=tmp_path,
    ) == downloaded
    assert len(calls) == call_count


def test_download_plm_open_inventory_detail_uses_inner_open_login_and_caches_json(monkeypatch, tmp_path) -> None:
    calls: list[tuple[str, dict, dict]] = []

    def fake_post_json(url: str, body: dict, headers: dict | None = None) -> dict:
        calls.append((url, body, headers or {}))
        if url.endswith("/system/innerOpen/login"):
            return {"success": True, "data": {"accessToken": "open-token"}}
        assert url.endswith("/open/inventory/overseas/listDetail")
        return {
            "success": True,
            "hasNext": False,
            "data": {
                "records": [
                    {"子SKU": "SKU-001", "集团": "集团八部", "最后一次入库时间": "2026-07-08 10:00:00"}
                ]
            },
        }

    monkeypatch.setattr(plm_download, "_post_json", fake_post_json)

    snapshot = plm_download.download_plm_open_inventory_detail(
        "2026-07-08",
        base_url="http://open.example/open",
        username="test-user",
        password="test-password",
        bloc_name="集团八部",
        cache_dir=tmp_path,
    )

    login_call = calls[0]
    list_call = calls[1]
    assert login_call[0] == "http://open.example/open/system/innerOpen/login"
    assert list_call[0] == "http://open.example/open/open/inventory/overseas/listDetail"
    assert list_call[1]["blocNameList"] == ["集团八部"]
    assert list_call[1]["latestStorageTimeStart"] == "2026-07-08 00:00:00"
    assert list_call[1]["latestStorageTimeEnd"] == "2026-07-08 23:59:59"
    assert list_call[2] == {"Authorization": "open-token"}

    payload = json.loads(snapshot.read_text(encoding="utf-8"))
    assert snapshot.name == "plm-open-2026-07-08.json"
    assert payload["row_count"] == 1
    assert payload["rows"][0]["子SKU"] == "SKU-001"

    call_count = len(calls)
    assert plm_download.download_plm_open_inventory_detail(
        "2026-07-08",
        base_url="http://open.example/open",
        username="test-user",
        password="test-password",
        bloc_name="集团八部",
        cache_dir=tmp_path,
    ) == snapshot
    assert len(calls) == call_count


def test_download_plm_open_inventory_detail_uses_total_over_false_has_next(monkeypatch, tmp_path) -> None:
    seen_pages: list[int] = []

    def fake_post_json(url: str, body: dict, headers: dict | None = None) -> dict:
        if url.endswith("/system/innerOpen/login"):
            return {"success": True, "data": {"accessToken": "open-token"}}
        page_num = body["pageNum"]
        seen_pages.append(page_num)
        return {
            "success": True,
            "total": 2,
            "hasNext": False,
            "data": [{"子SKU": f"SKU-{page_num}"}],
        }

    monkeypatch.setattr(plm_download, "_post_json", fake_post_json)

    snapshot = plm_download.download_plm_open_inventory_detail(
        "2026-07-08",
        base_url="http://open.example/open",
        username="test-user",
        password="test-password",
        bloc_name="集团八部",
        cache_dir=tmp_path,
        page_size=1,
    )

    payload = json.loads(snapshot.read_text(encoding="utf-8"))
    assert seen_pages == [1, 2]
    assert [row["子SKU"] for row in payload["rows"]] == ["SKU-1", "SKU-2"]
