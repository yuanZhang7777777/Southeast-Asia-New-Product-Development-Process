import os
import sys
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_historical_selection1.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook  # noqa: E402

from app import models  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.historical_selection1_import import (  # noqa: E402
    SOURCE_TYPE,
    apply_selection1_rows,
    normalize_selection1_history_rows,
    parse_selection1_workbook,
    plan_selection1_rows,
    revert_selection1_import,
)
from app.selection1_importer import SOURCE_TYPE as CURRENT_SELECTION1_SOURCE_TYPE  # noqa: E402


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def build_selection1_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "开发0623期"
    sheet.append(
        ["站点", "开发部门", "开发员", "一级类目", "关键词", "产品图片", "主SKU名称", "主SKU",
         "子SKU名称", "子SKU", "产品类型", "开品理由", "月销次高", None, None, "月销第三高", None, None]
    )
    sheet.append(
        [None] * 12 + ["月销次高链接", "售价(PHP）", "月销", "月销第三高链接", "售价(PHP）", "月销"]
    )
    # 核价列落在 BY（旧 cells 快照 A:BX+CC:CH 的缺口区），验证全列保真。
    sheet.cell(row=1, column=77, value="供应链核价")
    sheet.cell(row=2, column=77, value="采购核价人")
    sheet.append(
        ["菲律宾", "开发一部", "开发员A", "家居厨卫", "hook", None, "挂钩", "HXG15G",
         "挂钩-灰", "HXG15GD", "引流", "理由A", "https://a.example/1", 11.5, 300, "https://b.example/3", 9.9, 120]
    )
    sheet.cell(row=3, column=77, value="核价员A")
    sheet.append(["合计", None, None, None, None, None, None, None, None, "合计", None, None])
    sheet.append(["站点", "开发部门", "开发员", "一级类目", "关键词", "产品图片", "主SKU名称", "主SKU",
                  "子SKU名称", "子SKU", "产品类型", "开品理由"])
    sheet.append(
        ["泰国", "开发二部", "开发员B", "收纳", "box", None, "收纳盒", "TAB10K",
         "收纳盒-蓝", "TAB10KB", "利润", "理由B"]
    )
    pre_sheet = workbook.create_sheet("开发0407期")
    pre_sheet.append(["站点", "开发部门", "开发员", "一级类目", "关键词", "产品图片", "主SKU名称", "主SKU",
                      "子SKU名称", "子SKU", "产品类型", "开品理由"])
    pre_sheet.append([None] * 12)
    pre_sheet.append(["菲律宾", "部", "员", "类", "kw", None, "名", "OLD407", "名B", "OLD407A", "引流", "r"])
    # 旧世代双层结构（仿真实 0820 期）：R1 字段+分组表头（含重复子SKU列、无站点列），
    # R2 分组子字段，R3 公式行（无子SKU，应跳过），R4 起数据。
    old_sheet = workbook.create_sheet("开发0820期--表格容易错行")
    old_sheet.append(["序列", "需求日期", "开发部门", "开发员", "一级类目", "关键词", "产品图片", "主SKU名称",
                      "主SKU", "是否进入核价流程", "子SKU名称", "子SKU", "子SKU",
                      "竞品单价\n（链接1）\n(比索）", "竞品月销\n（链接1）", "参考定价   （比索）", "开发询价", None])
    old_sheet.append([None] * 16 + ["国内参考链接", "供应商名称"])
    old_sheet.append(["公式行", "2025-08-13", None, None, None, None, None, None,
                      None, None, None, None, None, 66, None, 75, None, None])
    old_sheet.append([None, "2025-08-13", "产品开发四部", "金彩", "宠物用品", "pet bed", None, "宠物窝",
                      "OLD820", "是", "蓝色招财猫S码", "OLD820A", None,
                      130, 392, 130, "https://detail.1688.com/offer/1", "唐山盈好"])
    note_sheet = workbook.create_sheet("说明文档")
    note_sheet.append(["这是说明文字"])
    workbook.save(path)


def build_old_generation_workbook(path: Path) -> None:
    workbook = Workbook()
    # 单层表头 + 空行 + 数据（含旧世代别名：关键词组 / 主SKU名称（33））
    sheet = workbook.active
    sheet.title = "开发0903期"
    sheet.append(["序列", "需求日期", "开发部门", "开发员", "一级类目", "关键词组", "产品图片",
                  "主SKU名称（33）", "主SKU", "子SKU名称", "子SKU", "参考定价   （比索）", "备注"])
    sheet.append([None] * 13)
    sheet.append([1, "2025-09-01", "产品开发五部", "开发员E", "收纳", "storage box", None,
                  "收纳箱", "OLD903", "收纳箱-绿", "OLD903A", 88, "旧世代备注"])
    # 单层表头 + 数据紧跟（仿真实 0827 期，无空行、无第二层表头）
    direct = workbook.create_sheet("开发0827期")
    direct.append(["需求日期", "一级类目", "产品图片", "主SKU名称", "主SKU", "子SKU名称", "子SKU",
                   "合计单销", "是否认领", "销售反馈"])
    direct.append(["2025-08-25", "玩具", None, "分数学习器", "OLD827", "全套磁性分数学习器", "OLD827A1",
                   0.5, "是", "反馈1"])
    direct.append(["2025-08-25", "玩具", None, None, None, "新款磁性分数盘", "OLD827A2", 0.5, None, None])
    # 前 3 行找不到主SKU/子SKU表头 → 整表跳过并报告
    broken = workbook.create_sheet("开发0908期")
    broken.append(["随便", "什么"])
    broken.append(["数据", "行"])
    broken.append(["还是", "数据"])
    workbook.save(path)


def parse_fixture(tmp_path: Path) -> dict:
    source = tmp_path / "选品1：海外仓开发部门开发新品认领-反馈-测试.xlsx"
    build_selection1_workbook(source)
    return parse_selection1_workbook(source)


def test_parse_scopes_current_periods_and_keeps_full_snapshot(tmp_path: Path) -> None:
    report = parse_fixture(tmp_path)
    assert report["sheets"] == [
        {"sheet": "开发0623期", "period": "开发0623期", "rows": 2, "skipped": 2},
    ]
    skipped = {entry["sheet"]: entry["reason"] for entry in report["skipped_sheets"]}
    assert skipped == {
        "开发0407期": "早于0414期",
        "开发0820期--表格容易错行": "排除旧期",
        "说明文档": "非期数sheet",
    }
    assert report["row_count"] == 2
    first = report["rows"][0]
    assert first["sub_sku"] == "HXG15GD"
    assert first["main_sku"] == "HXG15G"
    assert first["batch"] == "开发0623期"
    assert first["country"] == "PH"
    assert first["site"] == "菲律宾"
    assert first["developer_name"] == "开发员A"
    assert first["product_type"] == "引流"
    cells = first["snapshot"]["fields_by_cell"]
    assert cells["N"]["value"] == 11.5 and cells["N"]["group"] == "月销次高"
    assert cells["Q"]["value"] == 9.9 and cells["Q"]["group"] == "月销第三高"
    assert cells["O"]["value"] == 300
    assert cells["R"]["value"] == 120
    assert cells["BY"] == {"header": "采购核价人", "group": "供应链核价", "value": "核价员A"}


def test_parse_can_limit_to_requested_sheets(tmp_path: Path) -> None:
    source = tmp_path / "selection1.xlsx"
    build_selection1_workbook(source)
    report = parse_selection1_workbook(source, selected_sheets={"开发0623期"})
    assert report["row_count"] == 2
    assert [row["sub_sku"] for row in report["rows"]] == ["HXG15GD", "TAB10KB"]
    assert {entry["sheet"] for entry in report["skipped_sheets"]} >= {"开发0820期--表格容易错行"}


def test_long_positioning_note_stays_only_in_source_snapshot(tmp_path: Path) -> None:
    source = tmp_path / "selection1.xlsx"
    build_selection1_workbook(source)
    workbook = load_workbook(source)
    workbook["开发0623期"]["K3"] = "x" * 65
    workbook.save(source)

    row = parse_selection1_workbook(source)["rows"][0]
    assert row["product_type"] is None
    assert row["snapshot"]["fields_by_cell"]["K"]["value"] == "x" * 65

def test_parse_excludes_old_generation_sheets(tmp_path: Path) -> None:
    report = parse_fixture(tmp_path)

    assert all(row["batch"] != "开发0820期" for row in report["rows"])
    assert {entry["sheet"]: entry["reason"] for entry in report["skipped_sheets"]}["开发0820期--表格容易错行"] == "排除旧期"


def test_parse_excludes_old_generation_single_layer_variants(tmp_path: Path) -> None:
    source = tmp_path / "选品1：旧世代单层变体.xlsx"
    build_old_generation_workbook(source)
    report = parse_selection1_workbook(source)

    assert report["rows"] == []
    assert report["sheets"] == []
    assert {entry["sheet"]: entry["reason"] for entry in report["skipped_sheets"]} == {
        "开发0903期": "排除旧期",
        "开发0827期": "排除旧期",
        "开发0908期": "排除旧期",
    }

def test_dry_run_writes_nothing(tmp_path: Path) -> None:
    report = parse_fixture(tmp_path)
    with SessionLocal() as db:
        plan = plan_selection1_rows(db, report["rows"])
    assert plan["would_create"] == 2
    assert plan["by_period"] == {"开发0623期": {"rows": 2, "would_create": 2, "skipped": 0}}
    with SessionLocal() as db:
        assert db.query(models.NewProductOpportunity).count() == 0
        assert db.query(models.ImportBatch).count() == 0
        assert db.query(models.SourceRecordSnapshot).count() == 0
        assert db.query(models.AuditLog).count() == 0

def test_apply_creates_archive_rows_without_tasks_or_notifications(tmp_path: Path) -> None:
    report = parse_fixture(tmp_path)
    with SessionLocal() as db:
        counts = apply_selection1_rows(
            db, report["rows"], source_label=report["source_file"], batch_tag="t-apply",
            imported_by="test", source_sha256=report["source_sha256"],
        )
        db.commit()
    assert counts["created"] == 2
    assert counts["batch_already_imported"] is False
    with SessionLocal() as db:
        opportunities = db.query(models.NewProductOpportunity).all()
        assert {item.source_type for item in opportunities} == {SOURCE_TYPE}
        assert {item.current_status for item in opportunities} == {"historical_archive"}
        assert {item.sub_sku for item in opportunities} == {"HXG15GD", "TAB10KB"}
        assert db.query(models.SourceRecordSnapshot).count() == 2
        assert db.query(models.FlowTask).count() == 0
        assert db.query(models.FlowInstance).count() == 0
        assert db.query(models.SalesClaimForecast).count() == 0
        assert db.query(models.NotificationLog).count() == 0
        batch = db.query(models.ImportBatch).one()
        assert batch.status == "completed"
        assert batch.created_count == 2

def test_rerun_is_idempotent(tmp_path: Path) -> None:
    report = parse_fixture(tmp_path)
    with SessionLocal() as db:
        apply_selection1_rows(db, report["rows"], source_label=report["source_file"], batch_tag="t1", source_sha256=report["source_sha256"])
        db.commit()
    with SessionLocal() as db:
        same_tag = apply_selection1_rows(db, report["rows"], source_label=report["source_file"], batch_tag="t1", source_sha256=report["source_sha256"])
        db.commit()
    assert same_tag["batch_already_imported"] is True
    assert same_tag["created"] == 0
    with SessionLocal() as db:
        second = apply_selection1_rows(db, report["rows"], source_label=report["source_file"], batch_tag="t2", source_sha256=report["source_sha256"])
        db.commit()
    assert second["batch_already_imported"] is False
    assert second["created"] == 0
    assert second["skipped_existing_archive"] == 2
    with SessionLocal() as db:
        third = apply_selection1_rows(db, report["rows"], source_label=report["source_file"], batch_tag="t3", source_sha256=None)
        db.commit()
    assert third["created"] == 0
    assert third["skipped_existing_archive"] == 2
    with SessionLocal() as db:
        assert db.query(models.NewProductOpportunity).count() == 2
        assert db.query(models.SourceRecordSnapshot).count() == 2

def test_existing_current_flow_row_receives_source_backfill_without_state_change(tmp_path: Path) -> None:
    with SessionLocal() as db:
        db.add(models.NewProductOpportunity(
            source_type=CURRENT_SELECTION1_SOURCE_TYPE,
            source_sheet="开发0623期", batch="开发0623期", main_sku="HXG15G", sub_sku="HXG15GD",
            site="PH", country="PH", current_status="pending_assignment", snapshot={"marker": "keep-me"},
        ))
        db.commit()
    report = parse_fixture(tmp_path)
    source_row = next(row for row in report["rows"] if row["sub_sku"] == "HXG15GD")
    source_row["image_url"] = "https://oss.example/HXG15GD.png"
    with SessionLocal() as db:
        counts = apply_selection1_rows(db, report["rows"], source_label=report["source_file"], batch_tag="t-current", source_sha256=report["source_sha256"])
        db.commit()
    assert counts["created"] == 1
    assert counts["skipped_existing_current_flow"] == 1
    assert counts["backfilled_existing_current_flow"] == 1
    with SessionLocal() as db:
        current = db.query(models.NewProductOpportunity).filter_by(source_type=CURRENT_SELECTION1_SOURCE_TYPE).one()
        assert current.current_status == "pending_assignment"
        assert current.main_sku_name == "挂钩"
        assert current.image_url == "https://oss.example/HXG15GD.png"
        assert current.snapshot["marker"] == "keep-me"
        assert current.snapshot["historical_selection1"] == source_row["snapshot"]
        assert db.query(models.SourceRecordSnapshot).filter_by(opportunity_id=current.id).count() == 1
        assert db.query(models.NewProductOpportunity).filter_by(source_type=SOURCE_TYPE, sub_sku="HXG15GD").count() == 0
        created = db.query(models.NewProductOpportunity).filter_by(source_type=SOURCE_TYPE).all()
        assert {item.sub_sku for item in created} == {"TAB10KB"}


def test_existing_archive_row_receives_image_and_latest_source_snapshot(tmp_path: Path) -> None:
    report = parse_fixture(tmp_path)
    source_row = next(row for row in report["rows"] if row["sub_sku"] == "HXG15GD")
    source_row["image_url"] = "https://oss.example/HXG15GD.png"
    with SessionLocal() as db:
        db.add(models.NewProductOpportunity(
            source_type=SOURCE_TYPE, source_sheet="开发0623期", batch="开发0623期",
            main_sku="HXG15G", sub_sku="HXG15GD", country="PH", site="PH",
            current_status="historical_archive",
            snapshot={"development_source": {"supplier": "keep"}, "fields_by_cell": {"A": {"value": "old"}}},
        ))
        db.commit()
    with SessionLocal() as db:
        counts = apply_selection1_rows(db, report["rows"], source_label=report["source_file"], batch_tag="t-archive", source_sha256=report["source_sha256"])
        db.commit()
    assert counts["created"] == 1
    assert counts["skipped_existing_archive"] == 1
    assert counts["backfilled_existing_archive"] == 1
    with SessionLocal() as db:
        archive = db.query(models.NewProductOpportunity).filter_by(source_type=SOURCE_TYPE, sub_sku="HXG15GD").one()
        assert archive.current_status == "historical_archive"
        assert archive.image_url == "https://oss.example/HXG15GD.png"
        assert archive.snapshot["development_source"] == {"supplier": "keep"}
        assert archive.snapshot["fields_by_cell"] == source_row["snapshot"]["fields_by_cell"]
        assert db.query(models.SourceRecordSnapshot).filter_by(opportunity_id=archive.id).count() == 1

def test_revert_removes_batch_and_allows_reimport(tmp_path: Path) -> None:
    report = parse_fixture(tmp_path)
    with SessionLocal() as db:
        apply_selection1_rows(db, report["rows"], source_label=report["source_file"], batch_tag="t-revert", source_sha256=report["source_sha256"])
        db.commit()
    with SessionLocal() as db:
        reverted = revert_selection1_import(db, "t-revert", actor="test")
        db.commit()
    assert reverted == {"reverted_opportunities": 2, "reverted_snapshots": 2}
    with SessionLocal() as db:
        assert db.query(models.NewProductOpportunity).count() == 0
        assert db.query(models.SourceRecordSnapshot).count() == 0
        assert db.query(models.ImportBatch).one().status == "reverted"
    with SessionLocal() as db:
        again = apply_selection1_rows(db, report["rows"], source_label=report["source_file"], batch_tag="t-after-revert", source_sha256=report["source_sha256"])
        db.commit()
    assert again["batch_already_imported"] is False
    assert again["created"] == 2
    with SessionLocal() as db:
        assert db.query(models.NotificationLog).count() == 0


def test_same_sheet_and_sub_sku_with_different_site_or_main_sku_create_distinct_archives() -> None:
    rows = [
        {"source_file": "selection1.xlsx", "source_sheet": "开发0623期", "source_row": 10, "batch": "开发新品0623期", "country": "PH", "site": "菲律宾", "main_sku": "MAIN-PH", "sub_sku": "SHARED-SUB", "snapshot": {"source_reference": {"source_row": 10}}},
        {"source_file": "selection1.xlsx", "source_sheet": "开发0623期", "source_row": 11, "batch": "开发新品0623期", "country": "TH", "site": "泰国", "main_sku": "MAIN-TH", "sub_sku": "SHARED-SUB", "snapshot": {"source_reference": {"source_row": 11}}},
    ]
    with SessionLocal() as db:
        assert plan_selection1_rows(db, rows)["would_create"] == 2
        counts = apply_selection1_rows(db, rows, source_label="selection1.xlsx", batch_tag="t-distinct-identity")
        db.commit()
    assert counts["created"] == 2
    with SessionLocal() as db:
        archives = db.query(models.NewProductOpportunity).filter_by(source_type=SOURCE_TYPE).all()
        assert {(item.country, item.main_sku, item.sub_sku) for item in archives} == {
            ("PH", "MAIN-PH", "SHARED-SUB"), ("TH", "MAIN-TH", "SHARED-SUB"),
        }


def test_missing_first_main_sku_falls_back_to_sub_sku_with_source_marker(tmp_path: Path) -> None:
    source = tmp_path / "selection1-missing-main.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "开发0623期"
    sheet.append(["站点", "主SKU", "子SKU"])
    sheet.append([None, None, None])
    sheet.append(["菲律宾", None, "SUB-NO-MAIN"])
    workbook.save(source)

    row = parse_selection1_workbook(source)["rows"][0]
    assert row["main_sku"] == "SUB-NO-MAIN"
    assert row["snapshot"]["backfilled_fields"] == {"main_sku": {"source": "sub_sku", "value": "SUB-NO-MAIN"}}
    with SessionLocal() as db:
        assert plan_selection1_rows(db, [row])["would_create"] == 1


def test_parse_excludes_last_year_legacy_period(tmp_path: Path) -> None:
    source = tmp_path / "selection1-legacy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "开发0924期"
    sheet.append(["站点", "主SKU", "子SKU"])
    sheet.append(["菲律宾", "LEGACY", "LEGACY-A"])
    workbook.save(source)

    report = parse_selection1_workbook(source)
    assert report["rows"] == []
    assert report["skipped_sheets"] == [{"sheet": "开发0924期", "reason": "排除旧期"}]


def test_parse_0707_uses_leading_site_column_without_header(tmp_path: Path) -> None:
    source = tmp_path / "selection1-0707.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "开发0707期"
    sheet.append([None, "开发部门", "主SKU", "子SKU"])
    sheet.append(["菲律宾", "海外仓", "MAIN-0707", "SUB-0707"])
    workbook.save(source)

    row = parse_selection1_workbook(source)["rows"][0]
    assert row["country"] == "PH"
    assert row["site"] == "菲律宾"
    assert row["snapshot"]["site_resolution"] == "leading_site_column"


def test_parse_0414_inherits_one_checked_site_within_main_sku_group(tmp_path: Path) -> None:
    source = tmp_path / "selection1-0414.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "开发0414期"
    sheet.append(["主SKU", "子SKU", "Shopee"])
    sheet.append([None, None, "PH"])
    sheet.append(["MAIN-0414", "SUB-0414-A", "✔"])
    sheet.append([None, "SUB-0414-B", None])
    workbook.save(source)

    rows = parse_selection1_workbook(source)["rows"]
    inherited = next(row for row in rows if row["sub_sku"] == "SUB-0414-B")
    assert inherited["country"] == "PH"
    assert inherited["site"] == "PH"
    assert inherited["snapshot"]["site_resolution"] == "checkbox_inherited"


def test_parse_0414_keeps_site_empty_when_group_has_no_checkmark(tmp_path: Path) -> None:
    source = tmp_path / "selection1-0414-no-check.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "开发0414期"
    sheet.append(["主SKU", "子SKU", "Shopee"])
    sheet.append([None, None, "PH"])
    sheet.append(["MAIN-EMPTY", "SUB-EMPTY-A", None])
    sheet.append([None, "SUB-EMPTY-B", None])
    workbook.save(source)

    rows = parse_selection1_workbook(source)["rows"]
    assert {(row["country"], row["site"]) for row in rows} == {(None, None)}
    assert {row["snapshot"]["site_resolution"] for row in rows} == {"checkbox_group_unchecked"}


def test_parse_backfills_missing_main_sku_and_sub_sku_name(tmp_path: Path) -> None:
    source = tmp_path / "selection1-fallbacks.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "开发0623期"
    sheet.append(["站点", "主SKU名称", "主SKU", "子SKU名称", "子SKU"])
    sheet.append(["菲律宾", "主商品名", None, None, "SUB-ONLY"])
    workbook.save(source)

    row = parse_selection1_workbook(source)["rows"][0]
    assert row["main_sku"] == "SUB-ONLY"
    assert row["sub_sku_name"] == "主商品名"
    assert row["snapshot"]["backfilled_fields"] == {
        "main_sku": {"source": "sub_sku", "value": "SUB-ONLY"},
        "sub_sku_name": {"source": "main_sku_name", "value": "主商品名"},
    }

def test_parse_excludes_period_outside_current_historical_migration_scope(tmp_path: Path) -> None:
    source = tmp_path / "selection1-current-future-period.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "开发1001期"
    sheet.append(["站点", "主SKU", "子SKU"])
    sheet.append(["菲律宾", "CURRENT-1001", "CURRENT-1001-A"])
    workbook.save(source)

    report = parse_selection1_workbook(source)
    assert report["rows"] == []
    assert report["skipped_sheets"] == [{"sheet": "开发1001期", "reason": "排除非本次历史期"}]


def test_normalize_same_identity_merges_empty_complement_and_preserves_sources() -> None:
    rows = [
        {
            "source_file": "selection1.xlsx", "source_sheet": "开发0526期", "source_row": 129,
            "source_type": SOURCE_TYPE, "batch": "开发0526期", "country": "TH", "site": "泰国",
            "main_sku": "WHHG972", "sub_sku": "WHHG972A2", "main_sku_name": "同款商品",
            "sub_sku_name": None, "keyword": "原关键词",
            "snapshot": {"source_reference": {"source_row": 129}},
        },
        {
            "source_file": "selection1.xlsx", "source_sheet": "开发0526期", "source_row": 130,
            "source_type": SOURCE_TYPE, "batch": "开发0526期", "country": "TH", "site": "泰国",
            "main_sku": "WHHG972", "sub_sku": "WHHG972A2", "main_sku_name": "同款商品",
            "sub_sku_name": "WHHG972A2规格", "keyword": None,
            "snapshot": {"source_reference": {"source_row": 130}},
        },
    ]

    normalized = normalize_selection1_history_rows(rows)

    assert len(normalized["rows"]) == 1
    row = normalized["rows"][0]
    assert row["keyword"] == "原关键词"
    assert row["sub_sku_name"] == "WHHG972A2规格"
    assert row["snapshot"]["merged_source_rows"] == [129, 130]
    assert normalized["merged_groups"] == [{"identity": ["开发0526期", "th", "WHHG972", "WHHG972A2"], "source_rows": [129, 130]}]


def test_normalize_same_main_with_different_children_keeps_both_rows() -> None:
    rows = [
        {
            "source_file": "selection1.xlsx", "source_sheet": "开发0526期", "source_row": 129,
            "source_type": SOURCE_TYPE, "batch": "开发0526期", "country": "TH", "site": "泰国",
            "main_sku": "HFL045", "sub_sku": "HFL045S", "snapshot": {"source_reference": {"source_row": 129}},
        },
        {
            "source_file": "selection1.xlsx", "source_sheet": "开发0526期", "source_row": 130,
            "source_type": SOURCE_TYPE, "batch": "开发0526期", "country": "TH", "site": "泰国",
            "main_sku": "HFL045", "sub_sku": "HFL045L", "snapshot": {"source_reference": {"source_row": 130}},
        },
    ]

    normalized = normalize_selection1_history_rows(rows)

    assert [row["sub_sku"] for row in normalized["rows"]] == ["HFL045S", "HFL045L"]
    assert normalized["merged_groups"] == []
    assert normalized["business_repair_rows"] == []


def test_normalize_same_identity_with_conflicting_business_value_requires_business_fix() -> None:
    rows = [
        {
            "source_file": "selection1.xlsx", "source_sheet": "开发0526期", "source_row": 129,
            "source_type": SOURCE_TYPE, "batch": "开发0526期", "country": "TH", "site": "泰国",
            "main_sku": "WHHG972", "sub_sku": "WHHG972A2", "product_type": "利润款",
            "snapshot": {"source_reference": {"source_row": 129}},
        },
        {
            "source_file": "selection1.xlsx", "source_sheet": "开发0526期", "source_row": 130,
            "source_type": SOURCE_TYPE, "batch": "开发0526期", "country": "TH", "site": "泰国",
            "main_sku": "WHHG972", "sub_sku": "WHHG972A2", "product_type": "淘汰款",
            "snapshot": {"source_reference": {"source_row": 130}},
        },
    ]

    normalized = normalize_selection1_history_rows(rows)

    assert normalized["rows"] == []
    assert normalized["business_repair_rows"][0]["conflicting_fields"] == {
        "product_type": ["利润款", "淘汰款"],
    }

def test_apply_normalizes_duplicate_identity_before_archive_write() -> None:
    rows = [
        {
            "source_file": "selection1.xlsx", "source_sheet": "开发0526期", "source_row": 129,
            "source_type": SOURCE_TYPE, "batch": "开发0526期", "country": "TH", "site": "泰国",
            "main_sku": "WHHG972", "sub_sku": "WHHG972A2", "main_sku_name": "同款商品",
            "keyword": "原关键词", "snapshot": {"source_reference": {"source_row": 129}},
        },
        {
            "source_file": "selection1.xlsx", "source_sheet": "开发0526期", "source_row": 130,
            "source_type": SOURCE_TYPE, "batch": "开发0526期", "country": "TH", "site": "泰国",
            "main_sku": "WHHG972", "sub_sku": "WHHG972A2", "main_sku_name": "同款商品",
            "sub_sku_name": "WHHG972A2规格", "snapshot": {"source_reference": {"source_row": 130}},
        },
    ]

    with SessionLocal() as db:
        counts = apply_selection1_rows(db, rows, source_label="selection1.xlsx", batch_tag="t-normalize-before-apply")
        db.commit()

    assert counts["created"] == 1
    with SessionLocal() as db:
        archive = db.query(models.NewProductOpportunity).filter_by(source_type=SOURCE_TYPE).one()
        assert archive.keyword == "原关键词"
        assert archive.sub_sku_name == "WHHG972A2规格"
        assert archive.snapshot["merged_source_rows"] == [129, 130]
