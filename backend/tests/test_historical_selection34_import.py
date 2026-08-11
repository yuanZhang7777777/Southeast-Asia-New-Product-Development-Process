import json
import os
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

import pytest

os.environ["APP_ENV"] = "testing"
os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_workflow.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from app import models  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app import historical_selection34_import as selection34_import  # noqa: E402
from app.historical_selection34_import import (  # noqa: E402
    AUTHORITATIVE_ROWS_SHA256,
    AUTHORITATIVE_SHA256,
    AUTHORITATIVE_SHEETS,
    SOURCE_TYPE,
    apply_selection34_rows,
    main,
    parse_historical_selection34_row,
    parse_selection34_workbook,
    selection34_rows_sha256,
    validate_selection34_apply_payload,
)


AUTHORITY_WORKBOOK = Path(r"E:\soft\dingding\历史数据--选品3&选品4表.xlsx")
EXPECTED_SHEETS = [
    "小货老品-4月底",
    "开发高投入推荐-4月底",
    "销售自选0511期",
    "销售自选0518期",
    "销售自选0525期",
    "销售自选0601期",
    "销售自选0608期",
    "销售自选0613期",
    "销售自选0615期",
    "销售自选0622期",
    "销售自选0629期",
    "销售自选0706期",
    "直发热销转0615期",
    "直发热销转0630期",
]

assert tuple(EXPECTED_SHEETS) == AUTHORITATIVE_SHEETS


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def _row(values_by_column: dict[int, object]) -> tuple[object | None, ...]:
    row = [None] * 83
    for column, value in values_by_column.items():
        row[column - 1] = value
    return tuple(row)


def _headers() -> dict[str, object]:
    headers = {get_column_letter(index): f"字段{index}" for index in range(1, 84)}
    headers.update(
        {
            "A": "站点",
            "B": "开发部门",
            "C": "开发员",
            "D": "一级类目",
            "E": "二级类目",
            "F": "关键词",
            "G": "产品图片",
            "H": "主SKU名称",
            "I": "主SKU",
            "J": "子SKU名称",
            "K": "子SKU",
            "L": "产品类型",
            "M": "开品理由",
            "BZ": "不认领理由",
            "CA": "主销售员",
            "CB": "是否认领",
            "CC": "认领单销",
            "CD": "销售反馈\n总结",
            "CE": "备注",
        }
    )
    return headers


def _parsed_row(
    *,
    source_row: int = 2,
    main_sku: str = "MAIN-1",
    sub_sku: str = "CHILD-1",
    salesperson: object = "销售A",
    claim_flag: object = "是",
    daily_sales: object = "1",
    reject_reason: object = None,
    feedback: object = "反馈",
    note: object = "备注",
):
    return parse_historical_selection34_row(
        _row(
            {
                1: "菲律宾",
                2: "开发一部",
                3: "开发A",
                4: "一级",
                5: "二级",
                6: "关键词",
                7: "https://image.example/item.jpg",
                8: "主商品",
                9: main_sku,
                10: "子商品",
                11: sub_sku,
                12: "利润款",
                13: "开品理由",
                78: reject_reason,
                79: salesperson,
                80: claim_flag,
                81: daily_sales,
                82: feedback,
                83: note,
            }
        ),
        headers=_headers(),
        source_file="selection34.xlsx",
        source_sheet="销售自选0511期",
        source_row=source_row,
    )


def _valid_apply_payload() -> dict:
    rows = []
    for index in range(596):
        sheet = EXPECTED_SHEETS[min(index, len(EXPECTED_SHEETS) - 1)]
        rows.append(
            {
                "source_type": SOURCE_TYPE,
                "source_file": "authority.xlsx",
                "source_sheet": sheet,
                "source_row": index + 2,
                "batch": sheet,
                "country": "PH",
                "site": "PH",
                "main_sku": f"MAIN-{index}",
                "sub_sku": f"SUB-{index}",
                "claims": [{"salesperson_name": f"销售{index}", "claim_result": "claim"}],
                "rejected_sources": [],
            }
        )
    for index in range(8):
        duplicate = deepcopy(rows[index])
        duplicate["source_row"] = 1000 + index
        duplicate["claims"] = [{"salesperson_name": f"重复销售{index}", "claim_result": "claim"}]
        rows.append(duplicate)
    payload = {
        "source_sha256": AUTHORITATIVE_SHA256.lower(),
        "sheets": [{"sheet": sheet} for sheet in EXPECTED_SHEETS],
        "row_count": 604,
        "unique_opportunity_count": 596,
        "claim_count": 604,
        "reject_count": 0,
        "source_snapshot_count": 604,
        "rows": rows,
    }
    payload["rows_sha256"] = selection34_rows_sha256(rows)
    return payload


def test_parse_row_uses_selection1_layout_and_keeps_source_payload() -> None:
    parsed = _parsed_row(main_sku=" MAIN \n 1 ", sub_sku="CHILD\n 1", daily_sales="1")

    assert parsed is not None
    assert parsed["source_type"] == SOURCE_TYPE
    assert parsed["batch"] == "销售自选0511期"
    assert parsed["current_status"] == "historical_archive"
    assert parsed["claim_pool_open"] is False
    assert parsed["country"] == "PH"
    assert parsed["site"] == "PH"
    assert parsed["main_sku"] == "MAIN1"
    assert parsed["sub_sku"] == "CHILD1"
    assert parsed["developer_department"] == "开发一部"
    assert parsed["developer_name"] == "开发A"
    assert parsed["category_level1"] == "一级"
    assert parsed["category_level2"] == "二级"
    assert parsed["keyword"] == "关键词"
    assert parsed["image_url"] == "https://image.example/item.jpg"
    assert parsed["main_sku_name"] == "主商品"
    assert parsed["sub_sku_name"] == "子商品"
    assert parsed["product_type"] == "利润款"
    assert parsed["reason"] == "开品理由"
    assert parsed["snapshot"]["fields_by_cell"]["I"]["value"] == " MAIN \n 1 "
    assert parsed["snapshot"]["fields_by_cell"]["K"]["value"] == "CHILD\n 1"
    assert len(parsed["snapshot"]["fields_by_cell"]) == 83
    assert parsed["snapshot"]["has_embedded_image"] is False
    assert parsed["claims"] == [
        {
            "salesperson_name": "销售A",
            "claim_result": "claim",
            "claim_daily_sales": 1.0,
            "reject_reason": None,
            "feedback_summary": "反馈",
            "source_column": "BZ:CE:2",
            "source_payload": {
                "source_file": "selection34.xlsx",
                "source_sheet": "销售自选0511期",
                "source_row": 2,
                "source_columns": {
                    "reject_reason": "BZ",
                    "salesperson": "CA",
                    "claim_flag": "CB",
                    "daily_sales": "CC",
                    "feedback_summary": "CD",
                    "note": "CE",
                },
                "reject_reason": None,
                "salesperson": "销售A",
                "claim_flag": "是",
                "daily_sales": "1",
                "feedback_summary": "反馈",
                "note": "备注",
            },
        }
    ]
    assert parsed["rejected_sources"] == []


@pytest.mark.parametrize(
    ("claim_flag", "daily_sales", "reject_reason", "expected_reason"),
    [
        ("否", 1, None, "来源标记不认领"),
        ("是", None, None, "来源表未填写认领单销"),
        ("是", "  ", None, "来源表未填写认领单销"),
        ("是", 0, None, "来源表认领单销为0"),
        ("是", "销量不足", None, "销量不足"),
        ("是", 1, "竞品过多", "竞品过多"),
    ],
)
def test_parse_row_rejects_named_nonpositive_or_explicit_rejection(
    claim_flag: object,
    daily_sales: object,
    reject_reason: object,
    expected_reason: str,
) -> None:
    parsed = _parsed_row(
        claim_flag=claim_flag,
        daily_sales=daily_sales,
        reject_reason=reject_reason,
    )

    assert parsed is not None
    assert parsed["claims"] == []
    assert parsed["rejected_sources"][0]["claim_result"] == "reject"
    assert parsed["rejected_sources"][0]["reject_reason"] == expected_reason


def test_parse_row_without_salesperson_creates_no_relation() -> None:
    parsed = _parsed_row(salesperson=None, claim_flag="是", daily_sales=2, reject_reason="仍不建关系")

    assert parsed is not None
    assert parsed["claims"] == []
    assert parsed["rejected_sources"] == []


def test_workbook_resets_bad_dimensions_and_keeps_sheet_name(tmp_path: Path) -> None:
    workbook_path = tmp_path / "bad-dimension.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售自选0511期"
    sheet.append([_headers()[get_column_letter(index)] for index in range(1, 84)])
    sheet.append(_row({1: "菲律宾", 9: "MAIN", 11: "CHILD", 79: "销售A", 80: "是", 81: "1"}))
    workbook.save(workbook_path)

    rewritten = tmp_path / "rewritten.xlsx"
    with zipfile.ZipFile(workbook_path) as source, zipfile.ZipFile(rewritten, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<dimension ref="A1:CE2"/>', b'<dimension ref="A1:A1"/>')
            target.writestr(item, data)
    rewritten.replace(workbook_path)

    report = parse_selection34_workbook(workbook_path)

    assert report["row_count"] == 1
    assert report["unique_opportunity_count"] == 1
    assert report["sheets"] == [
        {
            "sheet": "销售自选0511期",
            "batch": "销售自选0511期",
            "rows": 1,
            "claims": 1,
            "rejects": 0,
            "unclaimed": 0,
        }
    ]
    assert report["rows"][0]["batch"] == "销售自选0511期"


def test_apply_merges_multi_operator_identity_and_is_idempotent_without_downstream_rows() -> None:
    first_row = _parsed_row(source_row=2, salesperson="销售A", daily_sales=1, feedback="反馈A", note="备注A")
    second_row = _parsed_row(source_row=3, salesperson="销售B", daily_sales=2, feedback="反馈B", note="备注B")
    assert first_row is not None
    assert second_row is not None
    downstream_models = (
        models.FlowInstance,
        models.FlowTask,
        models.ReviewRecord,
        models.StockingRequest,
        models.MarketResearchItem,
        models.ListingRecord,
        models.ItemObservationPeriod,
        models.NotificationLog,
    )

    with SessionLocal() as db:
        before = {model: db.query(model).count() for model in downstream_models}
        first = apply_selection34_rows(db, [first_row, second_row], source_label="selection34.xlsx")
        db.commit()
        second = apply_selection34_rows(db, [first_row, second_row], source_label="selection34.xlsx")
        db.commit()

        opportunity = db.query(models.NewProductOpportunity).one()
        claims = db.query(models.SalesClaimForecast).order_by(models.SalesClaimForecast.salesperson_name).all()
        snapshots = db.query(models.SourceRecordSnapshot).order_by(models.SourceRecordSnapshot.source_row).all()
        after = {model: db.query(model).count() for model in downstream_models}

        assert opportunity.source_type == SOURCE_TYPE
        assert opportunity.batch == "销售自选0511期"
        assert opportunity.current_status == "historical_archive"
        assert opportunity.claim_pool_open is False
        assert [(claim.salesperson_name, claim.claim_source, claim.claim_daily_sales) for claim in claims] == [
            ("销售A", SOURCE_TYPE, 1.0),
            ("销售B", SOURCE_TYPE, 2.0),
        ]
        assert [snapshot.source_row for snapshot in snapshots] == [2, 3]
        assert all(snapshot.column_range == "A:CE" for snapshot in snapshots)
        assert json.loads(claims[1].note)["source_payload"]["note"] == "备注B"
        assert before == after

    assert first["created_opportunities"] == 1
    assert first["created_source_snapshots"] == 2
    assert first["created_claims"] == 2
    assert second["created_opportunities"] == 0
    assert second["created_source_snapshots"] == 0
    assert second["created_claims"] == 0


@pytest.mark.parametrize(
    ("field", "bad_value"),
    [
        ("source_sha256", "not-authoritative"),
        ("sheets", [{"sheet": sheet} for sheet in reversed(EXPECTED_SHEETS)]),
        ("row_count", 603),
        ("unique_opportunity_count", 595),
        ("claim_count", 603),
        ("reject_count", 1),
        ("source_snapshot_count", 603),
    ],
)
def test_validate_apply_payload_rejects_authority_mismatch(
    field: str,
    bad_value: object,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payload = _valid_apply_payload()
    monkeypatch.setattr(selection34_import, "AUTHORITATIVE_ROWS_SHA256", payload["rows_sha256"])
    payload[field] = bad_value

    with pytest.raises(ValueError, match=field):
        validate_selection34_apply_payload(payload)


def test_validate_apply_payload_recomputes_counts_from_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    payload = _valid_apply_payload()
    monkeypatch.setattr(selection34_import, "AUTHORITATIVE_ROWS_SHA256", payload["rows_sha256"])
    payload["rows"].pop()

    with pytest.raises(ValueError, match="rows.row_count"):
        validate_selection34_apply_payload(payload)


def test_validate_apply_payload_rejects_row_tampering_with_unchanged_counts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payload = _valid_apply_payload()
    monkeypatch.setattr(selection34_import, "AUTHORITATIVE_ROWS_SHA256", payload["rows_sha256"])
    payload["rows"][0]["claims"][0]["salesperson_name"] = "篡改运营"

    with pytest.raises(ValueError, match="rows_sha256"):
        validate_selection34_apply_payload(payload)


def test_apply_cli_rejects_invalid_payload_before_opening_database_session(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payload = _valid_apply_payload()
    payload["source_sha256"] = "not-authoritative"
    payload_path = tmp_path / "invalid-selection34.json"
    payload_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    session_opened = False

    def forbidden_session():
        nonlocal session_opened
        session_opened = True
        raise AssertionError("database session must not open for an invalid authority payload")

    monkeypatch.setattr("app.db.SessionLocal", forbidden_session)
    monkeypatch.setattr(
        sys,
        "argv",
        ["historical_selection34_import", "--apply-rows", str(payload_path), "--apply-dev"],
    )

    with pytest.raises(SystemExit, match="source_sha256"):
        main()

    assert session_opened is False


@pytest.mark.skipif(not AUTHORITY_WORKBOOK.exists(), reason="权威工作簿不在当前机器")
def test_authoritative_workbook_audit_counts() -> None:
    report = parse_selection34_workbook(AUTHORITY_WORKBOOK)

    assert report["source_sha256"].upper() == AUTHORITATIVE_SHA256
    assert report["source_sha256_matches_authority"] is True
    assert report["rows_sha256"] == AUTHORITATIVE_ROWS_SHA256
    assert [sheet["sheet"] for sheet in report["sheets"]] == EXPECTED_SHEETS
    assert report["row_count"] == 604
    assert report["unique_opportunity_count"] == 596
    assert report["multi_claim_identity_count"] == 8
    assert report["claim_count"] == 604
    assert report["reject_count"] == 0
    assert report["unclaimed_count"] == 0
    assert report["source_snapshot_count"] == 604
    assert report["image_stats"] == {"embedded_images": 0, "rows_with_image": 0}
    assert sum("\n" in str(row["snapshot"]["fields_by_cell"]["K"]["value"] or "") for row in report["rows"]) == 4
    assert all(not any(character.isspace() for character in row["sub_sku"]) for row in report["rows"])
