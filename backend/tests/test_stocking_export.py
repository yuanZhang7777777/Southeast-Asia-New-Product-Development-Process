import os
import sys
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_workflow.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from io import BytesIO
from datetime import date, datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app import models, schemas, services  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402


client = TestClient(app)


STOCKING_HEADERS = [
    "操作状态",
    "申请日期",
    "备货类型",
    "选品数据源",
    "销售员",
    "主SKU",
    "子sku",
    "成本价",
    "单个体积",
    "备货单销",
    "备货数量",
    "备货国家",
    "备货仓库",
    "货值",
    "体积",
    "补货原因",
]


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def test_stocking_export_matches_official_workbook_structure() -> None:
    prepare_approved_claims([("销售A", 2.5)])

    response = export_all_available()

    assert response.status_code == 200
    assert "filename*=UTF-8''%E6%B5%B7%E5%A4%96%E4%BB%93%E5%A4%87%E8%B4%A7%E7%94%B3%E8%AF%B7%E8%A1%A8.xlsx" in response.headers[
        "content-disposition"
    ]
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook["PH"]
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, [cell.value for cell in sheet[2]]))

    assert headers == STOCKING_HEADERS
    assert row["销售员"] == "销售A"
    assert row["操作状态"] is None
    assert row["主SKU"] == "MAIN-EXPORT"
    assert row["子sku"] == "SUB-EXPORT"
    assert row["备货单销"] == 2.5
    assert row["备货数量"] == 75
    assert row["补货原因"] is None


def test_stocking_export_splits_sheets_by_country_and_uses_application_dates() -> None:
    prepare_approved_claims([("销售PH", 1)], country="PH", site="PH", main_sku="MAIN-PH", sub_sku="SUB-PH", source_row=1)
    prepare_approved_claims([("销售TH", 1)], country="TH", site="泰国", main_sku="MAIN-TH", sub_sku="SUB-TH", source_row=2)

    response = export_all_available()

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["PH", "TH"]
    assert workbook["PH"]["F2"].value == "MAIN-PH"
    assert workbook["TH"]["F2"].value == "MAIN-TH"
    assert workbook["PH"]["B2"].value.date() == date(2026, 7, 17)
    assert workbook["PH"]["B2"].number_format == "yyyy-mm-dd"


def test_multiple_approved_claims_for_one_child_sku_export_as_multiple_rows() -> None:
    prepare_approved_claims([("赵钰婷", 1.5), ("李干", 1)])

    response = client.get("/stocking/available-list")

    assert response.status_code == 200
    rows = sorted(response.json(), key=lambda item: item["salesperson_name"])
    assert [(row["salesperson_name"], row["sub_sku"], row["claim_daily_sales"], row["quantity"]) for row in rows] == [
        ("李干", "SUB-EXPORT", 1, 30),
        ("赵钰婷", "SUB-EXPORT", 1.5, 45),
    ]


def test_stocking_export_excludes_rows_after_later_status() -> None:
    opportunity_id = prepare_approved_claims([("销售A", 2.5)])

    first = export_all_available("business_period=BATCH-EXPORT")
    assert first.status_code == 200

    with SessionLocal() as db:
        claim = db.query(models.SalesClaimForecast).filter_by(opportunity_id=opportunity_id).one()
        opportunity = db.get(models.NewProductOpportunity, opportunity_id)
        claim.downstream_status = "waiting_secondary_research"
        opportunity.current_status = "已确认不认领"
        db.commit()

    second = client.get("/stocking/available-list?business_period=BATCH-EXPORT")
    assert second.status_code == 200
    assert second.json() == []

    with SessionLocal() as db:
        claim = db.query(models.SalesClaimForecast).filter_by(opportunity_id=opportunity_id).one()
        opportunity = db.get(models.NewProductOpportunity, opportunity_id)
        rows = db.query(models.ExportRow).filter_by(claim_record_id=claim.id).all()
        transition_audits = db.query(models.AuditLog).filter_by(
            action="opportunity.waiting_arrival",
            entity_id=opportunity_id,
        ).all()
        assert claim.downstream_status == "waiting_secondary_research"
        assert opportunity.current_status == "已确认不认领"
        assert len(rows) == 1
        assert transition_audits == []


def test_legacy_null_review_only_applies_to_claims_existing_when_reviewed() -> None:
    reviewed_at = datetime(2026, 7, 1, 8, tzinfo=timezone.utc)
    with SessionLocal() as db:
        opportunity = models.NewProductOpportunity(
            source_type="selection2_caigen_claim_feedback",
            source_file="选品2.xlsx",
            source_sheet="5.26期",
            source_row=1,
            batch="BATCH-EXPORT",
            country="PH",
            site="PH",
            main_sku="MAIN-LEGACY-REVIEW",
            sub_sku="SUB-LEGACY-REVIEW",
            current_status="ready_for_stocking",
        )
        db.add(opportunity)
        db.flush()
        existing_claim = models.SalesClaimForecast(
            opportunity_id=opportunity.id,
            salesperson_name="销售已复核",
            claim_result="claim",
            claim_daily_sales=1,
            source_column="platform",
            downstream_status="waiting_export",
            created_at=reviewed_at - timedelta(minutes=1),
        )
        later_claim = models.SalesClaimForecast(
            opportunity_id=opportunity.id,
            salesperson_name="销售后认领",
            claim_result="claim",
            claim_daily_sales=2,
            source_column="platform",
            created_at=reviewed_at + timedelta(minutes=1),
        )
        db.add_all(
            [
                existing_claim,
                later_claim,
                models.ReviewRecord(
                    opportunity_id=opportunity.id,
                    reviewer_name="历史主管",
                    review_status="approved",
                    created_at=reviewed_at,
                ),
            ]
        )
        db.flush()
        existing_claim_id = existing_claim.id
        db.commit()

    response = client.get("/stocking/available-list?business_period=BATCH-EXPORT")

    assert response.status_code == 200
    assert response.json() == []


def test_repeat_export_only_includes_new_waiting_export_rows() -> None:
    prepare_approved_claims([("销售A", 1)], main_sku="MAIN-FIRST", sub_sku="SUB-FIRST", source_row=1)
    first = export_all_available("business_period=BATCH-EXPORT")
    assert first.status_code == 200

    prepare_approved_claims([("销售B", 2)], main_sku="MAIN-LATER", sub_sku="SUB-LATER", source_row=2)
    second = export_all_available("business_period=BATCH-EXPORT")

    workbook = load_workbook(BytesIO(second.content), data_only=True)
    assert {workbook["PH"][f"F{row}"].value for row in range(2, workbook["PH"].max_row + 1)} == {"MAIN-LATER"}



def test_repeatable_exports_still_exclude_disabled_opportunities() -> None:
    opportunity_id = prepare_approved_claims([("销售A", 1)])
    with SessionLocal() as db:
        opportunity = db.get(models.NewProductOpportunity, opportunity_id)
        assert opportunity is not None
        opportunity.current_status = "disabled"
        db.commit()

    response = client.get("/stocking/available-list?business_period=BATCH-EXPORT")

    assert response.status_code == 200
    assert response.json() == []


def prepare_approved_claims(
    claims: list[tuple[str, float]],
    country: str = "PH",
    site: str = "PH",
    main_sku: str = "MAIN-EXPORT",
    sub_sku: str = "SUB-EXPORT",
    source_row: int = 1,
) -> str:
    with SessionLocal() as db:
        opportunity = models.NewProductOpportunity(
            source_type="selection2_caigen_claim_feedback",
            source_file="选品2.xlsx",
            source_sheet="5.26期",
            source_row=source_row,
            batch="BATCH-EXPORT",
            country=country,
            site=site,
            developer_department="产品开发六部",
            developer_name="开发A",
            category_level1="汽摩配",
            keyword="motor cover",
            image_url="https://example.com/image.jpg",
            main_sku_name="摩托车配件",
            main_sku=main_sku,
            sub_sku_name="黑色",
            sub_sku=sub_sku,
            product_type="利润款",
            reason="新品开发",
        )
        db.add(opportunity)
        db.flush()
        for salesperson_name, daily_sales in claims:
            services.submit_claim(
                db,
                schemas.ClaimCreate(
                    opportunity_id=opportunity.id,
                    salesperson_name=salesperson_name,
                    claim_result="claim",
                    claim_daily_sales=daily_sales,
                    claim_source="caigen_self_claim",
                ),
            )
        services.submit_review(
            db,
            schemas.ReviewCreate(
                opportunity_id=opportunity.id,
                reviewer_name="练玉君",
                review_status="approved",
                review_comment="通过",
            ),
        )
        for claim in db.query(models.SalesClaimForecast).filter_by(opportunity_id=opportunity.id):
            request = db.query(models.StockingRequest).filter_by(claim_record_id=claim.id).one()
            request.application_date = date(2026, 7, 17)
            request.status = "submitted"
            claim.downstream_status = "waiting_export"
        db.commit()
        return opportunity.id


# Task 4 selected-request export contract.
def prepare_export_request(
    *, status: str = "submitted", request_type: str = "initial",
    source_type: str = "selection1_developer_claim_feedback", source_sheet: str = "开发0623期",
    business_period: str = "BATCH-SELECTED", import_batch_id: str | None = None,
    application_date: date = date(2026, 7, 17), main_sku: str = "MAIN-SELECTED",
    sub_sku: str = "SUB-SELECTED", source_row: int = 10, cost_price: float = 10,
    unit_volume: float = 0.001, daily_sales: float = 2, quantity: int = 60,
    country: str = "PH", warehouse: str | None = None, amount: float = 600,
    volume: float = 0.06, reason: str | None = None,
) -> tuple[str, str, str]:
    with SessionLocal() as db:
        opportunity = models.NewProductOpportunity(
            source_type=source_type, import_batch_id=import_batch_id, source_file="source.xlsx",
            source_sheet=source_sheet, source_row=source_row, batch=business_period,
            country="PH", site="PH", main_sku=main_sku, sub_sku=sub_sku,
            current_status="ready_for_stocking",
        )
        db.add(opportunity)
        db.flush()
        claim = models.SalesClaimForecast(
            opportunity_id=opportunity.id, salesperson_name="销售A", claim_result="claim",
            claim_daily_sales=daily_sales, source_column="platform",
            downstream_status="waiting_export" if status == "submitted" else "waiting_stocking_request",
        )
        db.add(claim)
        db.flush()
        request = models.StockingRequest(
            opportunity_id=opportunity.id, claim_record_id=claim.id, application_date=application_date,
            request_type=request_type, salesperson_name="销售A", main_sku=main_sku, sub_sku=sub_sku,
            cost_price=cost_price, unit_volume=unit_volume, daily_sales=daily_sales, quantity=quantity,
            country=country, warehouse=warehouse, amount=amount, volume=volume, reason=reason, status=status,
        )
        db.add(request)
        db.commit()
        return opportunity.id, claim.id, request.id


def export_auth_headers(name: str, role: str, dingtalk_user_id: str) -> dict[str, str]:
    with SessionLocal() as db:
        if db.query(models.RoleMapping).filter_by(dingtalk_user_id=dingtalk_user_id).first() is None:
            db.add(models.RoleMapping(name=name, role=role, dingtalk_user_id=dingtalk_user_id, enabled=True))
            db.commit()
    login = client.post("/auth/dingtalk/login", json={"dingtalk_user_id": dingtalk_user_id, "name": name})
    assert login.status_code == 200
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def export_all_available(query: str = ""):
    suffix = f"?{query}" if query else ""
    request_ids = [row["request_id"] for row in client.get(f"/stocking/available-list{suffix}").json()]
    return client.post(
        "/stocking/available-list/export",
        headers=export_auth_headers("主管A", "manager", "dt-manager"),
        json={"request_ids": request_ids},
    )


def test_available_list_uses_submitted_request_snapshot_without_review_and_preserves_filters() -> None:
    selected = prepare_export_request(
        request_type="replenishment", source_type="selection2_caigen_claim_feedback", source_sheet="5.26期",
        import_batch_id="batch-selected", application_date=date(2026, 7, 18), cost_price=12.5,
        unit_volume=0.002, daily_sales=2.01, quantity=61, country="TH", warehouse="TH仓",
        amount=762.5, volume=0.122, reason="销量增长",
    )
    prepare_export_request(status="draft", sub_sku="SUB-DRAFT", source_row=11)
    prepare_export_request(sub_sku="SUB-OTHER", source_row=12, import_batch_id="batch-other")

    response = client.get("/stocking/available-list?import_batch_id=batch-selected")

    assert response.status_code == 200
    assert response.json() == [{
        "opportunity_id": selected[0], "request_id": selected[2], "claim_record_id": selected[1],
        "business_period": "BATCH-SELECTED", "operation_status": "未操作", "time": None,
        "application_date": "2026-07-18", "stocking_type": "补货", "selection_source": "选品2/财根",
        "salesperson_name": "销售A", "main_sku": "MAIN-SELECTED", "sub_sku": "SUB-SELECTED",
        "site": "PH", "claim_daily_sales": 2.01, "quantity": 61, "stocking_country": "TH",
        "warehouse": "TH仓", "cost_price": 12.5, "unit_volume": 0.002, "amount": 762.5,
        "volume": 0.122, "replenishment_reason": "销量增长", "needs_launch_email": None,
        "launch_email_status": None, "review_status": None, "status": "submitted",
    }]


def test_selected_export_uses_exact_workbook_contract_and_persists_request_snapshots() -> None:
    opportunity_id, claim_id, request_id = prepare_export_request(
        request_type="replenishment", source_type="selection2_caigen_claim_feedback", source_sheet="5.26期",
        application_date=date(2026, 7, 18), cost_price=12.5, unit_volume=0.002, daily_sales=2.01,
        quantity=61, country="TH", warehouse="TH仓", amount=762.5, volume=0.122, reason="销量增长",
    )
    other_id = prepare_export_request(main_sku="MAIN-OTHER", sub_sku="SUB-OTHER", source_row=11)[2]
    headers = export_auth_headers("主管A", "manager", "dt-manager")

    response = client.post("/stocking/available-list/export", headers=headers, json={"request_ids": [request_id]})

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook["TH"]
    assert [cell.value for cell in sheet[1]] == [
        "操作状态", "申请日期", "备货类型", "选品数据源", "销售员", "主SKU", "子sku", "成本价",
        "单个体积", "备货单销", "备货数量", "备货国家", "备货仓库", "货值", "体积", "补货原因",
    ]
    assert sheet["B2"].value.date() == date(2026, 7, 18)
    assert sheet["B2"].number_format == "yyyy-mm-dd"
    assert sheet["C2"].value == "补货"
    assert sheet["D2"].value == "选品2/财根"
    assert "5.26期" not in sheet["D2"].value
    assert [sheet.cell(2, column).value for column in range(8, 17)] == [
        12.5, 0.002, 2.01, 61, "TH", "TH仓", 762.5, 0.122, "销量增长",
    ]
    with SessionLocal() as db:
        batch = db.query(models.ExportBatch).one()
        row = db.query(models.ExportRow).one()
        assert batch.exported_by == "主管A"
        assert batch.scope == "stocking_available"
        assert (row.opportunity_id, row.claim_record_id, row.stocking_request_id) == (
            opportunity_id, claim_id, request_id,
        )
        assert (row.application_date, row.stocking_type, row.cost_price, row.unit_volume, row.amount,
                row.volume, row.replenishment_reason) == (
            date(2026, 7, 18), "replenishment", 12.5, 0.002, 762.5, 0.122, "销量增长",
        )
        assert row.selection_source == "选品2/财根"
        assert db.get(models.StockingRequest, request_id).status == "exported"
        assert db.get(models.SalesClaimForecast, claim_id).downstream_status == "waiting_arrival"
        assert db.get(models.NewProductOpportunity, opportunity_id).current_status == "ready_for_stocking"
    assert [row["request_id"] for row in client.get("/stocking/available-list").json()] == [other_id]


def test_selected_export_only_advances_selected_claim_when_opportunity_has_siblings() -> None:
    opportunity_id = prepare_approved_claims([("销售A", 1), ("销售B", 2)])
    headers = export_auth_headers("主管A", "manager", "dt-manager")
    with SessionLocal() as db:
        claims = db.query(models.SalesClaimForecast).filter_by(opportunity_id=opportunity_id).order_by(
            models.SalesClaimForecast.salesperson_name
        ).all()
        selected_claim, sibling_claim = claims
        selected_request = db.query(models.StockingRequest).filter_by(claim_record_id=selected_claim.id).one()
        sibling_request = db.query(models.StockingRequest).filter_by(claim_record_id=sibling_claim.id).one()
        selected_ids = (selected_claim.id, selected_request.id)
        sibling_ids = (sibling_claim.id, sibling_request.id)

    response = client.post(
        "/stocking/available-list/export", headers=headers, json={"request_ids": [selected_ids[1]]},
    )

    assert response.status_code == 200
    with SessionLocal() as db:
        assert db.get(models.StockingRequest, selected_ids[1]).status == "exported"
        assert db.get(models.SalesClaimForecast, selected_ids[0]).downstream_status == "waiting_arrival"
        assert db.get(models.StockingRequest, sibling_ids[1]).status == "submitted"
        assert db.get(models.SalesClaimForecast, sibling_ids[0]).downstream_status == "waiting_export"
        assert db.get(models.NewProductOpportunity, opportunity_id).current_status == "ready_for_stocking"


def test_selected_export_rejects_invalid_or_stale_request_sets_without_mutation() -> None:
    request_id = prepare_export_request()[2]
    draft_id = prepare_export_request(status="draft", sub_sku="SUB-DRAFT", source_row=11)[2]
    headers = export_auth_headers("主管A", "manager", "dt-manager")

    assert client.get("/stocking/available-list/export").status_code == 405
    assert client.post("/stocking/available-list/export", headers=headers, json={"request_ids": []}).status_code == 422
    assert client.post(
        "/stocking/available-list/export", headers=headers, json={"request_ids": [request_id, request_id]},
    ).status_code == 422
    assert client.post(
        "/stocking/available-list/export", headers=headers, json={"request_ids": ["missing-request"]},
    ).status_code == 404
    assert client.post(
        "/stocking/available-list/export", headers=headers, json={"request_ids": [request_id, draft_id]},
    ).status_code == 409
    with SessionLocal() as db:
        assert db.query(models.ExportBatch).count() == 0
        assert db.query(models.ExportRow).count() == 0
        assert db.get(models.StockingRequest, request_id).status == "submitted"
        assert db.get(models.StockingRequest, draft_id).status == "draft"


@pytest.mark.parametrize("ineligible_reason", ["disabled", "rejected", "non_platform"])
def test_selected_export_rejects_requests_absent_from_available_list(ineligible_reason: str) -> None:
    opportunity_id, claim_id, request_id = prepare_export_request()
    headers = export_auth_headers("主管A", "manager", "dt-manager")
    with SessionLocal() as db:
        opportunity = db.get(models.NewProductOpportunity, opportunity_id)
        claim = db.get(models.SalesClaimForecast, claim_id)
        if ineligible_reason == "disabled":
            opportunity.current_status = "disabled"
        elif ineligible_reason == "rejected":
            claim.claim_result = "reject"
        else:
            claim.source_column = "CC:CH"
        db.commit()

    assert client.get("/stocking/available-list").json() == []
    response = client.post(
        "/stocking/available-list/export", headers=headers, json={"request_ids": [request_id]},
    )

    assert response.status_code == 409
    with SessionLocal() as db:
        assert db.query(models.ExportBatch).count() == 0
        assert db.get(models.StockingRequest, request_id).status == "submitted"
        assert db.get(models.SalesClaimForecast, claim_id).downstream_status == "waiting_export"


def test_selected_export_build_failure_rolls_back_and_reexport_is_conflict(monkeypatch: pytest.MonkeyPatch) -> None:
    _, claim_id, request_id = prepare_export_request()
    headers = export_auth_headers("主管A", "manager", "dt-manager")
    original_build = services.build_available_stocking_workbook

    def fail_build(*_args, **_kwargs):
        raise RuntimeError("workbook build failed")

    monkeypatch.setattr(services, "build_available_stocking_workbook", fail_build)
    with pytest.raises(RuntimeError, match="workbook build failed"):
        client.post("/stocking/available-list/export", headers=headers, json={"request_ids": [request_id]})
    with SessionLocal() as db:
        assert db.query(models.ExportBatch).count() == 0
        assert db.get(models.StockingRequest, request_id).status == "submitted"
        assert db.get(models.SalesClaimForecast, claim_id).downstream_status == "waiting_export"

    monkeypatch.setattr(services, "build_available_stocking_workbook", original_build)
    assert client.post(
        "/stocking/available-list/export", headers=headers, json={"request_ids": [request_id]},
    ).status_code == 200
    assert client.post(
        "/stocking/available-list/export", headers=headers, json={"request_ids": [request_id]},
    ).status_code == 409
    with SessionLocal() as db:
        assert db.query(models.ExportBatch).count() == 1
        assert db.query(models.ExportRow).count() == 1


def test_operator_cannot_export_selected_requests() -> None:
    request_id = prepare_export_request()[2]
    headers = export_auth_headers("销售A", "operator", "dt-operator")

    response = client.post("/stocking/available-list/export", headers=headers, json={"request_ids": [request_id]})

    assert response.status_code == 403
