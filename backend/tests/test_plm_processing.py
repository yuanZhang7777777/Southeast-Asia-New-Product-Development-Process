import os
import sys
from datetime import date, timezone
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_plm_processing.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from app import models  # noqa: E402
from app.db import Base  # noqa: E402
from app import plm_processing  # noqa: E402
from app.plm_processing import process_plm_arrival_workbook  # noqa: E402

GROUP_EIGHT = "\u96c6\u56e2\u516b\u90e8"
GROUP_ONE = "\u96c6\u56e2\u4e00\u90e8"
SALES_A = "\u9500\u552eA"
SALES_B = "\u9500\u552eB"
PH = "\u83f2\u5f8b\u5bbe"
TH = "\u6cf0\u56fd"

engine = create_engine(os.environ["DATABASE_URL"])
SessionLocal = sessionmaker(bind=engine)


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def test_plm_processing_persists_all_rows_and_dry_run_does_not_open_secondary_research(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    with SessionLocal() as db:
        opportunity, claim = make_claim(" sub-a ", "PH", SALES_A, "waiting_arrival")
        db.add_all([opportunity, claim])
        db.flush()
        add_stocking_export(db, opportunity, claim)

        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=False,
        )

        assert result["row_count"] == 3
        assert result["new_arrival_count"] == 3
        assert result["restock_count"] == 0
        assert result["unknown_count"] == 0
        assert result["matched_count"] == 1
        assert result["arrival_record_count"] == 0
        assert result["planned_responsibilities"] == [
            {
                "claim_record_id": claim.id,
                "opportunity_id": opportunity.id,
                "business_period": "2026-07-12",
                "salesperson_name": SALES_A,
                "site": "PH",
                "country": "PH",
                "main_sku": "MAIN-1",
                "sub_sku": " sub-a ",
                "product_name": "new",
            }
        ]
        assert db.query(models.PlmArrivalBatch).count() == 1
        items = db.query(models.PlmArrivalItem).order_by(models.PlmArrivalItem.source_row).all()
        assert [item.arrival_type for item in items] == ["new_arrival", "new_arrival", "new_arrival"]
        assert [item.match_status for item in items] == ["matched_dry_run", "duplicate_claim_in_batch", "unmatched"]
        assert items[0].product_name == "new"
        assert items[0].raw_payload["_matched_claim_record_ids"] == [claim.id]
        assert db.get(models.SalesClaimForecast, claim.id).downstream_status == "waiting_arrival"
        assert db.query(models.ArrivalRecord).count() == 0


def test_plm_processing_automation_uses_utc_arrival_time_and_dedupes_source_hash(tmp_path: Path, monkeypatch) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    opened_arrivals = []

    original_open_secondary_research = plm_processing.services.open_secondary_research

    def capture_open_secondary_research(db, claim_record_id, arrived_at=None):
        opened_arrivals.append(arrived_at)
        return original_open_secondary_research(db, claim_record_id, arrived_at)

    monkeypatch.setattr(plm_processing.services, "open_secondary_research", capture_open_secondary_research)

    with SessionLocal() as db:
        opportunity, claim = make_claim("SUB-A", "PH", SALES_A, "waiting_arrival")
        db.add_all([opportunity, claim])
        db.flush()
        add_stocking_export(db, opportunity, claim)

        db.commit()

        first = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )
        second = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        assert first["arrival_record_count"] == 1
        assert first["matched_count"] == 1
        assert second["status"] == "duplicate"
        assert second["planned_responsibilities"][0]["claim_record_id"] == claim.id
        assert db.query(models.PlmArrivalBatch).count() == 1
        assert db.query(models.PlmArrivalItem).count() == 3
        assert db.query(models.ArrivalRecord).count() == 1
        assert opened_arrivals[0].tzinfo == timezone.utc
        assert opened_arrivals[0].isoformat() == "2026-07-12T02:00:00+00:00"
        # SQLite drops tzinfo on roundtrip; the monkeypatch above asserts the pre-flush value is UTC-aware.
        assert db.query(models.ArrivalRecord).one().arrived_at.isoformat() == "2026-07-12T02:00:00"
        saved_claim = db.get(models.SalesClaimForecast, claim.id)
        assert saved_claim.downstream_status == "waiting_secondary_research"
        assert saved_claim.arrival_detected_at is not None


def test_plm_processing_allows_same_workbook_for_different_first_listing_dates(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm-multi-day.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "商品名称",
            "国家",
            "子SKU",
            "主SKU",
            "销售员",
            "集团",
            "海外仓",
            "最后一次入库时间",
            "首次上架时间",
        ]
    )
    sheet.append(["day1", PH, "SUB-A", "MAIN-1", SALES_A, GROUP_EIGHT, "PH仓", "2026-07-12 10:00:00", "2026-07-12 00:00:00"])
    sheet.append(["day2", PH, "SUB-B", "MAIN-2", SALES_A, GROUP_EIGHT, "PH仓", "2026-07-12 10:00:00", "2026-07-13 00:00:00"])
    workbook.save(workbook_path)

    with SessionLocal() as db:
        day1 = process_plm_arrival_workbook(db, workbook_path, "2026-07-12", source_file="plm.xlsx", bloc_name=GROUP_EIGHT)
        day2 = process_plm_arrival_workbook(db, workbook_path, "2026-07-13", source_file="plm.xlsx", bloc_name=GROUP_EIGHT)

        assert day1["status"] == "processed"
        assert day2["status"] == "processed"
        assert db.query(models.PlmArrivalBatch).count() == 2
        items = (
            db.query(models.PlmArrivalItem)
            .join(models.PlmArrivalBatch, models.PlmArrivalBatch.id == models.PlmArrivalItem.batch_id)
            .order_by(models.PlmArrivalBatch.arrival_date, models.PlmArrivalItem.source_row)
            .all()
        )
        assert [item.product_name for item in items] == [
            "day1",
            "day2",
            "day1",
            "day2",
        ]


def test_plm_processing_does_not_duplicate_arrival_record_for_same_claim_across_batches(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    updated_path = tmp_path / "plm-updated.xlsx"
    build_workbook(updated_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "商品名称",
            "国家",
            "子SKU",
            "主SKU",
            "销售员",
            "集团",
            "海外仓",
            "最后一次入库时间",
            "首次上架时间",
        ]
    )
    sheet.append(["same claim", PH, "SUB-A", "MAIN-1", SALES_A, GROUP_EIGHT, "PH仓", "2026-07-12 10:00:00", "2026-07-12 00:00:00"])
    sheet.append(["different bytes", PH, "SUB-X", "MAIN-X", SALES_A, GROUP_ONE, "PH仓", "2026-07-12", "2026-07-12"])
    workbook.save(updated_path)

    with SessionLocal() as db:
        opportunity, claim = make_claim("SUB-A", "PH", SALES_A, "waiting_arrival")
        db.add_all([opportunity, claim])
        db.flush()
        add_stocking_export(db, opportunity, claim)
        db.commit()

        first = process_plm_arrival_workbook(db, workbook_path, "2026-07-12", source_file="plm.xlsx", bloc_name=GROUP_EIGHT, workflow_automation_enabled=True)
        second = process_plm_arrival_workbook(db, updated_path, "2026-07-12", source_file="plm-updated.xlsx", bloc_name=GROUP_EIGHT, workflow_automation_enabled=True)

        assert first["arrival_record_count"] == 1
        assert second["status"] == "processed"
        assert db.query(models.ArrivalRecord).filter_by(claim_record_id=claim.id).count() == 1


def test_plm_processing_queues_missing_system_sku_for_manager_assignment(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm-discovery.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "商品名称",
            "国家",
            "子SKU",
            "主SKU",
            "销售员",
            "集团",
            "海外仓",
            "最后一次入库时间",
            "首次上架时间",
            "海外仓可发",
            "真仓库存",
        ]
    )
    sheet.append(["missing system sku", PH, "MISS-SUB", "MISS-MAIN", SALES_A, GROUP_EIGHT, "PH仓", "2026-07-12 10:00:00", "2026-07-12 00:00:00", 2, 2])
    workbook.save(workbook_path)

    with SessionLocal() as db:
        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm-discovery.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).one()
        assert result["matched_count"] == 0
        assert result["arrival_record_count"] == 0
        assert item.match_status == "pending_assignment"
        assert item.matched_claim_record_id is None
        assert item.raw_payload["_plm_assignment"]["status"] == "pending_assignment"
        assert item.raw_payload["_plm_assignment"]["salesperson_name"] == SALES_A
        assert db.query(models.NewProductOpportunity).count() == 0
        assert db.query(models.SalesClaimForecast).count() == 0
        assert db.query(models.ArrivalRecord).count() == 0
        assert db.query(models.FlowTask).count() == 0


def test_plm_processing_does_not_duplicate_pending_assignment_across_batches(tmp_path: Path) -> None:
    first_path = tmp_path / "plm-first.xlsx"
    second_path = tmp_path / "plm-second.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "商品名称",
            "国家",
            "子SKU",
            "主SKU",
            "销售员",
            "集团",
            "海外仓",
            "最后一次入库时间",
            "首次上架时间",
        ]
    )
    sheet.append(["missing system sku", TH, "MISS-SUB", "MISS-MAIN", SALES_A, GROUP_ONE, "TH仓", "2026-08-10 10:00:00", "2026-08-10 08:00:00"])
    workbook.save(first_path)
    sheet.append(["different sku", TH, "NEW-SUB", "NEW-MAIN", SALES_A, GROUP_ONE, "TH仓", "2026-08-10 11:00:00", "2026-08-10 08:00:00"])
    workbook.save(second_path)

    with SessionLocal() as db:
        first = process_plm_arrival_workbook(
            db,
            first_path,
            "2026-08-10",
            source_file="plm-first.xlsx",
            workflow_automation_enabled=True,
        )
        second = process_plm_arrival_workbook(
            db,
            second_path,
            "2026-08-10",
            source_file="plm-second.xlsx",
            workflow_automation_enabled=True,
        )

        assert first["status"] == "processed"
        assert second["status"] == "processed"
        statuses = [
            item.match_status
            for item in (
                db.query(models.PlmArrivalItem)
                .join(models.PlmArrivalBatch, models.PlmArrivalBatch.id == models.PlmArrivalItem.batch_id)
                .order_by(models.PlmArrivalBatch.source_file, models.PlmArrivalItem.source_row)
                .all()
            )
        ]
        assert statuses == ["pending_assignment", "duplicate_plm_arrival_item", "pending_assignment"]
        pending_items = db.query(models.PlmArrivalItem).filter_by(match_status="pending_assignment").all()
        assert {item.sub_sku for item in pending_items} == {"MISS-SUB", "NEW-SUB"}
        assert db.query(models.NewProductOpportunity).count() == 0
        assert db.query(models.SalesClaimForecast).count() == 0
        assert db.query(models.ArrivalRecord).count() == 0


def test_plm_processing_routes_missing_sku_to_plm_salesperson_when_they_own_site(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm-discovery-site-owner.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "商品名称",
            "国家",
            "子SKU",
            "主SKU",
            "销售员",
            "集团",
            "海外仓",
            "最后一次入库时间",
            "首次上架时间",
            "海外仓可发",
            "真仓库存",
        ]
    )
    sheet.append(["site owner sku", PH, "OWN-SUB", "OWN-MAIN", SALES_A, GROUP_EIGHT, "PH仓", "2026-07-12 10:00:00", "2026-07-12 00:00:00", 2, 2])
    workbook.save(workbook_path)

    with SessionLocal() as db:
        user = models.User(name=SALES_A, enabled=True)
        db.add(user)
        db.flush()
        db.add_all(
            [
                models.RoleMapping(user_id=user.id, name=SALES_A, role="operator", enabled=True, notification_enabled=True),
                models.OperatorAssignmentProfile(operator_name=SALES_A, key_site="PH", enabled=True),
            ]
        )
        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm-discovery-site-owner.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).one()
        opportunity = db.query(models.NewProductOpportunity).filter_by(source_type="plm_arrival_discovery").one()
        claim = db.query(models.SalesClaimForecast).one()
        arrival = db.query(models.ArrivalRecord).one()

        assert result["matched_count"] == 1
        assert result["arrival_record_count"] == 1
        assert item.match_status == "matched_discovered"
        assert item.matched_claim_record_id == claim.id
        assert item.raw_payload["_plm_discovery"]["status"] == "matched_discovered"
        assert item.raw_payload["_plm_discovery"]["salesperson_name"] == SALES_A
        assert opportunity.batch == "PLM新增到货"
        assert opportunity.main_sku == "OWN-MAIN"
        assert opportunity.sub_sku == "OWN-SUB"
        assert claim.salesperson_name == SALES_A
        assert claim.downstream_status == "waiting_secondary_research"
        assert arrival.salesperson_name == SALES_A
        assert db.query(models.FlowTask).count() == 0


def test_plm_processing_queues_missing_sku_when_plm_salesperson_does_not_own_site(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm-discovery-site-mismatch.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "商品名称",
            "国家",
            "子SKU",
            "主SKU",
            "销售员",
            "集团",
            "海外仓",
            "最后一次入库时间",
            "首次上架时间",
            "海外仓可发",
            "真仓库存",
        ]
    )
    sheet.append(["site mismatch sku", PH, "OWN-SUB", "OWN-MAIN", SALES_A, GROUP_EIGHT, "PH仓", "2026-07-12 10:00:00", "2026-07-12 00:00:00", 2, 2])
    workbook.save(workbook_path)

    with SessionLocal() as db:
        user = models.User(name=SALES_A, enabled=True)
        db.add(user)
        db.flush()
        db.add_all(
            [
                models.RoleMapping(user_id=user.id, name=SALES_A, role="operator", enabled=True, notification_enabled=True),
                models.OperatorAssignmentProfile(operator_name=SALES_A, key_site="TH", enabled=True),
            ]
        )
        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm-discovery-site-mismatch.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).one()

        assert result["matched_count"] == 0
        assert result["arrival_record_count"] == 0
        assert item.match_status == "pending_assignment"
        assert item.matched_claim_record_id is None
        assert item.raw_payload["_plm_assignment"]["status"] == "pending_assignment"
        assert item.raw_payload["_plm_assignment"]["salesperson_name"] == SALES_A
        assert "等待主管/超管指派" in item.raw_payload["_plm_assignment"]["reason"]
        assert db.query(models.NewProductOpportunity).filter_by(source_type="plm_arrival_discovery").count() == 0
        assert db.query(models.SalesClaimForecast).count() == 0
        assert db.query(models.ArrivalRecord).count() == 0
        assert db.query(models.FlowTask).count() == 0


def test_plm_processing_matches_current_claim_by_country_and_sub_sku_only(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm-country-sub.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "商品名称",
            "国家",
            "子SKU",
            "主SKU",
            "销售员",
            "集团",
            "海外仓",
            "最后一次入库时间",
            "首次上架时间",
            "海外仓可发",
        ]
    )
    sheet.append(["same country stock", PH, "SUB-A", "PLM-MAIN", SALES_B, GROUP_ONE, "PH仓", "2026-07-12 10:00:00", "2026-07-14 09:00:00", 8])
    workbook.save(workbook_path)

    with SessionLocal() as db:
        opportunity, claim = make_claim("SUB-A", "PH", SALES_A, "waiting_arrival")
        opportunity.main_sku = "SYSTEM-MAIN"
        db.add_all([opportunity, claim])
        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm-country-sub.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).one()
        saved_claim = db.get(models.SalesClaimForecast, claim.id)
        arrival = db.query(models.ArrivalRecord).one()
        assert result["row_count"] == 1
        assert result["matched_count"] == 1
        assert result["arrival_record_count"] == 1
        assert item.match_status == "matched"
        assert item.matched_claim_record_id == claim.id
        assert item.raw_payload["bloc_name"] == GROUP_ONE
        assert saved_claim.salesperson_name == SALES_A
        assert saved_claim.downstream_status == "waiting_secondary_research"
        assert arrival.salesperson_name == SALES_A


def test_plm_processing_exact_match_ignores_non_platform_claim(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    with SessionLocal() as db:
        opportunity, claim = make_claim("SUB-A", "PH", SALES_A, "waiting_arrival", source_column="CC:CH")
        db.add_all([opportunity, claim])
        db.flush()
        add_stocking_export(db, opportunity, claim)

        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=False,
        )

        assert result["matched_count"] == 0
        assert result["planned_responsibilities"] == []
        assert db.get(models.SalesClaimForecast, claim.id).downstream_status == "waiting_arrival"
        assert db.query(models.PlmArrivalItem).filter_by(matched_claim_record_id=claim.id).count() == 0


def test_plm_processing_matches_the_country_snapshotted_at_export() -> None:
    with SessionLocal() as db:
        opportunity, claim = make_claim("SUB-A", "PH", SALES_A, "waiting_arrival")
        db.add_all([opportunity, claim])
        db.flush()
        add_stocking_export(db, opportunity, claim)
        opportunity.country = "TH"
        opportunity.site = "TH"
        db.commit()

        ph_matches = plm_processing._exact_claim_matches(
            db,
            {"sub_sku": "SUB-A", "country": "PH", "salesperson_name": SALES_A},
        )
        th_matches = plm_processing._exact_claim_matches(
            db,
            {"sub_sku": "SUB-A", "country": "TH", "salesperson_name": SALES_A},
        )

        assert [matched_claim.id for matched_claim, _ in ph_matches] == [claim.id]
        assert th_matches == []


def test_plm_processing_queues_ambiguous_country_sub_sku_matches(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    with SessionLocal() as db:
        claims = []
        for period in ["2026-W28", "2026-W29"]:
            opportunity, claim = make_claim("SUB-A", "PH", SALES_A, "waiting_arrival", business_period=period)
            db.add_all([opportunity, claim])
            db.flush()
            add_stocking_export(db, opportunity, claim)

            claims.append(claim)
        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).filter_by(source_row=2).one()
        assert result["arrival_record_count"] == 0
        assert result["matched_count"] == 0
        assert result["planned_responsibilities"] == []
        assert item.match_status == "existing_system_sku"
        assert len(item.raw_payload["_matching_current_products"]) == 2
        assert db.query(models.ArrivalRecord).count() == 0
        assert {db.get(models.SalesClaimForecast, claim.id).downstream_status for claim in claims} == {
            "waiting_arrival"
        }


def build_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "\u5546\u54c1\u540d\u79f0",
            "\u56fd\u5bb6",
            "\u5b50SKU",
            "\u4e3bSKU",
            "\u9500\u552e\u5458",
            "\u96c6\u56e2",
            "\u6d77\u5916\u4ed3",
            "\u6700\u540e\u4e00\u6b21\u5165\u5e93\u65f6\u95f4",
            "\u9996\u6b21\u4e0a\u67b6\u65f6\u95f4",
            "\u6d77\u5916\u4ed3\u53ef\u53d1",
            "\u771f\u4ed3\u5e93\u5b58",
        ]
    )
    sheet.append(["new", PH, " SUB-A ", "MAIN-1", f" {SALES_A} ", GROUP_EIGHT, "PH\u4ed3", "2026-07-12 10:00:00", "2026-07-12 00:00:00", 12, 30])
    sheet.append(["restock", PH, "SUB-B", "MAIN-2", SALES_A, GROUP_EIGHT, "PH\u4ed3", "2026-07-12 11:00:00", "2026-07-01", 3, 8])
    sheet.append(["unknown", TH, "SUB-C", "MAIN-3", SALES_B, GROUP_EIGHT, "TH\u4ed3", "2026-07-12 12:00:00", None, 4, 9])
    sheet.append(["wrong sales", PH, "SUB-A", "MAIN-1", SALES_B, GROUP_EIGHT, "PH\u4ed3", "2026-07-12 13:00:00", "2026-07-12", 5, 10])
    sheet.append(["other group", PH, "SUB-D", "MAIN-4", SALES_A, GROUP_ONE, "PH\u4ed3", "2026-07-12", "2026-07-12", 1, 1])
    workbook.save(path)


def make_claim(
    sub_sku: str,
    site: str,
    salesperson_name: str,
    downstream_status: str,
    source_column: str = "platform",
    business_period: str = "2026-07-12",
) -> tuple[models.NewProductOpportunity, models.SalesClaimForecast]:
    opportunity = models.NewProductOpportunity(
        id=models.new_id(),
        source_type="selection1_developer_claim_feedback",
        source_file="selection.xlsx",
        source_sheet="sheet",
        source_row=1,
        batch=business_period,
        country=site,
        site=site,
        main_sku="MAIN-1",
        sub_sku=sub_sku,
        current_status="claim_submitted",
        snapshot={},
    )
    claim = models.SalesClaimForecast(
        opportunity_id=opportunity.id,
        salesperson_name=salesperson_name,
        claim_result="claim",
        source_column=source_column,
        downstream_status=downstream_status,
    )
    return opportunity, claim


def add_stocking_export(db, opportunity, claim) -> None:
    request = models.StockingRequest(
        opportunity_id=opportunity.id, claim_record_id=claim.id, application_date=date(2026, 7, 17),
        request_type="initial", salesperson_name=claim.salesperson_name, main_sku=opportunity.main_sku,
        sub_sku=opportunity.sub_sku, daily_sales=1, quantity=30, status="exported",
    )
    db.add(request)
    db.flush()
    batch = models.ExportBatch(file_name="stocking.xlsx", scope="stocking_available", row_count=1)
    db.add(batch)
    db.flush()
    db.add(models.ExportRow(
        export_batch_id=batch.id, opportunity_id=opportunity.id, claim_record_id=claim.id,
        stocking_request_id=request.id, salesperson_name=claim.salesperson_name,
        main_sku=opportunity.main_sku, sub_sku=opportunity.sub_sku,
        claim_daily_sales=1, stocking_quantity=30, country=opportunity.country,
    ))

def test_plm_matches_current_platform_claims_with_or_without_stocking_export() -> None:
    row = {"sub_sku": "SUB-A", "country": "PH", "salesperson_name": SALES_A}
    with SessionLocal() as db:
        waiting_export = add_export_candidate(db, "waiting_export", "stocking_available")
        no_export_row = add_export_candidate(db, "waiting_arrival", None)
        traceability_only = add_export_candidate(db, "waiting_arrival", "traceability")
        sales_self = add_export_candidate(db, "waiting_arrival", "stocking_available", source_type="sales_self_selection")
        historical_selection1_opportunity, historical_selection1 = make_claim(
            "SUB-A",
            "PH",
            SALES_A,
            "waiting_secondary_research",
            source_column="history_selection1",
        )
        db.add_all([historical_selection1_opportunity, historical_selection1])
        db.commit()

        matches = plm_processing._exact_claim_matches(db, row)

        assert {claim.id for claim, _ in matches} == {
            waiting_export.id,
            no_export_row.id,
            traceability_only.id,
            sales_self.id,
            historical_selection1.id,
        }
        assert db.query(models.ReviewRecord).count() == 0


def test_plm_processing_opens_secondary_for_existing_system_claim_without_export(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    with SessionLocal() as db:
        opportunity, claim = make_claim("SUB-A", "PH", SALES_A, "waiting_export")
        db.add_all([opportunity, claim])
        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).filter_by(source_row=2).one()
        saved_claim = db.get(models.SalesClaimForecast, claim.id)
        assert result["matched_count"] == 1
        assert result["arrival_record_count"] == 1
        assert item.match_status == "matched"
        assert saved_claim.downstream_status == "waiting_secondary_research"


def add_export_candidate(
    db, downstream_status: str, scope: str | None, *, source_type: str = "selection1_developer_claim_feedback",
) -> models.SalesClaimForecast:
    opportunity, claim = make_claim("SUB-A", "PH", SALES_A, downstream_status)
    opportunity.id = models.new_id()
    opportunity.source_type = source_type
    claim.opportunity_id = opportunity.id
    db.add_all([opportunity, claim])
    db.flush()
    request = models.StockingRequest(
        opportunity_id=opportunity.id, claim_record_id=claim.id, application_date=date(2026, 7, 17),
        request_type="initial", salesperson_name=SALES_A, main_sku=opportunity.main_sku,
        sub_sku=opportunity.sub_sku, daily_sales=1, quantity=30, status="exported",
    )
    db.add(request)
    db.flush()
    if scope:
        batch = models.ExportBatch(file_name=f"{scope}.xlsx", scope=scope, row_count=1)
        db.add(batch)
        db.flush()
        db.add(models.ExportRow(
            export_batch_id=batch.id, opportunity_id=opportunity.id, claim_record_id=claim.id,
            stocking_request_id=request.id, salesperson_name=SALES_A, main_sku=opportunity.main_sku,
            sub_sku=opportunity.sub_sku, claim_daily_sales=1, stocking_quantity=30, country="PH",
        ))
    return claim


def test_plm_processing_skips_secondary_research_when_active_item_exists(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    with SessionLocal() as db:
        opportunity, claim = make_claim("SUB-A", "PH", SALES_A, "waiting_arrival")
        db.add_all([opportunity, claim])
        db.flush()
        add_stocking_export(db, opportunity, claim)
        listing = models.ListingRecord(
            source_group_key="history:MAIN-1",
            source_claim_ids=[],
            source_type="history_listing",
            business_period=opportunity.batch,
            country="PH",
            site="PH",
            main_sku=opportunity.main_sku,
            salesperson_name=SALES_A,
            shop="Shopee-PH",
            item="ITEM-1",
            listing_strategy="history",
            status="active",
            tracking_status="active",
        )
        db.add(listing)
        db.flush()
        db.add(models.ListingSkuBinding(
            listing_record_id=listing.id,
            main_sku=opportunity.main_sku,
            sub_sku=opportunity.sub_sku,
            binding_source="history_listing",
        ))
        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).filter_by(source_row=2).one()
        assert result["matched_count"] == 0
        assert result["arrival_record_count"] == 0
        assert item.match_status == "already_listed"
        assert db.get(models.SalesClaimForecast, claim.id).downstream_status == "waiting_arrival"


def test_plm_processing_matches_selection34_business_period_claim(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    with SessionLocal() as db:
        archived = models.NewProductOpportunity(
            source_type="history_selection34",
            source_file="history.xlsx",
            source_sheet="直发热销转0630期",
            source_row=74,
            batch="直发热销转0630期",
            country=PH,
            site=PH,
            main_sku="MAIN-1",
            sub_sku="SUB-A",
            current_status="historical_archive",
            snapshot={},
        )
        db.add(archived)
        db.flush()
        db.add(
            models.SalesClaimForecast(
                opportunity_id=archived.id,
                salesperson_name=SALES_A,
                claim_result="claim",
                claim_daily_sales=1,
                source_column="history_selection34",
                claim_source="history_selection34",
            )
        )
        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).filter_by(source_row=2).one()
        assert result["matched_count"] == 1
        assert result["arrival_record_count"] == 1
        assert item.match_status == "matched"
        assert item.matched_claim_record_id is not None
        assert db.query(models.NewProductOpportunity).filter_by(source_type="plm_arrival_discovery").count() == 0
        assert db.query(models.SalesClaimForecast).count() == 1
        assert db.query(models.ArrivalRecord).count() == 1
        saved_claim = db.query(models.SalesClaimForecast).one()
        saved_opportunity = db.get(models.NewProductOpportunity, archived.id)
        assert saved_claim.downstream_status == "waiting_secondary_research"
        assert saved_opportunity.current_status == "waiting_secondary_research"


def test_activate_existing_system_sku_discoveries_reopens_items_blocked_by_archive() -> None:
    with SessionLocal() as db:
        batch = models.PlmArrivalBatch(
            arrival_date="2026-08-10",
            source_file="plm-2026-08-10.xlsx",
            source_hash="hash-0810",
            bloc_name=GROUP_EIGHT,
            row_count=1,
        )
        db.add(batch)
        db.flush()
        item = models.PlmArrivalItem(
            batch_id=batch.id,
            source_sheet="汇总表格",
            source_row=850,
            arrival_type="new_arrival",
            product_name="旧档案挡住的新品",
            salesperson_name=SALES_A,
            country=PH,
            warehouse="PH仓",
            main_sku="MAIN-1",
            sub_sku="SUB-A",
            match_status="existing_system_sku",
            raw_payload={},
        )
        archived = models.NewProductOpportunity(
            source_type="history_selection34",
            source_file="history.xlsx",
            source_sheet="直发热销转0630期",
            source_row=74,
            batch="直发热销转0630期",
            country=PH,
            site=PH,
            main_sku="MAIN-1",
            sub_sku="SUB-A",
            current_status="historical_archive",
            snapshot={},
        )
        db.add_all([item, archived])
        db.commit()

        dry_run = plm_processing.activate_existing_system_sku_discoveries(db, arrival_date="2026-08-10", apply=False)
        assert dry_run["would_activate"] == 0
        assert dry_run["skipped"] == 1
        assert dry_run["items"][0]["status"] == "still_existing_system_sku"
        assert db.query(models.NewProductOpportunity).filter_by(source_type="plm_arrival_discovery").count() == 0

        result = plm_processing.activate_existing_system_sku_discoveries(db, arrival_date="2026-08-10", apply=True)

        assert db.query(models.NewProductOpportunity).filter_by(source_type="plm_arrival_discovery").count() == 0
        saved_item = db.get(models.PlmArrivalItem, item.id)
        assert result["activated"] == 0
        assert result["queued_for_assignment"] == 0
        assert result["skipped"] == 1
        assert saved_item.match_status == "existing_system_sku"
        assert saved_item.matched_claim_record_id is None
        assert db.query(models.SalesClaimForecast).count() == 0
        assert db.query(models.ArrivalRecord).count() == 0

def test_active_listing_keys_include_shared_item_binding() -> None:
    with SessionLocal() as db:
        listing = models.ListingRecord(
            source_group_key="history:REP-MAIN",
            source_claim_ids=[],
            source_type="history_listing",
            business_period="history",
            country="PH",
            site="PH",
            main_sku="REP-MAIN",
            salesperson_name=SALES_A,
            shop="Shopee-PH",
            item="ITEM-SHARED",
            listing_strategy="history",
            status="active",
            tracking_status="active",
        )
        db.add(listing)
        db.flush()
        db.add(models.ListingSkuBinding(
            listing_record_id=listing.id,
            main_sku="MAIN-1",
            sub_sku="SUB-A",
            binding_source="history_listing",
        ))
        db.commit()

        assert ("MAIN-1", "PH") in plm_processing._active_listing_keys(db)

def test_plm_processing_activates_historical_claim_after_verified_new_arrival(tmp_path: Path) -> None:
    workbook_path = tmp_path / "plm.xlsx"
    build_workbook(workbook_path)
    with SessionLocal() as db:
        user = models.User(name=SALES_A, dingtalk_user_id="dt-history-sales", enabled=True)
        db.add(user)
        db.flush()
        db.add(models.RoleMapping(
            user_id=user.id,
            name=SALES_A,
            dingtalk_user_id=user.dingtalk_user_id,
            role="operator",
            enabled=True,
        ))
        opportunity = models.NewProductOpportunity(
            source_type="history_selection1",
            source_file="selection1.xlsx",
            source_sheet="开发0623期",
            source_row=2,
            batch="开发0623期",
            country="PH",
            site="PH",
            main_sku="MAIN-1",
            sub_sku="SUB-A",
            current_status="historical_archive",
            snapshot={},
        )
        db.add(opportunity)
        db.flush()
        db.add(models.SalesClaimForecast(
            opportunity_id=opportunity.id,
            salesperson_name=SALES_A,
            claim_result="claim",
            claim_daily_sales=1.5,
            source_column="history_selection1",
            claim_source="history_selection1",
        ))
        db.commit()

        result = process_plm_arrival_workbook(
            db,
            workbook_path,
            "2026-07-12",
            source_file="plm.xlsx",
            bloc_name=GROUP_EIGHT,
            workflow_automation_enabled=True,
        )

        item = db.query(models.PlmArrivalItem).filter_by(source_row=2).one()
        active_claim = db.query(models.SalesClaimForecast).filter_by(source_column="platform").one()
        assert result["matched_count"] == 1
        assert item.match_status == "matched_historical"
        assert active_claim.salesperson_name == SALES_A
        assert active_claim.downstream_status == "waiting_secondary_research"
        assert db.query(models.ArrivalRecord).filter_by(claim_record_id=active_claim.id).count() == 1
