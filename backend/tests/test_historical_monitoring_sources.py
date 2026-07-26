from pathlib import Path
import sys

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.historical_monitoring_sources import audit_historical_monitoring_sources


def build_market_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PH精品"
    worksheet.append(["产品信息", None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, "最低价", None, None, "最热销", None, None, "新晋", None, None, "锚定", None, None, "开发新品填写", None, None, None, None, None, None, None, "二次调研--跟首次调研结论有变化的填写"])
    headers = [None] * 49
    for index, title in {
        1: "调研时间",
        2: "货品",
        3: "国家",
        4: "销售员",
        5: "主SKU",
        6: "子sku",
        7: "商品名称",
        8: "一级类目",
        9: "二级类目",
        16: "关键词",
        18: "平台综合推荐(前三页）最低价竞品链接1",
        19: "竞品单价（链接1）",
        20: "竞品子sku月销（链接1）",
        21: "平台综合推荐（前三页）月销量最多竞品链接2",
        22: "竞品单价（链接2）",
        23: "竞品子sku月销（链接2）",
        24: "平台综合推荐（前三页近3个月上架的）新晋竞品链接3",
        25: "竞品单价（链接3）",
        26: "竞品子sku月销（链接3）",
        27: "竟对参考单销(=子sku月销/30)",
        28: "竟对参考售价",
        30: "产品类型",
        31: "稳定期定价 （PHP）",
        38: "认领单销",
        39: "竞对链接",
        40: "调研结论--价格/月销/市场趋势变化等，没有变化就填“无变化”-可以不用填竞对链接",
        41: "产品定位（引流款/利润款/淘汰款）",
        42: "到货通知",
    }.items():
        headers[index - 1] = title
    worksheet.append(headers)
    worksheet.append(["7.8", "开发新品0707期", "菲律宾", "陆俊全", "MAIN1", "SUB1", "商品A", "类目1", "类目2", None, None, None, None, None, None, "kw", None, "url1", 10, 90, "url2", 12, 120, "url3", 13, 10, 4, 99, None, "利润", 88, None, None, None, None, None, None, 5, "anchor", "", "", "已到货"])
    worksheet.append(["7.8", "小货老品", "菲律宾", "陆俊全", "MAIN2", "SUB2", "商品B"])
    worksheet.append(["7.8", "直发转海外仓", "菲律宾", "陆俊全", "MAIN3", "SUB3", "商品C"])
    workbook.create_sheet("TH精品")
    workbook.create_sheet("VN精品")
    workbook.save(path)


def build_listing_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "精品流程"
    worksheet.append(["基础信息", None, None, "销售策略", None, None, None, None, None, None, None, None, "数据跟进一", None, None, None, "销售一次复盘", None, None, "数据跟进二", None, None, None, "销售二次复盘", None, None, "数据跟进三", None, None, None, "销售三次复盘", None, None, "数据跟进四", None, None, None, "总结"])
    worksheet.append([None] * 12 + ["第一周"] + [None] * 6 + ["第二周"] + [None] * 6 + ["第三周"] + [None] * 6 + ["第四周"])
    worksheet.append(["操作日期", "主SKU", "国家", "销售员", "店铺", "Item", "优化", "刷单", "广告", "联盟", "活动", "其他", "订单量", "总收入", "一次毛利额", "一次毛利率", "复盘时间", "产品定位", "优化操作", "订单量", "总收入", "一次毛利额", "一次毛利率", "复盘时间", "产品定位", "优化操作", "订单量", "总收入", "一次毛利额", "一次毛利率", "复盘时间", "产品定位", "优化操作", "订单量", "总收入", "一次毛利额", "一次毛利率", "总结"])
    worksheet.append(["2026-07-23", "MAIN1", "菲律宾", "陆俊全", "Shopee-406TH", "123456", "是", "否", "是", "是", None, "主图+广告", 0, 0, 0, 0, "2026-07-30", "利润款", "加视频"])
    worksheet.append(["2026-07-23", "MAIN2", "菲律宾", "陆俊全", "Shopee-406TH", "#VALUE!", "是", "否", "是", "是", None, "主图+广告"])
    workbook.save(path)


def test_audit_keeps_only_development_new_products_and_preserves_blank_fields() -> None:
    root = Path(".codex_tmp/test_historical_monitoring_sources")
    root.mkdir(parents=True, exist_ok=True)
    market = root / "market.xlsx"
    listing = root / "listing.xlsx"
    build_market_workbook(market)
    build_listing_workbook(listing)

    report = audit_historical_monitoring_sources(market, listing)

    market_report = report["market_monitor"]
    assert market_report["summary"]["raw_row_count"] == 3
    assert market_report["summary"]["development_new_product_row_count"] == 1
    assert market_report["summary"]["excluded_non_development_row_count"] == 2
    assert market_report["summary"]["missing_secondary_conclusion_count"] == 1
    assert market_report["summary"]["missing_positioning_count"] == 1
    assert market_report["summary"]["missing_target_daily_sales_count"] == 1
    assert market_report["summary"]["missing_selling_points_count"] == 1
    assert market_report["records"][0]["positioning"] is None
    assert market_report["records"][0]["target_daily_sales"] is None
    assert market_report["records"][0]["selling_points"] is None


def test_audit_listing_monitor_counts_anomalies_and_missing_finebi_metrics() -> None:
    root = Path(".codex_tmp/test_historical_monitoring_sources")
    root.mkdir(parents=True, exist_ok=True)
    market = root / "market.xlsx"
    listing = root / "listing.xlsx"
    build_market_workbook(market)
    build_listing_workbook(listing)

    report = audit_historical_monitoring_sources(market, listing)

    listing_report = report["listing_monitor"]
    assert listing_report["summary"]["raw_row_count"] == 2
    assert listing_report["summary"]["matched_development_main_sku_count"] == 1
    assert listing_report["summary"]["anomaly_item_count"] == 1
    assert listing_report["summary"]["missing_week_metric_rows"] == 1
    assert listing_report["summary"]["finebi_fillable_rows"] == 0
    assert listing_report["anomalies"][0]["item"] == "#VALUE!"
