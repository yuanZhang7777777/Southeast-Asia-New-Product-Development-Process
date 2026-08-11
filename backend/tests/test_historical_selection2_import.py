import os
import sys
from pathlib import Path

os.environ["DATABASE_URL"] = f"sqlite:///{Path(__file__).with_name('test_historical_selection2.db')}"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from app import models  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.historical_selection2_import import (  # noqa: E402
    SOURCE_TYPE,
    apply_selection2_rows,
    inspect_selection2_workbook,
    normalize_selection2_history_rows,
    parse_historical_selection2_row,
    parse_selection2_workbook,
)

OUTPUT_DIR = Path(__file__).resolve().parents[1] / ".test_outputs" / "historical_selection2_import"


def setup_function() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)



def _row(values_by_column: dict[int, object], width: int = 48) -> tuple[object | None, ...]:
    row = [None] * width
    for column, value in values_by_column.items():
        row[column - 1] = value
    return tuple(row)


def _headers(values_by_column: dict[int, str]) -> dict[str, str]:
    return {get_column_letter(column): title for column, title in values_by_column.items()}


def test_parse_row_creates_claims_only_for_numeric_daily_sales() -> None:
    parsed = parse_historical_selection2_row(
        _row(
            {
                2: "SPU-1",
                3: "SKU-1",
                5: "商品A",
                6: "红色",
                38: "冯卓宏",
                39: "是",
                40: 2,
                41: "李桂敏",
                42: "市场销量不足",
                43: "赵钰婷",
                44: "1.5",
            }
        ),
        headers=_headers(
            {
                2: "SPU",
                3: "SKU",
                5: "产品名称",
                6: "产品规格属性（材质、大小、颜色）",
                38: "开发表格认领情况--主销售员",
                39: "是否认领",
                40: "认领单销",
                41: "主销售员1",
                42: "认领单销/不认领原因",
                43: "主销售员2",
                44: "认领单销/不认领原因",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="7.27期",
        source_row=2,
    )

    assert parsed is not None
    assert parsed["source_type"] == SOURCE_TYPE
    assert parsed["current_status"] == "historical_archive"
    assert parsed["main_sku"] == "SPU-1"
    assert parsed["sub_sku"] == "SKU-1"
    assert parsed["task_policy"] == "none"
    assert [(claim["salesperson_name"], claim["claim_daily_sales"], claim["source_column"]) for claim in parsed["claims"]] == [
        ("冯卓宏", 2.0, "AL:AN"),
        ("赵钰婷", 1.5, "AQ:AR"),
    ]
    assert parsed["rejected_sources"] == [
        {"salesperson_name": "李桂敏", "raw_value": "市场销量不足", "source_column": "AO:AP"}
    ]


def test_parse_row_with_infringement_marker_keeps_archive_only() -> None:
    parsed = parse_historical_selection2_row(
        _row(
            {
                2: "SPU-2",
                3: "SKU-2",
                5: "商品B",
                38: "冯卓宏",
                40: 2,
                45: "侵权商品",
            }
        ),
        headers=_headers(
            {
                2: "SPU",
                3: "SKU",
                5: "产品名称",
                38: "开发表格认领情况--主销售员",
                40: "认领单销",
                45: "销售员3",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="7.27期",
        source_row=3,
    )

    assert parsed is not None
    assert parsed["claims"] == []
    assert parsed["operator_match_policy"] == "skip"
    assert parsed["archive_only_reason"] == "infringing_product"
    assert parsed["snapshot"]["fields_by_cell"]["AS"]["value"] == "侵权商品"


def test_parse_row_with_ip_infringement_marker_does_not_create_numeric_owner() -> None:
    parsed = parse_historical_selection2_row(
        _row({1: "SPU-IP", 2: "SKU-IP", 43: 0, 45: "IP侵权"}),
        headers=_headers(
            {
                1: "SPU",
                2: "SKU",
                43: "主认领销售1",
                44: "认领单销",
                45: "不认领原因",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="6.22期",
        source_row=11,
    )

    assert parsed is not None
    assert parsed["claims"] == []
    assert parsed["rejected_sources"] == []
    assert parsed["operator_match_policy"] == "skip"
    assert parsed["archive_only_reason"] == "infringing_product"


def test_parse_row_repairs_shifted_primary_salesperson_and_rejection_reason() -> None:
    parsed = parse_historical_selection2_row(
        _row({1: "SPU-SHIFT", 2: "SKU-SHIFT", 44: "李干", 46: "需要资质"}),
        headers=_headers(
            {
                1: "SPU",
                2: "SKU",
                43: "主认领销售1",
                44: "认领单销",
                45: "不认领原因",
                46: "认领销售2",
                47: "认领单销",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="6.30期",
        source_row=52,
    )

    assert parsed is not None
    assert parsed["claims"] == []
    assert parsed["rejected_sources"] == [
        {"salesperson_name": "李干", "raw_value": "需要资质", "source_column": "AR:AT"}
    ]


def test_parse_row_repairs_shifted_primary_salesperson_and_positive_daily_sales() -> None:
    parsed = parse_historical_selection2_row(
        _row({1: "SPU-SHIFT", 2: "SKU-SHIFT", 44: "冯卓宏", 45: 0.5}),
        headers=_headers(
            {
                1: "SPU",
                2: "SKU",
                43: "主认领销售1",
                44: "认领单销",
                45: "不认领原因",
                46: "认领销售2",
                47: "认领单销",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="6.30期",
        source_row=58,
    )

    assert parsed is not None
    assert parsed["claims"] == [
        {
            "salesperson_name": "冯卓宏",
            "claim_result": "claim",
            "claim_daily_sales": 0.5,
            "source_column": "AR:AS",
        }
    ]
    assert parsed["rejected_sources"] == []


def test_parse_empty_claim_area_is_historical_unclaimed_not_pending_task() -> None:
    parsed = parse_historical_selection2_row(
        _row({2: "SPU-3", 3: "SKU-3", 5: "商品C"}),
        headers=_headers({2: "SPU", 3: "SKU", 5: "产品名称"}),
        source_file="selection2.xlsx",
        source_sheet="7.27期",
        source_row=4,
    )

    assert parsed is not None
    assert parsed["claims"] == []
    assert parsed["operator_match_policy"] == "none"
    assert parsed["task_policy"] == "none"
    assert parsed["archive_only_reason"] == "historical_unclaimed"


def test_parse_row_deduplicates_same_claimant_and_daily_sales() -> None:
    parsed = parse_historical_selection2_row(
        _row({1: "SPU-4", 2: "SKU-4", 43: "冯卓宏", 44: 1, 46: "冯卓宏", 47: 1}),
        headers=_headers(
            {
                1: "SPU",
                2: "SKU",
                43: "主认领销售1",
                44: "认领单销",
                46: "认领销售2",
                47: "认领单销",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="7.11期",
        source_row=5,
    )

    assert parsed is not None
    assert parsed["claims"] == [
        {
            "salesperson_name": "冯卓宏",
            "claim_result": "claim",
            "claim_daily_sales": 1.0,
            "source_column": "AQ:AR",
            "source_columns": ["AQ:AR", "AT:AU"],
        }
    ]


def test_parse_row_treats_zero_as_rejection_and_flags_conflicting_claim_values() -> None:
    parsed = parse_historical_selection2_row(
        _row(
            {1: "SPU-5", 2: "SKU-5", 43: "冯卓宏", 44: 1, 46: "冯卓宏", 47: 2, 48: "李干", 49: 0},
            width=49,
        ),
        headers=_headers(
            {
                1: "SPU",
                2: "SKU",
                43: "主认领销售1",
                44: "认领单销",
                46: "认领销售2",
                47: "认领单销",
                48: "认领销售3",
                49: "认领单销",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="7.11期",
        source_row=6,
    )

    assert parsed is not None
    assert parsed["claims"] == []
    assert parsed["rejected_sources"] == [
        {"salesperson_name": "李干", "raw_value": 0, "source_column": "AV:AW"}
    ]
    assert parsed["archive_only_reason"] is None
    assert parsed["quality_issues"] == [
        {
            "code": "conflicting_claim_values",
            "salesperson_name": "冯卓宏",
            "raw_values": [1, 2],
            "source_columns": ["AQ:AR", "AT:AU"],
        }
    ]


def test_parse_row_treats_named_blank_slot_as_rejection() -> None:
    parsed = parse_historical_selection2_row(
        _row({1: "SPU-BLANK", 2: "SKU-BLANK", 43: "销售A", 44: None}),
        headers=_headers({1: "SPU", 2: "SKU", 43: "主认领销售1", 44: "认领单销"}),
        source_file="selection2.xlsx",
        source_sheet="7.11期",
        source_row=7,
    )

    assert parsed is not None
    assert parsed["claims"] == []
    assert parsed["rejected_sources"] == [
        {"salesperson_name": "销售A", "raw_value": None, "source_column": "AQ:AR"}
    ]


def test_parse_workbook_reports_latest_period_and_rows() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / "selection2-history.xlsx"
    workbook = Workbook()
    old = workbook.active
    old.title = "5.26期"
    _append_headers(old)
    old.append(_row({2: "OLD", 3: "OLD-A", 5: "旧商品"}))
    latest = workbook.create_sheet("7.27期")
    _append_headers(latest)
    latest.append(_row({2: "NEW", 3: "NEW-A", 5: "新商品", 41: "销售A", 42: 1}))
    note = workbook.create_sheet("说明")
    note.append(["说明"])
    workbook.save(path)

    report = parse_selection2_workbook(path)

    assert report["latest_sheet"] == "7.27期"
    assert report["latest_period"] == "0727"
    assert report["row_count"] == 2
    assert report["sheets"] == [
        {"sheet": "5.26期", "period": "0526", "rows": 1, "skipped": 0},
        {"sheet": "7.27期", "period": "0727", "rows": 1, "skipped": 0},
    ]
    assert report["skipped_sheets"] == [{"sheet": "说明", "reason": "非期数sheet"}]


def test_parse_workbook_uses_headers_for_711_claims_and_source_evidence() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / "selection2-header-driven-711.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "7.11期"
    headers = [None] * 48
    for column, title in {
        1: "SPU",
        2: "SKU",
        4: "产品名称",
        5: "产品规格属性（材质、大小、颜色）",
        6: "图片",
        43: "主认领销售1",
        44: "认领单销",
        46: "认领销售2",
        47: "认领单销",
    }.items():
        headers[column - 1] = title
    sheet.append(headers)
    sheet.append(
        _row(
            {
                1: "SPU-711",
                2: "SKU-711",
                4: "商品711",
                5: "黑色",
                6: "https://image.example/711.jpg",
                43: "冯卓宏",
                44: 1,
                46: "李干",
                47: 2,
            }
        )
    )
    workbook.save(path)

    report = parse_selection2_workbook(path)

    assert report["row_count"] == 1
    parsed = report["rows"][0]
    assert parsed["batch"] == "选品2-财根0711期"
    assert (parsed["main_sku"], parsed["sub_sku"]) == ("SPU-711", "SKU-711")
    assert [(claim["salesperson_name"], claim["claim_daily_sales"]) for claim in parsed["claims"]] == [
        ("冯卓宏", 1.0),
        ("李干", 2.0),
    ]
    assert parsed["snapshot"]["headers_by_cell"]["AQ"] == "主认领销售1"
    assert parsed["snapshot"]["fields_by_cell"]["AT"] == {"header": "认领销售2", "value": "李干"}


def test_inspect_workbook_reports_unreadable_zip_without_writing() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    broken = OUTPUT_DIR / "broken.xlsx"
    broken.write_bytes(b"PK\x03\x04not-a-complete-xlsx")

    report = inspect_selection2_workbook(broken)

    assert report["readable"] is False
    assert report["sheets"] == []
    assert "could not open workbook" in report["error"]


def _append_headers(sheet) -> None:
    headers = [None] * 48
    for column_index, title in {
        2: "SPU",
        3: "SKU",
        5: "产品名称",
        6: "产品规格属性（材质、大小、颜色）",
        38: "开发表格认领情况--主销售员",
        39: "是否认领",
        40: "认领单销",
        41: "销售员1",
        42: "认领单销/不认领原因",
        43: "销售员2",
        44: "认领单销/不认领原因",
        45: "销售员3",
        46: "认领单销/不认领原因",
    }.items():
        headers[column_index - 1] = title
    sheet.append(headers)


def test_parse_row_backfills_missing_main_sku_from_sub_sku() -> None:
    parsed = parse_historical_selection2_row(
        _row({2: "SKU-NO-MAIN", 5: "商品D"}),
        headers=_headers({1: "SPU", 2: "SKU", 5: "产品名称"}),
        source_file="selection2.xlsx",
        source_sheet="7.11期",
        source_row=7,
    )

    assert parsed is not None
    assert parsed["main_sku"] == "SKU-NO-MAIN"
    assert parsed["snapshot"]["backfilled_fields"] == {
        "main_sku": {"value": "SKU-NO-MAIN", "source": "sub_sku"}
    }


def test_parse_row_normalizes_pricing_aliases_and_source_salesperson() -> None:
    parsed = parse_historical_selection2_row(
        _row(
            {
                1: "SPU-PRICE",
                2: "SKU-PRICE",
                3: "商品价格",
                4: "黑色",
                5: "开发员A",
                6: 19.5,
                7: 22,
            }
        ),
        headers=_headers(
            {
                1: "SPU",
                2: "SKU",
                3: "产品名称",
                4: "产品规格属性（材质、大小、颜色）",
                5: "开发姓名",
                6: "SP上家₱",
                7: "SP上架₱",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="6.30期",
        source_row=8,
    )

    assert parsed is not None
    assert parsed["country"] == "PH"
    assert parsed["normalized_fields"]["Shopee稳定期定价"] == {
        "value": 19.5,
        "source_column": "F",
        "source_header": "SP上家₱",
    }
    assert parsed["normalized_fields"]["Shopee推广期定价"] == {
        "value": 22,
        "source_column": "G",
        "source_header": "SP上架₱",
    }
    assert parsed["normalized_fields"]["来源销售员"]["value"] == "开发员A"


def test_normalize_same_identity_nonempty_conflict_for_business_repair() -> None:
    def parsed_row(source_row: int, price: float):
        return parse_historical_selection2_row(
            _row({1: "SPU-CONFLICT", 2: "SKU-CONFLICT", 3: price}),
            headers=_headers({1: "SPU", 2: "SKU", 3: "SP上家₱"}),
            source_file="selection2.xlsx",
            source_sheet="6.30期",
            source_row=source_row,
        )

    first = parsed_row(9, 20)
    second = parsed_row(10, 21)

    report = normalize_selection2_history_rows([first, second])

    assert report["unified_rows"] == []
    assert report["business_repair_rows"] == [
        {
            "source_file": "selection2.xlsx",
            "source_sheet": "6.30期",
            "business_period": "选品2-财根0630期",
            "country": "PH",
            "main_sku": "SPU-CONFLICT",
            "sub_sku": "SKU-CONFLICT",
            "source_rows": [9, 10],
            "conflicting_fields": {"Shopee稳定期定价": [20, 21]},
        }
    ]


def test_normalize_keeps_variants_under_one_main_sku_as_independent_rows() -> None:
    def parsed_row(sub_sku: str, source_row: int):
        return parse_historical_selection2_row(
            _row({1: "SPU-VARIANT", 2: sub_sku, 3: 20}),
            headers=_headers({1: "SPU", 2: "SKU", 3: "SP上家₱"}),
            source_file="selection2.xlsx",
            source_sheet="6.30期",
            source_row=source_row,
        )

    report = normalize_selection2_history_rows([parsed_row("SKU-S", 11), parsed_row("SKU-L", 12)])

    assert report["business_repair_rows"] == []
    assert [row["sub_sku"] for row in report["unified_rows"]] == ["SKU-S", "SKU-L"]


def test_apply_rows_is_idempotent_and_creates_no_workflow_tasks() -> None:
    parsed = parse_historical_selection2_row(
        _row({1: "SPU-APPLY", 2: "SKU-APPLY", 4: "商品", 43: "销售A", 44: 1.5, 46: "销售B", 47: 0}),
        headers=_headers(
            {
                1: "SPU",
                2: "SKU",
                4: "产品名称",
                43: "主认领销售1",
                44: "认领单销",
                46: "认领销售2",
                47: "认领单销",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="7.11期",
        source_row=2,
    )
    assert parsed is not None

    with SessionLocal() as db:
        first = apply_selection2_rows(db, [parsed], source_label="selection2.xlsx")
        db.commit()
        second = apply_selection2_rows(db, [parsed], source_label="selection2.xlsx")
        db.commit()

        opportunity = db.query(models.NewProductOpportunity).one()
        claims = db.query(models.SalesClaimForecast).order_by(models.SalesClaimForecast.salesperson_name).all()
        assert opportunity.current_status == "historical_archive"
        assert opportunity.claim_pool_open is False
        assert [(claim.salesperson_name, claim.claim_result, claim.claim_daily_sales, claim.reject_reason) for claim in claims] == [
            ("销售A", "claim", 1.5, None),
            ("销售B", "reject", None, "来源认领单销为 0"),
        ]
        assert db.query(models.FlowInstance).count() == 0
        assert db.query(models.FlowTask).count() == 0
        assert db.query(models.ReviewRecord).count() == 0
        assert db.query(models.StockingRequest).count() == 0
        assert db.query(models.ListingRecord).count() == 0

    assert first["created_opportunities"] == 1
    assert first["created_claims"] == 2
    assert second["created_opportunities"] == 0
    assert second["updated_opportunities"] == 1
    assert second["created_claims"] == 0
    assert second["updated_claims"] == 2


def test_normalize_reparses_stale_payload_claims_from_raw_snapshot() -> None:
    parsed = parse_historical_selection2_row(
        _row({1: "SPU-CLEAN", 2: "SKU-CLEAN", 44: "李干", 46: "需要资质"}),
        headers=_headers(
            {
                1: "SPU",
                2: "SKU",
                43: "主认领销售1",
                44: "认领单销",
                45: "不认领原因",
                46: "认领销售2",
                47: "认领单销",
            }
        ),
        source_file="selection2.xlsx",
        source_sheet="6.30期",
        source_row=52,
    )
    assert parsed is not None
    parsed["rejected_sources"] = [
        {"salesperson_name": "需要资质", "raw_value": None, "source_column": "AT:AU"}
    ]
    parsed["snapshot"]["rejected_sources"] = parsed["rejected_sources"]

    normalized = normalize_selection2_history_rows([parsed])

    assert normalized["unified_rows"][0]["rejected_sources"] == [
        {"salesperson_name": "李干", "raw_value": "需要资质", "source_column": "AR:AT"}
    ]
