from pathlib import Path
import shutil
import sys

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import models  # noqa: E402
from app.db import Base  # noqa: E402
from app.historical_archive_import import (  # noqa: E402
    HISTORICAL_ARCHIVE_SOURCE_TYPE,
    HISTORICAL_ARCHIVE_STATUS,
    apply_archive_to_db,
    build_archive_rows,
    write_archive_review_outputs,
)


def sample_report() -> dict:
    return {
        "summary": {"record_count": 5},
        "records": [
            {
                "main_sku": "MAIN-WAIT",
                "child_sku": "SUB-WAIT",
                "country": "PH",
                "historical_salesperson": "销售A",
                "plm_salesperson": None,
                "owner_for_action": "销售A",
                "arrival_time": None,
                "has_secondary_research": False,
                "secondary_research_complete": False,
                "has_listing_item": False,
                "listing_source": None,
                "shop": None,
                "item": None,
                "finebi_fillable": False,
                "listing_evidence_status": "none",
                "multi_item_candidates": [],
                "classification_status": "到货监控中",
                "suggested_action": "继续等PLM到货，不要求运营补二次调研",
                "conflict_reason": "",
                "secondary_conclusion": None,
                "positioning": None,
                "target_daily_sales": None,
                "selling_points": None,
                "source_reference": {"source_file": "market.xlsx", "source_sheet": "PH精品", "source_row": 3},
            },
            {
                "main_sku": "MAIN-SECONDARY",
                "child_sku": "SUB-SECONDARY",
                "country": "PH",
                "historical_salesperson": "销售B",
                "plm_salesperson": "销售B",
                "owner_for_action": "销售B",
                "arrival_time": "2026-07-10 09:00:00",
                "has_secondary_research": False,
                "secondary_research_complete": False,
                "has_listing_item": False,
                "listing_source": None,
                "shop": None,
                "item": None,
                "finebi_fillable": False,
                "listing_evidence_status": "none",
                "multi_item_candidates": [],
                "classification_status": "已到货待补二次调研",
                "suggested_action": "要求运营补二次调研",
                "conflict_reason": "",
                "secondary_conclusion": None,
                "positioning": None,
                "target_daily_sales": None,
                "selling_points": None,
                "source_reference": {"source_file": "market.xlsx", "source_sheet": "PH精品", "source_row": 4},
            },
            {
                "main_sku": "MAIN-CONFLICT",
                "child_sku": "SUB-CONFLICT",
                "country": "TH",
                "historical_salesperson": "销售C",
                "plm_salesperson": "销售D",
                "owner_for_action": "销售D",
                "arrival_time": "2026-07-11 09:00:00",
                "has_secondary_research": True,
                "secondary_research_complete": False,
                "has_listing_item": True,
                "listing_source": "finebi",
                "shop": "TH店铺",
                "item": "123456",
                "finebi_fillable": True,
                "listing_evidence_status": "hard_conflict",
                "multi_item_candidates": [{"shop": "TH店铺", "item_id": "123456"}],
                "classification_status": "已刊登但二次调研缺失",
                "suggested_action": "不回退刊登，补二次调研并继续补观察数据",
                "conflict_reason": "PLM销售员与历史销售员不一致",
                "secondary_conclusion": "",
                "positioning": "",
                "target_daily_sales": None,
                "selling_points": None,
                "source_reference": {"source_file": "market.xlsx", "source_sheet": "TH精品", "source_row": 5},
            },
            {
                "main_sku": "MAIN-SHARED",
                "child_sku": "SUB-SHARED",
                "country": "PH",
                "historical_salesperson": "销售E",
                "plm_salesperson": "销售E",
                "owner_for_action": "销售E",
                "arrival_time": "2026-07-12 09:00:00",
                "has_secondary_research": True,
                "secondary_research_complete": False,
                "has_listing_item": True,
                "listing_source": "finebi",
                "shop": "Shopee-1PH",
                "item": "ITEM-1",
                "finebi_fillable": True,
                "listing_evidence_status": "shared_item_binding",
                "multi_item_candidates": [{"shop": "Shopee-1PH", "item_id": "ITEM-1"}],
                "shared_item_owners": [
                    {"country": "PH", "main_sku": "MAIN-SHARED", "shop": "Shopee-1PH", "item_id": "ITEM-1"},
                    {"country": "PH", "main_sku": "MAIN-OTHER", "shop": "Shopee-1PH", "item_id": "ITEM-1"},
                ],
                "metric_dimension": "item_summary",
                "metric_owner_main_sku": "MAIN-SHARED",
                "metrics_are_item_summary": True,
                "classification_status": "已刊登但二次调研缺失",
                "suggested_action": "不回退刊登，补二次调研并继续补观察数据",
                "conflict_reason": "",
                "secondary_conclusion": "",
                "positioning": "",
                "target_daily_sales": None,
                "selling_points": None,
                "source_reference": {"source_file": "market.xlsx", "source_sheet": "PH精品", "source_row": 6},
            },
            {
                "main_sku": "MAIN-MISSING-SUB",
                "child_sku": None,
                "country": "PH",
                "historical_salesperson": "销售F",
                "plm_salesperson": None,
                "owner_for_action": "销售F",
                "arrival_time": None,
                "has_secondary_research": False,
                "secondary_research_complete": False,
                "has_listing_item": False,
                "listing_source": None,
                "shop": None,
                "item": None,
                "finebi_fillable": False,
                "listing_evidence_status": "none",
                "multi_item_candidates": [],
                "classification_status": "到货监控中",
                "suggested_action": "继续等PLM到货，不要求运营补二次调研",
                "conflict_reason": "",
                "secondary_conclusion": None,
                "positioning": None,
                "target_daily_sales": None,
                "selling_points": None,
                "source_reference": {"source_file": "market.xlsx", "source_sheet": "PH精品", "source_row": 7},
            },
        ],
    }


def test_build_archive_rows_keeps_blanks_and_samples_by_status() -> None:
    rows = build_archive_rows(sample_report(), sample_per_status=1)

    assert len(rows) == 3
    wait = rows[0]
    assert wait["source_type"] == HISTORICAL_ARCHIVE_SOURCE_TYPE
    assert wait["current_status"] == HISTORICAL_ARCHIVE_STATUS
    assert wait["snapshot"]["secondary"]["target_daily_sales"] is None
    assert wait["snapshot"]["secondary"]["selling_points"] is None
    assert wait["snapshot"]["classification_status"] == "到货监控中"
    assert wait["snapshot"]["listing"]["evidence_status"] == "none"
    shared = next(row for row in build_archive_rows(sample_report()) if row["main_sku"] == "MAIN-SHARED")
    assert shared["snapshot"]["listing"]["evidence_status_label"] == "共享Item汇总"
    assert shared["snapshot"]["listing"]["metric_dimension"] == "item_summary"
    assert shared["snapshot"]["listing"]["metric_owner_main_sku"] == "MAIN-SHARED"
    assert shared["snapshot"]["listing"]["metrics_are_item_summary"] is True


def test_writes_conflict_and_pending_secondary_workbooks() -> None:
    rows = build_archive_rows(sample_report())
    output_dir = Path(__file__).resolve().parents[1] / ".test_outputs" / "historical_archive_import"
    shutil.rmtree(output_dir, ignore_errors=True)
    try:
        output = write_archive_review_outputs(rows, output_dir, "20260725")

        assert output.summary["total_rows"] == 5
        assert output.summary["conflict_rows"] == 1
        assert output.summary["pending_secondary_rows"] == 3
        assert output.summary["hard_conflict_rows"] == 1
        assert output.summary["shared_item_binding_rows"] == 1
        assert output.summary["skipped_rows"] == 1
        assert output.summary["plm_arrival_rows"] == 3
        assert output.summary["finebi_fillable_rows"] == 2
        assert output.hard_conflict_workbook.exists()
        assert output.multi_item_workbook.exists()
        assert output.missing_finebi_workbook.exists()
        assert output.plm_arrival_workbook.exists()
        assert output.finebi_metrics_workbook.exists()
        assert output.pending_secondary_workbook.exists()
        assert output.shared_item_workbook.exists()
        assert output.skipped_workbook.exists()

        conflict_wb = load_workbook(output.hard_conflict_workbook, read_only=True)
        try:
            headers = [cell.value for cell in conflict_wb.active[1]]
            assert conflict_wb.active.max_row == 2
            assert "是否有刊登证据" in headers
            assert "刊登证据状态" in headers
            assert "指标口径" in headers
            assert "代表主SKU" in headers
            assert "Item汇总" in headers
            assert "候选Item数" in headers
            assert "候选店铺Item" in headers
        finally:
            conflict_wb.close()

        shared_wb = load_workbook(output.shared_item_workbook, read_only=True)
        try:
            shared_rows = list(shared_wb.active.iter_rows(values_only=True))
            shared_headers = list(shared_rows[0])
            evidence_index = shared_headers.index("刊登证据状态")
            dimension_index = shared_headers.index("指标口径")
            owner_index = shared_headers.index("代表主SKU")
            assert shared_rows[1][evidence_index] == "共享Item汇总"
            assert shared_rows[1][dimension_index] == "item_summary"
            assert shared_rows[1][owner_index] == "MAIN-SHARED"
        finally:
            shared_wb.close()
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


def test_apply_archive_only_creates_opportunities_and_snapshots_idempotently() -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    rows = build_archive_rows(sample_report())

    with Session(engine) as db:
        first = apply_archive_to_db(db, rows, source_file="classification.json", imported_by="tester")
        db.commit()
        second = apply_archive_to_db(db, rows, source_file="classification.json", imported_by="tester")
        db.commit()
        existing = db.scalar(
            select(models.NewProductOpportunity).where(models.NewProductOpportunity.source_row == 3)
        )
        assert existing is not None
        existing.snapshot = {
            **existing.snapshot,
            "development_source": {"source_file": "东南亚海外仓新品表-PH.xlsx", "source_row": 10},
        }
        db.commit()
        rows[0]["snapshot"] = {**rows[0]["snapshot"], "conflict_reason": "人工备注"}
        third = apply_archive_to_db(db, rows, source_file="classification.json", imported_by="tester")
        db.commit()

        opportunities = db.scalars(select(models.NewProductOpportunity)).all()
        snapshots = db.scalars(select(models.SourceRecordSnapshot)).all()
        flow_tasks = db.scalars(select(models.FlowTask)).all()
        claims = db.scalars(select(models.SalesClaimForecast)).all()
        updated_snapshot = db.scalar(
            select(models.SourceRecordSnapshot).where(models.SourceRecordSnapshot.source_row == 3)
        )

    assert first.created_count == 4
    assert first.skipped_count == 1
    assert second.created_count == 0
    assert third.updated_count == 1
    assert len(opportunities) == 4
    assert len(snapshots) == 4
    assert flow_tasks == []
    assert claims == []
    assert {item.source_type for item in opportunities} == {HISTORICAL_ARCHIVE_SOURCE_TYPE}
    assert {item.current_status for item in opportunities} == {HISTORICAL_ARCHIVE_STATUS}
    assert opportunities[0].snapshot["secondary"]["target_daily_sales"] is None
    assert opportunities[0].snapshot["development_source"]["source_file"] == "东南亚海外仓新品表-PH.xlsx"
    assert updated_snapshot is not None
    assert updated_snapshot.payload["conflict_reason"] == "人工备注"
    assert updated_snapshot.payload["development_source"]["source_row"] == 10
