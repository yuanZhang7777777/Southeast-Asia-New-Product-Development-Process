"""FineBI 周数据自动拉取：登录 -> 创建导出 -> 下载 -> 校验表头 -> 入库。

协议来自生产自动化实测：
1. GET /webroot/decision/login 取登录页，正则解析 Dec.system = JSON.parse('...') 配置；
   要求 encryptionType=0 且带 encryptionKey（RSA 公钥），密码按每 50 字符分块做
   RSA PKCS1v15 加密 + base64，块间用 "---" 连接；再向同一 URL POST JSON
   {username, password: 加密串, validity: -2, sliderToken: "", origin: "", encrypted: true}，
   headers 必带 X-Requested-With / Origin / Referer / transEncryptLevel: "1" / Chrome 系 UA。
   成功判定 = HTTP 200 且会话出现 fine_auth_token Cookie（精确名）；
   失败时取响应 JSON 的 errorMsg/message/errorCode，并附登录页关键片段便于运维排查。
2. 先 GET /webroot/decision/view?entryType=5&reportId=..（报表预热，非 200 不视为失败），再
   POST /webroot/decision/v5/design/report/data/export?reportId=..&entryType=6&operationId=..，
   body = 配置文件（FINEBI_PAYLOAD_FILE）里的 JSON，其中所有 MMDD-MMDD 周期串（时间周期-SKU /
   时间周期-Order 过滤条件，生产实测 95 处）必须整体替换为目标周，否则报表按过时周期出空表。
   headers 必带 Chrome 系 UA / X-Requested-With / Origin / Referer(报表页)。
   创建成功也可能 Content-Length: 0，真正成败以下载结果为准。
3. GET /webroot/decision/v5/design/report/data/export/download/<operationId>?link=&sessionID=..&form=true，
   成功判定 = 返回 Excel/ZIP 文件头（PK..）而非 HTML，且 openpyxl 能打开、表头与既有 finebi_live 一致。
operationId / sessionID 每次运行用 uuid 重新生成，绝不写死。
"""

from __future__ import annotations

import argparse
import base64
import json
import os
import re
import uuid
from dataclasses import dataclass
from datetime import date
from http.cookiejar import CookieJar
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, Request, build_opener

from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric import padding
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.config import Settings, get_settings
from app.historical_finebi_import import (
    KEY_COLUMNS,
    METRIC_COLUMNS,
    apply_plan,
    build_plan,
    period_window,
    read_finebi_periods,
)
from app.historical_monitoring_sources import text_value

LOGIN_PATH = "/webroot/decision/login"
VIEW_PATH = "/webroot/decision/view"
EXPORT_PATH = "/webroot/decision/v5/design/report/data/export"
DOWNLOAD_PATH = "/webroot/decision/v5/design/report/data/export/download"
AUTH_COOKIE_NAME = "fine_auth_token"
DEC_SYSTEM_PATTERN = re.compile(r"Dec\.system\s*=\s*JSON\.parse\('((?:\\.|[^'])*)'\)")
PASSWORD_CHUNK_SIZE = 50
CHROME_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"
)
DEFAULT_TARGET_DIR = Path("outputs/historical_data/finebi_live")
WEEK_LABEL_PATTERN = re.compile(r"\d{4}-\d{4}")
REQUIRED_HEADER_COLUMNS = (*KEY_COLUMNS, *METRIC_COLUMNS)
REQUIRED_SETTINGS = (
    ("FINEBI_BASE_URL", "finebi_base_url"),
    ("FINEBI_USERNAME", "finebi_username"),
    ("FINEBI_PASSWORD", "finebi_password"),
    ("FINEBI_REPORT_ID", "finebi_report_id"),
    ("FINEBI_PAYLOAD_FILE", "finebi_payload_file"),
)
XLSX_MAGIC = b"PK\x03\x04"


class FineBIPullError(RuntimeError):
    """FineBI 协议层失败（登录 / 导出 / 下载 / 文件校验）。"""


@dataclass
class HttpResponse:
    status: int
    headers: dict[str, str]
    body: bytes

    def snippet(self, limit: int = 300) -> str:
        text = re.sub(r"\s+", " ", self.body.decode("utf-8", errors="replace")).strip()
        return text[:limit] or "空"


class UrllibSession:
    """带 Cookie 会话的最小 HTTP 客户端；测试中整体替换为假会话，绝不发起真实请求。"""

    def __init__(self) -> None:
        self.jar = CookieJar()
        self._opener = build_opener(HTTPCookieProcessor(self.jar))

    def request(
        self,
        method: str,
        url: str,
        *,
        data: bytes | None = None,
        headers: dict[str, str] | None = None,
        timeout: int = 120,
    ) -> HttpResponse:
        request = Request(url, data=data, headers=headers or {}, method=method)
        try:
            with self._opener.open(request, timeout=timeout) as response:
                return HttpResponse(
                    int(response.status),
                    {key.lower(): value for key, value in response.headers.items()},
                    response.read(),
                )
        except HTTPError as exc:
            return HttpResponse(
                int(exc.code),
                {key.lower(): value for key, value in (exc.headers or {}).items()},
                exc.read() or b"",
            )
        except (URLError, TimeoutError) as exc:
            raise FineBIPullError(f"FineBI 请求失败：{url} 不可达（{type(exc).__name__}）") from exc

    def cookie_names(self) -> set[str]:
        return {cookie.name for cookie in self.jar}


def validate_week_label(week_label: str, year: int) -> str:
    label = (week_label or "").strip()
    if not WEEK_LABEL_PATTERN.fullmatch(label):
        raise ValueError("week_label 格式应为 MMDD-MMDD，例如 0723-0729")
    try:
        period_window(label, year)
    except ValueError as exc:
        raise ValueError(f"week_label 不是合法日期区间：{label}") from exc
    return label


def pull_week(
    week_label: str,
    *,
    settings: Settings | None = None,
    target_dir: Path | str | None = None,
    session: UrllibSession | None = None,
    year: int | None = None,
) -> Path:
    settings = settings or get_settings()
    label = validate_week_label(week_label, year or date.today().year)
    _require_configuration(settings)
    payload_bytes = _load_payload(settings.finebi_payload_file)
    payload_bytes = _replace_payload_periods(payload_bytes, label)
    session = session or UrllibSession()
    base = settings.finebi_base_url.strip().rstrip("/")
    report_id = settings.finebi_report_id.strip()
    operation_id = uuid.uuid4().hex
    session_id = uuid.uuid4().hex
    _login(session, base, settings)
    _warmup_report(session, base, report_id)
    _create_export(session, base, report_id, operation_id, payload_bytes)
    content = _download_export(session, base, report_id, operation_id, session_id)
    _validate_headers(content, label)

    directory = Path(target_dir) if target_dir else DEFAULT_TARGET_DIR
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / f"{label}.xlsx"
    if target.exists():
        os.replace(target, target.with_name(f"{target.name}.bak"))
    partial = target.with_name(f"{target.name}.part")
    partial.write_bytes(content)
    os.replace(partial, target)
    return target


def pull_and_import(
    db: Session,
    week_label: str,
    imported_by: str | None = None,
    *,
    settings: Settings | None = None,
    target_dir: Path | str | None = None,
    session: UrllibSession | None = None,
    year: int | None = None,
) -> dict[str, object]:
    settings = settings or get_settings()
    year = year or date.today().year
    directory = Path(target_dir) if target_dir else DEFAULT_TARGET_DIR
    path = pull_week(week_label, settings=settings, target_dir=directory, session=session, year=year)
    label = path.stem
    finebi, periods, skipped = read_finebi_periods(directory, year, include_period=label)
    plan = build_plan(db, {}, finebi, year)
    counts = apply_plan(db, plan, source_label=path.name, imported_by=imported_by)
    return {
        "week_label": label,
        "file": str(path),
        "periods": periods,
        "excluded_periods": skipped,
        "summary": plan["summary"],
        "apply_counts": counts,
    }


def _require_configuration(settings: Settings) -> None:
    missing = [name for name, attribute in REQUIRED_SETTINGS if not str(getattr(settings, attribute)).strip()]
    if missing:
        raise ValueError(f"FineBI 拉取配置不完整，缺少：{', '.join(missing)}")


def _load_payload(payload_file: str) -> bytes:
    path = Path(payload_file)
    if not path.is_file():
        raise ValueError(f"FINEBI_PAYLOAD_FILE 不存在：{path}")
    content = path.read_bytes()
    try:
        json.loads(content.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"FINEBI_PAYLOAD_FILE 不是合法 JSON：{path}") from exc
    return content


def _replace_payload_periods(payload_bytes: bytes, target_label: str) -> bytes:
    """把 payload 里所有 MMDD-MMDD 周期串替换为目标周（时间周期-SKU / 时间周期-Order 等过滤条件）。

    不替换则 FineBI 按 payload 内过时周期出空表——2026-07 生产实测空导出的根因。
    """
    text = payload_bytes.decode("utf-8")
    detected = sorted(set(re.findall(r"(?<!\d)\d{4}-\d{4}(?!\d)", text)))
    for value in detected:
        if value != target_label:
            text = text.replace(value, target_label)
    if target_label not in text:
        raise FineBIPullError(
            "FINEBI_PAYLOAD_FILE 中未发现任何 MMDD-MMDD 周期串，无法定位周期过滤条件；"
            "请核对 payload 是否仍是浏览器抓取的导出请求原文"
        )
    return text.encode("utf-8")


def _report_url(base: str, report_id: str) -> str:
    return f"{base}{VIEW_PATH}?{urlencode({'entryType': 5, 'reportId': report_id})}"


def _warmup_report(session: UrllibSession, base: str, report_id: str) -> None:
    # 预热报表会话；生产自动化同款步骤，非 200 只影响后续下载（下载会自行报错），不在此失败。
    session.request(
        "GET",
        _report_url(base, report_id),
        headers={"User-Agent": CHROME_USER_AGENT, "Referer": f"{base}{LOGIN_PATH}"},
    )


def _login(session: UrllibSession, base: str, settings: Settings) -> None:
    login_url = f"{base}{LOGIN_PATH}"
    page = session.request("GET", login_url)
    if page.status != 200:
        raise FineBIPullError(f"FineBI 登录页获取失败：HTTP {page.status}，响应片段：{page.snippet()}")
    page_html = page.body.decode("utf-8", errors="replace")
    dec_system = _parse_dec_system(page_html)
    payload: dict[str, object] = {
        "username": settings.finebi_username,
        "password": _encrypt_password(settings.finebi_password, dec_system),
        "validity": -2,
        "sliderToken": "",
        "origin": "",
        "encrypted": True,
    }
    response = session.request(
        "POST",
        login_url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "Content-Type": "application/json",
            "X-Requested-With": "XMLHttpRequest",
            "Origin": base,
            "Referer": login_url,
            "transEncryptLevel": "1",
            "User-Agent": CHROME_USER_AGENT,
        },
    )
    if response.status != 200 or AUTH_COOKIE_NAME not in session.cookie_names():
        raise FineBIPullError(
            f"FineBI 登录失败：HTTP {response.status}，未获得 {AUTH_COOKIE_NAME} Cookie，"
            f"错误信息：{_login_error_message(response)}；登录页片段：{_page_snippet(page_html)}"
        )


def _parse_dec_system(html: str) -> dict[str, object]:
    match = DEC_SYSTEM_PATTERN.search(html)
    if not match:
        raise FineBIPullError(f"FineBI 登录页未找到 Dec.system 配置；登录页片段：{_page_snippet(html)}")
    try:
        unescaped = json.loads(f'"{match.group(1)}"')
        data = json.loads(unescaped)
    except (json.JSONDecodeError, ValueError) as exc:
        raise FineBIPullError(
            f"FineBI 登录页 Dec.system 配置解析失败：{exc}；登录页片段：{_page_snippet(html)}"
        ) from exc
    if not isinstance(data, dict):
        raise FineBIPullError(f"FineBI 登录页 Dec.system 配置不是对象；登录页片段：{_page_snippet(html)}")
    return data


def _encrypt_password(password: str, dec_system: dict[str, object]) -> str:
    encryption_type = dec_system.get("encryptionType")
    encryption_key = str(dec_system.get("encryptionKey") or "").strip()
    if encryption_type not in (0, "0") or not encryption_key:
        raise FineBIPullError(
            f"Unsupported encryption type: encryptionType={encryption_type!r}，"
            f"encryptionKey {'存在' if encryption_key else '缺失'}"
        )
    public_key = _load_public_key(encryption_key)
    chunks = [password[index:index + PASSWORD_CHUNK_SIZE] for index in range(0, len(password), PASSWORD_CHUNK_SIZE)]
    return "---".join(
        base64.b64encode(_rsa_encrypt(public_key, chunk.encode("utf-8"))).decode("ascii") for chunk in chunks
    )


def _public_key_pem(encryption_key: str) -> str:
    body = re.sub(r"\s+", "", encryption_key)
    lines = [body[index:index + 64] for index in range(0, len(body), 64)]
    return "-----BEGIN PUBLIC KEY-----\n" + "\n".join(lines) + "\n-----END PUBLIC KEY-----\n"


def _load_public_key(encryption_key: str):
    try:
        return serialization.load_pem_public_key(_public_key_pem(encryption_key).encode("ascii"))
    except Exception as exc:
        raise FineBIPullError(f"FineBI encryptionKey 不是合法 RSA 公钥：{exc}") from exc


def _rsa_encrypt(public_key, chunk: bytes) -> bytes:
    return public_key.encrypt(chunk, padding.PKCS1v15())


def _login_error_message(response: HttpResponse) -> str:
    try:
        data = json.loads(response.body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return response.snippet()
    if isinstance(data, dict):
        parts = [str(data[key]) for key in ("errorMsg", "message", "errorCode") if data.get(key)]
        if parts:
            return "；".join(parts)
    return response.snippet()


def _page_snippet(html: str, limit: int = 400) -> str:
    return re.sub(r"\s+", " ", html).strip()[:limit] or "空"


def _create_export(
    session: UrllibSession, base: str, report_id: str, operation_id: str, payload_bytes: bytes
) -> None:
    params = urlencode({"reportId": report_id, "entryType": 6, "operationId": operation_id})
    response = session.request(
        "POST",
        f"{base}{EXPORT_PATH}?{params}",
        data=payload_bytes,
        headers={
            "Content-Type": "application/json;charset=UTF-8",
            "User-Agent": CHROME_USER_AGENT,
            "X-Requested-With": "XMLHttpRequest",
            "Origin": base,
            "Referer": _report_url(base, report_id),
        },
    )
    if response.status != 200:
        raise FineBIPullError(f"FineBI 创建导出失败：HTTP {response.status}，响应片段：{response.snippet()}")
    # 生产实测：创建成功也可能返回 Content-Length: 0，因此空 body 不算失败，成败以下载结果为准。


def _download_export(
    session: UrllibSession, base: str, report_id: str, operation_id: str, session_id: str
) -> bytes:
    params = urlencode({"link": "", "sessionID": session_id, "form": "true"})
    response = session.request(
        "GET",
        f"{base}{DOWNLOAD_PATH}/{operation_id}?{params}",
        headers={"User-Agent": CHROME_USER_AGENT, "Referer": _report_url(base, report_id)},
    )
    content_type = response.headers.get("content-type", "")
    if (
        response.status != 200
        or "text/html" in content_type.lower()
        or not response.body.startswith(XLSX_MAGIC)
    ):
        raise FineBIPullError(
            f"FineBI 下载失败：HTTP {response.status}，Content-Type：{content_type or '未知'}，"
            f"返回的不是 Excel/ZIP 文件；响应片段：{response.snippet()}"
        )
    return response.body


def _validate_headers(content: bytes, week_label: str) -> None:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise FineBIPullError(f"FineBI 下载的 {week_label} 文件无法被 openpyxl 打开：{exc}") from exc
    try:
        seen: list[str] = []
        for worksheet in workbook.worksheets:
            # FineBI 自动导出的 xlsx 维度元数据是坏的（标称 A1:A1），read_only 模式必须重置
            worksheet.reset_dimensions()
            first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = {text_value(value) for value in first_row if text_value(value)}
            if all(column in headers for column in REQUIRED_HEADER_COLUMNS):
                return
            seen.extend(sorted(headers))
        raise FineBIPullError(
            f"FineBI 下载的 {week_label} 表头与既有 finebi_live 文件不一致，不入库："
            f"需要 {list(REQUIRED_HEADER_COLUMNS)}，实际 {seen or '无表头'}"
        )
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="FineBI 周数据自动拉取（默认仅下载，不入库）")
    parser.add_argument("--week-label", required=True, help="周标签，形如 0723-0729")
    parser.add_argument("--import-db", action="store_true", help="下载后解析并写入数据库")
    parser.add_argument("--imported-by", default="finebi_auto_pull")
    args = parser.parse_args()

    settings = get_settings()
    if not settings.finebi_auto_pull_enabled:
        raise SystemExit("FINEBI_AUTO_PULL_ENABLED=false，请先在部署环境开启")
    if args.import_db:
        from app.db import SessionLocal

        with SessionLocal() as db:
            report = pull_and_import(db, args.week_label, imported_by=args.imported_by, settings=settings)
            db.commit()
    else:
        report = {"week_label": args.week_label, "file": str(pull_week(args.week_label, settings=settings))}
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
