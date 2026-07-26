from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app import models, schemas, services
from app.excel_images import images_by_row, save_product_image
from app.field_mapping import normalize_header, number_value, text_value
from app.workbook_sheets import resolve_sheet_name
from app.workflow_status import OPPORTUNITY_ASSIGNED, OPPORTUNITY_PENDING_ASSIGNMENT, TASK_PENDING


SOURCE_TYPE = "selection2_caigen_claim_feedback"
DEFAULT_SHEET = "5.26期"
MAX_SOURCE_COLUMN = column_index_from_string("AV")
SNAPSHOT_COLUMNS = tuple(get_column_letter(index) for index in range(1, MAX_SOURCE_COLUMN + 1))
CLAIM_SOURCE_COLUMNS = ("AL:AN", "AO:AP", "AQ:AR", "AS:AT", "AU:AV")


def import_selection2_workbook(db: Session, payload: schemas.Selection2ImportRequest) -> schemas.Selection2ImportResponse:
    source_path = resolve_source_file(payload.source_file)
    import_batch = models.ImportBatch(
        source_type=SOURCE_TYPE,
        source_file=source_path.name,
        source_sheet=payload.source_sheet,
        status="running",
    )
    db.add(import_batch)
    db.flush()

    workbook = load_workbook(source_path, read_only=False, data_only=True)
    source_sheet = resolve_sheet_name(workbook, payload.source_sheet)
    import_batch.source_sheet = source_sheet

    worksheet = workbook[source_sheet]
    try:
        worksheet.reset_dimensions()
    except AttributeError:
        pass
    source_max_column = max(worksheet.max_column or 0, MAX_SOURCE_COLUMN)
    product_images = images_by_row(worksheet, "G")
    headers_by_column = source_headers_by_column(worksheet, source_max_column)

    created_count = 0
    updated_count = 0
    skipped_count = 0
    prefill_claim_count = 0
    task_count = 0
    processed_count = 0

    for source_row, row in enumerate(worksheet.iter_rows(min_row=2, max_col=source_max_column, values_only=True), start=2):
        if payload.max_rows is not None and processed_count >= payload.max_rows:
            break
        parsed = parse_selection2_row(row, headers_by_column)
        if parsed is None:
            skipped_count += 1
            continue
        attach_product_image(parsed, product_images.get(source_row), source_row)
        processed_count += 1

        opportunity, created = upsert_opportunity(db, parsed, source_path.name, source_sheet, source_row, import_batch.id)
        created_count += int(created)
        updated_count += int(not created)
        prefill_claim_count += replace_source_claims(db, opportunity.id, parsed)
        if ensure_import_claim_task(db, opportunity, parsed):
            task_count += 1

    import_batch.created_count = created_count
    import_batch.updated_count = updated_count
    import_batch.skipped_count = skipped_count
    import_batch.status = "completed"
    services.audit(
        db,
        "selection2.imported",
        "new_product_opportunity",
        None,
        {
            "source_file": source_path.name,
            "source_sheet": source_sheet,
            "imported_count": created_count + updated_count,
            "created_count": created_count,
            "updated_count": updated_count,
            "skipped_count": skipped_count,
        },
    )

    return schemas.Selection2ImportResponse(
        import_batch_id=import_batch.id,
        source_file=source_path.name,
        source_sheet=source_sheet,
        business_period=source_sheet,
        imported_count=created_count + updated_count,
        created_count=created_count,
        updated_count=updated_count,
        skipped_count=skipped_count,
        market_research_count=0,
        prefill_claim_count=prefill_claim_count,
        task_count=task_count,
    )


def attach_product_image(parsed: dict[str, Any], image: Any | None, source_row: int) -> None:
    if parsed["main"].get("image_url"):
        return
    image_url = save_product_image(image, SOURCE_TYPE, source_row)
    if image_url:
        parsed["main"]["image_url"] = image_url
        parsed["snapshot"]["extracted_image_url"] = image_url


def resolve_source_file(source_file: str | None) -> Path:
    if source_file:
        path = Path(source_file)
        if not path.is_absolute():
            path = Path.cwd() / path
        if not path.exists():
            raise FileNotFoundError(str(path))
        return path
    for root in [Path.cwd(), Path.cwd().parent]:
        matches = sorted(root.glob("选品2：海外仓财根团队开发新品认领-反馈*.xlsx"))
        if matches:
            return matches[0]
    raise FileNotFoundError("selection2 source workbook not found")


def parse_selection2_row(row: tuple[Any, ...], headers_by_column: dict[str, list[str]] | None = None) -> dict[str, Any] | None:
    raw_values = {get_column_letter(index): cell_value(row, get_column_letter(index)) for index in range(1, len(row) + 1)}
    values = {column: cell_value(row, column) for column in SNAPSHOT_COLUMNS}
    sub_sku = text_value(values["C"])
    if not sub_sku:
        return None
    main_sku = text_value(values["B"]) or sub_sku
    snapshot = {
        "source_type": SOURCE_TYPE,
        "allowed_columns": list(raw_values),
        "cells": values,
        "headers_by_column": headers_by_column or {},
        "fields_by_column": fields_by_column(raw_values),
        "fields_by_header": fields_by_header(headers_by_column or {}, raw_values),
        "claim_prefill": parse_claims(values),
    }
    return {
        "main": {
            "main_sku": main_sku,
            "sub_sku": sub_sku,
            "main_sku_name": text_value(values["E"]),
            "sub_sku_name": text_value(values["F"]),
            "image_url": text_value(values["G"]),
        },
        "claim_prefill": snapshot["claim_prefill"],
        "snapshot": snapshot,
    }


def upsert_opportunity(
    db: Session, parsed: dict[str, Any], source_file: str, source_sheet: str, source_row: int, import_batch_id: str
) -> tuple[models.NewProductOpportunity, bool]:
    existing = db.scalar(
        select(models.NewProductOpportunity).where(
            models.NewProductOpportunity.source_type == SOURCE_TYPE,
            models.NewProductOpportunity.source_sheet == source_sheet,
            models.NewProductOpportunity.sub_sku == parsed["main"]["sub_sku"],
        )
    )
    data = {
        **parsed["main"],
        "source_type": SOURCE_TYPE,
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
        "import_batch_id": import_batch_id,
        "batch": source_sheet,
        "snapshot": parsed["snapshot"],
    }
    if existing:
        for field, value in data.items():
            if field != "current_status":
                setattr(existing, field, value)
        add_source_snapshot(db, existing, import_batch_id)
        return existing, False

    opportunity = models.NewProductOpportunity(**data)
    db.add(opportunity)
    db.flush()
    add_source_snapshot(db, opportunity, import_batch_id)
    services.audit(
        db,
        "opportunity.created",
        "new_product_opportunity",
        opportunity.id,
        {"main_sku": opportunity.main_sku, "sub_sku": opportunity.sub_sku, "source_type": SOURCE_TYPE},
    )
    return opportunity, True


def add_source_snapshot(db: Session, opportunity: models.NewProductOpportunity, import_batch_id: str) -> None:
    db.add(
        models.SourceRecordSnapshot(
            import_batch_id=import_batch_id,
            opportunity_id=opportunity.id,
            source_file=opportunity.source_file,
            source_sheet=opportunity.source_sheet,
            source_row=opportunity.source_row,
            column_range="A:AV",
            payload=opportunity.snapshot,
        )
    )


def source_headers_by_column(worksheet: Any, max_col: int) -> dict[str, list[str]]:
    headers: dict[str, list[str]] = {}
    for index in range(1, max_col + 1):
        column = get_column_letter(index)
        top = text_value(worksheet.cell(row=1, column=index).value)
        candidates = [top] if top else []
        if candidates:
            headers[column] = candidates
    return headers


def fields_by_header(headers_by_column: dict[str, list[str]], values: dict[str, Any]) -> dict[str, Any]:
    fields: dict[str, Any] = {}
    for column, headers in headers_by_column.items():
        value = values.get(column)
        if value is None:
            continue
        for header in headers:
            key = normalize_header(header)
            if key and key not in fields:
                fields[key] = value
    return fields


def fields_by_column(values: dict[str, Any]) -> dict[str, Any]:
    return {column: value for column, value in values.items() if value not in (None, "")}


def parse_claims(values: dict[str, Any]) -> list[dict[str, Any]]:
    claims: list[dict[str, Any]] = []
    main_salesperson = text_value(values["AL"])
    if main_salesperson:
        claim_result = normalize_claim_result(values["AM"])
        claims.append(
            {
                "salesperson_name": main_salesperson,
                "claim_result": claim_result,
                "claim_daily_sales": number_value(values["AN"]) if claim_result == "claim" else None,
                "reject_reason": None,
                "source_column": "AL:AN",
            }
        )
    for name_column, value_column, source_column in [("AO", "AP", "AO:AP"), ("AQ", "AR", "AQ:AR"), ("AS", "AT", "AS:AT"), ("AU", "AV", "AU:AV")]:
        salesperson = text_value(values[name_column])
        raw_value = values[value_column]
        if not salesperson or raw_value in (None, ""):
            continue
        daily_sales = number_value(raw_value)
        if daily_sales and daily_sales > 0:
            claims.append(
                {
                    "salesperson_name": salesperson,
                    "claim_result": "claim",
                    "claim_daily_sales": daily_sales,
                    "reject_reason": None,
                    "source_column": source_column,
                }
            )
        else:
            claims.append(
                {
                    "salesperson_name": salesperson,
                    "claim_result": "reject",
                    "claim_daily_sales": None,
                    "reject_reason": text_value(raw_value),
                    "source_column": source_column,
                }
            )
    return claims


def replace_source_claims(db: Session, opportunity_id: str, parsed: dict[str, Any]) -> int:
    db.execute(
        delete(models.SalesClaimForecast).where(
            models.SalesClaimForecast.opportunity_id == opportunity_id,
            models.SalesClaimForecast.source_column.in_(CLAIM_SOURCE_COLUMNS),
        )
    )
    count = 0
    for claim in parsed["claim_prefill"]:
        db.add(models.SalesClaimForecast(opportunity_id=opportunity_id, platform="Shopee", **claim))
        count += 1
    return count


def ensure_import_claim_task(db: Session, opportunity: models.NewProductOpportunity, parsed: dict[str, Any]) -> bool:
    task = db.scalar(
        select(models.FlowTask)
        .join(models.FlowInstance)
        .where(models.FlowInstance.opportunity_id == opportunity.id, models.FlowTask.task_type == "sales_claim")
        .order_by(models.FlowTask.created_at.desc())
    )
    assignee_name = parsed["claim_prefill"][0]["salesperson_name"] if parsed["claim_prefill"] else None
    if task:
        if task.status == TASK_PENDING and assignee_name and not task.assignee_name:
            task.assignee_name = assignee_name
        if opportunity.current_status == OPPORTUNITY_PENDING_ASSIGNMENT and assignee_name:
            opportunity.current_status = OPPORTUNITY_ASSIGNED
        return False
    flow = models.FlowInstance(
        opportunity_id=opportunity.id,
        current_node="sales_claim",
        current_status=OPPORTUNITY_ASSIGNED if assignee_name else "open_claim_pool",
        owner_role="sales",
    )
    db.add(flow)
    db.flush()
    db.add(
        models.FlowTask(
            flow_instance_id=flow.id,
            node_code="sales_claim",
            task_type="sales_claim",
            assignee_name=assignee_name,
            assignee_role="sales",
        )
    )
    if assignee_name:
        opportunity.current_status = OPPORTUNITY_ASSIGNED
    return True


def normalize_claim_result(value: Any) -> str | None:
    text = text_value(value)
    if text in {"是", "认领", "claim", "CLAIM", "yes", "YES"}:
        return "claim"
    if text in {"否", "不认领", "reject", "REJECT", "no", "NO"}:
        return "reject"
    return None


def cell_value(row: tuple[Any, ...], column: str) -> Any:
    index = column_index_from_string(column) - 1
    return row[index] if index < len(row) else None
