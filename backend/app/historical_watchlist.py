from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


COUNTRY_ALIASES = {"TH": "TH", "泰国": "TH", "VN": "VN", "越南": "VN", "PH": "PH", "菲律宾": "PH"}
REJECTED_TEXT = ("不认领", "易碎品", "拒绝", "未认领", "不接受", "备注", "说明")


def build_historical_watchlist(source_files: Iterable[Path]) -> list[dict[str, Any]]:
    merged: dict[tuple[str, str], dict[str, Any]] = {}
    for source_path in sorted((Path(path) for path in source_files), key=lambda path: path.name):
        source_country = normalize_country(source_path.stem)
        source_sha256 = sha256_file(source_path)
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                worksheet.reset_dimensions()
                rows = worksheet.iter_rows(values_only=True)
                first_header = tuple(next(rows, ()))
                second_header = tuple(next(rows, ()))
                main_column = find_sku_column(first_header, second_header, "主SKU", "MAINSKU")
                child_column = find_sku_column(first_header, second_header, "子SKU", "SUBSKU")
                all_claim_columns = find_claim_columns(first_header, second_header)
                claim_columns = [
                    column for column in all_claim_columns if source_country is None or column[0] == source_country
                ]
                if main_column is None or child_column is None or not claim_columns:
                    continue
                for source_row, row in enumerate(rows, start=3):
                    main_sku = normalize_sku(cell(row, main_column))
                    child_sku = normalize_sku(cell(row, child_column))
                    if not main_sku or not child_sku:
                        continue
                    for country, claimant_column in claim_columns:
                        claimants = extract_person_names(normalize_text(cell(row, claimant_column)))
                        if not claimants:
                            continue
                        record = merged.setdefault(
                            (country, child_sku),
                            {"country": country, "child_sku": child_sku, "main_skus": set(), "historical_claimants": set(), "source_references": set()},
                        )
                        record["main_skus"].add(main_sku)
                        record["historical_claimants"].update(claimants)
                        record["source_references"].add((source_path.name, source_sha256, worksheet.title, source_row, claimant_column + 1))
        finally:
            workbook.close()
    return [serialize_record(record) for _, record in sorted(merged.items())]


def watchlist_payload(source_files: Iterable[Path]) -> dict[str, Any]:
    records = build_historical_watchlist(source_files)
    return {
        "key_count": len(records),
        "ambiguous_main_sku_key_count": sum(len(record["main_skus"]) > 1 for record in records),
        "records": records,
    }


def serialize_record(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "country": record["country"],
        "child_sku": record["child_sku"],
        "main_skus": sorted(record["main_skus"]),
        "historical_claimants": sorted(record["historical_claimants"]),
        "source_references": [
            {"source_file": source_file, "source_sha256": source_sha256, "source_sheet": source_sheet, "source_row": source_row, "source_column": source_column}
            for source_file, source_sha256, source_sheet, source_row, source_column in sorted(record["source_references"])
        ],
    }


def normalize_country(value: object) -> str | None:
    text = normalize_text(value).upper()
    return next((country for alias, country in COUNTRY_ALIASES.items() if alias in text), None)


def normalize_sku(value: object) -> str | None:
    text = normalize_text(value).upper()
    return text or None


def find_sku_column(first_header: tuple[Any, ...], second_header: tuple[Any, ...], *aliases: str) -> int | None:
    normalized_aliases = {normalize_text(alias).upper() for alias in aliases}
    for index in range(max(len(first_header), len(second_header))):
        if normalize_text(cell(first_header, index)).upper() in normalized_aliases:
            return index
        if normalize_text(cell(second_header, index)).upper() in normalized_aliases:
            return index
    return None


def find_claim_columns(first_header: tuple[Any, ...], second_header: tuple[Any, ...]) -> list[tuple[str, int]]:
    columns: list[tuple[str, int]] = []
    for index in range(max(len(first_header), len(second_header)) - 1):
        country = normalize_country(cell(first_header, index)) or normalize_country(cell(second_header, index))
        sales_header = normalize_text(cell(second_header, index + 1)) or normalize_text(cell(first_header, index + 1))
        if country and "预估单销" in sales_header:
            columns.append((country, index))
    return columns


def extract_person_names(value: str) -> list[str]:
    if not value or any(text in value for text in REJECTED_TEXT):
        return []
    if not re.fullmatch(r"[\u4e00-\u9fff]{2,4}(?:\s+[\u4e00-\u9fff]{2,4})?", value):
        return []
    return value.split()


def normalize_text(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split())


def cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a traceable historical claimed-SKU watchlist.")
    parser.add_argument("source_files", nargs="+", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    encoded = json.dumps(watchlist_payload(args.source_files), ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    if args.output:
        args.output.write_text(encoded, encoding="utf-8")
    else:
        print(encoded, end="")


if __name__ == "__main__":
    main()
