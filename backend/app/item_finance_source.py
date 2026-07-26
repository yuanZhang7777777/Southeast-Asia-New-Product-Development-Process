from __future__ import annotations

import argparse
import json
import math
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from app.historical_watchlist import sha256_file


REQUIRED_COLUMNS = ("ITEMID", "主SKU", "店铺", "审核时间")
SHOP_COUNTRY = re.compile(r"(TH|VN|PH)$", re.I)
DATA_SHEET = "ItemID财务数据八部"
FILTER_SHEET = "ItemID财务数据八部_过滤条件"


def read_item_finance_period(period: str, source_path: Path) -> list[dict[str, Any]]:
    source_path = Path(source_path)
    source_sha256 = sha256_file(source_path)
    records: list[dict[str, Any]] = []
    found_data_sheet = False
    workbook = load_workbook(source_path, read_only=False, data_only=True, keep_links=False)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows()
            header_row, header_cells = next(enumerate(rows, start=1), (0, ()))
            header = tuple(cell.value for cell in header_cells)
            columns = required_columns(header)
            if columns is None:
                continue
            found_data_sheet = True
            merged_values = merged_cell_values(worksheet, set(columns.values()))
            for source_row, row in enumerate(rows, start=header_row + 1):
                item_id = item_id_text(display_cell(row, columns["ITEMID"], source_row, merged_values))
                main_sku = normalize_main_sku(display_cell(row, columns["主SKU"], source_row, merged_values))
                shop = text_value(display_cell(row, columns["店铺"], source_row, merged_values))
                country = country_from_shop(shop)
                if not item_id or not main_sku or not shop or not country:
                    continue
                records.append(
                    {
                        "country": country,
                        "main_sku": main_sku,
                        "shop": shop,
                        "item_id": item_id,
                        "audit_time": time_value(display_cell(row, columns["审核时间"], source_row, merged_values)),
                        "source_reference": {
                            "period": period,
                            "source_file": source_path.name,
                            "source_sha256": source_sha256,
                            "source_sheet": worksheet.title,
                            "source_row": source_row,
                        },
                    }
                )
    finally:
        workbook.close()
    if not found_data_sheet:
        raise ValueError(f"{source_path.name} missing required columns: {', '.join(REQUIRED_COLUMNS)}")
    return records


def reconcile_item_candidates(watchlist: dict[str, Any], period_inputs: Iterable[tuple[str, Path]]) -> dict[str, Any]:
    candidates = merge_candidates(period_inputs)
    candidates_by_main: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for candidate in candidates:
        candidates_by_main[candidate["country"], candidate["main_sku"]].append(candidate)

    records = watchlist.get("records")
    if not isinstance(records, list):
        raise ValueError("watchlist records must be a list")
    matches = []
    for record in sorted(records, key=lambda item: (str(item.get("country", "")), str(item.get("child_sku", "")))):
        country = text_value(record.get("country"))
        main_skus = sorted({normalize_main_sku(value) for value in record.get("main_skus", []) if normalize_main_sku(value)})
        found = {
            (candidate["country"], candidate["main_sku"], candidate["shop"], candidate["item_id"]): candidate
            for main_sku in main_skus
            for candidate in candidates_by_main.get((country or "", main_sku), [])
        }
        selected = [found[key] for key in sorted(found)]
        classification = "unique" if len(selected) == 1 else "multiple" if selected else "unmatched"
        matches.append({**record, "country": country, "main_skus": main_skus, "classification": classification, "candidates": selected})

    return {
        "summary": {
            "watchlist_record_count": len(matches),
            "candidate_count": len(candidates),
            "matched_key_count": sum(match["classification"] != "unmatched" for match in matches),
            "unique_count": sum(match["classification"] == "unique" for match in matches),
            "multiple_count": sum(match["classification"] == "multiple" for match in matches),
            "unmatched_count": sum(match["classification"] == "unmatched" for match in matches),
        },
        "candidates": candidates,
        "matches": matches,
        "unique_candidates": [match for match in matches if match["classification"] == "unique"],
        "multiple_candidates": [match for match in matches if match["classification"] == "multiple"],
        "unmatched": [match for match in matches if match["classification"] == "unmatched"],
        "anomalies": conflicts(candidates),
    }


def merge_candidates(period_inputs: Iterable[tuple[str, Path]]) -> list[dict[str, Any]]:
    merged: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for period, source_path in period_inputs:
        for record in read_item_finance_period(period, source_path):
            key = record["country"], record["main_sku"], record["shop"], record["item_id"]
            candidate = merged.setdefault(
                key,
                {
                    "country": record["country"],
                    "main_sku": record["main_sku"],
                    "shop": record["shop"],
                    "item_id": record["item_id"],
                    "periods": set(),
                    "audit_times": set(),
                    "source_references": set(),
                    "row_count": 0,
                },
            )
            candidate["periods"].add(record["source_reference"]["period"])
            if record["audit_time"]:
                candidate["audit_times"].add(record["audit_time"])
            candidate["source_references"].add(tuple(record["source_reference"].items()))
            candidate["row_count"] += 1
    return [serialize_candidate(merged[key]) for key in sorted(merged)]


def serialize_candidate(candidate: dict[str, Any]) -> dict[str, Any]:
    periods = sorted(candidate["periods"])
    return {
        "country": candidate["country"],
        "main_sku": candidate["main_sku"],
        "shop": candidate["shop"],
        "item_id": candidate["item_id"],
        "first_period": periods[0],
        "last_period": periods[-1],
        "row_count": candidate["row_count"],
        "audit_times": sorted(candidate["audit_times"]),
        "source_references": [dict(reference) for reference in sorted(candidate["source_references"])],
    }


def conflicts(candidates: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    item_owners: dict[str, set[tuple[str, str, str]]] = defaultdict(set)
    shop_items: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    for candidate in candidates:
        owner = candidate["country"], candidate["main_sku"], candidate["shop"]
        item_owners[candidate["item_id"]].add(owner)
        shop_items[owner].add(candidate["item_id"])
    return {
        "item_multiple_owners": [
            {"item_id": item_id, "owners": [{"country": country, "main_sku": main_sku, "shop": shop} for country, main_sku, shop in sorted(owners)]}
            for item_id, owners in sorted(item_owners.items())
            if len(owners) > 1
        ],
        "main_sku_shop_multiple_items": [
            {"country": country, "main_sku": main_sku, "shop": shop, "item_ids": sorted(item_ids)}
            for (country, main_sku, shop), item_ids in sorted(shop_items.items())
            if len(item_ids) > 1
        ],
    }


def required_columns(header: tuple[Any, ...]) -> dict[str, int] | None:
    columns = {normalized_header(value): index for index, value in enumerate(header) if normalized_header(value)}
    return {name: columns[normalized_header(name)] for name in REQUIRED_COLUMNS} if all(normalized_header(name) in columns for name in REQUIRED_COLUMNS) else None


def merged_cell_values(worksheet: Any, zero_based_columns: set[int]) -> dict[tuple[int, int], Any]:
    values: dict[tuple[int, int], Any] = {}
    one_based_columns = {column + 1 for column in zero_based_columns}
    for merged_range in worksheet.merged_cells.ranges:
        columns = [column for column in one_based_columns if merged_range.min_col <= column <= merged_range.max_col]
        if not columns:
            continue
        value = worksheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in columns:
                values[row, column] = value
    return values


def display_cell(row: tuple[Any, ...], index: int, source_row: int, merged_values: dict[tuple[int, int], Any]) -> Any:
    value = row[index].value if index < len(row) else None
    return value if value is not None else merged_values.get((source_row, index + 1))


def normalized_header(value: object) -> str:
    return "".join((text_value(value) or "").split()).upper()


def item_id_text(value: object) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        try:
            return format(Decimal(str(value)), "f").rstrip("0").rstrip(".") or "0"
        except InvalidOperation:
            return None
    return text_value(value)


def normalize_main_sku(value: object) -> str | None:
    text = text_value(value)
    return text.upper() if text else None


def country_from_shop(shop: str | None) -> str | None:
    match = SHOP_COUNTRY.search(shop or "")
    return match.group(1).upper() if match else None


def text_value(value: object) -> str | None:
    text = " ".join(str(value or "").replace("\xa0", " ").split())
    return text or None


def time_value(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return text_value(value)


def cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def parse_input(value: str) -> tuple[str, Path]:
    period, separator, source = value.partition("=")
    if not separator or not period.strip() or not source.strip():
        raise argparse.ArgumentTypeError("--input must be PERIOD=PATH")
    return period.strip(), Path(source.strip())



def validate_period_inputs(
    period_inputs: Iterable[tuple[str, Path]], manifests: dict[str, Path], allow_unmanifested: bool = False
) -> None:
    inputs = list(period_inputs)
    input_periods = [period for period, _ in inputs]
    if len(input_periods) != len(set(input_periods)):
        raise ValueError("duplicate input period")
    unexpected = sorted(set(manifests) - set(input_periods))
    if unexpected:
        raise ValueError(f"manifest without input period: {', '.join(unexpected)}")
    for period, source_path in inputs:
        manifest_path = manifests.get(period)
        if manifest_path is None:
            if allow_unmanifested:
                continue
            raise ValueError(f"missing manifest for input period: {period}")
        validate_period_manifest(period, Path(source_path), manifest_path)


def validate_period_manifest(period: str, source_path: Path, manifest_path: Path) -> None:
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ValueError(f"invalid manifest {manifest_path}: {error}") from error
    if not isinstance(manifest, dict):
        raise ValueError(f"manifest must be a JSON object: {manifest_path}")
    if manifest.get("status") != "ok":
        raise ValueError(f"manifest status is not ok: {manifest_path}")
    if manifest.get("period") != period:
        raise ValueError(f"manifest period mismatch: {manifest_path}")
    if manifest.get("payload_widget_name") != DATA_SHEET:
        raise ValueError(f"manifest widget mismatch: {manifest_path}")
    replaced = manifest.get("period_replaced_count")
    if isinstance(replaced, bool) or not isinstance(replaced, int) or replaced <= 0:
        raise ValueError(f"manifest period_replaced_count must be positive: {manifest_path}")
    if Path(str(manifest.get("file") or "")).name != source_path.name:
        raise ValueError(f"manifest file mismatch: {manifest_path}")
    if manifest.get("size") != source_path.stat().st_size:
        raise ValueError(f"manifest size mismatch: {manifest_path}")
    if manifest.get("sha256") != sha256_file(source_path):
        raise ValueError(f"manifest sha256 mismatch: {manifest_path}")
    sheets = manifest.get("sheets")
    if not isinstance(sheets, list) or [sheet.get("name") for sheet in sheets if isinstance(sheet, dict)] != [DATA_SHEET, FILTER_SHEET]:
        raise ValueError(f"manifest sheets mismatch: {manifest_path}")
    row_count = sheets[0].get("row_count")
    if isinstance(row_count, bool) or not isinstance(row_count, int) or row_count <= 0:
        raise ValueError(f"manifest data sheet has no rows: {manifest_path}")
    workbook = load_workbook(source_path, read_only=True, data_only=True, keep_links=False)
    try:
        if workbook.sheetnames != [DATA_SHEET, FILTER_SHEET]:
            raise ValueError(f"workbook sheets mismatch: {source_path}")
        filters = workbook[FILTER_SHEET]
        filters.reset_dimensions()
        if f'"{period}"' not in str(filters.cell(row=1, column=2).value or ""):
            raise ValueError(f"workbook filter period mismatch: {source_path}")
    finally:
        workbook.close()

def main() -> None:

    parser = argparse.ArgumentParser(description="Read-only historical FineBI Item candidate reconciliation.")
    parser.add_argument("--input", action="append", required=True, type=parse_input, metavar="PERIOD=PATH")
    parser.add_argument("--manifest", action="append", default=[], type=parse_input, metavar="PERIOD=PATH")
    parser.add_argument("--allow-unmanifested", action="store_true")
    parser.add_argument("--watchlist", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    manifests = {}
    for period, manifest_path in args.manifest:
        if period in manifests:
            parser.error(f"duplicate manifest period: {period}")
        manifests[period] = manifest_path
    try:
        validate_period_inputs(args.input, manifests, args.allow_unmanifested)
    except ValueError as error:
        parser.error(str(error))
    result = reconcile_item_candidates(json.loads(args.watchlist.read_text(encoding="utf-8")), args.input)
    encoded = json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    if args.output:
        args.output.write_text(encoded, encoding="utf-8")
    else:
        print(encoded, end="")


if __name__ == "__main__":
    main()
