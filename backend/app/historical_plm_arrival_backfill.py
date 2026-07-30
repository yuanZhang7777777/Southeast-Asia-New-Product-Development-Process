from __future__ import annotations

import argparse
import json
import re
from datetime import date, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models
from app.historical_archive_import import APPLY_ALLOWED_ENVS
from app.historical_watchlist import sha256_file
from app.plm_arrivals import DEFAULT_BLOC_NAME, HEADERS, _date_value, _number, _text
from app.plm_processing import _datetime_value
from app.services import audit, normalize_site_code

EPOCH_PATTERN = re.compile(r"(\d{13})")


def snapshot_date_from_name(name: str) -> str | None:
    match = EPOCH_PATTERN.search(name)
    if not match:
        return None
    from datetime import datetime

    return datetime.fromtimestamp(int(match.group(1)) / 1000, tz=timezone.utc).date().isoformat()


def parse_plm_snapshot(workbook_path: Path, bloc_name: str = DEFAULT_BLOC_NAME) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header_values = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        indexes = {field: header_values.index(title) for field, title in HEADERS.items() if title in header_values}
        items: list[dict[str, Any]] = []
        for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = {field: row[index] if index < len(row) else None for field, index in indexes.items()}
            if _text(values.get("bloc_name")) != bloc_name:
                continue
            latest_date = _date_value(values.get("latest_storage_time"))
            first_date = _date_value(values.get("first_listing_time"))
            if first_date is None:
                arrival_type = "unknown"
            elif latest_date is not None and first_date == latest_date:
                arrival_type = "new_arrival"
            else:
                arrival_type = "restock"
            items.append(
                {
                    "source_sheet": sheet.title,
                    "source_row": source_row,
                    "arrival_type": arrival_type,
                    "product_name": _text(values.get("product_name")),
                    "salesperson_name": _text(values.get("salesperson_name")),
                    "sub_sku": _text(values.get("sub_sku")),
                    "main_sku": _text(values.get("main_sku")),
                    "country": _text(values.get("country")),
                    "warehouse": _text(values.get("warehouse")),
                    "latest_storage_time": _text(values.get("latest_storage_time")),
                    "first_listing_time": _text(values.get("first_listing_time")),
                    "available_quantity": _number(values.get("available_quantity")),
                    "real_stock_quantity": _number(values.get("real_stock_quantity")),
                    "daily_sales": _number(values.get("daily_sales")),
                }
            )
    finally:
        workbook.close()
    return items


def _history_claim_matches(db: Session, row: dict[str, Any]) -> list[tuple[models.SalesClaimForecast, models.NewProductOpportunity]]:
    sub_sku = (row.get("sub_sku") or "").strip().upper()
    site = normalize_site_code(row.get("country")) or ""
    salesperson = (row.get("salesperson_name") or "").strip()
    if not sub_sku or not site or not salesperson:
        return []
    rows = db.execute(
        select(models.SalesClaimForecast, models.NewProductOpportunity)
        .join(models.NewProductOpportunity, models.NewProductOpportunity.id == models.SalesClaimForecast.opportunity_id)
        .where(
            models.SalesClaimForecast.claim_result == "claim",
            models.SalesClaimForecast.source_column == "platform",
            models.SalesClaimForecast.downstream_status.is_not(None),
            models.SalesClaimForecast.salesperson_name == salesperson,
        )
    ).all()
    return [
        (claim, opportunity)
        for claim, opportunity in rows
        if (opportunity.sub_sku or "").strip().upper() == sub_sku
        and (normalize_site_code(opportunity.site or opportunity.country) or "") == site
    ]


def apply_snapshot(
    db: Session,
    workbook_path: Path,
    snapshot_date: str,
    *,
    bloc_name: str = DEFAULT_BLOC_NAME,
    apply: bool = False,
    actor: str | None = None,
) -> dict[str, Any]:
    items = parse_plm_snapshot(workbook_path, bloc_name=bloc_name)
    source_hash = sha256_file(workbook_path)
    existing_batch = db.scalars(
        select(models.PlmArrivalBatch).where(models.PlmArrivalBatch.source_hash == source_hash)
    ).one_or_none()

    counts = {
        "rows": len(items),
        "matched_rows": 0,
        "arrival_created": 0,
        "claims_already_have_arrival": 0,
        "batch_already_imported": existing_batch is not None,
    }
    batch = existing_batch
    if apply and batch is None:
        batch = models.PlmArrivalBatch(
            arrival_date=snapshot_date,
            source_file=workbook_path.name,
            source_hash=source_hash,
            bloc_name=bloc_name,
            status="history_archive",
            row_count=len(items),
        )
        db.add(batch)
        db.flush()
        for row in items:
            db.add(
                models.PlmArrivalItem(
                    batch_id=batch.id,
                    source_sheet=row.get("source_sheet"),
                    source_row=row.get("source_row"),
                    arrival_type=row["arrival_type"],
                    product_name=row.get("product_name"),
                    salesperson_name=row.get("salesperson_name"),
                    country=row.get("country"),
                    warehouse=row.get("warehouse"),
                    main_sku=row.get("main_sku"),
                    sub_sku=row.get("sub_sku"),
                    latest_storage_time=_datetime_value(row.get("latest_storage_time")),
                    first_listing_time=_datetime_value(row.get("first_listing_time")),
                    available_quantity=row.get("available_quantity"),
                    real_stock_quantity=row.get("real_stock_quantity"),
                    daily_sales=row.get("daily_sales"),
                    match_status="history_archive",
                    raw_payload=row,
                )
            )

    for row in items:
        matches = _history_claim_matches(db, row)
        if not matches:
            continue
        counts["matched_rows"] += 1
        for claim, opportunity in matches:
            has_arrival = db.scalar(
                select(models.ArrivalRecord.id).where(models.ArrivalRecord.claim_record_id == claim.id)
            )
            if has_arrival:
                counts["claims_already_have_arrival"] += 1
                continue
            counts["arrival_created"] += 1
            if apply and batch is not None:
                db.add(
                    models.ArrivalRecord(
                        opportunity_id=opportunity.id,
                        claim_record_id=claim.id,
                        plm_arrival_batch_id=batch.id,
                        salesperson_name=claim.salesperson_name,
                        country=row.get("country") or opportunity.country,
                        warehouse=row.get("warehouse"),
                        arrived_quantity=int(row["available_quantity"]) if row.get("available_quantity") is not None else None,
                        arrived_at=_datetime_value(row.get("latest_storage_time")),
                        note="历史PLM快照回填",
                    )
                )
    if apply:
        audit(
            db,
            "history.plm_arrivals_backfilled",
            "plm_arrival_batch",
            batch.id if batch else None,
            {**{k: v for k, v in counts.items() if k != "batch_already_imported"}, "snapshot_date": snapshot_date},
            actor,
        )
        db.flush()
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description="历史 PLM 汇总快照：库存档案 + 到货时间回填（不动状态不发通知）")
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--snapshot-date", help="快照日期 YYYY-MM-DD；缺省从文件名毫秒时间戳推导")
    parser.add_argument("--apply-dev", action="store_true")
    parser.add_argument("--actor", default="history_plm_backfill")
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    snapshot_date = args.snapshot_date or snapshot_date_from_name(workbook_path.name)
    if not snapshot_date:
        raise SystemExit("无法从文件名推导快照日期，请传 --snapshot-date")

    from app.config import get_settings
    from app.db import SessionLocal

    if args.apply_dev:
        settings = get_settings()
        if settings.app_env not in APPLY_ALLOWED_ENVS:
            raise SystemExit(f"apply blocked: app_env={settings.app_env}")
    with SessionLocal() as db:
        report = apply_snapshot(
            db, workbook_path, snapshot_date, apply=args.apply_dev, actor=args.actor
        )
        if args.apply_dev:
            db.commit()
    print(json.dumps({"mode": "apply" if args.apply_dev else "dry-run", "snapshot_date": snapshot_date, **report}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
