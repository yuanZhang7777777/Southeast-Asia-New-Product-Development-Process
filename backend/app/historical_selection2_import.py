"""选品2历史档案只读解析器。

默认只解析或 dry-run；写库时也只恢复商品、历史认领/不认领和来源证据，不创建
当前认领、复核、二调、刊登、观察或钉钉任务。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models
from app.field_mapping import json_safe_value, text_value
from app.historical_archive_import import APPLY_ALLOWED_ENVS
from app.historical_watchlist import sha256_file
from app.services import audit

SOURCE_TYPE = "history_selection2"
ARCHIVE_STATUS = "historical_archive"
PERIOD_PATTERN = re.compile(r"(?P<month>\d{1,2})\.(?P<day>\d{1,2})期|(?P<compact>\d{3,4})期")
INFRINGEMENT_TEXT = "侵权"
HEADER_SEARCH_ROWS = 5
DEFAULT_COUNTRY = "PH"
AUDIT_ACTION = "history_selection2.imported"
DISPIMG_PATTERN = re.compile(r'DISPIMG\("(?P<id>ID_[A-F0-9]+)"', flags=re.IGNORECASE)
XDR_NS = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CELL_IMAGE_NS = "{http://www.wps.cn/officeDocument/2017/etCustomData}"

# 7.11期是选品2历史清洗的主表。旧表仅在列名/列位上有差异；未列入这里的
# 非空字段会进入历史补充字段，保留来源而不臆测业务含义。
CANONICAL_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "来源销售员": ("开发姓名", "销售员"),
    "首单备货数量": ("首单备货数量", "首单备货"),
    "产品名称": ("产品名称", "商品名称"),
    "产品规格属性": ("产品规格属性（材质、大小、颜色）", "产品规格属性"),
    "图片": ("图片", "图链"),
    "进价": ("进价",),
    "Shopee稳定期定价": (
        "Shopee稳定期定价",
        "SP上家₱",
        "稳定期定价",
        "SP上家₱稳定期定价",
        "SP稳定期定价₱",
    ),
    "Shopee推广期定价": (
        "Shopee推广期定价",
        "SP上架₱",
        "推广期定价",
        "SP上架₱推广期定价",
        "SP推广期定价₱",
    ),
    "SP上家¥": ("SP上家¥",),
    "海运": ("海运",),
    "广告刷单费率+退货率+上架仓储费+资金费率+库存损益": (
        "广告刷单费率+退货率+上架仓储费+资金费率+库存损益",
    ),
    "国外出库操作费": ("国外出库操作费",),
    "国内操作费": ("国内操作费",),
    "进货运费": ("进货运费",),
    "销售成本（比索）": ("销售成本（比索）",),
    "SP利润额（比索）": ("SP利润额（比索）",),
    "SP利润额（元）": ("SP利润额（元）",),
    "推广期SP利润率": ("推广期SP利润率",),
    "稳定期SP利润率": ("稳定期SP利润率",),
    "长(箱)": ("长(箱)",),
    "宽(箱)": ("宽(箱)",),
    "高(箱)": ("高(箱)",),
    "体积(箱)": ("体积(箱)",),
    "件数": ("件数",),
    "体积(单)": ("体积(单)",),
    "重量KG": ("重量KG", "重量kg"),
    "高价高消链接": ("高价高消链接",),
    "进货链接": ("进货链接",),
    "1688商家核价记录": ("1688商家核价记录",),
    "审核": ("审核",),
    "低价高消链接": ("低价高消链接",),
    "最新低价链接": ("最新低价链接",),
    "最低价链接": ("最低价链接",),
    "是否贴标包装及下单备注": ("是否贴标包装及下单备注",),
    "首单备货金额": ("首单备货金额",),
    "首单备货体积": ("首单备货体积",),
    "供应链核价后采购链接": ("供应链核价后采购链接",),
    "供应链核价后采购价": ("供应链核价后采购价",),
    "核价意见": ("核价意见",),
    "核价人": ("核价人",),
    "核价时间": ("核价时间",),
}


def period_from_sheet(sheet_name: str) -> str | None:
    match = PERIOD_PATTERN.search(sheet_name)
    if not match:
        return None
    if compact := match.group("compact"):
        return compact.zfill(4)
    return f"{int(match.group('month')):02d}{int(match.group('day')):02d}"


def business_period_from_sheet(sheet_name: str) -> str | None:
    period = period_from_sheet(sheet_name)
    return f"选品2-财根{period}期" if period else None


def parse_selection2_workbook(
    path: Path,
    max_rows_per_sheet: int | None = None,
    upload_images: bool = False,
) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    sheet_reports: list[dict[str, Any]] = []
    skipped_sheets: list[dict[str, str]] = []
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        periods: list[tuple[int, str, str]] = []
        for sheet_name in workbook.sheetnames:
            period = period_from_sheet(sheet_name)
            if period is None:
                skipped_sheets.append({"sheet": sheet_name, "reason": "非期数sheet"})
                continue

            periods.append((int(period), sheet_name, period))
            worksheet = workbook[sheet_name]
            try:
                worksheet.reset_dimensions()
            except AttributeError:
                pass

            headers, header_row = worksheet_headers(worksheet)
            sheet_rows = skipped = 0
            if not headers:
                sheet_reports.append({"sheet": sheet_name, "period": period, "rows": sheet_rows, "skipped": skipped})
                continue

            for source_row, row in enumerate(
                worksheet.iter_rows(
                    min_row=header_row + 1,
                    values_only=True,
                ),
                start=header_row + 1,
            ):
                if max_rows_per_sheet is not None and sheet_rows >= max_rows_per_sheet:
                    break
                parsed = parse_historical_selection2_row(
                    row,
                    headers=headers,
                    source_file=path.name,
                    source_sheet=sheet_name,
                    source_row=source_row,
                    source_header_row=header_row,
                    business_period=business_period_from_sheet(sheet_name),
                )
                if parsed is None:
                    skipped += 1
                    continue
                rows.append(parsed)
                sheet_rows += 1
            sheet_reports.append({"sheet": sheet_name, "period": period, "rows": sheet_rows, "skipped": skipped})
    finally:
        workbook.close()
    image_stats = attach_wps_cell_images(path, rows) if upload_images else {
        "references": sum(bool(row["snapshot"].get("image_reference")) for row in rows),
        "uploaded": 0,
        "rows_with_image": sum(bool(row.get("image_url")) for row in rows),
        "upload_failures": [],
    }
    latest = max(periods, default=None)
    return {
        "source_type": SOURCE_TYPE,
        "source_file": path.name,
        "source_sha256": sha256_file(path),
        "readable": True,
        "latest_sheet": latest[1] if latest else None,
        "latest_period": latest[2] if latest else None,
        "sheets": sheet_reports,
        "skipped_sheets": skipped_sheets,
        "row_count": len(rows),
        "image_stats": image_stats,
        "rows": rows,
    }


def inspect_selection2_workbook(path: Path) -> dict[str, Any]:
    try:
        report = parse_selection2_workbook(path, max_rows_per_sheet=0)
    except Exception as exc:
        return {
            "source_type": SOURCE_TYPE,
            "source_file": path.name,
            "readable": False,
            "latest_sheet": None,
            "latest_period": None,
            "sheets": [],
            "skipped_sheets": [],
            "row_count": 0,
            "error": f"could not open workbook: {exc}",
        }
    return {key: value for key, value in report.items() if key != "rows"}


def dispimg_id(value: Any) -> str | None:
    match = DISPIMG_PATTERN.search(text_value(value) or "")
    return match.group("id").upper() if match else None


def is_image_url(value: Any) -> bool:
    return bool(re.match(r"^(?:https?://|/uploaded-sources/)", text_value(value) or "", flags=re.IGNORECASE))


def wps_cell_images(archive: zipfile.ZipFile) -> dict[str, tuple[bytes, str]]:
    image_xml = "xl/cellimages.xml"
    rels_xml = "xl/_rels/cellimages.xml.rels"
    if image_xml not in archive.namelist() or rels_xml not in archive.namelist():
        return {}
    relationships = ElementTree.fromstring(archive.read(rels_xml))
    targets = {
        relation.get("Id"): f"xl/{relation.get('Target', '').lstrip('/')}"
        for relation in relationships.iter(f"{REL_NS}Relationship")
    }
    root = ElementTree.fromstring(archive.read(image_xml))
    images: dict[str, tuple[bytes, str]] = {}
    for cell_image in root.iter(f"{CELL_IMAGE_NS}cellImage"):
        properties = cell_image.find(f".//{XDR_NS}cNvPr")
        blip = cell_image.find(f".//{A_NS}blip")
        if properties is None or blip is None:
            continue
        image_id = (properties.get("name") or "").upper()
        media_path = targets.get(blip.get(f"{R_NS}embed", ""))
        if not image_id or not media_path or media_path not in archive.namelist():
            continue
        images[image_id] = (archive.read(media_path), Path(media_path).suffix.lstrip(".") or "png")
    return images


def attach_wps_cell_images(path: Path, rows: list[dict[str, Any]]) -> dict[str, Any]:
    from app.oss_storage import upload_product_image

    failures: list[dict[str, Any]] = []
    uploaded_by_reference: dict[str, str] = {}
    with zipfile.ZipFile(path) as archive:
        images = wps_cell_images(archive)
        for row in rows:
            image_reference = row["snapshot"].get("image_reference")
            if not image_reference:
                continue
            image = images.get(image_reference)
            if image is None:
                failures.append(image_failure(row, image_reference, "image_reference_not_found"))
                continue
            url = uploaded_by_reference.get(image_reference)
            if url is None:
                data, extension = image
                digest = hashlib.sha1(data).hexdigest()[:16]
                try:
                    url = upload_product_image(
                        data,
                        extension,
                        f"{SOURCE_TYPE}-{row['source_sheet']}",
                        row["source_row"],
                        digest,
                    )
                except Exception as exc:
                    failures.append(image_failure(row, image_reference, type(exc).__name__))
                    continue
                if not url:
                    failures.append(image_failure(row, image_reference, "upload_returned_empty"))
                    continue
                uploaded_by_reference[image_reference] = url
            row["image_url"] = url
            row["snapshot"]["extracted_image_url"] = url
            row["snapshot"]["image_origin"] = "selection2_wps_cell_image"
    return {
        "references": sum(bool(row["snapshot"].get("image_reference")) for row in rows),
        "uploaded": len(uploaded_by_reference),
        "rows_with_image": sum(bool(row.get("image_url")) for row in rows),
        "upload_failures": failures,
    }


def image_failure(row: dict[str, Any], image_reference: str, reason: str) -> dict[str, Any]:
    return {
        "source_sheet": row["source_sheet"],
        "source_row": row["source_row"],
        "main_sku": row["main_sku"],
        "sub_sku": row["sub_sku"],
        "image_reference": image_reference,
        "reason": reason,
    }


def parse_historical_selection2_row(
    row: tuple[Any, ...],
    *,
    headers: dict[str, Any],
    source_file: str,
    source_sheet: str,
    source_row: int,
    source_header_row: int | None = None,
    business_period: str | None = None,
) -> dict[str, Any] | None:
    values = {
        column: json_safe_value(cell_value(row, column))
        for column in ordered_columns(headers)
    }
    main_sku_column = find_main_sku_column(headers)
    sub_sku_column = find_sub_sku_column(headers)
    sub_sku = text_value(values.get(sub_sku_column)) if sub_sku_column else None
    if not sub_sku or _is_repeated_header(sub_sku):
        return None

    main_sku = text_value(values.get(main_sku_column)) if main_sku_column else None
    backfilled_fields: dict[str, dict[str, str]] = {}
    if not main_sku:
        main_sku = sub_sku
        backfilled_fields["main_sku"] = {"value": sub_sku, "source": "sub_sku"}

    resolved_business_period = business_period or business_period_from_sheet(source_sheet) or source_sheet
    normalized_fields, supplementary_fields, field_quality_issues = normalize_selection2_fields(values, headers)
    source_reference: dict[str, Any] = {
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
    }
    if source_header_row is not None:
        source_reference["header_row"] = source_header_row

    archive_only_reason = None
    operator_match_policy = "none"
    claims: list[dict[str, Any]] = []
    rejected_sources: list[dict[str, Any]] = []
    claim_quality_issues: list[dict[str, Any]] = []
    if has_infringement_marker(values):
        archive_only_reason = "infringing_product"
        operator_match_policy = "skip"
    else:
        claims, rejected_sources, claim_quality_issues = parse_claim_sources(values, headers)
        if not claims and not rejected_sources:
            archive_only_reason = "historical_claim_quality_issue" if claim_quality_issues else "historical_unclaimed"

    quality_issues = [*field_quality_issues, *claim_quality_issues]
    image_value = field_value(normalized_fields, "图片")
    image_reference = dispimg_id(image_value)
    snapshot = {
        "archive_type": "historical_selection2",
        "business_period": resolved_business_period,
        "headers_by_cell": {column: headers.get(column) for column in ordered_columns(headers)},
        "raw_cells_by_cell": values,
        "fields_by_cell": fields_by_cell(values, headers),
        "normalized_fields": normalized_fields,
        "historical_supplementary_fields": supplementary_fields,
        "source_reference": source_reference,
        "claims": claims,
        "rejected_sources": rejected_sources,
        "quality_issues": quality_issues,
        "archive_only_reason": archive_only_reason,
    }
    if image_reference:
        snapshot["image_reference"] = image_reference
    if backfilled_fields:
        snapshot["backfilled_fields"] = backfilled_fields

    product_name = field_value(normalized_fields, "产品名称")
    specification = field_value(normalized_fields, "产品规格属性")
    image_url = image_value if is_image_url(image_value) else None
    return {
        "source_type": SOURCE_TYPE,
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
        "batch": resolved_business_period,
        "current_status": ARCHIVE_STATUS,
        "country": DEFAULT_COUNTRY,
        "site": DEFAULT_COUNTRY,
        "main_sku": main_sku,
        "sub_sku": sub_sku,
        "main_sku_name": product_name,
        "sub_sku_name": specification or product_name,
        "image_url": image_url,
        "normalized_fields": normalized_fields,
        "historical_supplementary_fields": supplementary_fields,
        "snapshot": snapshot,
        "claims": claims,
        "rejected_sources": rejected_sources,
        "quality_issues": quality_issues,
        "archive_only_reason": archive_only_reason,
        "operator_match_policy": operator_match_policy,
        "task_policy": "none",
    }


def worksheet_headers(worksheet: Any) -> tuple[dict[str, Any], int]:
    for header_row in range(1, HEADER_SEARCH_ROWS + 1):
        row = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            ),
            (),
        )
        headers = {
            get_column_letter(index): json_safe_value(value)
            for index, value in enumerate(row, start=1)
        }
        if find_main_sku_column(headers) and find_sub_sku_column(headers):
            return headers, header_row
    return {}, 1


def parse_claim_sources(
    values: dict[str, Any],
    headers: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    entries_by_salesperson: dict[str, list[dict[str, Any]]] = {}
    source_order: list[str] = []
    quality_issues: list[dict[str, Any]] = []
    slots = claim_slots(headers)
    shifted_entry, consumed_columns = shifted_primary_claim_entry(values, slots)
    if shifted_entry:
        salesperson = shifted_entry.pop("salesperson_name")
        entries_by_salesperson[salesperson] = [shifted_entry]
        source_order.append(salesperson)

    for slot in slots:
        if slot["name_column"] in consumed_columns:
            continue
        salesperson = text_value(values.get(slot["name_column"]))
        if not salesperson:
            continue
        if not valid_claimant_name(salesperson):
            quality_issues.append(
                {
                    "code": "invalid_claimant_name",
                    "raw_value": salesperson,
                    "source_column": slot["name_column"],
                }
            )
            continue

        raw_value = first_slot_value(values, slot)
        entry = {
            "raw_value": raw_value,
            "source_column": slot["source_column"],
            "kind": claim_value_kind(raw_value, values.get(slot["status_column"]) if slot["status_column"] else None),
        }
        if salesperson not in entries_by_salesperson:
            entries_by_salesperson[salesperson] = []
            source_order.append(salesperson)
        entries_by_salesperson[salesperson].append(entry)

    claims: list[dict[str, Any]] = []
    rejected_sources: list[dict[str, Any]] = []
    for salesperson in source_order:
        entries = entries_by_salesperson[salesperson]
        value_keys = {claim_value_key(entry["raw_value"]) for entry in entries}
        if len(value_keys) > 1:
            quality_issues.append(
                {
                    "code": "conflicting_claim_values",
                    "salesperson_name": salesperson,
                    "raw_values": distinct_raw_values(entries),
                    "source_columns": [entry["source_column"] for entry in entries],
                }
            )
            continue

        entry = entries[0]
        if entry["kind"] == "claim":
            claim = {
                "salesperson_name": salesperson,
                "claim_result": "claim",
                "claim_daily_sales": strict_positive_number(entry["raw_value"]),
                "source_column": entry["source_column"],
            }
            if len(entries) > 1:
                claim["source_columns"] = [entry["source_column"] for entry in entries]
            claims.append(claim)
        elif entry["kind"] == "reject":
            rejected_sources.append(
                {
                    "salesperson_name": salesperson,
                    "raw_value": entry["raw_value"],
                    "source_column": entry["source_column"],
                }
            )
        else:
            quality_issues.append(
                {
                    "code": "invalid_claim_value",
                    "salesperson_name": salesperson,
                    "raw_value": entry["raw_value"],
                    "source_column": entry["source_column"],
                }
            )
    return claims, rejected_sources, quality_issues


def shifted_primary_claim_entry(
    values: dict[str, Any],
    slots: list[dict[str, str | None]],
) -> tuple[dict[str, Any] | None, set[str]]:
    if not slots:
        return None, set()
    primary = slots[0]
    if values.get(primary["name_column"]) not in (None, ""):
        return None, set()
    shifted_name_column = primary["daily_sales_column"]
    salesperson = text_value(values.get(shifted_name_column)) if shifted_name_column else None
    if not salesperson or not valid_claimant_name(salesperson):
        return None, set()

    value_column = primary["reason_column"] or shifted_name_column
    raw_value = values.get(value_column)
    consumed_columns = {shifted_name_column}
    if raw_value in (None, "") and len(slots) > 1:
        misplaced_reason_column = slots[1]["name_column"]
        misplaced_reason = text_value(values.get(misplaced_reason_column))
        if misplaced_reason and not valid_claimant_name(misplaced_reason):
            raw_value = values[misplaced_reason_column]
            value_column = misplaced_reason_column
            consumed_columns.add(misplaced_reason_column)

    return (
        {
            "salesperson_name": salesperson,
            "raw_value": raw_value,
            "source_column": source_range(shifted_name_column, value_column),
            "kind": claim_value_kind(raw_value, None),
        },
        consumed_columns,
    )


def claim_slots(headers: dict[str, Any]) -> list[dict[str, str | None]]:
    claimant_columns = [
        column
        for column in ordered_columns(headers)
        if is_claimant_header(headers.get(column))
    ]
    slots: list[dict[str, str | None]] = []
    for index, name_column in enumerate(claimant_columns):
        next_column = claimant_columns[index + 1] if index + 1 < len(claimant_columns) else None
        related_columns = columns_between(headers, name_column, next_column)
        daily_sales_column = first_header_column(headers, related_columns, is_daily_sales_header)
        reason_column = first_header_column(headers, related_columns, is_rejection_reason_header)
        status_column = first_header_column(headers, related_columns, is_claim_status_header)
        source_end = daily_sales_column or reason_column or status_column or name_column
        slots.append(
            {
                "name_column": name_column,
                "daily_sales_column": daily_sales_column,
                "reason_column": reason_column,
                "status_column": status_column,
                "source_column": source_range(name_column, source_end),
            }
        )
    return slots


def first_slot_value(values: dict[str, Any], slot: dict[str, str | None]) -> Any:
    for key in ("daily_sales_column", "reason_column", "status_column"):
        column = slot[key]
        if column and values.get(column) not in (None, ""):
            return values[column]
    return None


def claim_value_kind(raw_value: Any, status_value: Any) -> str:
    if is_explicit_rejection(status_value):
        return "reject"
    if strict_positive_number(raw_value) is not None:
        return "claim"
    if is_formula_error(raw_value):
        return "invalid"
    return "reject"


def is_zero_daily_sales(value: Any) -> bool:
    if isinstance(value, bool) or value in (None, ""):
        return False
    if isinstance(value, int | float):
        return float(value) == 0
    text = str(value).strip().replace(",", "")
    return bool(re.fullmatch(r"\d+(?:\.\d+)?", text)) and float(text) == 0


def strict_positive_number(value: Any) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, int | float):
        number = float(value)
        return number if number > 0 else None
    text = str(value).strip().replace(",", "")
    if not re.fullmatch(r"\d+(?:\.\d+)?", text):
        return None
    number = float(text)
    return number if number > 0 else None


def has_infringement_marker(values: dict[str, Any]) -> bool:
    return any(INFRINGEMENT_TEXT in (text_value(value) or "") for value in values.values())


def normalize_selection2_fields(
    values: dict[str, Any],
    headers: dict[str, Any],
) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    normalized_fields: dict[str, dict[str, Any]] = {}
    primary_columns: set[str] = set()
    alias_columns: set[str] = set()
    quality_issues: list[dict[str, Any]] = []
    for canonical_field, aliases in CANONICAL_FIELD_ALIASES.items():
        alias_keys = {header_key(alias) for alias in aliases}
        matched_columns = [
            column
            for column in ordered_columns(headers)
            if header_key(headers.get(column)) in alias_keys
        ]
        alias_columns.update(matched_columns)
        for column in matched_columns:
            value = values.get(column)
            if value in (None, ""):
                continue
            if is_formula_error(value):
                quality_issues.append(
                    formula_error_issue(canonical_field, column, headers.get(column))
                )
                continue
            if canonical_field not in normalized_fields:
                normalized_fields[canonical_field] = {
                    "value": value,
                    "source_column": column,
                    "source_header": headers.get(column),
                }
                primary_columns.add(column)

    supplementary_fields: list[dict[str, Any]] = []
    for column in ordered_columns(headers):
        value = values.get(column)
        if value in (None, ""):
            continue
        header = headers.get(column)
        if is_formula_error(value):
            if column not in alias_columns:
                quality_issues.append(formula_error_issue(text_value(header) or column, column, header))
            continue
        if column in primary_columns or is_identity_or_claim_header(header):
            continue
        supplementary_fields.append(
            {
                "field": text_value(header) or column,
                "value": value,
                "source_column": column,
                "source_header": header,
                "reason": "additional_alias_value" if column in alias_columns else "unmapped_history_field",
            }
        )
    return normalized_fields, supplementary_fields, quality_issues


def formula_error_issue(field: str, column: str, header: Any) -> dict[str, Any]:
    return {
        "code": "formula_error",
        "field": field,
        "source_column": column,
        "source_header": header,
    }


def field_value(normalized_fields: dict[str, dict[str, Any]], field: str) -> str | None:
    evidence = normalized_fields.get(field)
    return text_value(evidence.get("value")) if evidence else None


def is_identity_or_claim_header(header: Any) -> bool:
    return (
        is_main_sku_header(header)
        or is_sub_sku_header(header)
        or is_claimant_header(header)
        or is_daily_sales_header(header)
        or is_rejection_reason_header(header)
        or is_claim_status_header(header)
    )


def normalize_selection2_history_rows(rows: list[dict[str, Any] | None]) -> dict[str, Any]:
    groups: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for row in rows:
        if row is None:
            continue
        row = reparse_claim_evidence(row)
        key = (
            row.get("source_file"),
            row.get("source_sheet"),
            row.get("batch"),
            row.get("country"),
            row.get("main_sku"),
            row.get("sub_sku"),
        )
        groups.setdefault(key, []).append(row)

    unified_rows: list[dict[str, Any]] = []
    business_repair_rows: list[dict[str, Any]] = []
    supplementary_fields: list[dict[str, Any]] = []
    for group in groups.values():
        group = sorted(group, key=lambda item: int(item["source_row"]))
        source_rows = [item["source_row"] for item in group]
        values_by_field: dict[str, list[dict[str, Any]]] = {}
        for item in group:
            for field, evidence in item.get("normalized_fields", {}).items():
                values_by_field.setdefault(field, []).append(evidence)
            for field in item.get("historical_supplementary_fields", []):
                supplementary_fields.append(
                    {
                        **field,
                        "source_file": item["source_file"],
                        "source_sheet": item["source_sheet"],
                        "business_period": item["batch"],
                        "country": item["country"],
                        "main_sku": item["main_sku"],
                        "sub_sku": item["sub_sku"],
                        "source_row": item["source_row"],
                    }
                )

        conflicts = {
            field: distinct_evidence_values(evidence)
            for field, evidence in values_by_field.items()
            if len(distinct_evidence_values(evidence)) > 1
        }
        if conflicts:
            first = group[0]
            business_repair_rows.append(
                {
                    "source_file": first["source_file"],
                    "source_sheet": first["source_sheet"],
                    "business_period": first["batch"],
                    "country": first["country"],
                    "main_sku": first["main_sku"],
                    "sub_sku": first["sub_sku"],
                    "source_rows": source_rows,
                    "conflicting_fields": conflicts,
                }
            )
            continue

        merged = dict(group[0])
        merged_fields = {
            field: evidence[0]
            for field, evidence in values_by_field.items()
        }
        merged["normalized_fields"] = merged_fields
        merged["source_rows"] = source_rows
        merged["claims"] = unique_records(group, "claims")
        merged["rejected_sources"] = unique_records(group, "rejected_sources")
        merged["quality_issues"] = unique_records(group, "quality_issues")
        merged["historical_supplementary_fields"] = [
            field
            for item in group
            for field in item.get("historical_supplementary_fields", [])
        ]
        merged_snapshot = dict(merged["snapshot"])
        merged_snapshot["normalized_fields"] = merged_fields
        merged_snapshot["merged_source_rows"] = source_rows
        merged_snapshot["claims"] = merged["claims"]
        merged_snapshot["rejected_sources"] = merged["rejected_sources"]
        merged_snapshot["quality_issues"] = merged["quality_issues"]
        merged_snapshot["historical_supplementary_fields"] = merged["historical_supplementary_fields"]
        merged["snapshot"] = merged_snapshot
        unified_rows.append(merged)

    return {
        "unified_rows": unified_rows,
        "business_repair_rows": business_repair_rows,
        "supplementary_fields": supplementary_fields,
    }


def reparse_claim_evidence(row: dict[str, Any]) -> dict[str, Any]:
    snapshot = row.get("snapshot") or {}
    values = snapshot.get("raw_cells_by_cell")
    headers = snapshot.get("headers_by_cell")
    if not isinstance(values, dict) or not isinstance(headers, dict):
        return row

    reparsed = dict(row)
    reparsed_snapshot = dict(snapshot)
    field_issues = [
        issue
        for issue in row.get("quality_issues", [])
        if issue.get("code") not in {"conflicting_claim_values", "invalid_claim_value", "invalid_claimant_name"}
    ]
    if has_infringement_marker(values):
        claims: list[dict[str, Any]] = []
        rejected_sources: list[dict[str, Any]] = []
        claim_issues: list[dict[str, Any]] = []
        archive_only_reason = "infringing_product"
        reparsed["operator_match_policy"] = "skip"
    else:
        claims, rejected_sources, claim_issues = parse_claim_sources(values, headers)
        archive_only_reason = None
        if not claims and not rejected_sources:
            archive_only_reason = "historical_claim_quality_issue" if claim_issues else "historical_unclaimed"

    quality_issues = [*field_issues, *claim_issues]
    reparsed.update(
        {
            "claims": claims,
            "rejected_sources": rejected_sources,
            "quality_issues": quality_issues,
            "archive_only_reason": archive_only_reason,
        }
    )
    reparsed_snapshot.update(
        {
            "claims": claims,
            "rejected_sources": rejected_sources,
            "quality_issues": quality_issues,
            "archive_only_reason": archive_only_reason,
        }
    )
    reparsed["snapshot"] = reparsed_snapshot
    return reparsed


def selection2_row_identity(row: dict[str, Any]) -> tuple[str, str, str, str]:
    return (
        row.get("batch") or "",
        row.get("country") or DEFAULT_COUNTRY,
        row.get("main_sku") or "",
        row.get("sub_sku") or "",
    )


def find_history_opportunity(db: Session, row: dict[str, Any]) -> models.NewProductOpportunity | None:
    period, country, main_sku, sub_sku = selection2_row_identity(row)
    return db.scalar(
        select(models.NewProductOpportunity)
        .where(
            models.NewProductOpportunity.source_type == SOURCE_TYPE,
            models.NewProductOpportunity.batch == period,
            models.NewProductOpportunity.country == country,
            models.NewProductOpportunity.main_sku == main_sku,
            models.NewProductOpportunity.sub_sku == sub_sku,
        )
        .order_by(models.NewProductOpportunity.created_at.asc())
    )


def plan_selection2_rows(db: Session, rows: list[dict[str, Any]]) -> dict[str, Any]:
    normalization = normalize_selection2_history_rows(rows)
    counts = {
        "would_create_opportunities": 0,
        "would_update_opportunities": 0,
        "would_create_claims": 0,
        "would_update_claims": 0,
        "skipped_conflicting_rows": len(normalization["business_repair_rows"]),
    }
    by_period: dict[str, dict[str, int]] = {}
    for row in normalization["unified_rows"]:
        period = row["batch"]
        period_counts = by_period.setdefault(period, {"rows": 0, "claims": 0, "rejects": 0})
        period_counts["rows"] += 1
        period_counts["claims"] += len(row["claims"])
        period_counts["rejects"] += len(row["rejected_sources"])
        opportunity = find_history_opportunity(db, row)
        counts["would_update_opportunities" if opportunity else "would_create_opportunities"] += 1
        for relation in history_relations(row):
            existing = find_history_claim(db, opportunity.id, relation) if opportunity else None
            counts["would_update_claims" if existing else "would_create_claims"] += 1
    return {**counts, "by_period": by_period, "business_repair_rows": normalization["business_repair_rows"]}


def apply_selection2_rows(
    db: Session,
    rows: list[dict[str, Any]],
    *,
    source_label: str,
    imported_by: str | None = None,
    source_sha256: str | None = None,
) -> dict[str, Any]:
    normalization = normalize_selection2_history_rows(rows)
    counts = {
        "created_opportunities": 0,
        "updated_opportunities": 0,
        "created_claims": 0,
        "updated_claims": 0,
        "created_source_snapshots": 0,
        "updated_source_snapshots": 0,
        "skipped_conflicting_rows": len(normalization["business_repair_rows"]),
    }
    batch = models.ImportBatch(
        source_type=SOURCE_TYPE,
        source_file=source_label,
        source_sheet="all",
        business_period="历史全期",
        imported_by=imported_by,
        status="running",
    )
    db.add(batch)
    db.flush()

    for row in normalization["unified_rows"]:
        opportunity = find_history_opportunity(db, row)
        if opportunity is None:
            opportunity = models.NewProductOpportunity(
                source_type=SOURCE_TYPE,
                source_file=row["source_file"],
                source_sheet=row["source_sheet"],
                source_row=row["source_row"],
                import_batch_id=batch.id,
                batch=row["batch"],
                country=row.get("country") or DEFAULT_COUNTRY,
                site=row.get("site") or DEFAULT_COUNTRY,
                main_sku=row["main_sku"],
                sub_sku=row["sub_sku"],
                main_sku_name=row.get("main_sku_name"),
                sub_sku_name=row.get("sub_sku_name"),
                image_url=row.get("image_url"),
                current_status=ARCHIVE_STATUS,
                claim_pool_open=False,
                snapshot=row["snapshot"],
            )
            db.add(opportunity)
            db.flush()
            counts["created_opportunities"] += 1
        else:
            update_history_opportunity(opportunity, row, batch.id)
            counts["updated_opportunities"] += 1

        source_snapshot = db.scalar(
            select(models.SourceRecordSnapshot).where(
                models.SourceRecordSnapshot.opportunity_id == opportunity.id,
                models.SourceRecordSnapshot.source_file == row["source_file"],
                models.SourceRecordSnapshot.source_sheet == row["source_sheet"],
                models.SourceRecordSnapshot.source_row == row["source_row"],
            )
        )
        if source_snapshot is None:
            source_snapshot = models.SourceRecordSnapshot(
                import_batch_id=batch.id,
                opportunity_id=opportunity.id,
                source_file=row["source_file"],
                source_sheet=row["source_sheet"],
                source_row=row["source_row"],
                column_range="A:AW",
                payload=row["snapshot"],
            )
            db.add(source_snapshot)
            counts["created_source_snapshots"] += 1
        else:
            source_snapshot.import_batch_id = batch.id
            source_snapshot.column_range = "A:AW"
            source_snapshot.payload = row["snapshot"]
            counts["updated_source_snapshots"] += 1

        for relation in history_relations(row):
            claim = find_history_claim(db, opportunity.id, relation)
            created = claim is None
            if claim is None:
                claim = models.SalesClaimForecast(
                    opportunity_id=opportunity.id,
                    source_column=relation["source_column"],
                    claim_source=SOURCE_TYPE,
                )
                db.add(claim)
            claim.salesperson_name = relation["salesperson_name"]
            claim.claim_result = relation["claim_result"]
            claim.claim_daily_sales = relation["claim_daily_sales"]
            claim.reject_reason = relation["reject_reason"]
            claim.feedback_summary = relation.get("feedback_summary")
            claim.note = relation["note"]
            claim.task_id = None
            claim.downstream_status = None
            counts["created_claims" if created else "updated_claims"] += 1

    batch.created_count = counts["created_opportunities"]
    batch.updated_count = counts["updated_opportunities"]
    batch.skipped_count = counts["skipped_conflicting_rows"]
    batch.status = "completed"
    audit(
        db,
        AUDIT_ACTION,
        "new_product_opportunity",
        None,
        {
            "source_file": source_label,
            "source_sha256": source_sha256,
            "import_batch_id": batch.id,
            **counts,
        },
        imported_by,
    )
    db.flush()
    return {
        "import_batch_id": batch.id,
        **counts,
        "business_repair_rows": normalization["business_repair_rows"],
    }


def update_history_opportunity(
    opportunity: models.NewProductOpportunity,
    row: dict[str, Any],
    import_batch_id: str,
) -> None:
    opportunity.source_file = row["source_file"]
    opportunity.source_sheet = row["source_sheet"]
    opportunity.source_row = row["source_row"]
    opportunity.import_batch_id = import_batch_id
    opportunity.site = row.get("site") or DEFAULT_COUNTRY
    opportunity.current_status = ARCHIVE_STATUS
    opportunity.claim_pool_open = False
    opportunity.snapshot = row["snapshot"]
    for field in ("main_sku_name", "sub_sku_name", "image_url"):
        if row.get(field) not in (None, ""):
            setattr(opportunity, field, row[field])


def history_relations(row: dict[str, Any]) -> list[dict[str, Any]]:
    relations: list[dict[str, Any]] = []
    for claim in row.get("claims", []):
        relations.append(
            history_relation(
                row,
                claim["salesperson_name"],
                "claim",
                claim["source_column"],
                claim.get("claim_daily_sales"),
                None,
                claim.get("source_columns"),
            )
        )
    for rejected in row.get("rejected_sources", []):
        relations.append(
            history_relation(
                row,
                rejected["salesperson_name"],
                "reject",
                rejected["source_column"],
                None,
                rejection_reason(rejected.get("raw_value")),
                rejected.get("source_columns"),
                rejected.get("raw_value"),
            )
        )
    return relations


def history_relation(
    row: dict[str, Any],
    salesperson_name: str,
    claim_result: str,
    source_range: str,
    claim_daily_sales: float | None,
    reject_reason: str | None,
    source_ranges: list[str] | None = None,
    raw_value: Any = None,
) -> dict[str, Any]:
    source_column = f"{SOURCE_TYPE}:{source_range}"
    note = {
        "history_source": {
            "source_file": row["source_file"],
            "source_sheet": row["source_sheet"],
            "source_period": row["batch"],
            "source_row": row["source_row"],
            "source_column": source_range,
            **({"source_columns": source_ranges} if source_ranges else {}),
            **({"raw_value": raw_value} if claim_result == "reject" else {}),
        }
    }
    return {
        "salesperson_name": salesperson_name,
        "claim_result": claim_result,
        "claim_daily_sales": claim_daily_sales,
        "reject_reason": reject_reason,
        "source_column": source_column,
        "note": json.dumps(note, ensure_ascii=False, default=str),
    }


def rejection_reason(raw_value: Any) -> str:
    if raw_value in (None, ""):
        return "来源未填写认领单销"
    if is_zero_daily_sales(raw_value):
        return "来源认领单销为 0"
    return text_value(raw_value) or "来源标记不认领"


def find_history_claim(
    db: Session,
    opportunity_id: str,
    relation: dict[str, Any],
) -> models.SalesClaimForecast | None:
    return db.scalar(
        select(models.SalesClaimForecast)
        .where(
            models.SalesClaimForecast.opportunity_id == opportunity_id,
            models.SalesClaimForecast.claim_source == SOURCE_TYPE,
            models.SalesClaimForecast.source_column == relation["source_column"],
            models.SalesClaimForecast.salesperson_name == relation["salesperson_name"],
        )
        .order_by(models.SalesClaimForecast.created_at.asc())
    )


def distinct_evidence_values(evidence: list[dict[str, Any]]) -> list[Any]:
    values: list[Any] = []
    seen: set[str] = set()
    for item in evidence:
        value = item.get("value")
        key = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
        if key not in seen:
            seen.add(key)
            values.append(value)
    return values


def unique_records(rows: list[dict[str, Any]], field: str) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rows:
        for value in row.get(field, []):
            key = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
            if key not in seen:
                seen.add(key)
                result.append(value)
    return result


def fields_by_cell(values: dict[str, Any], headers: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        column: {"header": headers.get(column), "value": value}
        for column, value in values.items()
        if value not in (None, "")
    }


def find_main_sku_column(headers: dict[str, Any]) -> str | None:
    return first_header_column(headers, ordered_columns(headers), is_main_sku_header)


def find_sub_sku_column(headers: dict[str, Any]) -> str | None:
    columns = ordered_columns(headers)
    return first_header_column(headers, columns, lambda header: header_key(header) == "SKU") or first_header_column(
        headers, columns, is_sub_sku_header
    )


def find_product_name_column(headers: dict[str, Any]) -> str | None:
    return first_header_column(headers, ordered_columns(headers), lambda header: "产品名称" in header_key(header) or "商品名称" in header_key(header))


def find_specification_column(headers: dict[str, Any]) -> str | None:
    return first_header_column(headers, ordered_columns(headers), lambda header: "规格" in header_key(header))


def find_image_column(headers: dict[str, Any]) -> str | None:
    return first_header_column(headers, ordered_columns(headers), lambda header: "图片" in header_key(header) or "图链" in header_key(header))


def is_main_sku_header(header: Any) -> bool:
    key = header_key(header)
    return key == "SPU" or "主SKU" in key


def is_sub_sku_header(header: Any) -> bool:
    key = header_key(header)
    return key in {"SKU", "子SKU", "开发子SKU"} or key.endswith("子SKU")


def is_claimant_header(header: Any) -> bool:
    key = header_key(header)
    if key in {header_key("销售员"), header_key("开发姓名")}:
        return False
    if not key or "单销" in key or "原因" in key or "反馈" in key:
        return False
    if "销售员" not in key and "认领销售" not in key:
        return False
    return "认领" in key or "主销售员" in key or key.startswith("销售员")


def is_daily_sales_header(header: Any) -> bool:
    return "单销" in header_key(header)


def is_rejection_reason_header(header: Any) -> bool:
    return "不认领原因" in header_key(header)


def is_claim_status_header(header: Any) -> bool:
    return "是否认领" in header_key(header)


def valid_claimant_name(value: str) -> bool:
    text = value.strip()
    if not text or re.fullmatch(r"\d+(?:\.\d+)?", text):
        return False
    return not any(
        marker in text.lower()
        for marker in ("表格", "已认领", "侵权", "需要", "资质", "趋势", "单销低", "http://", "https://", "链接")
    )


def is_explicit_rejection(value: Any) -> bool:
    key = header_key(value)
    return key in {"否", "不认领", "NO", "N"}


def is_formula_error(value: Any) -> bool:
    return bool(re.fullmatch(r"#(?:REF|VALUE|N/A|NAME\?)!?", text_value(value) or "", flags=re.IGNORECASE))


def claim_value_key(value: Any) -> tuple[str, Any]:
    number = strict_positive_number(value)
    if number is not None:
        return ("positive_number", number)
    if isinstance(value, int | float) and not isinstance(value, bool):
        return ("number", float(value))
    return ("text", text_value(value) or "")


def distinct_raw_values(entries: list[dict[str, Any]]) -> list[Any]:
    values: list[Any] = []
    seen: set[tuple[str, Any]] = set()
    for entry in entries:
        key = claim_value_key(entry["raw_value"])
        if key not in seen:
            seen.add(key)
            values.append(entry["raw_value"])
    return values


def ordered_columns(headers: dict[str, Any]) -> list[str]:
    return sorted(headers, key=column_index_from_string)


def columns_between(headers: dict[str, Any], start: str, end: str | None) -> list[str]:
    start_index = column_index_from_string(start)
    end_index = column_index_from_string(end) if end else None
    return [
        column
        for column in ordered_columns(headers)
        if column_index_from_string(column) > start_index
        and (end_index is None or column_index_from_string(column) < end_index)
    ]


def first_header_column(
    headers: dict[str, Any],
    columns: list[str],
    predicate: Any,
) -> str | None:
    return next((column for column in columns if predicate(headers.get(column))), None)


def source_range(start: str, end: str) -> str:
    return start if start == end else f"{start}:{end}"


def header_key(value: Any) -> str:
    return re.sub(r"[\s_\-—–:：/\\()（）\[\]【】]+", "", text_value(value) or "").upper()


def cell_value(row: tuple[Any, ...], column: str) -> Any:
    index = column_index_from_string(column) - 1
    return row[index] if index < len(row) else None


def _is_repeated_header(value: str) -> bool:
    return header_key(value) in {"SKU", "子SKU", "SUBSKU"}


def main() -> None:
    parser = argparse.ArgumentParser(description="选品2历史来源解析与一次性迁移；默认不写库。")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--inspect-only", action="store_true")
    parser.add_argument("--upload-images", action="store_true")
    parser.add_argument("--rows-out")
    parser.add_argument("--max-rows-per-sheet", type=int, default=None)
    parser.add_argument("--apply-rows")
    parser.add_argument("--apply-dev", action="store_true")
    parser.add_argument("--imported-by", default="history_selection2_import")
    args = parser.parse_args()

    if args.apply_rows:
        from app.config import get_settings
        from app.db import SessionLocal

        payload = json.loads(Path(args.apply_rows).read_text(encoding="utf-8"))
        rows = payload["rows"]
        with SessionLocal() as db:
            if not args.apply_dev:
                report = plan_selection2_rows(db, rows)
                print(json.dumps({"mode": "dry-run", "rows": len(rows), **report}, ensure_ascii=False, indent=2))
                return
            settings = get_settings()
            if settings.app_env not in APPLY_ALLOWED_ENVS:
                raise SystemExit(f"apply blocked: app_env={settings.app_env}")
            report = apply_selection2_rows(
                db,
                rows,
                source_label=payload.get("source_file", "selection2-history"),
                imported_by=args.imported_by,
                source_sha256=payload.get("source_sha256"),
            )
            db.commit()
        print(json.dumps({"mode": "apply", **report}, ensure_ascii=False, indent=2))
        return

    if args.workbook is None:
        parser.error("--workbook is required unless --apply-rows is used")
    report = (
        inspect_selection2_workbook(args.workbook)
        if args.inspect_only
        else parse_selection2_workbook(
            args.workbook,
            max_rows_per_sheet=args.max_rows_per_sheet,
            upload_images=args.upload_images,
        )
    )
    if args.rows_out and not args.inspect_only:
        Path(args.rows_out).write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    summary = {key: value for key, value in report.items() if key != "rows"}
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
