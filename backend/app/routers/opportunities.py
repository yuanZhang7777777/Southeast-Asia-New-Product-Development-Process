import json
import re
import shutil
import uuid
from pathlib import Path
from urllib.parse import unquote
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Response, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session

from app import import_jobs, models, schemas, selection1_importer, selection2_importer, services
from app.auth import AuthContext, require_roles
from app.db import get_db
from app.excel_images import UPLOADED_SOURCES_ROOT
from app.site_codes import normalize_site_code
from app.workflow_status import OPPORTUNITY_DISABLED

router = APIRouter(prefix="/opportunities", tags=["opportunities"])
UPLOAD_ROOT = Path(__file__).resolve().parents[2] / ".private_uploads" / "source-workbooks"


@router.get("", response_model=list[schemas.OpportunityListRead])
def list_opportunities(
    limit: int = Query(300, ge=1, le=300),
    source_sheet: str | None = Query(None),
    business_period: str | None = Query(None),
    import_batch_id: str | None = Query(None),
    include_disabled: bool = Query(False),
    db: Session = Depends(get_db),
    auth: AuthContext | None = Depends(require_roles("operator", "manager")),
) -> list[models.NewProductOpportunity]:
    query = select(models.NewProductOpportunity)
    if not include_disabled or (auth and "super_admin" not in auth.role_keys):
        query = query.where(models.NewProductOpportunity.current_status != OPPORTUNITY_DISABLED)
    period = business_period or source_sheet
    if business_period:
        query = query.where(models.NewProductOpportunity.batch == business_period)
    elif source_sheet:
        query = query.where(or_(models.NewProductOpportunity.batch == period, models.NewProductOpportunity.source_sheet == source_sheet))
    if import_batch_id:
        query = query.where(models.NewProductOpportunity.import_batch_id == import_batch_id)
    if auth and auth.role_keys.isdisjoint({"manager", "super_admin"}):
        task_opportunity_ids = (
            select(models.FlowInstance.opportunity_id)
            .join(models.FlowTask, models.FlowTask.flow_instance_id == models.FlowInstance.id)
            .where(
                or_(
                    models.FlowTask.assignee_name == auth.operator_name,
                    models.FlowTask.assignee_user_id == auth.user.id,
                )
            )
        )
        claim_opportunity_ids = select(models.SalesClaimForecast.opportunity_id).where(
            models.SalesClaimForecast.salesperson_name == auth.operator_name
        )
        query = query.where(
            or_(
                models.NewProductOpportunity.id.in_(task_opportunity_ids),
                models.NewProductOpportunity.id.in_(claim_opportunity_ids),
                # 历史档案对运营开放只读可见（刊登工作台点击进详情需要；2026-07-27 用户要求）。
                models.NewProductOpportunity.current_status == "historical_archive",
                and_(
                    models.NewProductOpportunity.source_type == "selection2_caigen_claim_feedback",
                    models.NewProductOpportunity.claim_pool_open.is_(True),
                ),
            )
        )
    opportunities = list(db.scalars(query.order_by(models.NewProductOpportunity.created_at.desc()).limit(limit)))
    attach_latest_summaries(db, opportunities)
    return opportunities


@router.get("/export")
def export_opportunities(
    source_sheet: str | None = Query(None),
    business_period: str | None = Query(None),
    import_batch_id: str | None = Query(None),
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> Response:
    rows = services.list_source_snapshot_rows(db, source_sheet=source_sheet, business_period=business_period, import_batch_id=import_batch_id)
    content = services.build_source_snapshot_workbook(rows)
    file_name = period_file_name("source-opportunities", business_period or source_sheet, import_batch_id)
    filename = quote(file_name)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.get("/import-batches", response_model=list[schemas.ImportBatchSummary])
def list_import_batches(
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> list[models.ImportBatch]:
    return list(db.scalars(select(models.ImportBatch).order_by(models.ImportBatch.imported_at.desc())))


@router.post("/import", response_model=list[schemas.OpportunityRead])
def import_opportunities(
    payload: schemas.OpportunityImportRequest,
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> list[models.NewProductOpportunity]:
    items = [services.create_opportunity(db, item) for item in payload.items]
    db.commit()
    return items


# 注意：必须注册在 /export、/import-batches 等静态 GET 路由之后，避免路径参数抢占匹配。
@router.get("/{opportunity_id}", response_model=schemas.OpportunityRead)
def get_opportunity(
    opportunity_id: str,
    db: Session = Depends(get_db),
    auth: AuthContext | None = Depends(require_roles("operator", "manager")),
) -> models.NewProductOpportunity:
    opportunity = db.get(models.NewProductOpportunity, opportunity_id)
    if opportunity is None:
        raise HTTPException(status_code=404, detail="opportunity not found")
    if opportunity.current_status == OPPORTUNITY_DISABLED and auth and "super_admin" not in auth.role_keys:
        raise HTTPException(status_code=404, detail="opportunity not found")
    attach_latest_summaries(db, [opportunity], include_historical_claims=True)
    return opportunity


@router.patch("/{opportunity_id}", response_model=schemas.OpportunityRead)
def update_opportunity(
    opportunity_id: str,
    payload: schemas.OpportunityUpdateRequest,
    db: Session = Depends(get_db),
    auth: AuthContext | None = Depends(require_roles("operator", "manager")),
) -> models.NewProductOpportunity:
    target = db.get(models.NewProductOpportunity, opportunity_id)
    if target is None:
        raise HTTPException(status_code=404, detail="opportunity not found")
    if services.is_read_only_historical_opportunity(target):
        raise HTTPException(status_code=403, detail="historical source products are read-only")
    if auth and auth.role_keys.isdisjoint({"manager", "super_admin"}):
        if target.current_status != "historical_archive" and not services.operator_has_opportunity_access(
            db,
            opportunity_id,
            auth.operator_name,
            auth.user.id,
        ):
            raise HTTPException(status_code=403, detail="opportunity does not belong to current operator")
        locked_fields = {"main_sku", "sub_sku", "site", "country", "current_status"}
        if payload.model_fields_set.intersection(locked_fields):
            raise HTTPException(status_code=403, detail="SPU, SKU, site, country and status require manager access")
        locked_columns = {
            column
            for field, column in services.opportunity_field_columns(target).items()
            if field in locked_fields
        }
        submitted_columns = {column.strip().upper() for column in (payload.source_cells or {})}
        if locked_columns.intersection(submitted_columns):
            raise HTTPException(status_code=403, detail="locked source identity fields require manager access")
    try:
        item = services.update_opportunity(
            db,
            opportunity_id,
            payload,
            actor_name=auth.user.name if auth else None,
            actor_user_id=auth.user.id if auth else None,
        )
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.commit()
    db.refresh(item)
    attach_latest_summaries(db, [item], include_historical_claims=True)
    return item


@router.post("/{opportunity_id}/disable", response_model=schemas.MessageResponse)
def disable_opportunity(
    opportunity_id: str,
    payload: schemas.DisableRequest,
    db: Session = Depends(get_db),
    auth: AuthContext | None = Depends(require_roles("super_admin")),
) -> schemas.MessageResponse:
    opportunity = db.get(models.NewProductOpportunity, opportunity_id)
    if opportunity is None:
        raise HTTPException(status_code=404, detail="opportunity not found")
    count = services.set_opportunities_disabled(
        db,
        [opportunity],
        payload.disabled,
        payload.reason,
        actor_name=auth.user.name if auth else None,
        actor_user_id=auth.user.id if auth else None,
    )
    db.commit()
    return schemas.MessageResponse(message="disabled" if payload.disabled else "restored", id=str(count))


@router.post("/{opportunity_id}/disable-group", response_model=schemas.MessageResponse)
def disable_opportunity_group(
    opportunity_id: str,
    payload: schemas.DisableRequest,
    db: Session = Depends(get_db),
    auth: AuthContext | None = Depends(require_roles("super_admin")),
) -> schemas.MessageResponse:
    opportunity = db.get(models.NewProductOpportunity, opportunity_id)
    if opportunity is None:
        raise HTTPException(status_code=404, detail="opportunity not found")
    site_key = normalize_site_code(opportunity.site or opportunity.country) or ""
    query = select(models.NewProductOpportunity).where(
        models.NewProductOpportunity.source_type == opportunity.source_type,
        models.NewProductOpportunity.batch == opportunity.batch,
        models.NewProductOpportunity.main_sku == opportunity.main_sku,
    )
    opportunities = [
        item for item in db.scalars(query) if (normalize_site_code(item.site or item.country) or "") == site_key
    ]
    count = services.set_opportunities_disabled(
        db,
        opportunities,
        payload.disabled,
        payload.reason,
        actor_name=auth.user.name if auth else None,
        actor_user_id=auth.user.id if auth else None,
        scope="main_sku_group",
    )
    db.commit()
    return schemas.MessageResponse(message="disabled" if payload.disabled else "restored", id=str(count))


@router.post("/import-batches/{batch_id}/disable", response_model=schemas.MessageResponse)
def disable_import_batch(
    batch_id: str,
    payload: schemas.DisableRequest,
    db: Session = Depends(get_db),
    auth: AuthContext | None = Depends(require_roles("super_admin")),
) -> schemas.MessageResponse:
    batch = db.get(models.ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="import batch not found")
    count = services.set_import_batch_disabled(
        db,
        batch,
        payload.disabled,
        payload.reason,
        actor_name=auth.user.name if auth else None,
        actor_user_id=auth.user.id if auth else None,
    )
    db.commit()
    return schemas.MessageResponse(message="disabled" if payload.disabled else "restored", id=str(count))


@router.post("/excel-sheets/upload", response_model=schemas.ExcelSheetListResponse)
def list_excel_sheets_upload(
    file: UploadFile = File(...),
    _: object = Depends(require_roles("manager")),
) -> schemas.ExcelSheetListResponse:
    try:
        workbook = load_workbook(file.file, read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Excel 文件无法读取：{exc}") from exc
    try:
        sheets = list(workbook.sheetnames)
    finally:
        workbook.close()
    return schemas.ExcelSheetListResponse(sheets=sheets, default_sheet=sheets[0] if sheets else None)


@router.post("/import/selection1", response_model=schemas.Selection1ImportResponse)
def import_selection1(
    payload: schemas.Selection1ImportRequest,
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> schemas.Selection1ImportResponse:
    try:
        result = selection1_importer.import_selection1_workbook(db, payload)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.commit()
    return result


@router.post("/import/selection1/upload", response_model=schemas.Selection1ImportResponse)
def import_selection1_upload(
    source_sheet: str = Form(""),
    business_period: str = Form(""),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> schemas.Selection1ImportResponse:
    source_file = _save_upload(file)
    source_sheet = source_sheet.strip() or _first_sheet(source_file)
    return import_selection1(
        schemas.Selection1ImportRequest(
            source_file=str(source_file),
            source_sheet=source_sheet,
            business_period=business_period.strip() or None,
        ),
        db,
    )


@router.post("/import/selection1/upload-async", response_model=schemas.ImportJobRead, status_code=status.HTTP_202_ACCEPTED)
def import_selection1_upload_async(
    source_sheet: str = Form(""),
    business_period: str = Form(""),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> models.ImportJob:
    source_file = _save_upload(file)
    source_sheet = source_sheet.strip() or _first_sheet(source_file)
    job = import_jobs.create_import_job(
        kind="selection1",
        source_file=str(source_file),
        source_sheet=source_sheet,
        business_period=business_period.strip() or None,
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    return job


@router.post("/import/selection2", response_model=schemas.Selection2ImportResponse)
def import_selection2(
    payload: schemas.Selection2ImportRequest,
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> schemas.Selection2ImportResponse:
    try:
        result = selection2_importer.import_selection2_workbook(db, payload)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.commit()
    return result


@router.post("/import/selection2/upload", response_model=schemas.Selection2ImportResponse)
def import_selection2_upload(
    source_sheet: str = Form(""),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> schemas.Selection2ImportResponse:
    source_file = _save_upload(file)
    source_sheet = source_sheet.strip() or _first_sheet(source_file)
    return import_selection2(schemas.Selection2ImportRequest(source_file=str(source_file), source_sheet=source_sheet), db)


@router.post("/import/selection2/upload-async", response_model=schemas.ImportJobRead, status_code=status.HTTP_202_ACCEPTED)
def import_selection2_upload_async(
    source_sheet: str = Form(""),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> models.ImportJob:
    source_file = _save_upload(file)
    source_sheet = source_sheet.strip() or _first_sheet(source_file)
    job = import_jobs.create_import_job(kind="selection2", source_file=str(source_file), source_sheet=source_sheet)
    db.add(job)
    db.commit()
    db.refresh(job)
    return job


@router.get("/import-jobs/{job_id}", response_model=schemas.ImportJobRead)
def get_import_job(
    job_id: str,
    db: Session = Depends(get_db),
    _: object = Depends(require_roles("manager")),
) -> models.ImportJob:
    job = db.get(models.ImportJob, job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="import job not found")
    return job


def _save_upload(file: UploadFile) -> Path:
    clean_filename = _clean_upload_filename(file.filename or "upload.xlsx")
    suffix = Path(clean_filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm", ".xls"}:
        raise HTTPException(status_code=400, detail="only Excel files are supported")
    UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[<>:"/\\|?*]+', "_", Path(clean_filename).name).strip(" .")
    target = UPLOAD_ROOT / f"{uuid.uuid4().hex}-{safe_name}"
    with target.open("wb") as output:
        shutil.copyfileobj(file.file, output)
    return target


def _clean_upload_filename(filename: str) -> str:
    return unquote(filename).split("?", 1)[0].strip() or "upload.xlsx"


def _first_sheet(source_file: Path) -> str:
    workbook = load_workbook(source_file, read_only=True, data_only=True)
    try:
        if not workbook.sheetnames:
            raise HTTPException(status_code=400, detail="Excel 文件没有 sheet")
        return workbook.sheetnames[0]
    finally:
        workbook.close()


def period_file_name(prefix: str, source_sheet: str | None, import_batch_id: str | None) -> str:
    period = source_sheet or import_batch_id
    return f"{prefix}-{period}.xlsx" if period else f"{prefix}.xlsx"


def attach_latest_summaries(
    db: Session,
    opportunities: list[models.NewProductOpportunity],
    *,
    include_historical_claims: bool = False,
) -> None:
    opportunity_ids = [opportunity.id for opportunity in opportunities]
    latest_claims: dict[str, models.SalesClaimForecast] = {}
    archive_claims: dict[str, models.SalesClaimForecast] = {}
    latest_reviews: dict[str, models.ReviewRecord] = {}
    pending_review_claims: dict[str, list[dict[str, object]]] = {opportunity_id: [] for opportunity_id in opportunity_ids}
    historical_claims: dict[str, list[dict[str, object]]] = {opportunity_id: [] for opportunity_id in opportunity_ids}
    if opportunity_ids:
        platform_claims = list(db.scalars(
            select(models.SalesClaimForecast)
            .where(
                models.SalesClaimForecast.opportunity_id.in_(opportunity_ids),
                models.SalesClaimForecast.source_column == "platform",
            )
            .order_by(models.SalesClaimForecast.last_updated_at.desc(), models.SalesClaimForecast.created_at.desc())
        ))
        for item in platform_claims:
            latest_claims.setdefault(item.opportunity_id, item)
        standalone_history_ids = [
            opportunity.id
            for opportunity in opportunities
            if opportunity.source_type in {"history_selection2", "history_selection34"}
        ]
        if standalone_history_ids:
            for item in db.scalars(
                select(models.SalesClaimForecast)
                .where(
                    models.SalesClaimForecast.opportunity_id.in_(standalone_history_ids),
                    models.SalesClaimForecast.source_column != "platform",
                )
                .order_by(models.SalesClaimForecast.created_at.asc())
            ):
                current = archive_claims.get(item.opportunity_id)
                if current is None or (current.claim_result != "claim" and item.claim_result == "claim"):
                    archive_claims[item.opportunity_id] = item
        for item in platform_claims:
            if services.latest_review_for_claim(db, item.opportunity_id, item):
                continue
            pending_review_claims[item.opportunity_id].append(
                {
                    "claim_record_id": item.id,
                    "claim_result": item.claim_result,
                    "salesperson_name": item.salesperson_name,
                    "claim_daily_sales": item.claim_daily_sales,
                    "reject_reason": item.reject_reason,
                    "feedback_summary": item.feedback_summary,
                    "note": item.note,
                }
            )
        for item in db.scalars(
            select(models.ReviewRecord)
            .where(models.ReviewRecord.opportunity_id.in_(opportunity_ids))
            .order_by(models.ReviewRecord.created_at.desc())
        ):
            latest_reviews.setdefault(item.opportunity_id, item)
        if include_historical_claims:
            claims = list(
                db.scalars(
                    select(models.SalesClaimForecast)
                    .where(models.SalesClaimForecast.opportunity_id.in_(opportunity_ids))
                    .order_by(models.SalesClaimForecast.created_at.asc(), models.SalesClaimForecast.id.asc())
                )
            )
            historical = [item for item in claims if item.source_column != "platform"]
            reviews_by_claim_id: dict[str, models.ReviewRecord] = {}
            if historical:
                for review in db.scalars(
                    select(models.ReviewRecord)
                    .where(models.ReviewRecord.claim_record_id.in_([item.id for item in historical]))
                    .order_by(models.ReviewRecord.created_at.desc())
                ):
                    if review.claim_record_id:
                        reviews_by_claim_id.setdefault(review.claim_record_id, review)
            opportunities_by_id = {opportunity.id: opportunity for opportunity in opportunities}
            for claim in historical:
                opportunity = opportunities_by_id[claim.opportunity_id]
                review = reviews_by_claim_id.get(claim.id)
                metadata = _historical_claim_metadata(claim.note)
                historical_claims[claim.opportunity_id].append(
                    {
                        "id": claim.id,
                        "salesperson_name": claim.salesperson_name,
                        "claim_result": claim.claim_result,
                        "claim_daily_sales": claim.claim_daily_sales,
                        "reject_reason": claim.reject_reason,
                        "feedback_summary": claim.feedback_summary,
                        "source_column": claim.source_column,
                        "source_period": metadata.get("source_period") or opportunity.batch or opportunity.source_sheet,
                        "source_row": metadata.get("source_row") or opportunity.source_row,
                        "source_note": metadata.get("source_note"),
                        "evidence_images": metadata.get("evidence_images") or [],
                        "manager_review_status": review.review_status if review else None,
                        "manager_review_comment": review.review_comment if review else None,
                    }
                )
    for opportunity in opportunities:
        claim = latest_claims.get(opportunity.id) or archive_claims.get(opportunity.id)
        review = latest_reviews.get(opportunity.id)
        opportunity.latest_claim_record_id = claim.id if claim else None
        opportunity.latest_claim_result = claim.claim_result if claim else None
        opportunity.latest_claim_salesperson = claim.salesperson_name if claim else None
        opportunity.latest_claim_daily_sales = claim.claim_daily_sales if claim else None
        opportunity.latest_reject_reason = claim.reject_reason if claim else None
        opportunity.latest_feedback_summary = claim.feedback_summary if claim else None
        opportunity.latest_claim_note = claim.note if claim else None
        opportunity.latest_review_status = review.review_status if review else None
        opportunity.latest_review_comment = review.review_comment if review else None
        opportunity.pending_review_claims = pending_review_claims[opportunity.id]
        if include_historical_claims:
            opportunity.historical_claims = historical_claims[opportunity.id]


def _historical_claim_metadata(note: str | None) -> dict[str, object]:
    try:
        payload = json.loads(note or "{}")
    except (json.JSONDecodeError, TypeError):
        return {}
    if not isinstance(payload, dict):
        return {}
    source = payload.get("history_source")
    source = source if isinstance(source, dict) else {}
    source_payload = payload.get("source_payload")
    source_payload = source_payload if isinstance(source_payload, dict) else {}
    images = payload.get("evidence_images")
    return {
        "source_period": source.get("business_period"),
        "source_row": source.get("source_row"),
        "source_note": source_payload.get("note"),
        "evidence_images": images if isinstance(images, list) else [],
    }
