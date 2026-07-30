from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ERROR_VALUES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
KNOWN_OPTIONAL_HEADER_ALIASES = {"产品图片": ("产品图片",)}

REQUIRED_HEADER_ALIASES = {
    "站点/国家": ("站点", "国家"),
    "开发部门/部门": ("开发部门", "部门"),
    "开发员": ("开发员",),
    "一级类目": ("一级类目",),
    "关键词": ("关键词",),
    "主SKU名称": ("主SKU名称",),
    "主SKU": ("主SKU",),
    "子SKU名称": ("子SKU名称",),
    "子SKU": ("子SKU", "子sku"),
    "产品类型": ("产品类型", "产品类型 / 引流or绑定or利润", "引流or绑定or利润"),
    "开品理由": ("开品理由",),
}


def audit_source_workbooks(paths: Iterable[str | Path]) -> dict[str, Any]:
    files = []
    key_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for raw_path in paths:
        path = Path(raw_path)
        file_report: dict[str, Any] = {"file": str(path), "exists": path.exists(), "sheets": []}
        if path.exists():
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                for worksheet in workbook.worksheets:
                    sheet_report = audit_sheet(worksheet)
                    file_report["sheets"].append(sheet_report)
                    for row in sheet_report["keys"]:
                        key_rows[row["key"]].append({"file": path.name, "sheet": worksheet.title, "row": row["row"]})
            finally:
                workbook.close()
        files.append(file_report)
    return {
        "files": files,
        "duplicate_keys": [
            {"key": key, "rows": rows}
            for key, rows in sorted(key_rows.items())
            if len(rows) > 1
        ],
    }


def audit_sheet(worksheet: Any) -> dict[str, Any]:
    try:
        worksheet.reset_dimensions()
    except AttributeError:
        pass
    rows = list(worksheet.iter_rows(values_only=True))
    header_rows = rows[:2]
    headers_by_column = headers_from_rows(header_rows)
    validation = validate_required_headers([header for headers in headers_by_column.values() for header in headers])
    matched_columns = matched_required_columns(headers_by_column)
    missing_counts = {name: 0 for name in REQUIRED_HEADER_ALIASES}
    data_row_count = 0
    formula_error_count = 0
    keys = []
    for row_number, row in enumerate(rows[2:], start=3):
        if not any(text_value(value) for value in row):
            continue
        data_row_count += 1
        formula_error_count += sum(1 for value in row if is_formula_error(value))
        for required_name, column_index in matched_columns.items():
            if not text_value(cell(row, column_index)):
                missing_counts[required_name] += 1
        site = text_value(cell(row, matched_columns.get("站点/国家")))
        sub_sku = text_value(cell(row, matched_columns.get("子SKU")))
        if site and sub_sku:
            keys.append({"key": f"{site}|{sub_sku}", "row": row_number})
    return {
        "name": worksheet.title,
        "headers": {column: headers for column, headers in headers_by_column.items()},
        "missing_required_headers": validation["missing"],
        "unexpected_headers": unexpected_headers(headers_by_column),
        "data_row_count": data_row_count,
        "missing_required_counts": {key: count for key, count in missing_counts.items() if count},
        "formula_error_count": formula_error_count,
        "keys": keys,
    }


def validate_required_headers(headers: Iterable[str]) -> dict[str, Any]:
    normalized = {normalize_header(header): header for header in headers if normalize_header(header)}
    matched = {}
    missing = []
    for required_name, aliases in REQUIRED_HEADER_ALIASES.items():
        match = next((normalized[normalize_header(alias)] for alias in aliases if normalize_header(alias) in normalized), None)
        if match:
            matched[required_name] = match
        else:
            missing.append(required_name)
    return {"matched": matched, "missing": missing}


def unexpected_headers(headers_by_column: dict[str, list[str]]) -> list[dict[str, str]]:
    known = known_header_keys()
    result = []
    for column, headers in headers_by_column.items():
        for header in headers:
            if normalize_header(header) and normalize_header(header) not in known:
                result.append({"column": column, "header": header})
    return result


def known_header_keys() -> set[str]:
    return {
        normalize_header(alias)
        for aliases in (*REQUIRED_HEADER_ALIASES.values(), *KNOWN_OPTIONAL_HEADER_ALIASES.values())
        for alias in aliases
    }


def matched_required_columns(headers_by_column: dict[str, list[str]]) -> dict[str, int]:
    result = {}
    for required_name, aliases in REQUIRED_HEADER_ALIASES.items():
        alias_keys = {normalize_header(alias) for alias in aliases}
        for column, headers in headers_by_column.items():
            if any(normalize_header(header) in alias_keys for header in headers):
                result[required_name] = column_index(column)
                break
    return result


def headers_from_rows(rows: list[tuple[Any, ...]]) -> dict[str, list[str]]:
    max_len = max((len(row) for row in rows), default=0)
    headers = {}
    for index in range(max_len):
        top = text_value(cell(rows[0], index) if len(rows) > 0 else None)
        sub = text_value(cell(rows[1], index) if len(rows) > 1 else None)
        values = [value for value in (top, sub, f"{top} / {sub}" if top and sub and top != sub else None) if value]
        if values:
            headers[get_column_letter(index + 1)] = values
    return headers


def normalize_header(value: object) -> str:
    return "".join(text_value(value).replace("\xa0", " ").split()).upper() if text_value(value) else ""


def text_value(value: object) -> str | None:
    text = " ".join(str(value or "").replace("\xa0", " ").split())
    return text or None


def cell(row: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index < 0 or index >= len(row):
        return None
    return row[index]


def column_index(column: str) -> int:
    value = 0
    for character in column.upper():
        value = value * 26 + ord(character) - ord("A") + 1
    return value - 1


def is_formula_error(value: object) -> bool:
    text = text_value(value)
    return bool(text and text.upper() in ERROR_VALUES)


def main() -> None:
    parser = argparse.ArgumentParser(description="Read-only historical source workbook audit.")
    parser.add_argument("workbooks", nargs="+", type=Path)
    args = parser.parse_args()
    print(json.dumps(audit_source_workbooks(args.workbooks), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
