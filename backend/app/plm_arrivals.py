from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADERS = {
    "product_name": "\u5546\u54c1\u540d\u79f0",
    "sub_sku": "\u5b50SKU",
    "main_sku": "\u4e3bSKU",
    "salesperson_name": "\u9500\u552e\u5458",
    "bloc_name": "\u96c6\u56e2",
    "warehouse": "\u6d77\u5916\u4ed3",
    "country": "\u56fd\u5bb6",
    "latest_storage_time": "\u6700\u540e\u4e00\u6b21\u5165\u5e93\u65f6\u95f4",
    "first_listing_time": "\u9996\u6b21\u4e0a\u67b6\u65f6\u95f4",
    "available_quantity": "\u6d77\u5916\u4ed3\u53ef\u53d1",
    "real_stock_quantity": "\u771f\u4ed3\u5e93\u5b58",
    "daily_sales": "\u5355\u9500",
}
DEFAULT_BLOC_NAME = "\u96c6\u56e2\u516b\u90e8"
UNKNOWN_SALESPERSON = "\u672a\u5339\u914d\u9500\u552e\u5458"


def parse_plm_arrival_preview(workbook_path: str | Path, date_text: str, bloc_name: str = DEFAULT_BLOC_NAME) -> dict[str, Any]:
    target_date = date.fromisoformat(date_text)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        sheet_name = sheet.title
        header_values = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        indexes = {field: header_values.index(title) for field, title in HEADERS.items() if title in header_values}
        items: list[dict[str, Any]] = []
        for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = {field: row[index] if index < len(row) else None for field, index in indexes.items()}
            if _text(values.get("bloc_name")) != bloc_name:
                continue
            latest_date = _date_value(values.get("latest_storage_time"))
            if latest_date != target_date:
                continue
            first_date = _date_value(values.get("first_listing_time"))
            if first_date is None:
                arrival_type = "unknown"
            elif first_date == latest_date:
                arrival_type = "new_arrival"
            else:
                arrival_type = "restock"
            items.append(
                {
                    "source_sheet": sheet_name,
                    "source_row": source_row,
                    "arrival_type": arrival_type,
                    "product_name": _text(values.get("product_name")),
                    "salesperson_name": _text(values.get("salesperson_name")) or UNKNOWN_SALESPERSON,
                    "sub_sku": _text(values.get("sub_sku")),
                    "main_sku": _text(values.get("main_sku")),
                    "country": _text(values.get("country")),
                    "warehouse": _text(values.get("warehouse")),
                    "latest_storage_time": _text(values.get("latest_storage_time")),
                    "first_listing_time": _text(values.get("first_listing_time")),
                    "available_quantity": _number(values.get("available_quantity")),
                    "real_stock_quantity": _number(values.get("real_stock_quantity")),
                    "daily_sales": _number(values.get("daily_sales")),
                }
            )
    finally:
        workbook.close()

    return {
        "date": date_text,
        "bloc_name": bloc_name,
        "row_count": len(items),
        "new_arrival_count": _count(items, "new_arrival"),
        "restock_count": _count(items, "restock"),
        "unknown_count": _count(items, "unknown"),
        "by_salesperson": _summaries(items),
        "items": items,
    }


def _summaries(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_name: dict[str, dict[str, Any]] = {}
    for item in items:
        summary = by_name.setdefault(
            item["salesperson_name"],
            {
                "salesperson_name": item["salesperson_name"],
                "new_arrival_count": 0,
                "restock_count": 0,
                "unknown_count": 0,
                "total_count": 0,
            },
        )
        summary[f"{item['arrival_type']}_count"] += 1
        summary["total_count"] += 1
    return sorted(by_name.values(), key=lambda item: item["salesperson_name"])


def _count(items: list[dict[str, Any]], arrival_type: str) -> int:
    return sum(1 for item in items if item["arrival_type"] == arrival_type)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    text = text.replace("/", "-")
    for fmt, length in (("%Y-%m-%d %H:%M:%S", 19), ("%Y-%m-%d", 10)):
        try:
            return datetime.strptime(text[:length], fmt).date()
        except ValueError:
            continue
    return None


def _number(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None
