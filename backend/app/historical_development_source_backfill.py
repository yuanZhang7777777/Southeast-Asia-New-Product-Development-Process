from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models
from app.config import get_settings
from app.field_mapping import normalize_header
from app.selection1_importer import number_value, text_value

from app.historical_archive_import import APPLY_ALLOWED_ENVS
HISTORICAL_ARCHIVE_SOURCE_TYPE = "historical_market_monitor_archive"
MAX_SOURCE_COLUMN = 50
COUNTRY_ALIASES = {"菲律宾": "PH", "PH": "PH", "泰国": "TH", "TH": "TH", "越南": "VN", "VN": "VN"}


def normalize_country(value: Any) -> str | None:
    text = text_value(value)
    if not text:
        return None
    return COUNTRY_ALIASES.get(text.strip(), text.strip().upper())


def normalize_sku(value: Any) -> str | None:
    text = text_value(value)
    return text.strip().upper() if text else None

MAIN_ALIASES = {
    "country": ["站点", "国家"],
    "developer_department": ["开发部门", "部门"],
    "developer_name": ["开发员"],
    "category_level1": ["一级类目"],
    "keyword": ["关键词"],
    "main_sku_name": ["主SKU名称"],
    "main_sku": ["主SKU"],
    "sub_sku_name": ["子SKU名称"],
    "sub_sku": ["子SKU", "子sku"],
    "product_type": ["产品类型", "引流or绑定or利润"],
    "reason": ["开品理由"],
}

FIELD_ALIASES = {
    "product_spec": ["产品规格", "开发询价/产品规格"],
    "outer_package": ["产品外包装", "开发询价/产品外包装"],
    "final_package": ["末道包材", "开发询价/末道包材"],
    "supplier_url": ["供应商链接", "开发询价/供应商链接"],
    "supplier_name": ["供应商名称", "开发询价/供应商名称"],
    "tax_included_cost_rmb": ["商品成本-含税（元）", "开发询价/商品成本-含税（元）"],
    "add_on_shipping_total": ["预估单销*30的加购运费", "预估单销*30 的加购运费"],
    "add_on_shipping_unit": ["单子sku的加购运费分摊"],
    "package_weight_kg": ["包装重量(kg)"],
    "package_length_cm": ["产品包装后体积长(cm)"],
    "package_width_cm": ["产品包装后体积宽(cm)"],
    "package_height_cm": ["产品包装后体积高(cm)"],
    "package_volume": ["包装后体积"],
    "reference_daily_sales": ["参考单销"],
    "stable_price": ["稳定期定价", "稳定期定价（PHP）", "稳定期定价（PHP)", "参考定价（THB）", "稳定期参考定价（VND）"],
    "gross_profit_local": ["一次毛利额（PHP）", "一次毛利额（THB）", "一次毛利额（VND）"],
    "gross_profit_rmb": ["一次毛利额（人民币）"],
    "stable_margin": ["稳定期利润率", "一次毛利率"],
    "estimated_daily_sales": ["预估单销"],
    "promo_price": ["推广期定价"],
    "promo_margin": ["推广期利润率"],
    "stable_total_cost": ["稳定期总成本（PHP）（含头程+平台费+基础设施）", "稳定期总成本（THB）（含头程+平台费+基础设施）", "稳定期总成本（VND）（含头程+平台费+基础设施）"],
    "promo_total_cost": ["推广期总成本（PHP）（含头程+平台费+基础设施）", "推广期总成本（THB）（含头程+平台费+基础设施）", "推广期总成本（VND）（含头程+平台费+基础设施）"],
    "first_leg_fee_rmb": ["头程费用（元）"],
}

COMPETITOR_GROUPS = [
    ("lowest_price", ["最低价链接"], ["售价1", "售价1(PHP）"], ["月销1"]),
    ("most_orders", ["月销最高链接链接1", "月销最高链接", "mostorders链接", "most/orders链接"], ["售价2", "售价2(PHP）"], ["月销2"]),
    ("second_orders", ["月销次高链接链接2", "月销次高链接"], ["售价(PHP）"], ["月销"]),
    ("third_orders", ["月销第三高链接链接3", "月销第三高链接"], ["售价(PHP）"], ["月销"]),
    ("new_arrival", ["新晋链接"], ["售价3", "售价3(PHP）"], ["月销3"]),
]


@dataclass(frozen=True)
class SourceMatch:
    status: str
    source: dict[str, Any] | None
    candidates: list[dict[str, Any]]


def build_source_index(paths: list[Path]) -> dict[tuple[str | None, str | None, str | None], list[dict[str, Any]]]:
    index: dict[tuple[str | None, str | None, str | None], list[dict[str, Any]]] = defaultdict(list)
    for path in paths:
        for record in iter_development_source_records(path):
            index[(record["country"], record["main_sku"], record["sub_sku"])].append(record)
    return dict(index)


def iter_development_source_records(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    records: list[dict[str, Any]] = []
    try:
        for sheet_name in workbook.sheetnames:
            if "说明" in sheet_name:
                continue
            worksheet = workbook[sheet_name]
            worksheet.reset_dimensions()
            iterator = worksheet.iter_rows(max_col=MAX_SOURCE_COLUMN, values_only=True)
            first = next(iterator, None)
            second = next(iterator, None)
            if not first or not second:
                continue
            headers = source_headers(first, second)
            for source_row, row in enumerate(iterator, start=3):
                record = parse_source_row(path.name, sheet_name, source_row, row, headers)
                if record:
                    records.append(record)
    finally:
        workbook.close()
    return records


def source_headers(top_row: tuple[Any, ...], sub_row: tuple[Any, ...]) -> dict[str, list[str]]:
    headers: dict[str, list[str]] = {}
    group: str | None = None
    for index in range(1, MAX_SOURCE_COLUMN + 1):
        column = get_column_letter(index)
        top = text_value(top_row[index - 1] if index <= len(top_row) else None)
        sub = text_value(sub_row[index - 1] if index <= len(sub_row) else None)
        if top:
            group = top
        candidates = [value for value in (top, sub, f"{group} / {sub}" if group and sub else None) if value]
        if candidates:
            headers[column] = candidates
    return headers


def parse_source_row(source_file: str, sheet_name: str, source_row: int, row: tuple[Any, ...], headers: dict[str, list[str]]) -> dict[str, Any] | None:
    raw = {get_column_letter(index): row[index - 1] if index <= len(row) else None for index in range(1, MAX_SOURCE_COLUMN + 1)}
    main = {field: source_value(raw, headers, aliases) for field, aliases in MAIN_ALIASES.items()}
    country = normalize_country(main.get("country"))
    main_sku = normalize_sku(main.get("main_sku"))
    sub_sku = normalize_sku(main.get("sub_sku"))
    if not main_sku or not sub_sku or main_sku in {"主SKU", "MAINSKU"} or sub_sku in {"子SKU", "SUBSKU"}:
        return None
    fields = {field: source_value(raw, headers, aliases) for field, aliases in FIELD_ALIASES.items()}
    fields = {key: value for key, value in fields.items() if value not in (None, "")}
    competitors = competitor_rows(raw, headers)
    return {
        "source_reference": {"source_file": source_file, "source_sheet": sheet_name, "source_row": source_row},
        "country": country,
        "main_sku": main_sku,
        "sub_sku": sub_sku,
        "main_sku_name": text_value(main.get("main_sku_name")),
        "sub_sku_name": text_value(main.get("sub_sku_name")),
        "developer_department": text_value(main.get("developer_department")),
        "developer_name": text_value(main.get("developer_name")),
        "category_level1": text_value(main.get("category_level1")),
        "keyword": text_value(main.get("keyword")),
        "product_type": text_value(main.get("product_type")),
        "reason": text_value(main.get("reason")),
        "development_inquiry": typed_fields(fields),
        "competitors": competitors,
        "pricing_snapshot": pricing_fields(fields),
    }


def source_value(raw: dict[str, Any], headers: dict[str, list[str]], aliases: list[str]) -> Any:
    alias_keys = {normalize_header(alias) for alias in aliases}
    for column, names in headers.items():
        if any(normalize_header(name) in alias_keys for name in names):
            value = raw.get(column)
            if value not in (None, ""):
                return value
    return None


def competitor_rows(raw: dict[str, Any], headers: dict[str, list[str]]) -> list[dict[str, Any]]:
    rows = []
    for kind, url_aliases, price_aliases, sales_aliases in COMPETITOR_GROUPS:
        url = text_value(source_value(raw, headers, url_aliases))
        price = number_value(source_value(raw, headers, price_aliases))
        sales = number_value(source_value(raw, headers, sales_aliases))
        if url or price is not None or sales is not None:
            rows.append({"kind": kind, "url": url, "price": price, "monthly_sales": sales})
    return rows


def typed_fields(fields: dict[str, Any]) -> dict[str, Any]:
    numeric = {
        "tax_included_cost_rmb",
        "add_on_shipping_total",
        "add_on_shipping_unit",
        "package_weight_kg",
        "package_length_cm",
        "package_width_cm",
        "package_height_cm",
        "package_volume",
    }
    return {key: number_value(value) if key in numeric else text_value(value) for key, value in fields.items() if key not in pricing_keys()}


def pricing_fields(fields: dict[str, Any]) -> dict[str, Any]:
    return {key: number_value(value) for key, value in fields.items() if key in pricing_keys() and number_value(value) is not None}


def pricing_keys() -> set[str]:
    return {
        "reference_daily_sales",
        "stable_price",
        "gross_profit_local",
        "gross_profit_rmb",
        "stable_margin",
        "estimated_daily_sales",
        "promo_price",
        "promo_margin",
        "stable_total_cost",
        "promo_total_cost",
        "first_leg_fee_rmb",
    }


def match_source(index: dict[tuple[str | None, str | None, str | None], list[dict[str, Any]]], row: dict[str, Any]) -> SourceMatch:
    key = (normalize_country(row.get("country")), normalize_sku(row.get("main_sku")), normalize_sku(row.get("child_sku") or row.get("sub_sku")))
    candidates = index.get(key, [])
    if not candidates:
        return SourceMatch("missing", None, [])
    distinct = distinct_source_candidates(candidates)
    if len(distinct) == 1:
        return SourceMatch("unique", distinct[0], candidates)
    return SourceMatch("multiple", None, distinct)


def distinct_source_candidates(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[str] = set()
    result = []
    for item in candidates:
        signature = json.dumps(
            {
                "main": {key: item.get(key) for key in ("developer_department", "developer_name", "category_level1", "keyword", "product_type", "reason")},
                "development_inquiry": item.get("development_inquiry") or {},
                "pricing_snapshot": item.get("pricing_snapshot") or {},
                "competitors": item.get("competitors") or [],
            },
            ensure_ascii=False,
            sort_keys=True,
            default=str,
        )
        if signature in seen:
            continue
        seen.add(signature)
        result.append(item)
    return result


def build_backfill_rows(classification: dict[str, Any], source_index: dict[tuple[str | None, str | None, str | None], list[dict[str, Any]]]) -> list[dict[str, Any]]:
    rows = []
    for record in classification.get("records") or []:
        match = match_source(source_index, record)
        rows.append({"record": record, "match_status": match.status, "source": match.source, "candidates": match.candidates})
    return rows


def write_outputs(rows: list[dict[str, Any]], output_dir: Path, run_date: str) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    counts = Counter(row["match_status"] for row in rows)
    rows_json = output_dir / f"historical-development-source-backfill-rows-{run_date}.json"
    rows_json.write_text(json.dumps(rows, ensure_ascii=False, indent=2, sort_keys=True, default=str) + "\n", encoding="utf-8")
    summary = {
        "total_rows": len(rows),
        "unique_matches": counts.get("unique", 0),
        "multiple_matches": counts.get("multiple", 0),
        "missing_matches": counts.get("missing", 0),
        "rows_json": str(rows_json),
    }
    (output_dir / f"historical-development-source-backfill-summary-{run_date}.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    write_workbook([row for row in rows if row["match_status"] == "unique"], output_dir / f"历史开发表字段可补清单-{run_date}.xlsx")
    write_workbook([row for row in rows if row["match_status"] == "multiple"], output_dir / f"历史开发表字段多候选清单-{run_date}.xlsx")
    write_workbook([row for row in rows if row["match_status"] == "missing"], output_dir / f"历史开发表字段未命中清单-{run_date}.xlsx")
    return summary


def write_workbook(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "清单"
    worksheet.append(["状态", "主SKU", "子SKU", "国家", "销售员", "来源文件", "来源Sheet", "来源行", "开发询价", "参考单销", "参考售价/定价", "商品成本", "候选数"])
    for row in rows:
        record = row["record"]
        source = row.get("source") or (row.get("candidates") or [{}])[0]
        inquiry = source.get("development_inquiry") or {}
        pricing = source.get("pricing_snapshot") or {}
        ref = source.get("source_reference") or {}
        worksheet.append(
            [
                row["match_status"],
                record.get("main_sku"),
                record.get("child_sku"),
                record.get("country"),
                record.get("historical_salesperson"),
                ref.get("source_file"),
                ref.get("source_sheet"),
                ref.get("source_row"),
                inquiry.get("product_spec"),
                pricing.get("reference_daily_sales"),
                pricing.get("stable_price"),
                inquiry.get("tax_included_cost_rmb"),
                len(row.get("candidates") or []),
            ]
        )
    workbook.save(path)


def apply_backfill(db: Session, rows: list[dict[str, Any]]) -> dict[str, int]:
    updated = skipped = 0
    for row in rows:
        if row["match_status"] != "unique" or not row.get("source"):
            skipped += 1
            continue
        record = row["record"]
        opportunity = db.scalar(
            select(models.NewProductOpportunity).where(
                models.NewProductOpportunity.source_type == HISTORICAL_ARCHIVE_SOURCE_TYPE,
                models.NewProductOpportunity.source_file == (record.get("source_reference") or {}).get("source_file"),
                models.NewProductOpportunity.source_sheet == (record.get("source_reference") or {}).get("source_sheet"),
                models.NewProductOpportunity.source_row == (record.get("source_reference") or {}).get("source_row"),
                models.NewProductOpportunity.main_sku == record.get("main_sku"),
                models.NewProductOpportunity.sub_sku == record.get("child_sku"),
            )
        )
        if not opportunity:
            skipped += 1
            continue
        snapshot = dict(opportunity.snapshot or {})
        snapshot["development_source"] = row["source"]
        opportunity.snapshot = snapshot
        source_snapshot = db.scalar(
            select(models.SourceRecordSnapshot).where(
                models.SourceRecordSnapshot.opportunity_id == opportunity.id,
                models.SourceRecordSnapshot.source_file == opportunity.source_file,
                models.SourceRecordSnapshot.source_sheet == opportunity.source_sheet,
                models.SourceRecordSnapshot.source_row == opportunity.source_row,
            )
        )
        if source_snapshot:
            source_snapshot.payload = snapshot
        updated += 1
    return {"updated": updated, "skipped": skipped}


def ensure_apply_allowed() -> None:
    env = get_settings().app_env.strip().lower()
    if env not in APPLY_ALLOWED_ENVS:
        raise RuntimeError(f"refusing historical source backfill outside dev/local env: {env}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Backfill historical archive snapshots from development source workbooks.")
    parser.add_argument("--classification", type=Path)
    parser.add_argument("--source-workbook", type=Path, action="append")
    parser.add_argument("--backfill-json", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/historical_data/historical_development_source_backfill_20260726"))
    parser.add_argument("--run-date", default="20260726")
    parser.add_argument("--apply-dev", action="store_true")
    args = parser.parse_args()

    if args.backfill_json:
        rows = json.loads(args.backfill_json.read_text(encoding="utf-8"))
    else:
        if not args.classification or not args.source_workbook:
            raise SystemExit("--classification and --source-workbook are required unless --backfill-json is provided")
        classification = json.loads(args.classification.read_text(encoding="utf-8"))
        source_index = build_source_index(args.source_workbook)
        rows = build_backfill_rows(classification, source_index)
    summary = write_outputs(rows, args.output_dir, args.run_date)
    apply_result = None
    if args.apply_dev:
        ensure_apply_allowed()
        from app.db import SessionLocal

        with SessionLocal() as db:
            apply_result = apply_backfill(db, rows)
            db.commit()
    print(json.dumps({"review": summary, "apply": apply_result}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
