"""一次性选品3/4历史档案导入器；只恢复商品、来源行和历史认领事实。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models
from app.field_mapping import json_safe_value, text_value
from app.historical_archive_import import APPLY_ALLOWED_ENVS
from app.historical_selection2_import import is_explicit_rejection, is_zero_daily_sales, strict_positive_number
from app.historical_watchlist import sha256_file
from app.services import audit
from app.site_codes import normalize_site_code

SOURCE_TYPE = "history_selection34"
ARCHIVE_STATUS = "historical_archive"
AUDIT_ACTION = "history_selection34.imported"
AUTHORITATIVE_SHA256 = "3A85646D94C6CAFFA203854DB3A8F45F09D57B75D55DDDB05C0E98698F62DDA7"
AUTHORITATIVE_ROWS_SHA256 = "8D7B8C08DDEFAD6E88A9149286A75E527847F43EAC4A3078BB27B072238B3C36"
AUTHORITATIVE_SHEETS = (
    "小货老品-4月底",
    "开发高投入推荐-4月底",
    "销售自选0511期",
    "销售自选0518期",
    "销售自选0525期",
    "销售自选0601期",
    "销售自选0608期",
    "销售自选0613期",
    "销售自选0615期",
    "销售自选0622期",
    "销售自选0629期",
    "销售自选0706期",
    "直发热销转0615期",
    "直发热销转0630期",
)
AUTHORITATIVE_COUNTS = {
    "row_count": 604,
    "unique_opportunity_count": 596,
    "claim_count": 604,
    "reject_count": 0,
    "source_snapshot_count": 604,
}
LAST_COLUMN = 83
CLAIM_COLUMNS = {
    "reject_reason": "BZ",
    "salesperson": "CA",
    "claim_flag": "CB",
    "daily_sales": "CC",
    "feedback_summary": "CD",
    "note": "CE",
}


def normalize_sku(value: Any) -> str | None:
    text = text_value(value)
    return re.sub(r"\s+", "", text) if text else None


def parse_historical_selection34_row(
    row: tuple[Any, ...],
    *,
    headers: dict[str, Any],
    source_file: str,
    source_sheet: str,
    source_row: int,
) -> dict[str, Any] | None:
    sub_sku = normalize_sku(_cell(row, "K"))
    if not sub_sku or sub_sku.upper() in {"子SKU", "SUBSKU"}:
        return None
    main_sku = normalize_sku(_cell(row, "I")) or sub_sku
    site = normalize_site_code(_cell(row, "A"))
    source_payload = {
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
        "source_columns": CLAIM_COLUMNS,
        "reject_reason": json_safe_value(_cell(row, "BZ")),
        "salesperson": json_safe_value(_cell(row, "CA")),
        "claim_flag": json_safe_value(_cell(row, "CB")),
        "daily_sales": json_safe_value(_cell(row, "CC")),
        "feedback_summary": json_safe_value(_cell(row, "CD")),
        "note": json_safe_value(_cell(row, "CE")),
    }
    claims, rejected_sources = _claim_relations(source_payload)
    fields_by_cell = {
        column: {
            "header": json_safe_value(headers.get(column)),
            "value": json_safe_value(_cell(row, column)),
        }
        for column in _columns()
    }
    snapshot = {
        "archive_type": "historical_selection34",
        "business_period": source_sheet,
        "source_reference": {
            "source_file": source_file,
            "source_sheet": source_sheet,
            "source_row": source_row,
        },
        "fields_by_cell": fields_by_cell,
        "claims": claims,
        "rejected_sources": rejected_sources,
        "has_embedded_image": False,
    }
    image = text_value(_cell(row, "G"))
    image_url = image if image and re.match(r"^(?:https?://|/uploaded-sources/)", image, re.IGNORECASE) else None
    return {
        "source_type": SOURCE_TYPE,
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
        "batch": source_sheet,
        "current_status": ARCHIVE_STATUS,
        "claim_pool_open": False,
        "country": site,
        "site": site,
        "developer_department": text_value(_cell(row, "B")),
        "developer_name": text_value(_cell(row, "C")),
        "category_level1": text_value(_cell(row, "D")),
        "category_level2": text_value(_cell(row, "E")),
        "keyword": text_value(_cell(row, "F")),
        "image_url": image_url,
        "main_sku_name": text_value(_cell(row, "H")),
        "main_sku": main_sku,
        "sub_sku_name": text_value(_cell(row, "J")),
        "sub_sku": sub_sku,
        "product_type": text_value(_cell(row, "L")),
        "reason": text_value(_cell(row, "M")),
        "snapshot": snapshot,
        "claims": claims,
        "rejected_sources": rejected_sources,
    }


def _claim_relations(source_payload: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    salesperson = text_value(source_payload["salesperson"])
    if not salesperson:
        return [], []
    raw_daily_sales = source_payload["daily_sales"]
    daily_sales = strict_positive_number(raw_daily_sales)
    reject_reason = text_value(source_payload["reject_reason"])
    claim_flag = source_payload["claim_flag"]
    relation = {
        "salesperson_name": salesperson,
        "claim_result": "claim",
        "claim_daily_sales": daily_sales,
        "reject_reason": None,
        "feedback_summary": text_value(source_payload["feedback_summary"]),
        "source_column": f"BZ:CE:{source_payload['source_row']}",
        "source_payload": source_payload,
    }
    if daily_sales is not None and not reject_reason and not _explicit_rejection(claim_flag):
        return [relation], []
    relation["claim_result"] = "reject"
    relation["claim_daily_sales"] = None
    relation["reject_reason"] = _rejection_reason(
        raw_daily_sales,
        reject_reason=reject_reason,
        explicit_rejection=_explicit_rejection(claim_flag),
    )
    return [], [relation]


def _explicit_rejection(value: Any) -> bool:
    return is_explicit_rejection(value) or (text_value(value) or "").strip().upper() in {"拒绝", "FALSE"}


def _rejection_reason(
    raw_daily_sales: Any,
    *,
    reject_reason: str | None,
    explicit_rejection: bool,
) -> str:
    if reject_reason:
        return reject_reason
    if text_value(raw_daily_sales) is None:
        return "来源表未填写认领单销"
    if is_zero_daily_sales(raw_daily_sales):
        return "来源表认领单销为0"
    if strict_positive_number(raw_daily_sales) is None:
        return text_value(raw_daily_sales) or "来源标记不认领"
    if explicit_rejection:
        return "来源标记不认领"
    return "来源标记不认领"


def parse_selection34_workbook(
    path: Path,
    max_rows_per_sheet: int | None = None,
) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    sheet_reports: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for worksheet in workbook.worksheets:
            try:
                worksheet.reset_dimensions()
            except AttributeError:
                pass
            iterator = worksheet.iter_rows(min_row=1, max_col=LAST_COLUMN, values_only=True)
            header_row = next(iterator, ())
            headers = {
                get_column_letter(index): json_safe_value(value)
                for index, value in enumerate(header_row, start=1)
            }
            sheet_rows: list[dict[str, Any]] = []
            for source_row, values in enumerate(iterator, start=2):
                if max_rows_per_sheet is not None and len(sheet_rows) >= max_rows_per_sheet:
                    break
                parsed = parse_historical_selection34_row(
                    values,
                    headers=headers,
                    source_file=path.name,
                    source_sheet=worksheet.title,
                    source_row=source_row,
                )
                if parsed is not None:
                    sheet_rows.append(parsed)
            rows.extend(sheet_rows)
            sheet_reports.append(
                {
                    "sheet": worksheet.title,
                    "batch": worksheet.title,
                    "rows": len(sheet_rows),
                    "claims": sum(len(row["claims"]) for row in sheet_rows),
                    "rejects": sum(len(row["rejected_sources"]) for row in sheet_rows),
                    "unclaimed": sum(not row["claims"] and not row["rejected_sources"] for row in sheet_rows),
                }
            )
    finally:
        workbook.close()

    source_sha256 = sha256_file(path)
    groups = _group_rows(rows)
    return {
        "source_type": SOURCE_TYPE,
        "source_file": path.name,
        "source_sha256": source_sha256,
        "source_sha256_matches_authority": source_sha256.upper() == AUTHORITATIVE_SHA256,
        "rows_sha256": selection34_rows_sha256(rows),
        "sheets": sheet_reports,
        "row_count": len(rows),
        "unique_opportunity_count": len(groups),
        "multi_claim_identity_count": sum(
            sum(len(row["claims"]) + len(row["rejected_sources"]) for row in group) > 1
            for group in groups.values()
        ),
        "claim_count": sum(len(row["claims"]) for row in rows),
        "reject_count": sum(len(row["rejected_sources"]) for row in rows),
        "unclaimed_count": sum(not row["claims"] and not row["rejected_sources"] for row in rows),
        "source_snapshot_count": len(rows),
        "image_stats": {
            "embedded_images": _embedded_image_count(path),
            "rows_with_image": sum(bool(row["image_url"]) for row in rows),
        },
        "rows": rows,
    }


def selection34_row_identity(row: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        SOURCE_TYPE,
        (row.get("source_sheet") or "").strip().casefold(),
        normalize_site_code(row.get("country") or row.get("site")) or "",
        (normalize_sku(row.get("main_sku")) or "").upper(),
        (normalize_sku(row.get("sub_sku")) or "").upper(),
    )


def _group_rows(rows: list[dict[str, Any]]) -> dict[tuple[str, str, str, str, str], list[dict[str, Any]]]:
    groups: dict[tuple[str, str, str, str, str], list[dict[str, Any]]] = {}
    for row in rows:
        groups.setdefault(selection34_row_identity(row), []).append(row)
    return groups


def selection34_rows_sha256(rows: list[dict[str, Any]]) -> str:
    canonical = json.dumps(
        rows,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(canonical).hexdigest().upper()


def validate_selection34_apply_payload(payload: dict[str, Any]) -> None:
    rows = payload.get("rows")
    errors: list[str] = []
    source_sha256 = str(payload.get("source_sha256") or "").upper()
    if source_sha256 != AUTHORITATIVE_SHA256:
        errors.append(f"source_sha256 expected {AUTHORITATIVE_SHA256}, got {payload.get('source_sha256')!r}")

    sheets = payload.get("sheets")
    sheet_names = [item.get("sheet") for item in sheets] if isinstance(sheets, list) else None
    if sheet_names != list(AUTHORITATIVE_SHEETS):
        errors.append(f"sheets expected {list(AUTHORITATIVE_SHEETS)!r}, got {sheet_names!r}")

    for field, expected in AUTHORITATIVE_COUNTS.items():
        if payload.get(field) != expected:
            errors.append(f"{field} expected {expected}, got {payload.get(field)!r}")

    if not isinstance(rows, list):
        errors.append(f"rows expected list, got {type(rows).__name__}")
    else:
        rows_sha256 = selection34_rows_sha256(rows)
        if rows_sha256 != AUTHORITATIVE_ROWS_SHA256:
            errors.append(f"rows_sha256 expected {AUTHORITATIVE_ROWS_SHA256}, got {rows_sha256}")
        if str(payload.get("rows_sha256") or "").upper() != rows_sha256:
            errors.append(f"declared rows_sha256 does not match rows: {payload.get('rows_sha256')!r} != {rows_sha256}")
        row_sheet_names = list(dict.fromkeys(row.get("source_sheet") for row in rows))
        row_counts = {
            "row_count": len(rows),
            "unique_opportunity_count": len(_group_rows(rows)),
            "claim_count": sum(len(row.get("claims", [])) for row in rows),
            "reject_count": sum(len(row.get("rejected_sources", [])) for row in rows),
            "source_snapshot_count": len(
                {
                    (selection34_row_identity(row), row.get("source_sheet"), row.get("source_row"))
                    for row in rows
                }
            ),
        }
        if row_sheet_names != list(AUTHORITATIVE_SHEETS):
            errors.append(f"rows.sheets expected {list(AUTHORITATIVE_SHEETS)!r}, got {row_sheet_names!r}")
        for field, expected in AUTHORITATIVE_COUNTS.items():
            if row_counts[field] != expected:
                errors.append(f"rows.{field} expected {expected}, got {row_counts[field]!r}")

    if errors:
        raise ValueError("selection34 apply payload validation failed: " + "; ".join(errors))


def _find_opportunity(db: Session, row: dict[str, Any]) -> models.NewProductOpportunity | None:
    site = normalize_site_code(row.get("country") or row.get("site"))
    return db.scalar(
        select(models.NewProductOpportunity)
        .where(
            models.NewProductOpportunity.source_type == SOURCE_TYPE,
            models.NewProductOpportunity.source_sheet == row["source_sheet"],
            models.NewProductOpportunity.country == site,
            models.NewProductOpportunity.main_sku == normalize_sku(row.get("main_sku")),
            models.NewProductOpportunity.sub_sku == normalize_sku(row.get("sub_sku")),
        )
        .order_by(models.NewProductOpportunity.created_at.asc())
    )


def _find_snapshot(
    db: Session,
    opportunity_id: str,
    row: dict[str, Any],
) -> models.SourceRecordSnapshot | None:
    return db.scalar(
        select(models.SourceRecordSnapshot).where(
            models.SourceRecordSnapshot.opportunity_id == opportunity_id,
            models.SourceRecordSnapshot.source_sheet == row["source_sheet"],
            models.SourceRecordSnapshot.source_row == row["source_row"],
        )
    )


def _find_claim(
    db: Session,
    opportunity_id: str,
    relation: dict[str, Any],
) -> models.SalesClaimForecast | None:
    return db.scalar(
        select(models.SalesClaimForecast).where(
            models.SalesClaimForecast.opportunity_id == opportunity_id,
            models.SalesClaimForecast.claim_source == SOURCE_TYPE,
            models.SalesClaimForecast.source_column == relation["source_column"],
        )
    )


def plan_selection34_rows(db: Session, rows: list[dict[str, Any]]) -> dict[str, Any]:
    counts = {
        "would_create_opportunities": 0,
        "would_update_opportunities": 0,
        "would_create_source_snapshots": 0,
        "would_update_source_snapshots": 0,
        "would_create_claims": 0,
        "would_update_claims": 0,
    }
    for group in _group_rows(rows).values():
        opportunity = _find_opportunity(db, group[0])
        counts["would_update_opportunities" if opportunity else "would_create_opportunities"] += 1
        for row in group:
            snapshot = _find_snapshot(db, opportunity.id, row) if opportunity else None
            counts["would_update_source_snapshots" if snapshot else "would_create_source_snapshots"] += 1
            for relation in _relations(row):
                claim = _find_claim(db, opportunity.id, relation) if opportunity else None
                counts["would_update_claims" if claim else "would_create_claims"] += 1
    return counts


def apply_selection34_rows(
    db: Session,
    rows: list[dict[str, Any]],
    *,
    source_label: str,
    imported_by: str | None = None,
    source_sha256: str | None = None,
) -> dict[str, Any]:
    counts = {
        "created_opportunities": 0,
        "updated_opportunities": 0,
        "created_source_snapshots": 0,
        "updated_source_snapshots": 0,
        "created_claims": 0,
        "updated_claims": 0,
    }
    batch = models.ImportBatch(
        source_type=SOURCE_TYPE,
        source_file=source_label,
        source_sheet="all",
        business_period="历史选品3/4全Sheet",
        imported_by=imported_by,
        status="running",
    )
    db.add(batch)
    db.flush()

    for group in _group_rows(rows).values():
        first = group[0]
        opportunity = _find_opportunity(db, first)
        if opportunity is None:
            opportunity = models.NewProductOpportunity(
                source_type=SOURCE_TYPE,
                source_file=first.get("source_file"),
                source_sheet=first["source_sheet"],
                source_row=first.get("source_row"),
                batch=first["source_sheet"],
                country=normalize_site_code(first.get("country") or first.get("site")),
                site=normalize_site_code(first.get("country") or first.get("site")),
                developer_department=first.get("developer_department"),
                developer_name=first.get("developer_name"),
                category_level1=first.get("category_level1"),
                category_level2=first.get("category_level2"),
                keyword=first.get("keyword"),
                image_url=first.get("image_url"),
                main_sku_name=first.get("main_sku_name"),
                main_sku=normalize_sku(first.get("main_sku")),
                sub_sku_name=first.get("sub_sku_name"),
                sub_sku=normalize_sku(first.get("sub_sku")),
                product_type=first.get("product_type"),
                reason=first.get("reason"),
                current_status=ARCHIVE_STATUS,
                claim_pool_open=False,
                snapshot=first["snapshot"],
                import_batch_id=batch.id,
            )
            db.add(opportunity)
            db.flush()
            counts["created_opportunities"] += 1
        else:
            _update_opportunity(opportunity, first, batch.id)
            counts["updated_opportunities"] += 1

        for row in group:
            snapshot = _find_snapshot(db, opportunity.id, row)
            if snapshot is None:
                snapshot = models.SourceRecordSnapshot(
                    import_batch_id=batch.id,
                    opportunity_id=opportunity.id,
                    source_file=row.get("source_file"),
                    source_sheet=row["source_sheet"],
                    source_row=row.get("source_row"),
                    column_range="A:CE",
                    payload=row["snapshot"],
                )
                db.add(snapshot)
                counts["created_source_snapshots"] += 1
            else:
                snapshot.import_batch_id = batch.id
                snapshot.source_file = row.get("source_file")
                snapshot.column_range = "A:CE"
                snapshot.payload = row["snapshot"]
                counts["updated_source_snapshots"] += 1

            for relation in _relations(row):
                claim = _find_claim(db, opportunity.id, relation)
                if claim is None:
                    claim = models.SalesClaimForecast(
                        opportunity_id=opportunity.id,
                        claim_source=SOURCE_TYPE,
                        source_column=relation["source_column"],
                    )
                    db.add(claim)
                    counts["created_claims"] += 1
                else:
                    counts["updated_claims"] += 1
                _update_claim(claim, relation)

    batch.created_count = counts["created_opportunities"]
    batch.updated_count = counts["updated_opportunities"]
    batch.status = "completed"
    audit(
        db,
        AUDIT_ACTION,
        "new_product_opportunity",
        None,
        {
            "source_file": source_label,
            "source_sha256": source_sha256,
            "import_batch_id": batch.id,
            **counts,
        },
        imported_by,
    )
    db.flush()
    return {"import_batch_id": batch.id, **counts}


def _update_opportunity(
    opportunity: models.NewProductOpportunity,
    row: dict[str, Any],
    import_batch_id: str,
) -> None:
    opportunity.source_file = row.get("source_file")
    opportunity.source_row = row.get("source_row")
    opportunity.import_batch_id = import_batch_id
    opportunity.batch = row["source_sheet"]
    opportunity.country = normalize_site_code(row.get("country") or row.get("site"))
    opportunity.site = opportunity.country
    opportunity.current_status = ARCHIVE_STATUS
    opportunity.claim_pool_open = False
    opportunity.snapshot = row["snapshot"]
    for field in (
        "developer_department",
        "developer_name",
        "category_level1",
        "category_level2",
        "keyword",
        "image_url",
        "main_sku_name",
        "sub_sku_name",
        "product_type",
        "reason",
    ):
        if row.get(field) not in (None, ""):
            setattr(opportunity, field, row[field])


def _update_claim(claim: models.SalesClaimForecast, relation: dict[str, Any]) -> None:
    source = relation["source_payload"]
    claim.salesperson_name = relation["salesperson_name"]
    claim.claim_result = relation["claim_result"]
    claim.claim_daily_sales = relation["claim_daily_sales"]
    claim.reject_reason = relation["reject_reason"]
    claim.feedback_summary = relation["feedback_summary"]
    claim.task_id = None
    claim.downstream_status = None
    claim.note = json.dumps(
        {
            "history_source": {
                "source_file": source["source_file"],
                "source_sheet": source["source_sheet"],
                "source_row": source["source_row"],
                "business_period": source["source_sheet"],
                "claim_column": "BZ:CE",
            },
            "source_payload": source,
        },
        ensure_ascii=False,
        separators=(",", ":"),
        default=str,
    )


def _relations(row: dict[str, Any]) -> list[dict[str, Any]]:
    return [*row.get("claims", []), *row.get("rejected_sources", [])]


def _columns() -> list[str]:
    return [get_column_letter(index) for index in range(1, LAST_COLUMN + 1)]


def _cell(row: tuple[Any, ...], column: str) -> Any:
    index = _column_index(column) - 1
    return row[index] if index < len(row) else None


def _column_index(column: str) -> int:
    value = 0
    for character in column:
        value = value * 26 + ord(character.upper()) - ord("A") + 1
    return value


def _embedded_image_count(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        return sum(name.startswith("xl/media/") and not name.endswith("/") for name in archive.namelist())


def main() -> None:
    parser = argparse.ArgumentParser(description="选品3/4一次性历史只读导入器；默认仅 dry-run。")
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--rows-out", type=Path)
    parser.add_argument("--max-rows-per-sheet", type=int, default=None)
    parser.add_argument("--apply-rows", type=Path)
    parser.add_argument("--apply-dev", action="store_true")
    parser.add_argument("--imported-by", default="history_selection34_import")
    args = parser.parse_args()

    if args.apply_rows:
        payload = json.loads(args.apply_rows.read_text(encoding="utf-8"))
        try:
            validate_selection34_apply_payload(payload)
        except ValueError as exc:
            raise SystemExit(str(exc)) from exc
        rows = payload["rows"]
        from app.config import get_settings
        from app.db import SessionLocal

        with SessionLocal() as db:
            if not args.apply_dev:
                report = plan_selection34_rows(db, rows)
                print(json.dumps({"mode": "dry-run", "rows": len(rows), **report}, ensure_ascii=False, indent=2))
                return
            settings = get_settings()
            if settings.app_env not in APPLY_ALLOWED_ENVS:
                raise SystemExit(f"apply blocked: app_env={settings.app_env}")
            report = apply_selection34_rows(
                db,
                rows,
                source_label=payload.get("source_file", "selection34-history"),
                imported_by=args.imported_by,
                source_sha256=payload.get("source_sha256"),
            )
            db.commit()
        print(json.dumps({"mode": "apply", **report}, ensure_ascii=False, indent=2))
        return

    if args.workbook is None:
        parser.error("--workbook is required unless --apply-rows is used")
    report = parse_selection34_workbook(args.workbook, max_rows_per_sheet=args.max_rows_per_sheet)
    if args.rows_out:
        args.rows_out.write_text(json.dumps(report, ensure_ascii=False, default=str), encoding="utf-8")
    print(json.dumps({key: value for key, value in report.items() if key != "rows"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
