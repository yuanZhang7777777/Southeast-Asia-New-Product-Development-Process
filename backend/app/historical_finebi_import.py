from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models, services
from app.historical_monitoring_sources import (
    LISTING_HEADER_ROW,
    LISTING_SHEET,
    listing_record as parse_listing_row,
    text_value,
    valid_item,
)
from app.historical_watchlist import sha256_file
from app.item_finance_source import (
    country_from_shop,
    display_cell,
    item_id_text,
    merged_cell_values,
    normalize_main_sku,
)

SOURCE_TYPE = "history_finebi"
RECORD_SOURCE = "history_finebi"
BUSINESS_PERIOD = "历史归档"
EXCLUDED_PERIODS: tuple[str, ...] = ()
from app.historical_archive_import import APPLY_ALLOWED_ENVS
PERIOD_FILE_PATTERN = re.compile(r"(\d{4}-\d{4})\.xlsx$")
KEY_COLUMNS = ("ITEMID", "主SKU", "店铺")
METRIC_COLUMNS = ("总收入", "订单量", "一次毛利")


def period_window(period: str, year: int) -> tuple[date, date]:
    start_text, end_text = period.split("-")
    start = date(year, int(start_text[:2]), int(start_text[2:]))
    end = date(year, int(end_text[:2]), int(end_text[2:]))
    if end < start:
        end = date(year + 1, end.month, end.day)
    return start, end


def read_workbench(path: Path) -> tuple[dict[tuple[str, str], dict[str, Any]], list[dict[str, Any]]]:
    source_sha256 = sha256_file(path)
    merged: dict[tuple[str, str], dict[str, Any]] = {}
    warnings: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if LISTING_SHEET not in workbook.sheetnames:
            raise ValueError(f"{path.name} missing sheet {LISTING_SHEET}")
        worksheet = workbook[LISTING_SHEET]
        worksheet.reset_dimensions()
        rows = list(worksheet.iter_rows(values_only=True))
        for source_row, row in enumerate(rows[LISTING_HEADER_ROW:], start=LISTING_HEADER_ROW + 1):
            if not any(text_value(value) for value in row):
                continue
            record = parse_listing_row(path, source_sha256, source_row, row)
            shop = record["shop"]
            item = record["item"]
            if not shop or not valid_item(item):
                warnings.append({"source_row": source_row, "shop": shop, "item": item, "message": "invalid shop/item"})
                continue
            entry = merged.setdefault(
                (shop, item),
                {
                    "shop": shop,
                    "item": item,
                    "main_skus": set(),
                    "salespersons": set(),
                    "countries": set(),
                    "operation_dates": set(),
                    "strategy": {},
                    "weeks": {},
                    "summary": None,
                    "source_rows": [],
                },
            )
            if record["main_sku"]:
                entry["main_skus"].add(record["main_sku"])
            if record["salesperson"]:
                entry["salespersons"].add(record["salesperson"])
            if record["country"]:
                entry["countries"].add(record["country"])
            if record["operation_date"]:
                entry["operation_dates"].add(record["operation_date"])
            for key, value in record["strategy"].items():
                if value and not entry["strategy"].get(key):
                    entry["strategy"][key] = value
            for week in record["weeks"]:
                existing = entry["weeks"].get(week["week_number"])
                if existing is None or (week["has_metrics"] and not existing["has_metrics"]):
                    entry["weeks"][week["week_number"]] = week
                elif week["has_review"] and not existing["has_review"]:
                    for field in ("review_time", "positioning", "optimization_action", "has_review"):
                        existing[field] = week[field]
            if record["summary"] and not entry["summary"]:
                entry["summary"] = record["summary"]
            entry["source_rows"].append(source_row)
    finally:
        workbook.close()
    return merged, warnings


def read_finebi_periods(
    finebi_dir: Path,
    year: int,
    excluded: tuple[str, ...] = EXCLUDED_PERIODS,
    include_period: str | None = None,
) -> tuple[dict[tuple[str, str], dict[str, dict[str, Any]]], list[str], list[str]]:
    metrics: dict[tuple[str, str], dict[str, dict[str, Any]]] = defaultdict(dict)
    periods: list[str] = []
    skipped: list[str] = []
    for source_path in sorted(finebi_dir.glob("*.xlsx")):
        match = PERIOD_FILE_PATTERN.search(source_path.name)
        if not match:
            continue
        period = match.group(1)
        if period in excluded and period != include_period:
            skipped.append(period)
            continue
        periods.append(period)
        workbook = load_workbook(source_path, read_only=False, data_only=True, keep_links=False)
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows()
                header_index, header_cells = next(enumerate(rows, start=1), (0, ()))
                headers = {text_value(cell.value): index for index, cell in enumerate(header_cells) if text_value(cell.value)}
                if not all(column in headers for column in (*KEY_COLUMNS, *METRIC_COLUMNS)):
                    continue
                key_indexes = {headers[column] for column in KEY_COLUMNS}
                merged_values = merged_cell_values(worksheet, key_indexes)
                for source_row, row in enumerate(rows, start=header_index + 1):
                    item = item_id_text(display_cell(row, headers["ITEMID"], source_row, merged_values))
                    main_sku = normalize_main_sku(display_cell(row, headers["主SKU"], source_row, merged_values))
                    shop = text_value(display_cell(row, headers["店铺"], source_row, merged_values))
                    if not item or not shop:
                        continue
                    bucket = metrics[(shop, item)].setdefault(
                        period,
                        {"orders": 0.0, "revenue": 0.0, "gross_profit": 0.0, "row_count": 0, "main_skus": set(),
                         "source_file": source_path.name},
                    )
                    bucket["orders"] += _metric(row, headers["订单量"])
                    bucket["revenue"] += _metric(row, headers["总收入"])
                    bucket["gross_profit"] += _metric(row, headers["一次毛利"])
                    bucket["row_count"] += 1
                    if main_sku:
                        bucket["main_skus"].add(main_sku)
                break
        finally:
            workbook.close()
    return metrics, periods, skipped


def _metric(row: tuple[Any, ...], index: int) -> float:
    value = row[index].value if index < len(row) else None
    if isinstance(value, (int, float)):
        return float(value)
    return 0.0


def build_plan(
    db: Session,
    workbench: dict[tuple[str, str], dict[str, Any]],
    finebi: dict[tuple[str, str], dict[str, dict[str, Any]]],
    year: int,
) -> dict[str, Any]:
    known_main_skus = {
        value for value in db.scalars(select(models.NewProductOpportunity.main_sku).distinct()) if value
    }
    opportunity_index: dict[tuple[str, str | None], list[str]] = defaultdict(list)
    for opportunity_id, main_sku, country in db.execute(
        select(
            models.NewProductOpportunity.id,
            models.NewProductOpportunity.main_sku,
            models.NewProductOpportunity.country,
        )
    ):
        opportunity_index[(main_sku, country)].append(opportunity_id)

    existing_active = {
        (shop, item): listing_id
        for listing_id, shop, item in db.execute(
            select(
                models.ListingRecord.id,
                models.ListingRecord.shop,
                models.ListingRecord.item,
            ).where(models.ListingRecord.status == "active")
        )
    }

    scoped_keys = set(workbench)
    out_of_scope_finebi = 0
    for key, period_data in finebi.items():
        if key in scoped_keys:
            continue
        finebi_main_skus = set().union(*(bucket["main_skus"] for bucket in period_data.values()))
        if key in existing_active or finebi_main_skus & known_main_skus:
            scoped_keys.add(key)
        else:
            out_of_scope_finebi += 1

    items: list[dict[str, Any]] = []
    for key in sorted(scoped_keys):
        shop, item = key
        bench = workbench.get(key)
        period_data = finebi.get(key, {})
        main_skus = sorted(
            (bench["main_skus"] if bench else set())
            | set().union(*(bucket["main_skus"] for bucket in period_data.values()), set())
        )
        finebi_weeks = _latest_four_finebi_weeks(period_data, year)
        workbench_fallback_weeks = []
        review_weeks = []
        if bench:
            for week_number, week in sorted(bench["weeks"].items()):
                if week["has_review"]:
                    review_weeks.append(week)
        country = None
        if bench and bench["countries"]:
            country = sorted(bench["countries"])[0]
        country = country or country_from_shop(shop)
        items.append(
            {
                "shop": shop,
                "item": item,
                "main_skus": main_skus,
                "representative_main_sku": main_skus[0] if main_skus else None,
                "salesperson": sorted(bench["salespersons"])[0] if bench and len(bench["salespersons"]) == 1 else None,
                "country": country,
                "strategy": bench["strategy"] if bench else {},
                "summary": bench["summary"] if bench else None,
                "operation_dates": sorted(bench["operation_dates"]) if bench else [],
                "source_rows": bench["source_rows"] if bench else [],
                "finebi_weeks": finebi_weeks,
                "workbench_fallback_weeks": workbench_fallback_weeks,
                "review_weeks": review_weeks,
                "opportunity_by_main_sku": {
                    main_sku: opportunity_index.get((main_sku, country), [])
                    for main_sku in main_skus
                },
                "existing_listing": existing_active.get(key),
                "is_shared": len(main_skus) > 1,
            }
        )
    summary = {
        "scoped_items": len(items),
        "workbench_items": len(workbench),
        "finebi_items_total": len(finebi),
        "finebi_items_out_of_scope": out_of_scope_finebi,
        "items_without_main_sku": sum(1 for entry in items if not entry["main_skus"]),
        "listings_to_update": sum(1 for entry in items if entry["existing_listing"]),
        "listings_to_create": sum(1 for entry in items if not entry["existing_listing"] and entry["main_skus"]),
        "shared_items": sum(1 for entry in items if entry["is_shared"]),
        "bindings_planned": sum(len(entry["main_skus"]) for entry in items),
        "finebi_week_rows": sum(len(entry["finebi_weeks"]) for entry in items),
        "workbench_fallback_week_rows": sum(len(entry["workbench_fallback_weeks"]) for entry in items),
        "review_weeks": sum(len(entry["review_weeks"]) for entry in items),
        "opportunity_linked_items": sum(
            1 for entry in items if any(len(ids) == 1 for ids in entry["opportunity_by_main_sku"].values())
        ),
    }
    return {"summary": summary, "items": items}


def build_plan_from_candidate_json(db: Session, raw_plan: dict[str, Any]) -> dict[str, Any]:
    existing_active = {
        (shop, item): listing_id
        for listing_id, shop, item in db.execute(
            select(
                models.ListingRecord.id,
                models.ListingRecord.shop,
                models.ListingRecord.item,
            ).where(models.ListingRecord.status == "active")
        )
    }
    opportunity_by_period: dict[tuple[str, str | None, str | None], list[str]] = defaultdict(list)
    opportunity_by_country: dict[tuple[str, str | None], list[str]] = defaultdict(list)
    for opportunity_id, main_sku, country, batch in db.execute(
        select(
            models.NewProductOpportunity.id,
            models.NewProductOpportunity.main_sku,
            models.NewProductOpportunity.country,
            models.NewProductOpportunity.batch,
        )
    ):
        opportunity_by_period[(main_sku, country, batch)].append(opportunity_id)
        opportunity_by_country[(main_sku, country)].append(opportunity_id)

    entries: dict[tuple[str, str], dict[str, Any]] = {}
    for candidate in raw_plan.get("push_listing_candidates", []):
        if not candidate.get("can_enter_observation"):
            continue
        shop = candidate.get("shop")
        item = candidate.get("item")
        main_sku = normalize_main_sku(candidate.get("main_sku"))
        if not shop or not item or not main_sku:
            continue
        entry = entries.setdefault(
            (shop, item),
            {
                "shop": shop,
                "item": item,
                "main_skus": set(),
                "salespersons": set(),
                "countries": set(),
                "business_periods": set(),
                "source_rows": [],
                "strategy": {},
            },
        )
        entry["main_skus"].add(main_sku)
        if candidate.get("push_salesperson"):
            entry["salespersons"].add(candidate["push_salesperson"])
        if candidate.get("country_zh"):
            entry["countries"].add(candidate["country_zh"])
        if candidate.get("selected_platform_period"):
            entry["business_periods"].add(candidate["selected_platform_period"])
        source_row = (candidate.get("push_source") or {}).get("row")
        if source_row is not None:
            entry["source_rows"].append(source_row)
        for key, value in (candidate.get("strategy_raw") or {}).items():
            if value and not entry["strategy"].get(key):
                entry["strategy"][key] = value

    weeks_by_item: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in raw_plan.get("candidate_observation_periods", []):
        shop = row.get("shop")
        item = row.get("item")
        if not shop or not item:
            continue
        source = row.get("finebi_source") or {}
        weeks_by_item[(shop, item)].append(
            {
                "period": row["period_code"],
                "window": (date.fromisoformat(row["period_start"]), date.fromisoformat(row["period_end"])),
                "orders": row.get("orders") or 0,
                "revenue": row.get("revenue") or 0,
                "gross_profit": row.get("gross_profit") or 0,
                "row_count": source.get("source_row_count") or len(source.get("rows") or []),
                "source_file": source.get("file_name"),
                "main_skus": sorted(set(source.get("main_skus") or row.get("main_skus") or [row.get("main_sku")])),
                "binding_keys": sorted(set(row.get("binding_keys") or [row.get("binding_key")])),
                "source_rows": sorted(set(row.get("source_rows") or [])),
            }
        )

    items: list[dict[str, Any]] = []
    for key in sorted(entries):
        entry = entries[key]
        shop, item = key
        main_skus = sorted(entry["main_skus"])
        countries = sorted(entry["countries"])
        business_periods = sorted(entry["business_periods"])
        country = countries[0] if countries else country_from_shop(shop)
        business_period = business_periods[0] if len(business_periods) == 1 else BUSINESS_PERIOD
        opportunity_by_main_sku = {}
        for main_sku in main_skus:
            period_matches = opportunity_by_period.get((main_sku, country, business_period), [])
            opportunity_by_main_sku[main_sku] = (
                period_matches
                if len(period_matches) == 1
                else opportunity_by_country.get((main_sku, country), [])
            )
        finebi_weeks = sorted(weeks_by_item.get(key, []), key=lambda week: week["window"][0])
        items.append(
            {
                "shop": shop,
                "item": item,
                "main_skus": main_skus,
                "representative_main_sku": main_skus[0] if main_skus else None,
                "salesperson": sorted(entry["salespersons"])[0] if len(entry["salespersons"]) == 1 else None,
                "country": country,
                "business_period": business_period,
                "strategy": entry["strategy"],
                "summary": None,
                "operation_dates": [],
                "source_rows": sorted(set(entry["source_rows"])),
                "finebi_weeks": finebi_weeks,
                "workbench_fallback_weeks": [],
                "review_weeks": [],
                "opportunity_by_main_sku": opportunity_by_main_sku,
                "existing_listing": existing_active.get(key),
                "is_shared": len(main_skus) > 1,
            }
        )
    return {
        "summary": {
            "scoped_items": len(items),
            "workbench_items": len(items),
            "finebi_items_total": len(weeks_by_item),
            "finebi_items_out_of_scope": 0,
            "items_without_main_sku": sum(1 for entry in items if not entry["main_skus"]),
            "listings_to_update": sum(1 for entry in items if entry["existing_listing"]),
            "listings_to_create": sum(1 for entry in items if not entry["existing_listing"] and entry["main_skus"]),
            "shared_items": sum(1 for entry in items if entry["is_shared"]),
            "bindings_planned": sum(len(entry["main_skus"]) for entry in items),
            "finebi_week_rows": sum(len(entry["finebi_weeks"]) for entry in items),
            "workbench_fallback_week_rows": 0,
            "review_weeks": 0,
            "opportunity_linked_items": sum(
                1 for entry in items if any(len(ids) == 1 for ids in entry["opportunity_by_main_sku"].values())
            ),
        },
        "items": items,
    }



def _latest_four_finebi_weeks(period_data: dict[str, dict[str, Any]], year: int) -> list[dict[str, Any]]:
    selected = sorted(
        ((period, bucket) for period, bucket in period_data.items() if bucket["row_count"] > 0),
        key=lambda item: period_window(item[0], year)[0],
    )[-4:]
    return [
        {
            "period": period,
            "window": period_window(period, year),
            **{field: bucket[field] for field in ("orders", "revenue", "gross_profit", "row_count", "source_file")},
            "main_skus": sorted(bucket["main_skus"]),
        }
        for period, bucket in selected
    ]
def apply_plan(db: Session, plan: dict[str, Any], source_label: str, imported_by: str | None = None) -> dict[str, int]:
    batch = models.ImportBatch(
        source_type=SOURCE_TYPE,
        source_file=source_label,
        source_sheet=LISTING_SHEET,
        business_period=BUSINESS_PERIOD,
        imported_by=imported_by,
        status="running",
    )
    db.add(batch)
    db.flush()
    counts = {
        "listings_created": 0,
        "listings_reused": 0,
        "bindings_created": 0,
        "weeks_created": 0,
        "weeks_updated": 0,
        "items_skipped_no_sku": 0,
    }
    now = models.now_utc()
    for entry in plan["items"]:
        if not entry["main_skus"]:
            counts["items_skipped_no_sku"] += 1
            continue
        listing = db.scalars(
            select(models.ListingRecord).where(
                models.ListingRecord.shop == entry["shop"],
                models.ListingRecord.item == entry["item"],
                models.ListingRecord.status == "active",
            )
        ).one_or_none()
        if listing is None:
            first_window = entry["finebi_weeks"][0]["window"] if entry["finebi_weeks"] else None
            business_period = entry.get("business_period") or BUSINESS_PERIOD
            listing = models.ListingRecord(
                id=models.new_id(),
                source_group_key=f"history:finebi:{entry['shop']}:{entry['item']}",
                source_claim_ids=[],
                source_type=SOURCE_TYPE,
                business_period=business_period,
                country=entry["country"],
                site=entry["country"],
                main_sku=entry["representative_main_sku"],
                main_sku_name=None,
                salesperson_name=entry["salesperson"] or "",
                shop=entry["shop"],
                item=entry["item"],
                listing_strategy=_strategy_text(entry["strategy"]),
                first_period_start=first_window[0] if first_window else None,
                first_period_end=first_window[1] if first_window else None,
                representative_rule="single_binding",
            )
            db.add(listing)
            db.flush()
            counts["listings_created"] += 1
        else:
            counts["listings_reused"] += 1
        existing_bindings = {
            (binding.main_sku, binding.sub_sku)
            for binding in db.scalars(
                select(models.ListingSkuBinding).where(models.ListingSkuBinding.listing_record_id == listing.id)
            )
        }
        for main_sku in entry["main_skus"]:
            if (main_sku, None) in existing_bindings:
                continue
            opportunity_ids = entry["opportunity_by_main_sku"].get(main_sku, [])
            db.add(
                models.ListingSkuBinding(
                    id=models.new_id(),
                    listing_record_id=listing.id,
                    main_sku=main_sku,
                    sub_sku=None,
                    salesperson_name=entry["salesperson"],
                    opportunity_id=opportunity_ids[0] if len(opportunity_ids) == 1 else None,
                    binding_source=RECORD_SOURCE,
                )
            )
            counts["bindings_created"] += 1
        db.flush()
        services.recompute_listing_binding_state(db, listing)

        existing_weeks = {
            period.week_number: period
            for period in db.scalars(
                select(models.ItemObservationPeriod).where(
                    models.ItemObservationPeriod.listing_record_id == listing.id,
                    models.ItemObservationPeriod.record_source == RECORD_SOURCE,
                )
            )
        }
        week_number = 0
        for week in entry["finebi_weeks"]:
            week_number += 1
            _upsert_week(
                db,
                existing_weeks,
                listing,
                week_number,
                counts,
                period_start=week["window"][0],
                period_end=week["window"][1],
                orders=int(round(week["orders"])),
                revenue=week["revenue"],
                gross_profit=week["gross_profit"],
                metrics_origin="finebi_live",
                snapshot={
                    "period": week["period"],
                    "source_file": week["source_file"],
                    "row_count": week["row_count"],
                    "main_skus": week["main_skus"],
                    "binding_keys": week.get("binding_keys") or [],
                    "push_source_rows": week.get("source_rows") or [],
                },
                fetched_at=now,
            )
        for week in entry["workbench_fallback_weeks"]:
            _upsert_week(
                db,
                existing_weeks,
                listing,
                week["week_number"],
                counts,
                period_start=None,
                period_end=None,
                orders=int(week["orders"]) if week["orders"] is not None else None,
                revenue=week["revenue"],
                gross_profit=week["gross_profit"],
                metrics_origin="workbench_fallback",
                snapshot={"source": "刊登监控", "week_number": week["week_number"]},
                fetched_at=now,
            )
        for week in entry["review_weeks"]:
            period = existing_weeks.get(week["week_number"])
            if period is None:
                continue
            if week["positioning"] and not period.product_positioning:
                period.product_positioning = week["positioning"]
            if week["optimization_action"] and not period.optimization_action:
                period.optimization_action = week["optimization_action"]
        if entry["summary"]:
            week_four = existing_weeks.get(4) or existing_weeks.get(max(existing_weeks) if existing_weeks else 0)
            if week_four is not None and not week_four.four_week_summary:
                week_four.four_week_summary = entry["summary"]
    batch.created_count = counts["listings_created"]
    batch.updated_count = counts["listings_reused"]
    batch.skipped_count = counts["items_skipped_no_sku"]
    batch.status = "completed"
    services.audit(
        db,
        "history.finebi_imported",
        "listing_record",
        None,
        {**counts, "import_batch_id": batch.id},
        imported_by,
    )
    db.flush()
    return counts


def _upsert_week(
    db: Session,
    existing_weeks: dict[int, models.ItemObservationPeriod],
    listing: models.ListingRecord,
    week_number: int,
    counts: dict[str, int],
    *,
    period_start: date | None,
    period_end: date | None,
    orders: int | None,
    revenue: float | None,
    gross_profit: float | None,
    metrics_origin: str | None,
    snapshot: dict[str, Any],
    fetched_at,
) -> models.ItemObservationPeriod:
    period = existing_weeks.get(week_number)
    if period is None:
        period = models.ItemObservationPeriod(
            id=models.new_id(),
            listing_record_id=listing.id,
            week_number=week_number,
            period_start=period_start,
            period_end=period_end,
            status="completed",
            record_source=RECORD_SOURCE,
            metrics_origin=metrics_origin,
            order_count=orders,
            total_revenue=revenue,
            gross_profit_amount=gross_profit,
            source_snapshot=snapshot,
            metrics_fetched_at=fetched_at,
        )
        db.add(period)
        existing_weeks[week_number] = period
        counts["weeks_created"] += 1
        return period
    if metrics_origin is None:
        return period
    if period.metrics_origin == "finebi_live" and metrics_origin != "finebi_live":
        return period
    changed = False
    for field, value in (
        ("order_count", orders),
        ("total_revenue", revenue),
        ("gross_profit_amount", gross_profit),
    ):
        if value is not None and getattr(period, field) != value:
            setattr(period, field, value)
            changed = True
    if changed or period.metrics_origin != metrics_origin:
        period.metrics_origin = metrics_origin
        period.period_start = period_start
        period.period_end = period_end
        period.source_snapshot = snapshot
        period.metrics_fetched_at = fetched_at
        counts["weeks_updated"] += 1
    return period


def _strategy_text(strategy: dict[str, Any]) -> str:
    labels = {
        "optimization": "优化",
        "brush_order": "刷单",
        "advertising": "广告",
        "affiliate": "联盟",
        "campaign": "活动",
        "other": "其他",
    }
    parts = [f"{labels.get(key, key)}:{value}" for key, value in strategy.items() if value]
    return "；".join(parts)


def main() -> None:
    parser = argparse.ArgumentParser(description="历史 FineBI/刊登监控导入器（默认 dry-run）")
    parser.add_argument("--workbench", help="刊登监控7.25.xlsx 路径")
    parser.add_argument("--finebi-dir", help="finebi_live 目录路径")
    parser.add_argument("--plan-json", help="已审核的历史推品/FineBI JSON 计划；提供后不重新按源表匹配")
    parser.add_argument("--year", type=int, default=2026)
    parser.add_argument("--output-dir", default="outputs")
    parser.add_argument("--apply-dev", action="store_true", help="真正写入开发库（默认只 dry-run）")
    parser.add_argument("--imported-by", default="history_finebi_import")
    args = parser.parse_args()

    from app.config import get_settings
    from app.db import SessionLocal

    with SessionLocal() as db:
        if args.plan_json:
            raw_plan = json.loads(Path(args.plan_json).read_text(encoding="utf-8"))
            plan = build_plan_from_candidate_json(db, raw_plan)
            report = {
                "mode": "apply" if args.apply_dev else "dry-run",
                "source": "plan-json",
                "plan_json": args.plan_json,
                **plan["summary"],
            }
            source_label = Path(args.plan_json).name
        else:
            if not args.workbench or not args.finebi_dir:
                parser.error("--workbench and --finebi-dir are required unless --plan-json is provided")
            workbench, warnings = read_workbench(Path(args.workbench))
            finebi, periods, skipped = read_finebi_periods(Path(args.finebi_dir), args.year)
            plan = build_plan(db, workbench, finebi, args.year)
            report = {
                "mode": "apply" if args.apply_dev else "dry-run",
                "source": "workbench+finebi",
                "periods": periods,
                "excluded_periods": skipped,
                "workbench_warnings": warnings,
                **plan["summary"],
            }
            source_label = Path(args.workbench).name
        if args.apply_dev:
            settings = get_settings()
            if settings.app_env not in APPLY_ALLOWED_ENVS:
                raise SystemExit(f"apply blocked: app_env={settings.app_env} 不在允许环境 {sorted(APPLY_ALLOWED_ENVS)}")
            counts = apply_plan(db, plan, source_label=source_label, imported_by=args.imported_by)
            db.commit()
            report["apply_counts"] = counts
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / f"history-finebi-import-{'apply' if args.apply_dev else 'dryrun'}-{date.today().isoformat()}.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    print(f"report written to {report_path}")


if __name__ == "__main__":
    main()
