from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import models

HEADER_FIELDS = {
    "operator_name": "销售员",
    "role": "岗位",
    "operator_level": "运营分层",
    "business_type": "业务类型",
    "key_site": "重点站点",
    "key_category1": "重点品类1",
    "key_category2": "重点品类2",
}
DEFAULT_PERSONNEL_FILENAME = "集团八部销售员信息表.xlsx"


def resolve_personnel_config_file(source_file: str | Path | None = None) -> Path:
    if source_file:
        path = Path(source_file)
        if not path.is_absolute():
            path = Path.cwd() / path
        if not path.exists():
            raise FileNotFoundError(str(path))
        return path

    search_roots = [Path.cwd(), Path.cwd().parent, Path(__file__).resolve().parents[2]]
    for root in search_roots:
        path = root / DEFAULT_PERSONNEL_FILENAME
        if path.exists():
            return path
    raise FileNotFoundError(DEFAULT_PERSONNEL_FILENAME)


def import_personnel_config(db: Session, source_file: str | Path, source_sheet: str | None = None) -> dict[str, int]:
    workbook = load_workbook(resolve_personnel_config_file(source_file), data_only=True, read_only=True)
    try:
        worksheet = workbook[source_sheet] if source_sheet else workbook[workbook.sheetnames[0]]
        headers = [_text(cell.value) for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        column_indexes = {field: headers.index(title) for field, title in HEADER_FIELDS.items() if title in headers}

        created_count = 0
        updated_count = 0
        skipped_count = 0
        next_order = int(db.scalar(select(func.max(models.OperatorAssignmentProfile.display_order))) or 0) + 1
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            values = {field: _text(row[index]) for field, index in column_indexes.items() if index < len(row)}
            operator_name = values.get("operator_name")
            if not operator_name:
                skipped_count += 1
                continue

            profile = db.scalar(
                select(models.OperatorAssignmentProfile).where(
                    models.OperatorAssignmentProfile.operator_name == operator_name
                )
            )
            if profile is None:
                profile = models.OperatorAssignmentProfile(operator_name=operator_name, display_order=next_order)
                next_order += 1
                db.add(profile)
                created_count += 1
            else:
                updated_count += 1

            for field in HEADER_FIELDS:
                if field != "operator_name":
                    setattr(profile, field, values.get(field))
            profile.enabled = True
            upsert_operator_role_mapping(db, operator_name)

        return {"created_count": created_count, "updated_count": updated_count, "skipped_count": skipped_count}
    finally:
        workbook.close()


def upsert_operator_role_mapping(db: Session, operator_name: str) -> None:
    mapping = db.scalar(
        select(models.RoleMapping).where(
            models.RoleMapping.name == operator_name,
            models.RoleMapping.role.in_(("operator", "sales")),
        )
    )
    if mapping is None:
        db.add(models.RoleMapping(name=operator_name, role="operator", enabled=True))
        return
    mapping.role = "operator"
    mapping.enabled = True


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
