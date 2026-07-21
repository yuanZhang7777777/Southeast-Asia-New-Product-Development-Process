from __future__ import annotations

import json
import math
import time
from io import BytesIO
from typing import Any, Mapping
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlsplit, urlunsplit
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from app.config import Settings


PRODUCT_DIMENSIONS = (
    "产品体积长(cm)",
    "产品体积宽(cm)",
    "产品体积高(cm)",
)
PACKAGING_DIMENSIONS = (
    "产品包装后体积长(cm)",
    "产品包装后体积宽(cm)",
    "产品包装后体积高(cm)",
)
REQUIRED_HEADERS = ("sku", "主SKU", *PRODUCT_DIMENSIONS, *PACKAGING_DIMENSIONS)
PORTION_FIELDS = (
    "sku,mainsku,pretendlength,pretendwidth,pretendlheight,"
    "actualLength,actualWidth,actualHeight"
)
POLL_ATTEMPTS = 30
POLL_INTERVAL_SECONDS = 2


def calculate_unit_volume(row: Mapping[str, object]) -> float | None:
    packaging = _dimensions(row, PACKAGING_DIMENSIONS)
    product = _dimensions(row, PRODUCT_DIMENSIONS)
    selected = packaging or product
    if selected is None:
        return None
    return round(selected[0] * selected[1] * selected[2] / 1_000_000, 12)


def parse_product_list_workbook(content: bytes, requested_skus: list[str]) -> dict[str, float | None]:
    result = dict.fromkeys(requested_skus)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = tuple(next(rows, ()))
        if not set(REQUIRED_HEADERS).issubset(headers):
            raise ValueError("ERP product workbook has unexpected headers")
        for values in rows:
            row = dict(zip(headers, values))
            sku = str(row.get("sku") or "").strip()
            if sku in result:
                result[sku] = calculate_unit_volume(row)
        return result
    finally:
        workbook.close()


def fetch_product_volumes(skus: list[str], settings: Settings) -> dict[str, float | None]:
    fallback = dict.fromkeys(skus)
    if not skus or not all(
        str(value).strip()
        for value in (
            settings.erp_login_url,
            settings.erp_product_list_url,
            settings.erp_download_list_url,
            settings.erp_username,
            settings.erp_password,
        )
    ):
        return fallback

    try:
        login = _post_json(
            settings.erp_login_url,
            {"username": settings.erp_username, "password": settings.erp_password},
        )
        _assert_business_ok(login, "ERP login")
        token = _access_token(login)
        if not token:
            return fallback

        headers = {"Authorization": token}
        list_body = {"pageNum": 1, "pageSize": 20}
        before = _post_json(settings.erp_download_list_url, list_body, headers)
        _assert_business_ok(before, "ERP download list")
        existing_ids = {str(row["id"]) for row in _response_rows(before) if row.get("id") is not None}

        created = _post_json(
            settings.erp_product_list_url,
            {"skuList": skus, "isfile": "0", "portionFieldSet": PORTION_FIELDS},
            headers,
        )
        _assert_business_ok(created, "ERP product list creation")

        for attempt in range(POLL_ATTEMPTS):
            listing = _post_json(settings.erp_download_list_url, list_body, headers)
            _assert_business_ok(listing, "ERP download list")
            ready = _new_product_list_export(_response_rows(listing), existing_ids)
            if ready:
                return parse_product_list_workbook(_download_bytes(str(ready["downloadUrl"])), skus)
            if attempt + 1 < POLL_ATTEMPTS:
                time.sleep(POLL_INTERVAL_SECONDS)
    except Exception:
        return fallback
    return fallback


def _dimensions(
    row: Mapping[str, object], names: tuple[str, str, str]
) -> tuple[float, float, float] | None:
    values = tuple(_positive_number(row.get(name)) for name in names)
    if any(value is None for value in values):
        return None
    return values  # type: ignore[return-value]


def _positive_number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) and number > 0 else None


def _new_product_list_export(
    rows: list[dict[str, Any]], existing_ids: set[str]
) -> dict[str, Any] | None:
    for row in rows:
        row_id = row.get("id")
        if (
            row_id is not None
            and str(row_id) not in existing_ids
            and row.get("source") == "productList"
            and row.get("status") == "success"
            and row.get("downloadUrl")
        ):
            return row
    return None


def _post_json(url: str, body: dict[str, Any], headers: dict[str, str] | None = None) -> Any:
    request = Request(
        url,
        data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
        headers={"content-type": "application/json", **(headers or {})},
        method="POST",
    )
    try:
        with urlopen(request, timeout=30) as response:
            return json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        raise RuntimeError(f"ERP request failed with HTTP {exc.code}") from None
    except (URLError, TimeoutError, UnicodeError, json.JSONDecodeError):
        raise RuntimeError("ERP request failed") from None


def _download_bytes(url: str) -> bytes:
    parts = urlsplit(url)
    encoded_url = urlunsplit(
        (
            parts.scheme,
            parts.netloc,
            quote(parts.path, safe="/%"),
            quote(parts.query, safe="=&%+/:;,@$"),
            parts.fragment,
        )
    )
    try:
        with urlopen(encoded_url, timeout=60) as response:
            return response.read()
    except HTTPError as exc:
        raise RuntimeError(f"ERP file download failed with HTTP {exc.code}") from None
    except (URLError, TimeoutError):
        raise RuntimeError("ERP file download failed") from None


def _access_token(response: Any) -> str:
    if not isinstance(response, dict):
        return ""
    data = response.get("data") if isinstance(response.get("data"), dict) else {}
    return str(data.get("accessToken") or response.get("accessToken") or "")


def _assert_business_ok(response: Any, label: str) -> None:
    if not isinstance(response, dict):
        raise RuntimeError(f"{label} failed")
    try:
        code = int(response.get("code"))
    except (TypeError, ValueError):
        code = None
    if response.get("success") is False or (code is not None and code >= 400):
        raise RuntimeError(f"{label} failed")


def _response_rows(response: Any) -> list[dict[str, Any]]:
    if not isinstance(response, dict):
        return []
    data = response.get("data")
    if isinstance(data, list):
        return [row for row in data if isinstance(row, dict)]
    if not isinstance(data, dict):
        return []
    for key in ("list", "records", "rows"):
        if isinstance(data.get(key), list):
            return [row for row in data[key] if isinstance(row, dict)]
    return []
