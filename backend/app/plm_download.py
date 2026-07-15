from __future__ import annotations

import json
import os
import re
import time
from datetime import date, datetime, time as wall_time, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlsplit, urlunsplit
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


BEIJING = ZoneInfo("Asia/Shanghai")
DOWNLOAD_CENTER_PATH = "/api/hz-inventory/minioDownloadFile/selectPage"
EXPORT_PATH = "/api/hz-inventory/overseas/details/exportSummaryExcel"
LOGIN_PATH = "/api/system/login"
OPEN_LOGIN_PATH = "/system/innerOpen/login"
OPEN_LIST_DETAIL_PATH = "/open/inventory/overseas/listDetail"
POLL_ATTEMPTS = 60
POLL_INTERVAL_SECONDS = 10


def beijing_window_ms(date_text: str) -> tuple[int, int]:
    day = date.fromisoformat(date_text)
    start = int(datetime.combine(day, wall_time.min, tzinfo=BEIJING).timestamp() * 1000)
    return start, start + 86_400_000 - 1_000


def beijing_window_text(date_text: str) -> tuple[str, str]:
    day = date.fromisoformat(date_text)
    return f"{day.isoformat()} 00:00:00", f"{day.isoformat()} 23:59:59"


def previous_beijing_date(now: datetime | None = None) -> str:
    current = now or datetime.now(BEIJING)
    if current.tzinfo is None:
        current = current.replace(tzinfo=BEIJING)
    return (current.astimezone(BEIJING).date() - timedelta(days=1)).isoformat()


def select_new_export(
    rows: list[dict[str, Any]], existing_ids: set[str], created_after_ms: int
) -> dict[str, Any] | None:
    candidates: list[tuple[int, dict[str, Any]]] = []
    for row in rows:
        match = re.search(r"(\d{13})", str(row.get("fileName") or ""))
        stamp = int(match.group(1)) if match else 0
        if (
            row.get("id")
            and str(row["id"]) not in existing_ids
            and str(row.get("status")) == "2"
            and row.get("downloadUrl")
            and stamp >= created_after_ms
        ):
            candidates.append((stamp, row))
    return max(candidates, key=lambda item: item[0])[1] if candidates else None


def download_plm_export(
    date_text: str,
    *,
    base_url: str,
    username: str,
    password: str,
    bloc_name: str,
    cache_dir: str | Path,
    force: bool = False,
) -> Path:
    cache_path = Path(cache_dir) / f"plm-{date_text}.xlsx"
    if not force and _valid_workbook(cache_path):
        return cache_path

    required = {"PLM_BASE_URL": base_url, "PLM_USERNAME": username, "PLM_PASSWORD": password, "PLM_BLOC_NAME": bloc_name}
    missing = [name for name, value in required.items() if not str(value).strip()]
    if missing:
        raise ValueError(f"missing PLM settings: {', '.join(missing)}")

    base = base_url.rstrip("/")
    login = _post_json(f"{base}{LOGIN_PATH}", {"username": username, "password": password})
    _assert_business_ok(login, "PLM login")
    token = _first_token(login)
    if not token:
        raise RuntimeError("PLM login returned no token")

    headers = {"authorization": token}
    list_body = {"downloadType": "1", "pageNum": 1, "pageSize": 10, "orderBy": ""}
    before = _post_json(f"{base}{DOWNLOAD_CENTER_PATH}", list_body, headers)
    _assert_business_ok(before, "PLM download center query")
    existing_ids = {
        str(row["id"])
        for row in _response_rows(before)
        if row.get("id")
    }

    start, end = beijing_window_ms(date_text)
    created = _post_json(
        f"{base}{EXPORT_PATH}",
        {
            "skuList": [],
            "mskuList": [],
            "bindSkuList": [],
            "kcsl": None,
            "kcslStart": None,
            "kcslEnd": None,
            "sort": None,
            "pageNum": 1,
            "pageSize": 10,
            "blocNameList": [bloc_name],
            "queryType": "1",
            "bindType": 1,
            "mergeSaleName": 0,
            "hideZeroData": 0,
            "latestStorageTimeStart": start,
            "latestStorageTimeEnd": end,
        },
        headers,
    )
    _assert_business_ok(created, "PLM export creation")
    task_now_ms = _integer(created.get("now")) if isinstance(created, dict) else None
    created_after_ms = (task_now_ms or int(time.time() * 1000)) - 60_000

    for attempt in range(POLL_ATTEMPTS):
        listing = _post_json(f"{base}{DOWNLOAD_CENTER_PATH}", list_body, headers)
        _assert_business_ok(listing, "PLM download center query")
        ready = select_new_export(_response_rows(listing), existing_ids, created_after_ms)
        if ready:
            content = _download_bytes(str(ready["downloadUrl"]))
            _validate_workbook_bytes(content)
            cache_path.parent.mkdir(parents=True, exist_ok=True)
            partial_path = cache_path.with_name(f"{cache_path.name}.part")
            partial_path.write_bytes(content)
            os.replace(partial_path, cache_path)
            return cache_path
        if attempt + 1 < POLL_ATTEMPTS:
            time.sleep(POLL_INTERVAL_SECONDS)
    raise TimeoutError("PLM export was not ready within 10 minutes")


def download_plm_open_inventory_detail(
    date_text: str,
    *,
    base_url: str,
    username: str,
    password: str,
    bloc_name: str,
    cache_dir: str | Path,
    page_size: int = 300,
    force: bool = False,
) -> Path:
    cache_path = Path(cache_dir) / f"plm-open-{date_text}.json"
    if not force and _valid_json_snapshot(cache_path):
        return cache_path

    required = {"PLM_OPEN_BASE_URL": base_url, "PLM_USERNAME": username, "PLM_PASSWORD": password, "PLM_BLOC_NAME": bloc_name}
    missing = [name for name, value in required.items() if not str(value).strip()]
    if missing:
        raise ValueError(f"missing PLM open settings: {', '.join(missing)}")

    base = base_url.rstrip("/")
    login = _post_json(f"{base}{OPEN_LOGIN_PATH}", {"username": username, "password": password})
    _assert_business_ok(login, "PLM open login")
    token = _first_token(login)
    if not token:
        raise RuntimeError("PLM open login returned no token")

    start, end = beijing_window_text(date_text)
    rows: list[dict[str, Any]] = []
    page_num = 1
    while True:
        response = _post_json(
            f"{base}{OPEN_LIST_DETAIL_PATH}",
            {
                "pageNum": page_num,
                "pageSize": page_size,
                "mergeSaleName": 0,
                "blocNameList": [bloc_name],
                "latestStorageTimeStart": start,
                "latestStorageTimeEnd": end,
            },
            {"Authorization": token},
        )
        _assert_business_ok(response, "PLM open inventory detail query")
        batch = _response_rows(response)
        rows.extend(batch)
        if not _has_next(response, page_num, page_size, len(rows), len(batch)):
            break
        page_num += 1

    cache_path.parent.mkdir(parents=True, exist_ok=True)
    partial_path = cache_path.with_name(f"{cache_path.name}.part")
    partial_path.write_text(
        json.dumps({"date": date_text, "bloc_name": bloc_name, "row_count": len(rows), "rows": rows}, ensure_ascii=False),
        encoding="utf-8",
    )
    os.replace(partial_path, cache_path)
    return cache_path


def _post_json(url: str, body: dict[str, Any], headers: dict[str, str] | None = None) -> Any:
    request = Request(
        url,
        data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
        headers={"content-type": "application/json", **(headers or {})},
        method="POST",
    )
    try:
        with urlopen(request, timeout=30) as response:
            return json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        raise RuntimeError(f"PLM request failed with HTTP {exc.code}") from exc
    except (URLError, TimeoutError) as exc:
        raise RuntimeError("PLM request failed") from exc


def _download_bytes(url: str) -> bytes:
    try:
        parts = urlsplit(url)
        encoded_url = urlunsplit(
            (parts.scheme, parts.netloc, quote(parts.path, safe="/%"), quote(parts.query, safe="=&%+/:;,@$"), parts.fragment)
        )
        with urlopen(encoded_url, timeout=60) as response:
            return response.read()
    except HTTPError as exc:
        raise RuntimeError(f"PLM file download failed with HTTP {exc.code}") from exc
    except (URLError, TimeoutError) as exc:
        raise RuntimeError("PLM file download failed") from exc


def _first_token(response: Any) -> str:
    if not isinstance(response, dict):
        return ""
    if isinstance(response.get("data"), str):
        return response["data"]
    data = response.get("data") if isinstance(response.get("data"), dict) else {}
    return str(
        response.get("authorization")
        or response.get("accessToken")
        or response.get("token")
        or data.get("authorization")
        or data.get("accessToken")
        or data.get("token")
        or ""
    )


def _assert_business_ok(response: Any, label: str) -> None:
    if not isinstance(response, dict):
        return
    code = _integer(response.get("code"))
    if response.get("success") is False or (code is not None and code >= 400):
        message = response.get("msg") or response.get("errmsg") or code or "unknown error"
        raise RuntimeError(f"{label} failed: {message}")


def _response_rows(response: Any) -> list[dict[str, Any]]:
    if not isinstance(response, dict):
        return []
    data = response.get("data")
    if isinstance(data, list):
        return [row for row in data if isinstance(row, dict)]
    if not isinstance(data, dict):
        return []
    rows = None
    for key in ("list", "records", "rows"):
        if isinstance(data.get(key), list):
            rows = data[key]
            break
    return [row for row in rows if isinstance(row, dict)] if isinstance(rows, list) else []


def _has_next(response: Any, page_num: int, page_size: int, seen_count: int, batch_count: int) -> bool:
    if not isinstance(response, dict):
        return False
    data = response.get("data") if isinstance(response.get("data"), dict) else {}
    total = _integer(response.get("total") or data.get("total"))
    if total is not None:
        return seen_count < total
    for container in (response, response.get("data") if isinstance(response.get("data"), dict) else {}):
        if "hasNext" in container:
            return bool(container.get("hasNext"))
        if "has_next" in container:
            return bool(container.get("has_next"))
    return batch_count == page_size and page_num < 10_000


def _integer(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _valid_workbook(path: Path) -> bool:
    if not path.is_file():
        return False
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        workbook.close()
        return True
    except Exception:
        return False


def _valid_json_snapshot(path: Path) -> bool:
    if not path.is_file():
        return False
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        return isinstance(payload, dict) and isinstance(payload.get("rows"), list)
    except Exception:
        return False


def _validate_workbook_bytes(content: bytes) -> None:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        workbook.close()
    except Exception as exc:
        raise ValueError("PLM download is not a readable Excel workbook") from exc
