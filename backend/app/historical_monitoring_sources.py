from __future__ import annotations

import argparse
import json
import math
import re
import warnings
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from app.historical_watchlist import sha256_file

COUNTRY_ALIASES = {"菲律宾": "PH", "PH": "PH", "泰国": "TH", "TH": "TH", "越南": "VN", "VN": "VN"}
MARKET_SHEETS = ("PH精品", "TH精品", "VN精品")
MARKET_HEADER_ROW = 2
LISTING_SHEET = "精品流程"
LISTING_HEADER_ROW = 3
WEEK_METRIC_COLUMNS = ((13, 16), (20, 23), (27, 30), (34, 37))
WEEK_REVIEW_COLUMNS = ((17, 19), (24, 26), (31, 33), (38, 38))
ERROR_VALUES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}


def audit_historical_monitoring_sources(
    market_workbook: str | Path,
    listing_workbook: str | Path,
    finebi_candidates: str | Path | None = None,
) -> dict[str, Any]:
    market = audit_market_monitor(Path(market_workbook))
    finebi_keys = load_finebi_listing_keys(Path(finebi_candidates)) if finebi_candidates else set()
    listing = audit_listing_monitor(Path(listing_workbook), {record["main_sku"] for record in market["records"]}, finebi_keys)
    return {"market_monitor": market, "listing_monitor": listing}


def audit_market_monitor(path: Path) -> dict[str, Any]:
    source_sha256 = sha256_file(path)
    records: list[dict[str, Any]] = []
    excluded: Counter[str] = Counter()
    raw_row_count = 0
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet_name in MARKET_SHEETS:
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            worksheet.reset_dimensions()
            rows = list(worksheet.iter_rows(values_only=True))
            if len(rows) < MARKET_HEADER_ROW:
                continue
            headers = header_columns(rows[MARKET_HEADER_ROW - 1])
            for source_row, row in enumerate(rows[MARKET_HEADER_ROW:], start=MARKET_HEADER_ROW + 1):
                if not any(text_value(value) for value in row):
                    continue
                raw_row_count += 1
                product_bucket = value_by_header(row, headers, "货品") or ""
                if not product_bucket.startswith("开发新品"):
                    excluded[product_bucket or "空货品"] += 1
                    continue
                records.append(market_record(path, source_sha256, sheet_name, source_row, row, headers))
    finally:
        workbook.close()
    duplicate_keys = duplicate_records(records, ("country", "salesperson", "main_sku", "child_sku"))
    summary = {
        "raw_row_count": raw_row_count,
        "development_new_product_row_count": len(records),
        "excluded_non_development_row_count": sum(excluded.values()),
        "missing_secondary_conclusion_count": sum(record["secondary_conclusion"] is None for record in records),
        "missing_positioning_count": sum(record["positioning"] is None for record in records),
        "missing_target_daily_sales_count": sum(record["target_daily_sales"] is None for record in records),
        "missing_selling_points_count": sum(record["selling_points"] is None for record in records),
        "duplicate_key_count": len(duplicate_keys),
        "formula_error_count": sum(record["formula_error_count"] for record in records),
    }
    return {"summary": summary, "excluded_buckets": dict(sorted(excluded.items())), "duplicate_keys": duplicate_keys, "records": records}


def market_record(path: Path, source_sha256: str, sheet_name: str, source_row: int, row: tuple[Any, ...], headers: dict[str, int]) -> dict[str, Any]:
    country_raw = value_by_header(row, headers, "国家")
    return {
        "source_reference": {"source_file": path.name, "source_sha256": source_sha256, "source_sheet": sheet_name, "source_row": source_row},
        "country": normalize_country(country_raw),
        "country_raw": country_raw,
        "product_bucket": value_by_header(row, headers, "货品"),
        "salesperson": value_by_header(row, headers, "销售员"),
        "main_sku": normalize_sku(value_by_header(row, headers, "主SKU")),
        "child_sku": normalize_sku(value_by_header(row, headers, "子sku")),
        "product_name": value_by_header(row, headers, "商品名称"),
        "category_level1": value_by_header(row, headers, "一级类目"),
        "category_level2": value_by_header(row, headers, "二级类目"),
        "keyword": value_by_header(row, headers, "关键词"),
        "sales_note": value_by_header(row, headers, "销售备注"),
        "competitors": market_competitors(row, headers),
        "reference_daily_sales": number_value(value_by_header(row, headers, "竟对参考单销(=子sku月销/30)")),
        "reference_price": number_value(value_by_header(row, headers, "竟对参考售价")),
        "product_type": value_by_header(row, headers, "产品类型"),
        "stable_price": number_value(value_by_header_prefix(row, headers, "稳定期定价")),
        "claimed_daily_sales": number_value(value_by_header(row, headers, "认领单销")),
        "secondary_research_at": time_value(value_by_header(row, headers, "调研时间", occurrence=2)),
        "secondary_competitor_url": value_by_header_prefix(row, headers, "竞对链接") or value_by_header_prefix(row, headers, "锚定链接"),
        "secondary_conclusion": value_by_header(row, headers, "调研结论--价格/月销/市场趋势变化等，没有变化就填“无变化”-可以不用填竞对链接"),
        "positioning": value_by_header_prefix(row, headers, "产品定位"),
        "target_daily_sales": number_value(value_by_header(row, headers, "目标单销")),
        "selling_points": value_by_header(row, headers, "卖点总结"),
        "arrival_note": value_by_header(row, headers, "到货通知"),
        "formula_error_count": sum(1 for value in row if is_formula_error(value)),
    }


def market_competitors(row: tuple[Any, ...], headers: dict[str, int]) -> list[dict[str, Any]]:
    specs = [
        ("lowest_price", "平台综合推荐(前三页）最低价竞品链接1", "竞品单价（链接1）", "竞品子sku月销（链接1)"),
        ("most_orders", "平台综合推荐（前三页）月销量最多竞品链接2", "竞品单价（链接2）", "竞品子sku月销（链接2)"),
        ("new_arrival", "平台综合推荐（前三页近3个月上架的）新晋竞品链接3", "竞品单价（链接3）", "竞品子sku月销（链接3)"),
    ]
    result = []
    for kind, url_header, price_header, sales_header in specs:
        url = value_by_header(row, headers, url_header)
        price = number_value(value_by_header(row, headers, price_header))
        sales = number_value(value_by_header(row, headers, sales_header))
        if url or price is not None or sales is not None:
            result.append({"kind": kind, "url": url, "price": price, "monthly_sales": sales})
    return result


def audit_listing_monitor(path: Path, development_main_skus: set[str], finebi_keys: set[tuple[str | None, str | None, str | None, str | None]]) -> dict[str, Any]:
    source_sha256 = sha256_file(path)
    records: list[dict[str, Any]] = []
    anomalies: list[dict[str, Any]] = []
    raw_items = raw_numeric_column_values(path, LISTING_SHEET, "F")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=r"Cell .* is marked as a date but the serial value .* is outside the limits.*")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if LISTING_SHEET not in workbook.sheetnames:
            return {"summary": {"raw_row_count": 0}, "anomalies": [{"message": f"missing sheet {LISTING_SHEET}"}], "records": []}
        worksheet = workbook[LISTING_SHEET]
        worksheet.reset_dimensions()
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message=r"Cell .* is marked as a date but the serial value .* is outside the limits.*")
            rows = list(worksheet.iter_rows(values_only=True))
        for source_row, row in enumerate(rows[LISTING_HEADER_ROW:], start=LISTING_HEADER_ROW + 1):
            if not any(text_value(value) for value in row):
                continue
            record = listing_record(path, source_sha256, source_row, row, raw_items.get(source_row))
            records.append(record)
            if not valid_item(record["item"]):
                anomalies.append({"source_row": source_row, "item": record["item"], "message": "invalid item"})
    finally:
        workbook.close()
    matched_rows = [record for record in records if record["main_sku"] in development_main_skus]
    missing_week_metrics = [record for record in matched_rows if not record["has_any_week_metrics"]]
    finebi_fillable = [record for record in matched_rows if (record["country"], record["main_sku"], record["shop"], record["item"]) in finebi_keys]
    summary = {
        "raw_row_count": len(records),
        "matched_development_main_sku_count": len(matched_rows),
        "anomaly_item_count": len(anomalies),
        "missing_week_metric_rows": len(missing_week_metrics),
        "finebi_fillable_rows": len(finebi_fillable),
        "unique_item_count": len({record["item"] for record in records if valid_item(record["item"])}),
    }
    return {"summary": summary, "anomalies": anomalies, "records": records}


def listing_record(path: Path, source_sha256: str, source_row: int, row: tuple[Any, ...], raw_item: str | None = None) -> dict[str, Any]:
    country_raw = text_value(cell(row, 2))
    weeks = [week_record(index + 1, row, WEEK_METRIC_COLUMNS[index], WEEK_REVIEW_COLUMNS[index]) for index in range(4)]
    item = item_text(cell(row, 5))
    item_recovery = None
    if not valid_item(item) and raw_item:
        item = raw_item
        item_recovery = "raw_numeric_cell"
    return {
        "source_reference": {"source_file": path.name, "source_sha256": source_sha256, "source_sheet": LISTING_SHEET, "source_row": source_row},
        "operation_date": time_value(cell(row, 0)),
        "main_sku": normalize_sku(cell(row, 1)),
        "country": normalize_country(country_raw),
        "country_raw": country_raw,
        "salesperson": text_value(cell(row, 3)),
        "shop": text_value(cell(row, 4)),
        "item": item,
        "item_recovery": item_recovery,
        "strategy": {
            "optimization": text_value(cell(row, 6)),
            "brush_order": text_value(cell(row, 7)),
            "advertising": text_value(cell(row, 8)),
            "affiliate": text_value(cell(row, 9)),
            "campaign": text_value(cell(row, 10)),
            "other": text_value(cell(row, 11)),
        },
        "weeks": weeks,
        "has_any_week_metrics": any(week["has_metrics"] for week in weeks),
        "summary": text_value(cell(row, 37)),
    }


def week_record(week_number: int, row: tuple[Any, ...], metric_range: tuple[int, int], review_range: tuple[int, int]) -> dict[str, Any]:
    metric_values = [cell(row, column - 1) for column in range(metric_range[0], metric_range[1] + 1)]
    review_values = [cell(row, column - 1) for column in range(review_range[0], review_range[1] + 1)]
    return {
        "week_number": week_number,
        "orders": number_value(metric_values[0]) if len(metric_values) > 0 else None,
        "revenue": number_value(metric_values[1]) if len(metric_values) > 1 else None,
        "gross_profit": number_value(metric_values[2]) if len(metric_values) > 2 else None,
        "gross_margin": number_value(metric_values[3]) if len(metric_values) > 3 else None,
        "review_time": time_value(review_values[0]) if len(review_values) > 0 else None,
        "positioning": text_value(review_values[1]) if len(review_values) > 1 else None,
        "optimization_action": text_value(review_values[2]) if len(review_values) > 2 else None,
        "has_metrics": any(text_value(value) is not None for value in metric_values),
        "has_review": any(text_value(value) is not None for value in review_values),
    }


def load_finebi_listing_keys(path: Path) -> set[tuple[str | None, str | None, str | None, str | None]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    candidates = payload.get("candidates", []) if isinstance(payload, dict) else []
    return {
        (text_value(item.get("country")), normalize_sku(item.get("main_sku")), text_value(item.get("shop")), item_text(item.get("item_id")))
        for item in candidates
        if isinstance(item, dict)
    }


def header_columns(row: tuple[Any, ...]) -> dict[str, int]:
    counts: Counter[str] = Counter()
    result: dict[str, int] = {}
    for index, value in enumerate(row):
        header = normalize_header(value)
        if not header:
            continue
        counts[header] += 1
        key = header if counts[header] == 1 else f"{header}#{counts[header]}"
        result[key] = index
    return result


def value_by_header(row: tuple[Any, ...], headers: dict[str, int], header: str, occurrence: int = 1) -> Any:
    key = normalize_header(header)
    if occurrence > 1:
        key = f"{key}#{occurrence}"
    return text_value(cell(row, headers[key])) if key in headers else None


def value_by_header_prefix(row: tuple[Any, ...], headers: dict[str, int], prefix: str) -> Any:
    normalized_prefix = normalize_header(prefix)
    for header, index in headers.items():
        if header.split("#", 1)[0].startswith(normalized_prefix):
            return text_value(cell(row, index))
    return None


def raw_numeric_column_values(path: Path, sheet_name: str, column: str) -> dict[int, str]:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    document_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    try:
        with ZipFile(path) as archive:
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            sheet = next(
                item for item in workbook.findall(f"{{{main_ns}}}sheets/{{{main_ns}}}sheet")
                if item.get("name") == sheet_name
            )
            rel_id = sheet.get(f"{{{document_rel_ns}}}id")
            relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(
                item.get("Target") for item in relationships.findall(f"{{{package_rel_ns}}}Relationship")
                if item.get("Id") == rel_id
            )
            target = target.lstrip("/")
            if not target.startswith("xl/"):
                target = f"xl/{target}"
            sheet_xml = ElementTree.fromstring(archive.read(target))
    except (ElementTree.ParseError, KeyError, OSError, StopIteration, ValueError):
        return {}
    result: dict[int, str] = {}
    for item in sheet_xml.findall(f".//{{{main_ns}}}c"):
        reference = item.get("r") or ""
        if not re.fullmatch(rf"{re.escape(column)}(\d+)", reference) or item.get("t") not in {None, "n"}:
            continue
        value = item.find(f"{{{main_ns}}}v")
        if value is not None and value.text:
            result[int(reference[len(column):])] = item_text(value.text) or ""
    return result


def duplicate_records(records: list[dict[str, Any]], keys: tuple[str, ...]) -> list[dict[str, Any]]:
    groups: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for record in records:
        key = tuple(record.get(name) for name in keys)
        if all(key):
            groups.setdefault(key, []).append(record)
    return [
        {"key": "|".join(str(value) for value in key), "rows": [record["source_reference"] for record in rows]}
        for key, rows in sorted(groups.items())
        if len(rows) > 1
    ]


def normalize_header(value: object) -> str:
    return "".join((text_value(value) or "").replace("（", "(").replace("）", ")").split()).upper()


def normalize_country(value: object) -> str | None:
    text = text_value(value)
    if not text:
        return None
    upper = text.upper()
    return next((country for alias, country in COUNTRY_ALIASES.items() if alias.upper() in upper), None)


def normalize_sku(value: object) -> str | None:
    text = text_value(value)
    return text.upper() if text else None


def item_text(value: object) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        try:
            return format(Decimal(str(value)), "f").rstrip("0").rstrip(".") or "0"
        except InvalidOperation:
            return None
    return text_value(value)


def valid_item(value: object) -> bool:
    text = item_text(value)
    return bool(text and text.upper() not in ERROR_VALUES and not text.startswith("#"))


def number_value(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    text = text_value(value)
    if not text or text.upper() in ERROR_VALUES or text.startswith("=") or text in {"/", "-"}:
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def time_value(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return text_value(value)


def text_value(value: object) -> str | None:
    text = " ".join(str(value or "").replace("\xa0", " ").split())
    return text or None


def is_formula_error(value: object) -> bool:
    text = text_value(value)
    return bool(text and text.upper() in ERROR_VALUES)


def cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if 0 <= index < len(row) else None


def main() -> None:
    parser = argparse.ArgumentParser(description="Read-only historical market/listing monitor audit.")
    parser.add_argument("--market", required=True, type=Path)
    parser.add_argument("--listing", required=True, type=Path)
    parser.add_argument("--finebi-candidates", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = audit_historical_monitoring_sources(args.market, args.listing, args.finebi_candidates)
    encoded = json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded, encoding="utf-8")
    else:
        print(encoded, end="")


if __name__ == "__main__":
    main()
